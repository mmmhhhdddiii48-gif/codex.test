
def today_str():
    return datetime.now().strftime('%Y-%m-%d')

def sanitize_filename(name: str) -> str:
    bad = '<>:"/\\|?*\n\t'
    cleaned = ''.join('_' if ch in bad else ch for ch in str(name))
    cleaned = cleaned.strip().strip('.')
    return cleaned or 'report'
import sys, json, webbrowser, shutil, hashlib, os
from pathlib import Path
from datetime import datetime
from urllib.parse import quote
from copy import deepcopy
from openpyxl import load_workbook
from PySide6.QtWidgets import (
    QApplication, QWidget, QVBoxLayout, QHBoxLayout, QLabel, QLineEdit,
    QPushButton, QMessageBox, QGridLayout, QTableWidget, QTableWidgetItem,
    QDialog, QFormLayout, QSpinBox, QDoubleSpinBox, QHeaderView,
    QComboBox, QTextEdit, QTextBrowser, QDateEdit, QFrame, QAbstractItemView, QCheckBox, QFileDialog, QScrollArea, QTabWidget, QScrollArea,
    QListWidget, QListWidgetItem, QGraphicsOpacityEffect, QSizePolicy, QAbstractSpinBox, QStackedWidget,
    QGraphicsDropShadowEffect, QSlider, QToolButton
)
from PySide6.QtCore import Qt, QDate, QUrl, QTimer, QPoint, QEasingCurve, QPropertyAnimation, QParallelAnimationGroup, QLocale, Signal, QSizeF, QMarginsF
from PySide6.QtGui import QDesktopServices, QPixmap, QImage, QIcon, QFont, QPainter, QColor, QLinearGradient, QPainterPath, QTextDocument, QPageLayout, QPageSize
from PySide6.QtPrintSupport import QPrinter

def app_base_dir():
    if getattr(sys, 'frozen', False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def resource_base_dir():
    return Path(getattr(sys, '_MEIPASS', Path(__file__).resolve().parent))


BASE_DIR = app_base_dir()
RESOURCE_DIR = resource_base_dir()

def _can_write_dir(path: Path) -> bool:
    try:
        path.mkdir(parents=True, exist_ok=True)
        probe = path / '.nukhba_write_test'
        probe.write_text('ok', encoding='utf-8')
        probe.unlink(missing_ok=True)
        return True
    except Exception:
        return False


def _user_data_root() -> Path:
    app_name = 'Nukhba'
    if os.name == 'nt':
        base = os.environ.get('LOCALAPPDATA') or os.environ.get('APPDATA')
        if base:
            return Path(base) / app_name
    return Path.home() / f'.{app_name.lower()}'


def app_data_dir():
    """
    مسار بيانات التشغيل يجب أن يكون قابلاً للكتابة دائمًا.
    - في نسخة التطوير: نفضّل app_data بجانب المشروع إذا كان قابلاً للكتابة.
    - في النسخة المجمّعة أو عند تعذر الكتابة: نستخدم AppData الخاص بالمستخدم.
    """
    portable_dir = BASE_DIR / 'app_data'
    running_frozen = bool(getattr(sys, 'frozen', False))

    if (not running_frozen) and _can_write_dir(portable_dir):
        return portable_dir

    user_dir = _user_data_root() / 'app_data'
    user_dir.mkdir(parents=True, exist_ok=True)
    return user_dir


def data_dir():
    return app_data_dir()


APP_DATA_DIR = app_data_dir()


def safe_new_id(prefix='evt'):
    try:
        return new_id(prefix)
    except Exception:
        from datetime import datetime as _dt
        return f"{prefix}_{_dt.now().strftime('%Y%m%d%H%M%S%f')}"


def item_stock_rows(db):
    rows = []
    for row in db.get('items', []) or []:
        rows.append({
            'code': row.get('code', ''),
            'name': str(row.get('name', '') or '').strip(),
            'qty': float(row.get('qty', 0) or 0),
            'buy_price': float(row.get('buy_price', 0) or 0),
            'sell_price': float(row.get('sell_price', 0) or 0),
            'unit': row.get('unit', ''),
        })
    return rows
PASSWORD_FILE = APP_DATA_DIR / 'pass.txt'
DB_FILE = APP_DATA_DIR / 'data.json'
INVOICES_DIR = APP_DATA_DIR / 'invoices'
ATTACHMENTS_DIR = APP_DATA_DIR / 'attachments'
BACKUPS_DIR = APP_DATA_DIR / 'backups'
LOGO_PNG = RESOURCE_DIR / 'nokhba_logo.png'
APP_ICON = RESOURCE_DIR / 'nokhba_icon.ico'


def save_lines_pdf(title, lines, output_path: Path):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = QTextDocument()
    title_html = str(title).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')
    body_html = '<br>'.join(str(x).replace('&','&amp;').replace('<','&lt;').replace('>','&gt;') for x in lines)
    html_doc = f'''<html dir="rtl"><head><meta charset="utf-8"></head><body style="font-family:Arial; font-size:12pt; direction:rtl; text-align:right;"><h2 style="text-align:center;">{title_html}</h2><div style="white-space:pre-wrap; line-height:1.7;">{body_html}</div></body></html>'''
    doc.setHtml(html_doc)
    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(str(output_path))
    print_fn = getattr(doc, 'print', None) or getattr(doc, 'print_', None)
    if not callable(print_fn):
        raise RuntimeError('محرك PDF غير مدعوم في هذه البيئة')
    print_fn(printer)
    return output_path


def save_text_as_pdf(title, text_content: str, output_path: Path):
    return save_lines_pdf(title, str(text_content or '').splitlines(), output_path)

def ensure_single_data_location():
    """ننقل الداتا المبدئية مرة واحدة من المسارات المضمّنة إلى مسار المستخدم القابل للكتابة إذا لزم."""
    bundle_candidates = [
        RESOURCE_DIR / 'app_data',
        BASE_DIR / '_internal' / 'app_data',
        BASE_DIR / 'app_data',
    ]

    APP_DATA_DIR.mkdir(parents=True, exist_ok=True)

    selected = None
    for src_dir in bundle_candidates:
        try:
            if src_dir.exists():
                selected = src_dir
                break
        except Exception:
            pass

    if not selected:
        return

    files_to_seed = ['data.json', 'pass.txt']
    for name in files_to_seed:
        src = selected / name
        dst = APP_DATA_DIR / name
        try:
            if src.exists() and not dst.exists():
                dst.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(src, dst)
        except Exception:
            pass

    for sub in ['attachments', 'invoices', 'backups']:
        try:
            (APP_DATA_DIR / sub).mkdir(parents=True, exist_ok=True)
        except Exception:
            pass


def funders_tabs_style():
    return (
        f"QTabWidget::pane {{border:1px solid {BORDER}; background:{CARD}; border-radius:14px; margin-top:8px;}}"
        f"QTabBar::tab {{min-width:170px; min-height:42px; padding:8px 14px; margin:4px 6px; border-radius:10px; background:{DARK}; color:{TEXT}; font-weight:800;}}"
        f"QTabBar::tab:selected {{background:{ACCENT}; color:{TEXT_ON_ACCENT};}}"
    )


def style_funders_tabs(tab_widget):
    try:
        tab_widget.setDocumentMode(True)
    except Exception:
        pass
    tab_widget.setStyleSheet(funders_tabs_style())
    return tab_widget


def resolve_app_file(rel_path):
    if not rel_path:
        return None
    rel = Path(rel_path)
    if rel.is_absolute():
        return rel
    return APP_DATA_DIR / rel

def app_relative_path(path_value):
    try:
        return str(Path(path_value).resolve().relative_to(APP_DATA_DIR.resolve()))
    except Exception:
        return str(Path(path_value))


def apply_branding(widget):
    try:
        if APP_ICON.exists():
            widget.setWindowIcon(QIcon(str(APP_ICON)))
    except Exception:
        pass



def fix_date_edit_widget(widget):
    try:
        widget.setCalendarPopup(True)
        widget.setDisplayFormat('yyyy-MM-dd')
        widget.setLocale(QLocale(QLocale.Arabic, QLocale.Iraq))
        widget.setWrapping(False)
        widget.setKeyboardTracking(False)
        widget.setAccelerated(False)
        cal = widget.calendarWidget()
        if cal is not None:
            cal.setGridVisible(True)
            try:
                cal.setLocale(QLocale(QLocale.Arabic, QLocale.Iraq))
            except Exception:
                pass
            try:
                cal.setFirstDayOfWeek(Qt.Saturday)
            except Exception:
                pass
    except Exception:
        pass
    return widget



def tune_numeric_widget(widget):
    try:
        widget.setButtonSymbols(QAbstractSpinBox.NoButtons)
    except Exception:
        pass
    try:
        widget.setAlignment(Qt.AlignRight)
    except Exception:
        pass
    try:
        widget.setKeyboardTracking(False)
    except Exception:
        pass
    try:
        widget.setAccelerated(False)
    except Exception:
        pass
    try:
        widget.setMinimumHeight(46)
    except Exception:
        pass
    return widget

def make_logo_label(size=86):
    lbl = QLabel()
    lbl.setAlignment(Qt.AlignCenter)
    lbl.setFixedSize(size, size)
    if LOGO_PNG.exists():
        pm = QPixmap(str(LOGO_PNG))
        if not pm.isNull():
            lbl.setPixmap(pm.scaled(size, size, Qt.KeepAspectRatio, Qt.SmoothTransformation))
            return lbl
    lbl.setText('ن')
    lbl.setStyleSheet(f'font-size:{max(26, int(size*0.48))}px;font-weight:900;color:{TEXT};')
    return lbl

DEFAULT_DB = {
    'items': [],
    'customers': [],
    'suppliers': [],
    'funders': [],
    'inbound': [],
    'sales': [],
    'cash': [],
    'returns': [],
    'damaged': [],
    'inventory_ledger': [],
    'profit_distributions': [],
    'profit_events': [],
    'profit_entries': [],
    'expenses': [],
    'agents_custody': [],
    'reconciliation': {'current_actual_cash': 0.0, 'current_notes': '', 'records': []},
    'notifications_log': [],
    'operations_log': [],
    'opening': {
        'opening_cash': 0.0,
        'operating_cash': 0.0,
        'customers': [],
        'suppliers': [],
        'stock_items': [],
        'old_funders': [],
        'old_totals': {
            'old_sales_total': 0.0,
            'old_purchases_total': 0.0,
            'old_expenses_total': 0.0,
            'old_funders_withdrawals_total': 0.0,
            'old_hidab_withdrawals_total': 0.0,
            'old_mustafa_withdrawals_total': 0.0,
            'old_profit_clearing_total': 0.0
        },
        'start_date': '',
        'hidab_opening_balance': 0.0,
        'mustafa_opening_balance': 0.0,
        'opening_locked': False
    },
    'settings': {
        'company_name': 'مخزن النخبة',
        'company_phone': '',
        'company_address': '',
        'next_invoice_no': 1,
        'funders_profit_pool_pct': 33.3333,
        'owner_profit_pct': 50.0,
        'partner_profit_pct': 50.0,
        'low_stock_threshold': 5,
        'current_theme': 'dark_lux',
        'notification_sound_enabled': True,
        'toast_notifications_enabled': True,
        'critical_only_sound': False,
        'autosave_interval_sec': 60,
        'backup_keep_files': 20,
        'create_backup_on_manual_save': True,
        'create_backup_on_export': True,
        'autosave_backup_every': 5,
        'confirm_before_logout': True,
        'confirm_before_export': False,
        'confirm_before_app_exit': True,
    }
}

THEMES = {
    'dark_lux': {
        'name': 'داكن فاخر',
        'dark': '#08111f',
        'card': '#0f172a',
        'accent': '#6cd8ff',
        'accent2': '#7af0c1',
        'text': '#f8fafc',
        'border': '#334155',
        'muted': 'rgba(255,255,255,0.72)',
        'text_on_accent': '#ffffff',
    },
    'midnight_blue': {
        'name': 'ليلي هادئ',
        'dark': '#0f172a',
        'card': '#172033',
        'accent': '#2563eb',
        'accent2': '#3b82f6',
        'text': '#eff6ff',
        'border': '#334155',
        'muted': '#bfd3ff',
        'text_on_accent': '#ffffff',
    },
    'emerald_night': {
        'name': 'نعناعي مريح',
        'dark': '#0f1c17',
        'card': '#172720',
        'accent': '#0f9d73',
        'accent2': '#16c08d',
        'text': '#eefcf7',
        'border': '#2d4b40',
        'muted': '#b8ddd1',
        'text_on_accent': '#ffffff',
    },
    'light_clean': {
        'name': 'طبي فاتح',
        'dark': '#f4f7fb',
        'card': '#ffffff',
        'accent': '#2563eb',
        'accent2': '#3b82f6',
        'text': '#111827',
        'border': '#d0d7e2',
        'muted': '#5b6473',
        'text_on_accent': '#ffffff',
    },
    'sand_gold': {
        'name': 'رملي ذهبي',
        'dark': '#f7f2e8',
        'card': '#fffaf1',
        'accent': '#b7791f',
        'accent2': '#d69e2e',
        'text': '#2d2418',
        'border': '#d8c4a5',
        'muted': '#6d5d45',
        'text_on_accent': '#ffffff',
    },
    'rose_modern': {
        'name': 'وردي عصري',
        'dark': '#fff4f7',
        'card': '#ffffff',
        'accent': '#d9467a',
        'accent2': '#ec5f95',
        'text': '#2b1a22',
        'border': '#edc8d7',
        'muted': '#7a5362',
        'text_on_accent': '#ffffff',
    },
}

DARK = CARD = ACCENT = ACCENT2 = TEXT = BORDER = MUTED = TEXT_ON_ACCENT = ''
BUTTON_STYLE = SECONDARY_BUTTON = INPUT_STYLE = TABLE_STYLE = WINDOW_STYLE = CARD_FRAME_STYLE = ''


def build_styles(theme_key='dark_lux'):
    global DARK, CARD, ACCENT, ACCENT2, TEXT, BORDER, MUTED, TEXT_ON_ACCENT
    global BUTTON_STYLE, SECONDARY_BUTTON, INPUT_STYLE, TABLE_STYLE, WINDOW_STYLE, CARD_FRAME_STYLE
    theme = THEMES.get(theme_key, THEMES['dark_lux'])
    DARK = theme['dark']
    CARD = theme['card']
    ACCENT = theme['accent']
    ACCENT2 = theme['accent2']
    TEXT = theme['text']
    BORDER = theme['border']
    MUTED = theme['muted']
    TEXT_ON_ACCENT = theme.get('text_on_accent', '#ffffff')
    BUTTON_STYLE = f'''
    QPushButton {{
        background-color: {ACCENT};
        color: {TEXT_ON_ACCENT};
        font-size: 17px;
        font-weight: bold;
        border-radius: 18px;
        padding: 14px 16px;
        min-height: 22px;
        border: none;
    }}
    QPushButton:hover {{ background-color: {ACCENT2}; }}
    QPushButton:pressed {{ background-color: {ACCENT}; }}
    '''
    SECONDARY_BUTTON = f'''
    QPushButton {{
        background-color: {CARD};
        color: {TEXT};
        font-size: 15px;
        font-weight: 700;
        border: 1px solid {BORDER};
        border-radius: 12px;
        padding: 10px 14px;
        min-height: 20px;
    }}
    QPushButton:hover {{ background-color: {ACCENT2}; color: {TEXT_ON_ACCENT}; border: 1px solid {ACCENT2}; }}
    '''
    INPUT_STYLE = f'''
    QLineEdit, QComboBox, QDateEdit, QTextEdit, QSpinBox, QDoubleSpinBox {{
        background-color: {CARD};
        color: {TEXT};
        border: 1px solid {BORDER};
        border-radius: 8px;
        padding: 8px;
        font-size: 17px;
        min-height: 28px;
        padding: 14px 16px;
        selection-background-color: {ACCENT};
        selection-color: {TEXT_ON_ACCENT};
    }}
    QComboBox QAbstractItemView {{
        background-color: {CARD};
        color: {TEXT};
        border: 1px solid {BORDER};
        selection-background-color: {ACCENT};
        selection-color: {TEXT_ON_ACCENT};
    }}
    '''
    TABLE_STYLE = f'''
    QTableWidget {{
        background-color: {CARD};
        color: {TEXT};
        gridline-color: {BORDER};
        font-size: 14px;
        border: 1px solid {BORDER};
        alternate-background-color: {DARK};
    }}
    QHeaderView::section {{
        background-color: {ACCENT};
        color: {TEXT_ON_ACCENT};
        padding: 7px;
        border: none;
        font-weight: bold;
    }}
    '''
    WINDOW_STYLE = f'background-color: {DARK}; color: {TEXT};'
    CARD_FRAME_STYLE = f'QFrame{{background:{CARD};border:1px solid {BORDER};border-radius:16px;}}'
    return theme


build_styles('dark_lux')


def notification_severity(level):
    mapping = {
        'info': 1,
        'success': 2,
        'warning': 3,
        'danger': 4,
    }
    return mapping.get(str(level or 'info').lower(), 1)


def level_colors(level):
    mapping = {
        'success': ('#14532d', '#dcfce7', '#166534'),
        'warning': ('#78350f', '#fef3c7', '#b45309'),
        'danger': ('#7f1d1d', '#fee2e2', '#b91c1c'),
        'info': (ACCENT, TEXT_ON_ACCENT, ACCENT2),
    }
    return mapping.get(level, mapping['info'])


def fmt_datetime_text(value):
    if not value:
        return '—'
    return str(value).replace('-', '/').replace('T', ' ')


def fmt_relative_minutes(seconds_value):
    try:
        secs = max(0, int(seconds_value or 0))
    except Exception:
        secs = 0
    mins = secs // 60
    if mins <= 0:
        return 'الآن'
    if mins == 1:
        return 'منذ دقيقة'
    return f'منذ {mins} دقيقة'


def fmt_money(v):
    try:
        return f"{round(float(v)):,}"
    except Exception:
        return '0'

def fmt_pct(v):
    try:
        return f"{round(float(v))}%"
    except Exception:
        return '0%'

def rgba_from_hex(value, alpha):
    value = str(value).lstrip('#')
    if len(value) != 6:
        return f'rgba(255,255,255,{alpha})'
    r = int(value[0:2], 16)
    g = int(value[2:4], 16)
    b = int(value[4:6], 16)
    return f'rgba({r}, {g}, {b}, {alpha})'


def paint_app_background(widget, event=None):
    painter = QPainter(widget)
    painter.setRenderHint(QPainter.Antialiasing)
    rect = widget.rect()

    gradient = QLinearGradient(0, 0, rect.width(), rect.height())
    gradient.setColorAt(0.0, QColor(DARK))
    gradient.setColorAt(0.55, QColor(CARD))
    gradient.setColorAt(1.0, QColor(DARK))
    painter.fillRect(rect, gradient)

    glow_specs = [
        (rect.width() * 0.82, rect.height() * 0.18, max(120, int(rect.width() * 0.22)), ACCENT, 38),
        (rect.width() * 0.12, rect.height() * 0.78, max(100, int(rect.width() * 0.18)), ACCENT2, 28),
        (rect.width() * 0.55, rect.height() * 0.52, max(140, int(rect.width() * 0.16)), BORDER, 18),
    ]
    for cx, cy, radius, color_hex, alpha in glow_specs:
        path = QPainterPath()
        path.addEllipse(QPoint(int(cx), int(cy)), radius, radius)
        painter.fillPath(path, QColor(rgba_from_hex(color_hex, alpha / 100.0)))

    painter.setPen(QColor(rgba_from_hex(TEXT, 0.05)))
    step = max(42, int(min(rect.width(), rect.height()) / 14))
    for x in range(0, rect.width(), step):
        painter.drawLine(x, 0, x, rect.height())
    for y in range(0, rect.height(), step):
        painter.drawLine(0, y, rect.width(), y)


def now_str():
    return datetime.now().strftime('%Y-%m-%d %H:%M:%S')


def ensure_item_inventory_fields(item):
    qty = int(float(item.get('qty', 0) or 0))
    avg_cost = float(item.get('avg_cost', item.get('buy_price', 0) or 0) or 0)
    total_value = float(item.get('total_value', qty * avg_cost) or 0)
    if qty <= 0:
        qty = 0
        if total_value < 0:
            total_value = 0.0
    if qty > 0 and total_value <= 0 and avg_cost > 0:
        total_value = qty * avg_cost
    if qty > 0 and avg_cost <= 0 and total_value > 0:
        avg_cost = total_value / qty
    if qty > 0 and total_value > 0:
        avg_cost = total_value / qty
    item['qty'] = int(qty)
    item['avg_cost'] = round(max(0.0, avg_cost), 4)
    item['total_value'] = round(max(0.0, total_value), 2)
    item['buy_price'] = round(item['avg_cost'], 4)
    return item


def item_avg_cost(item):
    ensure_item_inventory_fields(item)
    return float(item.get('avg_cost', 0) or 0)


def item_inventory_value(item):
    ensure_item_inventory_fields(item)
    return float(item.get('total_value', 0) or 0)


def inventory_add_stock(item, qty, unit_cost):
    item = ensure_item_inventory_fields(item)
    qty = int(float(qty or 0))
    unit_cost = float(unit_cost or 0)
    if qty <= 0:
        return item
    item['total_value'] = round(float(item.get('total_value', 0) or 0) + (qty * unit_cost), 2)
    item['qty'] = int(item.get('qty', 0) or 0) + qty
    ensure_item_inventory_fields(item)
    return item


def inventory_remove_stock(item, qty, unit_cost=None):
    item = ensure_item_inventory_fields(item)
    qty = int(float(qty or 0))
    current_qty = int(item.get('qty', 0) or 0)
    if qty <= 0:
        return 0.0
    if qty > current_qty:
        raise ValueError('insufficient_stock')
    if unit_cost is None:
        unit_cost = item_avg_cost(item)
    unit_cost = float(unit_cost or 0)
    removed_value = round(qty * unit_cost, 2)
    item['qty'] = current_qty - qty
    item['total_value'] = round(max(0.0, float(item.get('total_value', 0) or 0) - removed_value), 2)
    ensure_item_inventory_fields(item)
    if item['qty'] == 0:
        item['total_value'] = 0.0
        item['avg_cost'] = 0.0
        item['buy_price'] = 0.0
    return removed_value


def add_inventory_movement(db, item, movement_type, qty_in=0, qty_out=0, unit_cost=0, sale_price=0, reference_type='', reference_id='', date='', notes='', movement_uid=''):
    db.setdefault('inventory_ledger', [])
    movement_uid = movement_uid or f"{movement_type}_{reference_id}_{item.get('name','')}"
    if any(x.get('movement_uid') == movement_uid for x in db.get('inventory_ledger', [])):
        return
    db['inventory_ledger'].append({
        'id': generate_id('mov'),
        'movement_uid': movement_uid,
        'date': date or datetime.now().strftime('%Y-%m-%d'),
        'item': item.get('name', ''),
        'item_code': item.get('code', ''),
        'unit': item.get('unit', ''),
        'movement_type': movement_type,
        'qty_in': int(float(qty_in or 0)),
        'qty_out': int(float(qty_out or 0)),
        'unit_cost': float(unit_cost or 0),
        'sale_price': float(sale_price or 0),
        'reference_type': reference_type,
        'reference_id': reference_id,
        'notes': notes or '',
        'created_at': now_str(),
    })


def remove_inventory_movement(db, movement_uid):
    db['inventory_ledger'] = [x for x in db.get('inventory_ledger', []) if x.get('movement_uid') != movement_uid]


def inventory_movements_for_item(db, item_name):
    rows = [x for x in db.get('inventory_ledger', []) if x.get('item') == item_name]
    def sort_key(row):
        return (row.get('date', ''), row.get('created_at', ''), row.get('movement_uid', ''))
    return sorted(rows, key=sort_key)


def ensure_inventory_baseline(db):
    db.setdefault('inventory_ledger', [])
    if db.get('_inventory_baseline_done'):
        return
    if db.get('inventory_ledger'):
        db['_inventory_baseline_done'] = True
        return
    for item in db.get('items', []):
        ensure_item_inventory_fields(item)
        qty = int(item.get('qty', 0) or 0)
        if qty <= 0:
            continue
        add_inventory_movement(
            db, item, 'opening', qty_in=qty, unit_cost=item_avg_cost(item),
            reference_type='opening', reference_id=item.get('code') or item.get('name',''),
            date=opening_data(db).get('start_date') or datetime.now().strftime('%Y-%m-%d'),
            notes='رصيد افتتاحي/حالي',
            movement_uid=f"opening_current_{item.get('code') or item.get('name','')}"
        )
    db['_inventory_baseline_done'] = True


def next_invoice_no(db):
    settings = db.setdefault('settings', {})
    current = int(settings.get('next_invoice_no', 1) or 1)
    settings['next_invoice_no'] = current + 1
    return current


def ensure_invoice_no(db, sale):
    if sale.get('invoice_no'):
        return sale['invoice_no']
    sale['invoice_no'] = int(db.setdefault('settings', {}).get('next_invoice_no', 1) or 1)
    db['settings']['next_invoice_no'] = sale['invoice_no'] + 1
    return sale['invoice_no']


def get_invoice_sales(db, sale):
    invoice_no = ensure_invoice_no(db, sale)
    group_id = sale.get('invoice_group_id')
    rows = []
    for s in db.get('sales', []):
        if group_id and s.get('invoice_group_id') == group_id:
            rows.append(s)
        elif not group_id and ensure_invoice_no(db, s) == invoice_no:
            rows.append(s)
    if not rows:
        rows = [sale]
    return rows


def invoice_html(db, sale):
    company = db.get('settings', {}).get('company_name', 'مخزن النخبة')
    phone = db.get('settings', {}).get('company_phone', '')
    address = db.get('settings', {}).get('company_address', '')
    invoice_no = ensure_invoice_no(db, sale)
    lines = get_invoice_sales(db, sale)
    sale0 = lines[0]
    customer = sale0.get('customer','')
    payment_type = sale0.get('payment_type','نقدي')
    total_amount = sum(float(x.get('total', 0)) for x in lines)
    current_due = current_sale_due(db, sale0)
    paid_amount = round(max(0.0, total_amount - current_due), 2)
    due_amount = round(current_due, 2)
    notes = sale0.get('notes', '') or '-'
    rows_html = []
    for i, line in enumerate(lines, 1):
        rows_html.append(f"""
        <tr>
          <td>{i}</td>
          <td>{line.get('item','')}</td>
          <td>{line.get('qty',0)}</td>
          <td>{fmt_money(line.get('unit_price',0))} د.ع</td>
          <td>{fmt_money(line.get('total',0))} د.ع</td>
        </tr>""")
    rows_html = ''.join(rows_html)
    return f"""<!doctype html>
<html lang='ar' dir='rtl'>
<head>
<meta charset='utf-8'>
<title>فاتورة بيع {invoice_no}</title>
<style>
body {{font-family: Tahoma, Arial, sans-serif; background:#eef2f7; color:#111827; margin:0; padding:24px;}}
.sheet {{max-width:950px; margin:auto; background:#fff; border:1px solid #d7dde6; border-radius:22px; padding:28px; box-shadow:0 14px 45px rgba(0,0,0,.08);}}
.header {{display:flex; justify-content:space-between; align-items:flex-start; gap:20px; border-bottom:3px solid #1f2937; padding-bottom:18px; margin-bottom:18px;}}
.brand h1 {{margin:0 0 8px; font-size:32px;}}
.brand .sub {{color:#4b5563; line-height:1.8;}}
.meta {{text-align:left;}}
.badge {{display:inline-block; background:#1f2937; color:#fff; padding:10px 18px; border-radius:999px; font-weight:bold;}}
.grid {{display:grid; grid-template-columns:repeat(3, 1fr); gap:14px; margin:20px 0;}}
.card {{background:#f9fafb; border:1px solid #e5e7eb; border-radius:14px; padding:14px;}}
.card .label {{color:#6b7280; font-size:13px; margin-bottom:4px;}}
.card .value {{font-size:18px; font-weight:bold;}}
.tablebox {{margin-top:18px; border:1px solid #d1d5db; border-radius:16px; overflow:hidden;}}
table {{width:100%; border-collapse:collapse;}}
th, td {{border-bottom:1px solid #e5e7eb; padding:13px 10px; text-align:center;}}
th {{background:#1f2937; color:#fff;}}
tr:last-child td {{border-bottom:none;}}
.totals {{margin-top:20px; width:420px; margin-right:auto; border:1px solid #d1d5db; border-radius:16px; overflow:hidden;}}
.totals td {{font-weight:bold; background:#fff;}}
.totals tr:nth-child(odd) td {{background:#f9fafb;}}
.note {{margin-top:22px; padding:14px; border:1px dashed #9ca3af; border-radius:12px; background:#fafafa;}}
.footer {{margin-top:24px; text-align:center; color:#6b7280; font-size:13px;}}
.printbar {{margin:0 auto 16px; max-width:950px; display:flex; justify-content:flex-start; gap:10px;}}
.btn {{background:#1f2937; color:#fff; border:none; border-radius:10px; padding:10px 14px; cursor:pointer; font-weight:bold;}}
@media print {{ body {{background:#fff; padding:0;}} .printbar {{display:none;}} .sheet {{box-shadow:none; border:none; padding:0;}} }}
</style>
</head>
<body>
<div class='printbar'><button class='btn' onclick='window.print()'>طباعة الفاتورة</button></div>
<div class='sheet'>
  <div class='header'>
    <div class='brand'>
      <h1>{company}</h1>
      <div class='sub'>{address}</div>
      <div class='sub'>{phone}</div>
    </div>
    <div class='meta'>
      <div class='badge'>فاتورة بيع #{invoice_no}</div>
      <div style='margin-top:10px;'>التاريخ: {sale0.get('date','')}</div>
      <div>وقت الإنشاء: {sale0.get('created_at','')}</div>
    </div>
  </div>
  <div class='grid'>
    <div class='card'><div class='label'>اسم الزبون</div><div class='value'>{customer}</div></div>
    <div class='card'><div class='label'>نوع الدفع</div><div class='value'>{payment_type}</div></div>
    <div class='card'><div class='label'>عدد الأصناف</div><div class='value'>{len(lines)}</div></div>
  </div>
  <div class='tablebox'>
  <table>
    <thead><tr><th>#</th><th>الصنف</th><th>الكمية</th><th>سعر الوحدة</th><th>المجموع</th></tr></thead>
    <tbody>{rows_html}</tbody>
  </table>
  </div>
  <table class='totals'>
    <tr><td>إجمالي الفاتورة</td><td>{fmt_money(total_amount)} د.ع</td></tr>
    <tr><td>المقبوض</td><td>{fmt_money(paid_amount)} د.ع</td></tr>
    <tr><td>الباقي</td><td>{fmt_money(due_amount)} د.ع</td></tr>
  </table>
  <div class='note'><strong>ملاحظات:</strong> {notes}</div>
  <div class='footer'>شكراً لتعاملكم مع {company}</div>
</div>
</body>
</html>"""


def write_invoice_file(db, sale):
    invoice_no = ensure_invoice_no(db, sale)
    INVOICES_DIR.mkdir(parents=True, exist_ok=True)
    path = INVOICES_DIR / f'invoice_{invoice_no}.html'
    path.write_text(invoice_html(db, sale), encoding='utf-8')
    return path


def show_invoice_dialog(parent, db, sale, print_after=False):
    invoice_no = ensure_invoice_no(db, sale)
    html_path = write_invoice_file(db, sale)
    try:
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(html_path.resolve())))
    except Exception as exc:
        QMessageBox.warning(parent, 'تعذر فتح الفاتورة', f'تعذر فتح ملف الفاتورة: {exc}')
        return
    if print_after:
        pdf_path = INVOICES_DIR / f'invoice_{invoice_no}.pdf'
        try:
            save_text_as_pdf(f'فاتورة بيع #{invoice_no}', html_path.read_text(encoding='utf-8'), pdf_path)
        except Exception:
            pass

def receipt_html(db, title, number, date, party, amount, method='نقد', account_no='', notes='', extra_rows=None):
    company = db.get('settings', {}).get('company_name', 'مخزن النخبة')
    phone = db.get('settings', {}).get('company_phone', '')
    address = db.get('settings', {}).get('company_address', '')
    extra_rows = extra_rows or []
    rows = [
        ('الطرف', party),
        ('التاريخ', date),
        ('طريقة الدفع', method),
        ('المبلغ', f"{fmt_money(amount)} د.ع"),
    ]
    if account_no:
        rows.append(('رقم الحساب', account_no))
    rows.extend(extra_rows)
    rows_html = ''.join([f"<tr><td>{k}</td><td>{v}</td></tr>" for k,v in rows])
    return f"""<!doctype html>
<html lang='ar' dir='rtl'>
<head>
<meta charset='utf-8'>
<title>{title} {number}</title>
<style>
body {{font-family: Tahoma, Arial, sans-serif; background:#eef2f7; color:#111827; margin:0; padding:24px;}}
.sheet {{max-width:820px; margin:auto; background:#fff; border:1px solid #d7dde6; border-radius:22px; padding:28px; box-shadow:0 14px 45px rgba(0,0,0,.08);}}
.header {{display:flex; justify-content:space-between; align-items:flex-start; gap:20px; border-bottom:3px solid #1f2937; padding-bottom:18px; margin-bottom:18px;}}
.brand h1 {{margin:0 0 8px; font-size:30px;}}
.brand .sub {{color:#4b5563; line-height:1.8;}}
.meta {{text-align:left;}}
.badge {{display:inline-block; background:#1f2937; color:#fff; padding:10px 18px; border-radius:999px; font-weight:bold;}}
.tablebox {{margin-top:18px; border:1px solid #d1d5db; border-radius:16px; overflow:hidden;}}
table {{width:100%; border-collapse:collapse;}}
th, td {{border-bottom:1px solid #e5e7eb; padding:13px 10px; text-align:center;}}
tr:last-child td {{border-bottom:none;}}
th {{background:#1f2937; color:#fff;}}
.note {{margin-top:22px; padding:14px; border:1px dashed #9ca3af; border-radius:12px; background:#fafafa;}}
.printbar {{margin:0 auto 16px; max-width:820px; display:flex; justify-content:flex-start; gap:10px;}}
.btn {{background:#1f2937; color:#fff; border:none; border-radius:10px; padding:10px 14px; cursor:pointer; font-weight:bold;}}
@media print {{ body {{background:#fff; padding:0;}} .printbar {{display:none;}} .sheet {{box-shadow:none; border:none; padding:0;}} }}
</style>
</head>
<body>
<div class='printbar'><button class='btn' onclick='window.print()'>طباعة / حفظ PDF</button></div>
<div class='sheet'>
  <div class='header'>
    <div class='brand'>
      <h1>{company}</h1>
      <div class='sub'>{address}</div>
      <div class='sub'>{phone}</div>
    </div>
    <div class='meta'>
      <div class='badge'>{title} #{number}</div>
      <div style='margin-top:10px;'>التاريخ: {date}</div>
      <div>وقت الإنشاء: {now_str()}</div>
    </div>
  </div>
  <div class='tablebox'>
    <table>
      <thead><tr><th>البيان</th><th>القيمة</th></tr></thead>
      <tbody>{rows_html}</tbody>
    </table>
  </div>
  <div class='note'><strong>ملاحظات:</strong> {notes or '-'}</div>
</div>
</body>
</html>"""


def write_receipt_file(db, receipt_type, number, title, date, party, amount, method='نقد', account_no='', notes='', extra_rows=None):
    INVOICES_DIR.mkdir(parents=True, exist_ok=True)
    safe_type = str(receipt_type).replace(' ', '_')
    path = INVOICES_DIR / f'{safe_type}_{number}.html'
    path.write_text(receipt_html(db, title, number, date, party, amount, method, account_no, notes, extra_rows), encoding='utf-8')
    return path




def build_customer_statement_data(db, customer_name):
    customer_name = str(customer_name or '').strip()
    snapshot = customer_receivable_snapshot(db, customer_name)
    old_debt = float(opening_customer_amount(db, customer_name) or 0)

    sales_rows = []
    sales_source = [s for s in db.get('sales', []) if s.get('customer') == customer_name]
    sales_source.sort(key=lambda s: (str(s.get('date', '') or ''), int(ensure_invoice_no(db, s) or 0), str(s.get('created_at', '') or '')))
    for s in sales_source:
        state = snapshot.get('sales_state', {}).get(s.get('id', ''), {})
        sales_rows.append({
            'date': str(s.get('date', '') or ''),
            'invoice_no': ensure_invoice_no(db, s),
            'item': str(s.get('item', '') or '—'),
            'qty': float(s.get('qty', 0) or 0),
            'unit_price': float(s.get('unit_price', 0) or 0),
            'total': float(s.get('total', 0) or 0),
            'initial_paid': float(s.get('paid_amount', 0) or 0),
            'allocated_paid': float(state.get('later_paid', 0) or 0),
            'remaining': float(state.get('remaining', max(0.0, float(s.get('total', 0) or 0) - float(s.get('paid_amount', 0) or 0))) or 0),
        })

    payment_rows = []
    for s in sales_rows:
        if s['initial_paid'] > 0:
            payment_rows.append({
                'date': s['date'],
                'amount': s['initial_paid'],
                'details': f"دفعة مع الفاتورة #{s['invoice_no']} - {s['item']}",
            })
    for p in db.get('cash', []):
        if p.get('party') != customer_name or p.get('source') not in ('customer_payment', 'opening_customer_payment'):
            continue
        details = str(p.get('notes', '') or ('تسديد رصيد سابق' if p.get('source') == 'opening_customer_payment' else 'مبلغ تسديد'))
        payment_rows.append({
            'date': str(p.get('date', '') or ''),
            'amount': float(p.get('amount', 0) or 0),
            'details': details,
        })
    payment_rows.sort(key=lambda r: (r['date'] or '9999-99-99', r['details']))

    for row in sales_rows:
        row['total_paid'] = round(row['initial_paid'] + row['allocated_paid'], 2)

    totals = {
        'sales_total': round(sum(r['total'] for r in sales_rows), 2),
        'initial_paid_total': round(sum(r['initial_paid'] for r in sales_rows), 2),
        'settled_total': round(sum(r['total_paid'] for r in sales_rows), 2),
        'remaining_total': round(sum(r['remaining'] for r in sales_rows) + float(snapshot.get('opening_remaining', 0) or 0), 2),
        'payments_total': round(sum(r['amount'] for r in payment_rows), 2),
        'remaining_opening': round(float(snapshot.get('opening_remaining', 0) or 0), 2),
        'old_debt': round(old_debt, 2),
        'total_payments': round(sum(r['amount'] for r in payment_rows), 2),
        'final_due': round(float(snapshot.get('final_due', 0) or 0), 2),
    }
    return {
        'customer_name': customer_name,
        'sales_rows': sales_rows,
        'payment_rows': payment_rows,
        'totals': totals,
    }

def customer_statement_html(db, customer_name):
    company = db.get('settings', {}).get('company_name', 'مخزن النخبة') or 'مخزن النخبة'
    data = build_customer_statement_data(db, customer_name)
    customer_name = data['customer_name']
    sales_rows = data['sales_rows']
    payment_rows = data['payment_rows']
    totals = data['totals']

    sales_html_parts = []
    for idx, row in enumerate(sales_rows, 1):
        qty_txt = str(int(row['qty'])) if float(row['qty']).is_integer() else f"{row['qty']:.2f}"
        sales_html_parts.append(f"""
        <tr>
          <td>{idx}</td>
          <td>{row['date'] or '—'}</td>
          <td class='item'>{row['item']}</td>
          <td>{qty_txt}</td>
          <td>{fmt_money(row['unit_price'])} د.ع</td>
          <td>{fmt_money(row['total'])} د.ع</td>
          <td>{fmt_money(row['initial_paid'])} د.ع</td>
          <td>{fmt_money(row['total_paid'])} د.ع</td>
          <td>{fmt_money(row['remaining'])} د.ع</td>
        </tr>""")
    if not sales_html_parts:
        sales_html_parts.append("<tr><td colspan='9'>لا توجد مواد أو فواتير محفوظة لهذا الزبون بعد.</td></tr>")
    sales_html_parts.append(f"""
        <tr class='totals-row'>
          <td colspan='5'>المجموع النهائي</td>
          <td>{fmt_money(totals['sales_total'])} د.ع</td>
          <td>{fmt_money(totals['initial_paid_total'])} د.ع</td>
          <td>{fmt_money(totals['settled_total'])} د.ع</td>
          <td>{fmt_money(totals['remaining_total'])} د.ع</td>
        </tr>""")
    sales_html = ''.join(sales_html_parts)

    payments_html_parts = []
    for idx, row in enumerate(payment_rows, 1):
        payments_html_parts.append(f"""
        <tr>
          <td>{idx}</td>
          <td>{fmt_money(row['amount'])} د.ع</td>
          <td>{row['date'] or '—'}</td>
          <td class='item'>{row['details'] or '—'}</td>
        </tr>""")
    if not payments_html_parts:
        payments_html_parts.append("<tr><td colspan='4'>لا توجد تسديدات محفوظة لهذا الزبون بعد.</td></tr>")
    payments_html_parts.append(f"""
        <tr class='totals-row'>
          <td colspan='1'>المجموع</td>
          <td>{fmt_money(totals['payments_total'])} د.ع</td>
          <td colspan='2'>المبلغ المتبقي النهائي: {fmt_money(totals['final_due'])} د.ع</td>
        </tr>""")
    payments_html = ''.join(payments_html_parts)

    return f"""<!doctype html>
<html lang='ar' dir='rtl'>
<head>
<meta charset='utf-8'>
<title>كشف حركة الزبون - {customer_name}</title>
<style>
body {{font-family: Tahoma, Arial, sans-serif; background:#eef2f7; color:#111827; margin:0; padding:24px;}}
.sheet {{max-width:1320px; margin:auto; background:#fff; border:1px solid #d7dde6; border-radius:22px; padding:28px; box-shadow:0 14px 45px rgba(0,0,0,.08);}}
.header {{display:flex; justify-content:space-between; align-items:center; gap:18px; border-bottom:3px solid #1f2937; padding-bottom:16px; margin-bottom:18px;}}
.brand {{font-size:34px; font-weight:900;}}
.meta {{display:flex; gap:28px; font-size:16px; font-weight:800; color:#243447; flex-wrap:wrap;}}
.meta span {{background:#f4f7fb; border:1px solid #d7dde6; border-radius:999px; padding:10px 16px;}}
.section-title {{margin:24px 0 12px; font-size:24px; font-weight:900; color:#10233d;}}
.tablebox {{border:1px solid #d1d5db; border-radius:16px; overflow:hidden;}}
table {{width:100%; border-collapse:collapse;}}
th, td {{border-bottom:1px solid #e5e7eb; padding:12px 8px; text-align:center;}}
th {{background:#1f2937; color:#fff; white-space:nowrap;}}
tr:last-child td {{border-bottom:none;}}
.item {{text-align:right; font-weight:700;}}
.totals-row td {{background:#eef4ff; color:#10233d; font-weight:900; border-top:2px solid #c9d6ea;}}
.printbar {{margin:0 auto 16px; max-width:1320px; display:flex; justify-content:flex-start; gap:10px;}}
.btn {{background:#1f2937; color:#fff; border:none; border-radius:10px; padding:10px 14px; cursor:pointer; font-weight:bold;}}
.note {{margin-top:20px; padding:14px; border:1px dashed #9ca3af; border-radius:12px; background:#fafafa; line-height:1.8;}}
@media print {{ body {{background:#fff; padding:0;}} .printbar {{display:none;}} .sheet {{box-shadow:none; border:none; padding:0;}} }}
</style>
</head>
<body>
<div class='printbar'><button class='btn' onclick='window.print()'>طباعة كشف الزبون</button></div>
<div class='sheet'>
  <div class='header'>
    <div class='brand'>{company}</div>
    <div class='meta'>
      <span>اسم الزبون: {customer_name or '—'}</span>
      <span>التاريخ: {today_str()}</span>
    </div>
  </div>

  <div class='section-title'>تفاصيل المواد</div>
  <div class='tablebox'>
    <table>
      <thead>
        <tr>
          <th>#</th><th>التاريخ</th><th>المادة</th><th>الكمية</th><th>سعر المفرد</th><th>المجموع الكلي</th><th>المبلغ المستلم</th><th>المبلغ المسدد الكلي</th><th>المبلغ المتبقي</th>
        </tr>
      </thead>
      <tbody>{sales_html}</tbody>
    </table>
  </div>

  <div class='section-title'>التسديدات</div>
  <div class='tablebox'>
    <table>
      <thead>
        <tr>
          <th>#</th><th>مبلغ التسديد</th><th>التاريخ</th><th>التفاصيل</th>
        </tr>
      </thead>
      <tbody>{payments_html}</tbody>
    </table>
  </div>

  <div class='note'>هذا الكشف يوضح تفاصيل المواد بشكل مبسط، مع قسم مستقل للتسديدات حتى يكون أوضح وأسهل للإرسال أو الطباعة.</div>
</div>
</body>
</html>"""


def write_customer_statement_file(db, customer_name):
    INVOICES_DIR.mkdir(parents=True, exist_ok=True)
    safe_name = sanitize_filename(customer_name)
    path = INVOICES_DIR / f'customer_statement_{safe_name}.html'
    path.write_text(customer_statement_html(db, customer_name), encoding='utf-8')
    return path


def customer_statement_share_html(db, customer_name):
    company = db.get('settings', {}).get('company_name', 'مخزن النخبة') or 'مخزن النخبة'
    data = build_customer_statement_data(db, customer_name)
    customer_name = data['customer_name']
    sales_rows = data['sales_rows']
    payment_rows = data['payment_rows']
    totals = data['totals']

    sales_html_parts = []
    for idx, row in enumerate(sales_rows, 1):
        qty_txt = str(int(row['qty'])) if float(row['qty']).is_integer() else f"{row['qty']:.2f}"
        sales_html_parts.append(f"""
        <tr>
          <td>{idx}</td>
          <td>{row['date'] or '—'}</td>
          <td class='item'>{row['item']}</td>
          <td>{qty_txt}</td>
          <td>{fmt_money(row['unit_price'])}</td>
          <td>{fmt_money(row['total'])}</td>
          <td>{fmt_money(row['initial_paid'])}</td>
          <td>{fmt_money(row['total_paid'])}</td>
          <td>{fmt_money(row['remaining'])}</td>
        </tr>
        """)
    if not sales_html_parts:
        sales_html_parts.append("<tr><td colspan='9'>لا توجد مواد أو فواتير محفوظة لهذا الزبون بعد.</td></tr>")
    sales_html_parts.append(f"""
        <tr class='totals-row'>
          <td colspan='5'>المجموع النهائي</td>
          <td>{fmt_money(totals['sales_total'])}</td>
          <td>{fmt_money(totals['initial_paid_total'])}</td>
          <td>{fmt_money(totals['settled_total'])}</td>
          <td>{fmt_money(totals['remaining_total'])}</td>
        </tr>
    """)
    sales_html = ''.join(sales_html_parts)

    payments_html_parts = []
    for idx, row in enumerate(payment_rows, 1):
        payments_html_parts.append(f"""
        <tr>
          <td>{idx}</td>
          <td>{fmt_money(row['amount'])}</td>
          <td>{row['date'] or '—'}</td>
          <td class='item'>{row['details'] or '—'}</td>
        </tr>
        """)
    if not payments_html_parts:
        payments_html_parts.append("<tr><td colspan='4'>لا توجد تسديدات محفوظة لهذا الزبون بعد.</td></tr>")
    payments_html_parts.append(f"""
        <tr class='totals-row'>
          <td>المجموع</td>
          <td>{fmt_money(totals['payments_total'])}</td>
          <td colspan='2'>المتبقي النهائي: {fmt_money(totals['final_due'])}</td>
        </tr>
    """)
    payments_html = ''.join(payments_html_parts)

    return f"""<!doctype html>
<html lang='ar' dir='rtl'>
<head>
<meta charset='utf-8'>
<title>كشف زبون للمشاركة - {customer_name}</title>
<style>
body {{font-family: Tahoma, Arial, sans-serif; background:#edf3fb; margin:0; padding:26px; color:#10233d;}}
.sheet {{max-width:1540px; margin:auto; background:#ffffff; border-radius:30px; overflow:hidden; box-shadow:0 24px 70px rgba(15,23,42,.14); border:1px solid #d7e3f1;}}
.hero {{background:linear-gradient(135deg,#214e9b 0%,#0f766e 52%,#16a34a 100%); color:#fff; padding:26px 30px 20px;}}
.hero-top {{display:flex; justify-content:space-between; align-items:center; gap:16px; flex-wrap:wrap;}}
.brand {{font-size:34px; font-weight:900;}}
.meta {{display:flex; gap:14px; flex-wrap:wrap; font-size:16px; font-weight:800;}}
.meta span {{background:rgba(255,255,255,.16); border:1px solid rgba(255,255,255,.28); padding:10px 16px; border-radius:999px;}}
.wrap {{padding:22px;}}
.section-title {{font-size:27px; font-weight:900; margin:0 0 12px; color:#17324d;}}
.tablebox {{background:#fff; border:1px solid #dbe7f3; border-radius:24px; overflow:hidden; box-shadow:0 8px 25px rgba(15,23,42,.05); margin-bottom:24px;}}
table {{width:100%; border-collapse:collapse; table-layout:auto;}}
th, td {{padding:14px 10px; text-align:center; border-bottom:1px solid #e8eef6; font-size:14px;}}
th {{background:#13243b; color:#fff; font-size:13px; font-weight:800; white-space:nowrap;}}
tr:nth-child(even) td {{background:#f8fbff;}}
.item {{text-align:right; font-weight:700; color:#1e3a5f;}}
.totals-row td {{background:#eef6ff !important; color:#10233d; font-weight:900; border-top:2px solid #c9d8ea;}}
.footer {{padding:0 22px 24px; color:#526277; font-size:13px; line-height:1.8;}}
</style>
</head>
<body>
<div class='sheet'>
  <div class='hero'>
    <div class='hero-top'>
      <div class='brand'>{company}</div>
      <div class='meta'>
        <span>اسم الزبون: {customer_name or '—'}</span>
        <span>التاريخ: {today_str()}</span>
      </div>
    </div>
  </div>
  <div class='wrap'>
    <div class='section-title'>تفاصيل المواد</div>
    <div class='tablebox'>
      <table>
        <thead>
          <tr>
            <th>#</th><th>التاريخ</th><th>المادة</th><th>الكمية</th><th>سعر المفرد</th><th>المجموع الكلي</th><th>المبلغ المستلم</th><th>المبلغ المسدد الكلي</th><th>المبلغ المتبقي</th>
          </tr>
        </thead>
        <tbody>{sales_html}</tbody>
      </table>
    </div>

    <div class='section-title'>التسديدات</div>
    <div class='tablebox'>
      <table>
        <thead>
          <tr>
            <th>#</th><th>مبلغ التسديد</th><th>التاريخ</th><th>التفاصيل</th>
          </tr>
        </thead>
        <tbody>{payments_html}</tbody>
      </table>
    </div>
  </div>
  <div class='footer'>نسخة مشاركة مرتبة وواضحة للواتساب، مقسومة بين تفاصيل المواد والتسديدات حتى تبين للزبون بشكل مباشر ومفهوم.</div>
</div>
</body>
</html>"""


def agent_custody_report_html(db, agent_name):
    company = db.get('settings', {}).get('company_name', 'مخزن النخبة') or 'مخزن النخبة'
    agent_name = str(agent_name or '').strip()
    rows = [x for x in agents_custody_rows(db) if (x.get('agent', '') or '').strip() == agent_name]
    rows.sort(key=lambda x: ((x.get('date', '') or ''), (x.get('created_at', '') or '')))
    incoming = transfers = expenses = settlements = balance = 0.0
    incoming_rows = []
    outgoing_rows = []
    settlement_parts = []

    for row in rows:
        amount = float(row.get('amount', 0) or 0)
        effect = agent_custody_effect(row)
        typ = (row.get('type', '') or '').strip()
        direction = (row.get('settlement_direction', '') or '').strip() or '—'
        party = row.get('party', '') or '—'
        notes = row.get('notes', '') or '—'
        date_val = row.get('date', '') or '—'

        if typ == 'وارد مستلم':
            incoming += amount
            incoming_rows.append({
                'date': date_val,
                'party': party,
                'amount': amount,
                'notes': notes,
            })
        elif typ == 'تحويل':
            transfers += amount
            outgoing_rows.append({
                'date': date_val,
                'type': typ,
                'party': party,
                'amount': amount,
                'notes': notes,
            })
        elif typ == 'مصروف':
            expenses += amount
            outgoing_rows.append({
                'date': date_val,
                'type': typ,
                'party': party,
                'amount': amount,
                'notes': notes,
            })
        elif typ == 'معالجة فرق':
            settlements += amount if (row.get('settlement_direction', '') or '').strip() == 'لصالح العهدة' else -amount
            settlement_parts.append(f"""
            <tr>
              <td>{len(settlement_parts)+1}</td>
              <td>{fmt_money(amount)}</td>
              <td>{date_val}</td>
              <td>{direction}</td>
              <td class='notes'>{notes}</td>
            </tr>
            """)

        balance += effect

    incoming_parts = []
    for idx, row in enumerate(incoming_rows, 1):
        incoming_parts.append(f"""
        <tr>
          <td>{idx}</td>
          <td>{row['date']}</td>
          <td class='item'>{row['party']}</td>
          <td>{fmt_money(row['amount'])}</td>
          <td class='notes'>{row['notes']}</td>
        </tr>
        """)
    if not incoming_parts:
        incoming_parts.append("<tr><td colspan='5'>لا توجد حركات وارد محفوظة لهذا المندوب بعد.</td></tr>")
    incoming_parts.append(f"""
        <tr class='totals-row'>
          <td colspan='3'>مجموع الوارد</td>
          <td>{fmt_money(incoming)}</td>
          <td>عدد الحركات: {len(incoming_rows)}</td>
        </tr>
    """)

    outgoing_total = transfers + expenses
    outgoing_parts = []
    for idx, row in enumerate(outgoing_rows, 1):
        outgoing_parts.append(f"""
        <tr>
          <td>{idx}</td>
          <td>{row['date']}</td>
          <td>{row['type']}</td>
          <td class='item'>{row['party']}</td>
          <td>{fmt_money(row['amount'])}</td>
          <td class='notes'>{row['notes']}</td>
        </tr>
        """)
    if not outgoing_parts:
        outgoing_parts.append("<tr><td colspan='6'>لا توجد حركات صادرة محفوظة لهذا المندوب بعد.</td></tr>")
    outgoing_parts.append(f"""
        <tr class='totals-row'>
          <td colspan='4'>مجموع الصادر</td>
          <td>{fmt_money(outgoing_total)}</td>
          <td>تحويلات: {fmt_money(transfers)} | مصاريف: {fmt_money(expenses)}</td>
        </tr>
    """)

    if not settlement_parts:
        settlement_parts.append("<tr><td colspan='5'>لا توجد تسويات محفوظة لهذا المندوب بعد.</td></tr>")
    settlement_parts.append(f"""
        <tr class='totals-row'>
          <td>المجموع</td>
          <td>{fmt_money(abs(settlements))}</td>
          <td colspan='3'>صافي التسويات: {fmt_money(settlements)}</td>
        </tr>
    """)

    incoming_html = ''.join(incoming_parts)
    outgoing_html = ''.join(outgoing_parts)
    settlements_html = ''.join(settlement_parts)
    final_balance = round(balance, 2)
    return f"""<!doctype html>
<html lang='ar' dir='rtl'>
<head>
<meta charset='utf-8'>
<title>كشف عهدة المندوب - {agent_name}</title>
<style>
body {{font-family: Tahoma, Arial, sans-serif; background:#edf3fb; margin:0; padding:26px; color:#10233d;}}
.sheet {{max-width:1540px; margin:auto; background:#ffffff; border-radius:30px; overflow:hidden; box-shadow:0 24px 70px rgba(15,23,42,.14); border:1px solid #d7e3f1;}}
.hero {{background:linear-gradient(135deg,#214e9b 0%,#0f766e 52%,#16a34a 100%); color:#fff; padding:26px 30px 20px;}}
.hero-top {{display:flex; justify-content:space-between; align-items:center; gap:16px; flex-wrap:wrap;}}
.brand {{font-size:34px; font-weight:900;}}
.meta {{display:flex; gap:14px; flex-wrap:wrap; font-size:16px; font-weight:800;}}
.meta span {{background:rgba(255,255,255,.16); border:1px solid rgba(255,255,255,.28); padding:10px 16px; border-radius:999px;}}
.wrap {{padding:22px;}}
.section-title {{font-size:27px; font-weight:900; margin:0 0 12px; color:#17324d;}}
.tablebox {{background:#fff; border:1px solid #dbe7f3; border-radius:24px; overflow:hidden; box-shadow:0 8px 25px rgba(15,23,42,.05); margin-bottom:24px;}}
table {{width:100%; border-collapse:collapse; table-layout:auto;}}
th, td {{padding:14px 10px; text-align:center; border-bottom:1px solid #e8eef6; font-size:14px;}}
th {{background:#13243b; color:#fff; font-size:13px; font-weight:800; white-space:nowrap;}}
tr:nth-child(even) td {{background:#f8fbff;}}
.item {{text-align:right; font-weight:700; color:#1e3a5f;}}
.notes {{text-align:right; font-weight:700; color:#294766;}}
.totals-row td {{background:#eef6ff !important; color:#10233d; font-weight:900; border-top:2px solid #c9d8ea;}}
.summary-box {{margin-top:12px; background:linear-gradient(135deg,#f7fbff 0%,#ecf7f5 100%); border:1px solid #dbe7f3; border-radius:22px; padding:14px 18px; box-shadow:0 8px 25px rgba(15,23,42,.05);}}
.summary-title {{font-size:20px; font-weight:900; margin:0 0 10px; color:#17324d;}}
.summary-row {{display:flex; justify-content:space-between; align-items:center; gap:16px; padding:8px 0; border-bottom:1px dashed #cfdceb; font-size:15px; font-weight:800;}}
.summary-row:last-child {{border-bottom:none; color:#0f5132;}}
.summary-value {{display:inline-block; min-width:190px; text-align:center; background:#ffffff; border:1px solid #cfe0ee; border-radius:14px; padding:8px 14px; font-size:17px; font-weight:900; color:#17324d; box-shadow:0 6px 18px rgba(15,23,42,.06);}}
.summary-row:last-child .summary-value {{background:#ecfdf3; border-color:#b7e3c8; color:#0f5132;}}
.footer {{padding:0 22px 24px; color:#526277; font-size:13px; line-height:1.8;}}
</style>
</head>
<body>
<div class='sheet'>
  <div class='hero'>
    <div class='hero-top'>
      <div class='brand'>{company}</div>
      <div class='meta'>
        <span>اسم المندوب: {agent_name or '—'}</span>
        <span>التاريخ: {today_str()}</span>
      </div>
    </div>
  </div>
  <div class='wrap'>
    <div class='section-title'>الوارد</div>
    <div class='tablebox'>
      <table>
        <thead>
          <tr>
            <th>#</th><th>التاريخ</th><th>الجهة</th><th>المبلغ الوارد</th><th>التفاصيل</th>
          </tr>
        </thead>
        <tbody>{incoming_html}</tbody>
      </table>
    </div>

    <div class='section-title'>الصادر</div>
    <div class='tablebox'>
      <table>
        <thead>
          <tr>
            <th>#</th><th>التاريخ</th><th>النوع</th><th>الجهة</th><th>المبلغ الصادر</th><th>التفاصيل</th>
          </tr>
        </thead>
        <tbody>{outgoing_html}</tbody>
      </table>
    </div>

    <div class='section-title'>التسويات</div>
    <div class='tablebox'>
      <table>
        <thead>
          <tr>
            <th>#</th><th>مبلغ التسوية</th><th>التاريخ</th><th>اتجاه المعالجة</th><th>التفاصيل</th>
          </tr>
        </thead>
        <tbody>{settlements_html}</tbody>
      </table>
    </div>

    <div class='summary-box'>
      <div class='summary-title'>الملخص النهائي</div>
      <div class='summary-row'><span>مجموع الوارد</span><span class='summary-value'>{fmt_money(incoming)} د.ع</span></div>
      <div class='summary-row'><span>مجموع الصادر</span><span class='summary-value'>{fmt_money(outgoing_total)} د.ع</span></div>
      <div class='summary-row'><span>المتبقي بالعهدة</span><span class='summary-value'>{fmt_money(final_balance)} د.ع</span></div>
    </div>
  </div>
  <div class='footer'>الكشف مفصول بين الوارد والصادر والتسويات حتى يكون واضحًا عند الإرسال والمراجعة، والملخص النهائي بالأسفل يوضح المتبقي الحالي بالعهدة بشكل مباشر.</div>
</div>
</body>
</html>"""


def _render_html_to_pdf(html, output_path, page_width_mm=297, landscape=False):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = QTextDocument()
    doc.setHtml(html)
    printer = QPrinter(QPrinter.HighResolution)
    printer.setOutputFormat(QPrinter.PdfFormat)
    printer.setOutputFileName(str(output_path))
    printer.setPageMargins(QMarginsF(12, 12, 12, 12), QPageLayout.Millimeter)
    printer.setPageSize(QPageSize(QPageSize.A4))
    printer.setPageOrientation(QPageLayout.Landscape if landscape else QPageLayout.Portrait)
    page_rect = printer.pageRect(QPrinter.Point)
    doc.setPageSize(QSizeF(page_rect.width(), page_rect.height()))
    doc.print_(printer)
    return output_path


def _render_html_to_image(html, output_path, width=1400, background='#ffffff'):
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc = QTextDocument()
    doc.setHtml(html)
    doc.setTextWidth(width)
    size = doc.size().toSize()
    img = QImage(max(width, size.width() + 24), max(800, size.height() + 24), QImage.Format_ARGB32)
    img.fill(QColor(background))
    painter = QPainter(img)
    doc.drawContents(painter)
    painter.end()
    img.save(str(output_path))
    return output_path


def export_customer_statement_pdf(db, customer_name, output_path=None):
    safe_name = sanitize_filename(customer_name)
    if output_path is None:
        output_path = INVOICES_DIR / f'customer_statement_{safe_name}.pdf'
    return _render_html_to_pdf(customer_statement_html(db, customer_name), output_path)


def export_customer_statement_share_image(db, customer_name, output_path=None):
    safe_name = sanitize_filename(customer_name)
    if output_path is None:
        output_path = INVOICES_DIR / f'customer_statement_share_{safe_name}.png'
    return _render_html_to_image(customer_statement_share_html(db, customer_name), output_path, width=1380, background='#f4f7fb')


def write_agent_custody_report_file(db, agent_name):
    reports_dir = data_dir() / 'reports'
    reports_dir.mkdir(parents=True, exist_ok=True)
    safe_name = sanitize_filename(agent_name)
    path = reports_dir / f'agent_custody_report_{safe_name}.html'
    path.write_text(agent_custody_report_html(db, agent_name), encoding='utf-8')
    return path


def export_agent_custody_report_pdf(db, agent_name, output_path=None):
    reports_dir = data_dir() / 'reports'
    reports_dir.mkdir(parents=True, exist_ok=True)
    safe_name = sanitize_filename(agent_name)
    if output_path is None:
        output_path = reports_dir / f'agent_custody_report_{safe_name}.pdf'
    return _render_html_to_pdf(agent_custody_report_html(db, agent_name), output_path, landscape=True)


def export_agent_custody_report_share_image(db, agent_name, output_path=None):
    reports_dir = data_dir() / 'reports'
    reports_dir.mkdir(parents=True, exist_ok=True)
    safe_name = sanitize_filename(agent_name)
    if output_path is None:
        output_path = reports_dir / f'agent_custody_report_share_{safe_name}.png'
    return _render_html_to_image(agent_custody_report_html(db, agent_name), output_path, width=1380, background='#eef2f7')


class AgentCustodyReportPreviewDialog(QDialog):
    def __init__(self, db, agent_name, parent=None):
        super().__init__(parent)
        self.db = db
        self.agent_name = str(agent_name or '').strip()
        self.reports_dir = data_dir() / 'reports'
        self.reports_dir.mkdir(parents=True, exist_ok=True)
        self.setWindowTitle(f'كشف عهدة المندوب - {self.agent_name}')
        self.setWindowFlag(Qt.WindowMaximizeButtonHint, True)
        self.resize(1220, 860)
        self.setLayoutDirection(Qt.RightToLeft)
        root = QVBoxLayout(self)
        root.setContentsMargins(14, 14, 14, 14)
        root.setSpacing(10)

        actions = QHBoxLayout()
        self.btn_open_html = QPushButton('🌐 فتح HTML')
        self.btn_pdf = QPushButton('💾 حفظ PDF')
        self.btn_image = QPushButton('🖼 حفظ صورة واتساب')
        self.btn_folder = QPushButton('📂 فتح مجلد الحفظ')
        self.btn_whatsapp = QPushButton('🟢 فتح واتساب ويب')
        self.btn_close = QPushButton('إغلاق')
        for btn in [self.btn_open_html, self.btn_pdf, self.btn_image, self.btn_folder, self.btn_whatsapp, self.btn_close]:
            btn.setMinimumHeight(42)
            actions.addWidget(btn)
        root.addLayout(actions)

        hint = QLabel('اختر الحفظ كـ PDF للأرشفة والطباعة، أو احفظ صورة واتساب مرتبة حتى ترسل كشف عهدة المندوب مباشرة.')
        hint.setWordWrap(True)
        hint.setAlignment(Qt.AlignRight)
        root.addWidget(hint)

        self.browser = QTextBrowser()
        self.browser.setOpenExternalLinks(True)
        self.browser.setHtml(agent_custody_report_html(self.db, self.agent_name))
        root.addWidget(self.browser, 1)

        self.btn_open_html.clicked.connect(self.open_html)
        self.btn_pdf.clicked.connect(self.save_pdf)
        self.btn_image.clicked.connect(self.save_share_image)
        self.btn_folder.clicked.connect(self.open_folder)
        self.btn_whatsapp.clicked.connect(self.open_whatsapp)
        self.btn_close.clicked.connect(self.accept)

    def open_html(self):
        path = write_agent_custody_report_file(self.db, self.agent_name)
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(path.resolve())))

    def save_pdf(self):
        safe_name = sanitize_filename(self.agent_name)
        default_path = str((self.reports_dir / f'agent_custody_report_{safe_name}.pdf').resolve())
        path, _ = QFileDialog.getSaveFileName(self, 'حفظ كشف عهدة المندوب PDF', default_path, 'PDF Files (*.pdf)')
        if not path:
            return
        try:
            saved = export_agent_custody_report_pdf(self.db, self.agent_name, path)
            QMessageBox.information(self, 'تم الحفظ', f"تم حفظ ملف PDF بنجاح\n{saved}")
        except Exception as exc:
            QMessageBox.warning(self, 'تعذر الحفظ', f"تعذر حفظ ملف PDF\n{exc}")

    def save_share_image(self):
        safe_name = sanitize_filename(self.agent_name)
        default_path = str((self.reports_dir / f'agent_custody_report_share_{safe_name}.png').resolve())
        path, _ = QFileDialog.getSaveFileName(self, 'حفظ صورة واتساب', default_path, 'PNG Images (*.png)')
        if not path:
            return
        try:
            saved = export_agent_custody_report_share_image(self.db, self.agent_name, path)
            QMessageBox.information(self, 'تم الحفظ', f"تم حفظ صورة واتساب بنجاح\n{saved}")
        except Exception as exc:
            QMessageBox.warning(self, 'تعذر الحفظ', f"تعذر حفظ صورة واتساب\n{exc}")

    def open_folder(self):
        self.reports_dir.mkdir(parents=True, exist_ok=True)
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(self.reports_dir.resolve())))

    def open_whatsapp(self):
        webbrowser.open('https://web.whatsapp.com/')
        QMessageBox.information(self, 'واتساب ويب', 'فتحنا واتساب ويب. احفظ صورة الكشف أو الـ PDF ثم أرفق الملف من واتساب.')


class CustomerStatementDialog(QDialog):
    def __init__(self, parent, db, customer_name):
        super().__init__(parent)
        self.db = db
        self.customer_name = str(customer_name or '').strip()
        self.setWindowTitle(f'كشف الزبون - {self.customer_name}')
        self.resize(1220, 860)
        self.setLayoutDirection(Qt.RightToLeft)
        root = QVBoxLayout(self)
        root.setContentsMargins(14, 14, 14, 14)
        root.setSpacing(10)

        actions = QHBoxLayout()
        self.btn_open_html = QPushButton('🌐 فتح HTML')
        self.btn_pdf = QPushButton('💾 حفظ PDF')
        self.btn_image = QPushButton('🖼 حفظ صورة واتساب')
        self.btn_folder = QPushButton('📂 فتح مجلد الحفظ')
        self.btn_whatsapp = QPushButton('🟢 فتح واتساب ويب')
        self.btn_close = QPushButton('إغلاق')
        for btn in [self.btn_open_html, self.btn_pdf, self.btn_image, self.btn_folder, self.btn_whatsapp, self.btn_close]:
            btn.setMinimumHeight(42)
            actions.addWidget(btn)
        root.addLayout(actions)

        hint = QLabel('اختر الحفظ كـ PDF للأرشفة والطباعة، أو احفظ صورة واتساب ملونة ومرتبة حتى ترسلها مباشرة للزبون.')
        hint.setWordWrap(True)
        hint.setAlignment(Qt.AlignRight)
        root.addWidget(hint)

        self.browser = QTextBrowser()
        self.browser.setOpenExternalLinks(True)
        self.browser.setHtml(customer_statement_html(self.db, self.customer_name))
        root.addWidget(self.browser, 1)

        self.btn_open_html.clicked.connect(self.open_html)
        self.btn_pdf.clicked.connect(self.save_pdf)
        self.btn_image.clicked.connect(self.save_share_image)
        self.btn_folder.clicked.connect(self.open_folder)
        self.btn_whatsapp.clicked.connect(self.open_whatsapp)
        self.btn_close.clicked.connect(self.accept)

    def open_html(self):
        path = write_customer_statement_file(self.db, self.customer_name)
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(path.resolve())))

    def save_pdf(self):
        safe_name = sanitize_filename(self.customer_name)
        default_path = str((INVOICES_DIR / f'customer_statement_{safe_name}.pdf').resolve())
        path, _ = QFileDialog.getSaveFileName(self, 'حفظ كشف الزبون PDF', default_path, 'PDF Files (*.pdf)')
        if not path:
            return
        try:
            saved = export_customer_statement_pdf(self.db, self.customer_name, path)
            QMessageBox.information(self, 'تم الحفظ', f'تم حفظ ملف PDF بنجاح\n{saved}')
        except Exception as exc:
            QMessageBox.warning(self, 'تعذر الحفظ', f'تعذر حفظ ملف PDF\n{exc}')

    def save_share_image(self):
        safe_name = sanitize_filename(self.customer_name)
        default_path = str((INVOICES_DIR / f'customer_statement_share_{safe_name}.png').resolve())
        path, _ = QFileDialog.getSaveFileName(self, 'حفظ صورة واتساب', default_path, 'PNG Images (*.png)')
        if not path:
            return
        try:
            saved = export_customer_statement_share_image(self.db, self.customer_name, path)
            QMessageBox.information(self, 'تم الحفظ', f'تم حفظ صورة واتساب بنجاح\n{saved}')
        except Exception as exc:
            QMessageBox.warning(self, 'تعذر الحفظ', f'تعذر حفظ صورة واتساب\n{exc}')

    def open_folder(self):
        INVOICES_DIR.mkdir(parents=True, exist_ok=True)
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(INVOICES_DIR.resolve())))

    def open_whatsapp(self):
        webbrowser.open('https://web.whatsapp.com/')
        QMessageBox.information(self, 'واتساب ويب', 'فتحنا واتساب ويب. احفظ صورة الكشف أو الـ PDF ثم أرفق الملف من واتساب.')


def next_receipt_no(db, bucket='next_receipt_no'):
    settings = db.setdefault('settings', {})
    current = int(settings.get(bucket, 1) or 1)
    settings[bucket] = current + 1
    return current



def customer_receivable_snapshot(db, customer_name=''):
    customer_name = str(customer_name or '').strip()
    opening_total = opening_customer_amount(db, customer_name)
    opening_remaining = round(max(0.0, opening_total), 2)

    sales_source = [s for s in db.get('sales', []) if s.get('customer') == customer_name]
    sales_source.sort(key=lambda s: (str(s.get('date', '') or ''), ensure_invoice_no(db, s), str(s.get('created_at', '') or ''), str(s.get('id', '') or '')))

    events = []
    for s in sales_source:
        events.append({
            'kind': 'sale',
            'date': str(s.get('date', '') or ''),
            'created_at': str(s.get('created_at', '') or ''),
            'sort_ref': str(s.get('id', '') or ''),
            'sale': s,
        })

    for ret in db.get('returns', []):
        if ret.get('customer') != customer_name:
            continue
        credit_amount = max(0.0, _safe_float(ret.get('credit_amount', ret.get('total', 0))) - _safe_float(ret.get('credit_used', 0)) - _safe_float(ret.get('cash_paid_out', 0)))
        events.append({
            'kind': 'return',
            'date': str(ret.get('date', '') or ''),
            'created_at': str(ret.get('created_at', '') or ''),
            'sort_ref': str(ret.get('id', '') or ''),
            'return': ret,
            'credit_amount': round(credit_amount, 2),
        })

    for p in db.get('cash', []):
        if p.get('party') != customer_name or p.get('source') not in ('customer_payment', 'opening_customer_payment'):
            continue
        events.append({
            'kind': 'payment',
            'date': str(p.get('date', '') or ''),
            'created_at': str(p.get('created_at', '') or ''),
            'sort_ref': str(p.get('receipt_no', '') or p.get('created_at', '') or ''),
            'payment': p,
            'amount': round(max(0.0, _safe_float(p.get('amount', 0))), 2),
        })

    kind_order = {'sale': 0, 'return': 1, 'payment': 2}
    events.sort(key=lambda e: (e.get('date', '') or '9999-99-99', e.get('created_at', '') or '', kind_order.get(e.get('kind', ''), 9), e.get('sort_ref', '')))

    sales_state = {}
    sales_order = []
    payment_events = []
    return_states = {}
    credit_buckets = []

    def _sale_holder(sale_row):
        total = round(max(0.0, _safe_float(sale_row.get('total', 0))), 2)
        initial_paid = round(max(0.0, _safe_float(sale_row.get('paid_amount', 0))), 2)
        remaining = round(max(0.0, total - initial_paid), 2)
        return {
            'sale': sale_row,
            'sale_id': sale_row.get('id', ''),
            'invoice_group_id': sale_row.get('invoice_group_id') or f"single-{ensure_invoice_no(db, sale_row)}",
            'invoice_no': ensure_invoice_no(db, sale_row),
            'customer': sale_row.get('customer', ''),
            'sale_date': sale_row.get('date', ''),
            'item': sale_row.get('item', ''),
            'line_total': total,
            'initial_paid': initial_paid,
            'later_paid': 0.0,
            'credit_applied': 0.0,
            'remaining': remaining,
        }

    def _apply_credit_to_holder(holder, amount):
        amount = round(max(0.0, amount), 2)
        if amount <= 0 or holder is None:
            return 0.0
        take = round(min(holder.get('remaining', 0.0), amount), 2)
        if take <= 0:
            return 0.0
        holder['remaining'] = round(max(0.0, holder.get('remaining', 0.0) - take), 2)
        holder['credit_applied'] = round(holder.get('credit_applied', 0.0) + take, 2)
        return take

    def _consume_credit_buckets(prefer_sale_id=''):
        nonlocal credit_buckets
        if prefer_sale_id and prefer_sale_id in sales_state:
            holder = sales_state.get(prefer_sale_id)
            for bucket in credit_buckets:
                if holder.get('remaining', 0.0) <= 0:
                    break
                take = _apply_credit_to_holder(holder, bucket.get('remaining', 0.0))
                if take <= 0:
                    continue
                bucket['remaining'] = round(max(0.0, bucket.get('remaining', 0.0) - take), 2)
                ret_state = return_states.get(bucket.get('return_id', ''), {})
                ret_state['remaining_credit'] = round(max(0.0, ret_state.get('remaining_credit', 0.0) - take), 2)
                allocs = ret_state.setdefault('allocations', [])
                allocs.append({'sale_id': prefer_sale_id, 'amount': take, 'auto_from_bucket': True})
            credit_buckets = [b for b in credit_buckets if b.get('remaining', 0.0) > 0.009]

    def _apply_new_return_credit(amount, prefer_sale_id=''):
        nonlocal opening_remaining
        remaining = round(max(0.0, amount), 2)
        opening_take = 0.0
        sale_allocs = []

        if prefer_sale_id and prefer_sale_id in sales_state:
            holder = sales_state.get(prefer_sale_id)
            take = _apply_credit_to_holder(holder, remaining)
            if take > 0:
                sale_allocs.append({'sale_id': prefer_sale_id, 'amount': take, 'preferred': True})
                remaining = round(max(0.0, remaining - take), 2)

        if remaining > 0 and opening_remaining > 0:
            opening_take = round(min(opening_remaining, remaining), 2)
            opening_remaining = round(max(0.0, opening_remaining - opening_take), 2)
            remaining = round(max(0.0, remaining - opening_take), 2)

        if remaining > 0:
            for sid in sales_order:
                if sid == prefer_sale_id:
                    continue
                holder = sales_state.get(sid)
                take = _apply_credit_to_holder(holder, remaining)
                if take <= 0:
                    continue
                sale_allocs.append({'sale_id': sid, 'amount': take, 'preferred': False})
                remaining = round(max(0.0, remaining - take), 2)
                if remaining <= 0:
                    break

        return {
            'opening_take': round(opening_take, 2),
            'sale_allocations': sale_allocs,
            'remaining_excess': round(max(0.0, remaining), 2),
        }

    for ev in events:
        if ev.get('kind') == 'sale':
            sale_row = ev.get('sale', {})
            holder = _sale_holder(sale_row)
            sale_id = holder.get('sale_id', '')
            sales_state[sale_id] = holder
            sales_order.append(sale_id)
            _consume_credit_buckets(sale_id)
            continue

        if ev.get('kind') == 'return':
            ret = ev.get('return', {})
            ret_id = ret.get('id', '')
            ret_state = return_states.setdefault(ret_id, {
                'return_row': ret,
                'own_sale_applied': 0.0,
                'opening_applied': 0.0,
                'remaining_credit': 0.0,
                'allocations': [],
            })
            credit_amount = round(max(0.0, ev.get('credit_amount', 0.0)), 2)
            if credit_amount <= 0:
                continue
            apply_result = _apply_new_return_credit(credit_amount, ret.get('sale_id', ''))
            own_sale_applied = sum(a.get('amount', 0.0) for a in apply_result.get('sale_allocations', []) if a.get('sale_id') == ret.get('sale_id', ''))
            ret_state['own_sale_applied'] = round(ret_state.get('own_sale_applied', 0.0) + own_sale_applied, 2)
            ret_state['opening_applied'] = round(ret_state.get('opening_applied', 0.0) + apply_result.get('opening_take', 0.0), 2)
            ret_state['allocations'].extend(apply_result.get('sale_allocations', []))
            excess = round(max(0.0, apply_result.get('remaining_excess', 0.0)), 2)
            if excess > 0:
                ret_state['remaining_credit'] = round(ret_state.get('remaining_credit', 0.0) + excess, 2)
                credit_buckets.append({
                    'return_id': ret_id,
                    'remaining': excess,
                    'date': str(ret.get('date', '') or ''),
                    'created_at': str(ret.get('created_at', '') or ''),
                })
            continue

        if ev.get('kind') == 'payment':
            p = ev.get('payment', {})
            amount = round(max(0.0, ev.get('amount', 0.0)), 2)
            remaining = amount
            opening_take = 0.0
            if remaining > 0 and opening_remaining > 0:
                opening_take = round(min(opening_remaining, remaining), 2)
                opening_remaining = round(max(0.0, opening_remaining - opening_take), 2)
                remaining = round(max(0.0, remaining - opening_take), 2)
            allocations = []
            for sid in sales_order:
                if remaining <= 0:
                    break
                holder = sales_state.get(sid)
                if not holder:
                    continue
                take = round(min(holder.get('remaining', 0.0), remaining), 2)
                if take <= 0:
                    continue
                holder['remaining'] = round(max(0.0, holder.get('remaining', 0.0) - take), 2)
                holder['later_paid'] = round(holder.get('later_paid', 0.0) + take, 2)
                remaining = round(max(0.0, remaining - take), 2)
                sale_row = holder.get('sale', {})
                allocations.append({
                    'sale_id': sid,
                    'invoice_group_id': holder.get('invoice_group_id', ''),
                    'invoice_no': holder.get('invoice_no', ''),
                    'customer': customer_name,
                    'sale_date': sale_row.get('date', ''),
                    'item': sale_row.get('item', ''),
                    'amount': take,
                    'sale_row': sale_row,
                    'payment_row': p,
                })
            payment_events.append({
                'payment_row': p,
                'customer': customer_name,
                'opening_take': opening_take,
                'allocations': allocations,
                'unapplied_amount': remaining,
            })

    invoice_due_map = {}
    for sid in sales_order:
        holder = sales_state.get(sid)
        if not holder:
            continue
        gid = holder.get('invoice_group_id', '')
        invoice_due_map[gid] = round(invoice_due_map.get(gid, 0.0) + holder.get('remaining', 0.0), 2)

    carry_credit_total = round(sum(max(0.0, b.get('remaining', 0.0)) for b in credit_buckets), 2)
    invoices_due_total = round(sum(invoice_due_map.values()), 2)

    return {
        'customer_name': customer_name,
        'opening_remaining': round(opening_remaining, 2),
        'invoice_due_map': invoice_due_map,
        'invoices_due_total': invoices_due_total,
        'sales_state': sales_state,
        'sales_order': sales_order,
        'payment_events': payment_events,
        'return_states': return_states,
        'carry_credit_total': carry_credit_total,
        'final_due': round(max(0.0, opening_remaining + invoices_due_total), 2),
    }


def customer_return_credit(db, customer_name):
    snapshot = customer_receivable_snapshot(db, customer_name)
    return round(float(snapshot.get('carry_credit_total', 0) or 0), 2)

def funder_cash_rows(db, name, source):
    return [x for x in db.get('cash', []) if x.get('source') == source and x.get('party') == name]


def funder_total_paid_profit(db, name):
    return sum(float(x.get('amount', 0) or 0) for x in funder_cash_rows(db, name, 'funder_profit_payment'))


def funder_total_deposit(db, name):
    return sum(float(x.get('amount', 0) or 0) for x in funder_cash_rows(db, name, 'funder_capital_in'))


def funder_total_withdraw_capital(db, name):
    return sum(float(x.get('amount', 0) or 0) for x in funder_cash_rows(db, name, 'funder_capital_out'))


def opening_old_funder_capital(db, name):
    for row in opening_data(db).get('old_funders', []) or []:
        if (row.get('name', '') or '').strip() == (name or '').strip():
            capital = float(row.get('capital', 0) or 0)
            withdrawals = float(row.get('withdrawals', 0) or 0)
            return max(0.0, capital - withdrawals)
    return 0.0


def ensure_funder_identity(row, fallback_name=''):
    row = row or {}
    name = (row.get('name', '') or fallback_name or '').strip()
    if not name:
        return row
    row.setdefault('name', name)
    # Keep legacy/opening funders anchored to their original timestamp when available.
    # Do not force a fresh now_str() here because that can move the base-capital event
    # after historical profit events and distort the movement timeline.
    row.setdefault('created_at', row.get('created_at', '') or '')
    row.setdefault('base_capital', float(row.get('base_capital', row.get('capital', 0)) or 0))
    history = list(row.get('status_history', []) or [])
    if not history:
        stamp = str(row.get('created_at', '') or '')
        history = [{'active': bool(row.get('active', True)), 'date': stamp[:10], 'created_at': stamp}]
    history.sort(key=lambda x: (str(x.get('date', '') or ''), str(x.get('created_at', '') or '')))
    row['status_history'] = history
    return row


def resolve_funder_record(db, funder_row_or_name, create_missing=False):
    if isinstance(funder_row_or_name, dict):
        name = (funder_row_or_name.get('name', '') or '').strip()
        template = funder_row_or_name
    else:
        name = (funder_row_or_name or '').strip()
        template = {}
    if not name:
        return None
    for row in db.get('funders', []):
        if (row.get('name', '') or '').strip() == name:
            return ensure_funder_identity(row, name)
    if not create_missing:
        return None
    stamp = now_str()
    base_capital = max(0.0, _safe_float(template.get('base_capital', template.get('capital', 0)) or opening_old_funder_capital(db, name)))
    row = {
        'name': name,
        'capital': base_capital,
        'base_capital': base_capital,
        'phone': template.get('phone', ''),
        'notes': template.get('notes', ''),
        'active': bool(template.get('active', True)),
        'is_owner_capital': bool(template.get('is_owner_capital', (name == 'هضاب'))),
        'created_at': str(template.get('created_at', '') or stamp),
        'status_history': list(template.get('status_history', []) or [{'active': bool(template.get('active', True)), 'date': str(template.get('created_at', '') or stamp)[:10], 'created_at': str(template.get('created_at', '') or stamp)}]),
    }
    db.setdefault('funders', []).append(row)
    return ensure_funder_identity(row, name)


def funder_capital_base(db, name):
    row = resolve_funder_record(db, name, create_missing=False)
    if row is not None:
        base_capital = _safe_float(row.get('base_capital', row.get('capital', 0)))
        if base_capital > 0:
            return base_capital
        current_capital = _safe_float(row.get('capital', 0))
        if current_capital > 0:
            return current_capital
    return opening_old_funder_capital(db, name)


def funder_current_ratio_pct(db, name):
    funder_row = next((x for x in normalized_funders(db) if (x.get('name', '') or '').strip() == (name or '').strip()), None)
    if not funder_row:
        funder_row = next((x for x in db.get('funders', []) if (x.get('name', '') or '').strip() == (name or '').strip()), None)
    if not funder_row:
        return 0.0
    return funder_ratio_pct(db, funder_row)


def funder_event_order(row):
    source = row.get('source', '')
    if source == 'funder_base_capital':
        return 0
    if source == 'funder_capital_in':
        return 1
    if source == 'funder_capital_out':
        return 2
    if source == 'funder_profit_payment':
        return 3
    return 9


def funder_movement_rows(db, funder_name=None):
    funders = {x.get('name', ''): ensure_funder_identity(dict(x), x.get('name', '')) for x in normalized_funders(db) if x.get('name', '')}
    build_profit_ledger(db)
    events = []
    for name, f in funders.items():
        base_cap = funder_capital_base(db, name)
        events.append({
            'date': str((f or {}).get('created_at', '') or '')[:10], 'created_at': str((f or {}).get('created_at', '') or ''), 'party': name, 'amount': base_cap, 'notes': 'رأس المال الأساسي',
            'source': 'funder_base_capital', 'category': 'رأس المال الأساسي', 'payment_method': '-', 'account_no': '', 'receipt_no': '',
        })
    for row in db.get('cash', []):
        if row.get('source') in ('funder_capital_in', 'funder_capital_out', 'funder_profit_payment') and row.get('party', '') in funders:
            events.append(dict(row))
    for row in db.get('profit_entries', []):
        if row.get('beneficiary_type') == 'funder' and row.get('beneficiary_name', '') in funders:
            events.append({
                'date': row.get('date', ''), 'created_at': row.get('created_at', ''), 'party': row.get('beneficiary_name', ''),
                'amount': row.get('amount', 0), 'notes': row.get('notes', ''),
                'source': 'profit_event', 'category': 'ربح ممول' if _safe_float(row.get('amount', 0)) >= 0 else 'عكس ربح ممول',
                'payment_method': '-', 'account_no': '', 'receipt_no': row.get('event_uid', ''),
            })
    events.sort(key=lambda x: (str(x.get('date', '') or ''), str(x.get('created_at', '') or ''), funder_event_order(x), str(x.get('receipt_no', '') or ''), str(x.get('party', '') or '')))

    running = {name: 0.0 for name in funders}
    earned_running = {name: 0.0 for name in funders}
    paid_running = {name: 0.0 for name in funders}
    rows = []
    for ev in events:
        name = ev.get('party', '')
        if not name:
            continue
        source = ev.get('source', '')
        amount = _safe_float(ev.get('amount', 0))
        if source in ('funder_base_capital', 'funder_capital_in'):
            running[name] = max(0.0, running.get(name, 0.0) + amount)
        elif source == 'funder_capital_out':
            running[name] = max(0.0, running.get(name, 0.0) - amount)
        elif source == 'funder_profit_payment':
            paid_running[name] = round(paid_running.get(name, 0.0) + amount, 2)
        elif source == 'profit_event':
            earned_running[name] = round(earned_running.get(name, 0.0) + amount, 2)
        total_active = 0.0
        for fname, frow in funders.items():
            if funder_is_active_as_of(frow, ev.get('date', ''), ev.get('created_at', '')) and running.get(fname, 0.0) > 0:
                total_active += running.get(fname, 0.0)
        ratio_after = (running.get(name, 0.0) / total_active * 100.0) if total_active and funder_is_active_as_of(funders[name], ev.get('date', ''), ev.get('created_at', '')) and running.get(name, 0.0) > 0 else 0.0
        rows.append({
            'party': name, 'date': ev.get('date', ''), 'created_at': ev.get('created_at', ''),
            'movement': ev.get('category', ev.get('source', '')), 'amount': amount, 'source': source,
            'payment_method': ev.get('payment_method', '-'), 'notes': ev.get('notes', ''),
            'account_no': ev.get('account_no', ''), 'receipt_no': ev.get('receipt_no', ''),
            'capital_after': running.get(name, 0.0), 'ratio_after': ratio_after,
            'earned_profit': earned_running.get(name, 0.0), 'paid_profit': paid_running.get(name, 0.0),
            'pending_profit': round(earned_running.get(name, 0.0) - paid_running.get(name, 0.0), 2),
        })
    if funder_name:
        rows = [x for x in rows if x.get('party') == funder_name]
    return rows

def rename_funder_references(db, old_name, new_name):
    old_name = (old_name or '').strip()
    new_name = (new_name or '').strip()
    if not old_name or not new_name or old_name == new_name:
        return
    for row in db.get('cash', []):
        if row.get('party', '') == old_name and row.get('source') in ('funder_capital_in', 'funder_capital_out', 'funder_profit_payment'):
            row['party'] = new_name
    for row in opening_data(db).get('old_funders', []):
        if row.get('name', '') == old_name:
            row['name'] = new_name


def save_db(data):
    DB_FILE.parent.mkdir(parents=True, exist_ok=True)
    DB_FILE.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')


def hash_password(pwd):
    return hashlib.sha256((pwd or '').encode('utf-8')).hexdigest()


def password_record_for_storage(pwd):
    return f"sha256${hash_password(pwd)}"


def verify_password_input(input_pwd, saved_value):
    saved_value = (saved_value or '').strip()
    if saved_value.startswith('sha256$'):
        return hash_password(input_pwd) == saved_value.split('$', 1)[1]
    return (input_pwd or '') == saved_value


def create_backup_file(data, reason='manual', keep_files=20):
    BACKUPS_DIR.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    safe_reason = ''.join(ch if ch.isalnum() or ch in ('_', '-') else '_' for ch in str(reason or 'backup'))
    backup_path = BACKUPS_DIR / f'nokhba_{safe_reason}_{stamp}.json'
    backup_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
    try:
        keep_files = max(3, int(keep_files or 20))
    except Exception:
        keep_files = 20
    files = sorted(BACKUPS_DIR.glob('nokhba_*.json'), key=lambda p: p.stat().st_mtime, reverse=True)
    for old in files[keep_files:]:
        try:
            old.unlink()
        except Exception:
            pass
    return backup_path


def _safe_float(v):
    try:
        return float(v or 0)
    except Exception:
        return 0.0


def _safe_int(v):
    try:
        return int(float(v or 0))
    except Exception:
        return 0


def _event_ts_tuple(date_value='', created_at=''):
    return (str(date_value or ''), str(created_at or ''), '')



def _funder_cash_events_as_of(db, date_value='', created_at=''):
    cutoff = (str(date_value or ''), str(created_at or ''), '')
    rows = []

    for f in normalized_funders(db):
        name = (f.get('name', '') or '').strip()
        if not name:
            continue
        amount = max(0.0, _safe_float(funder_capital_base(db, name)))
        if amount <= 0:
            continue
        stamp = str(f.get('created_at', '') or '')
        key = (stamp[:10], stamp, '')
        if key <= cutoff:
            rows.append({
                'date': stamp[:10],
                'created_at': stamp,
                'party': name,
                'amount': amount,
                'source': 'funder_base_capital',
            })

    for row in db.get('cash', []):
        if row.get('source') in ('funder_capital_in', 'funder_capital_out'):
            key = (str(row.get('date', '') or ''), str(row.get('created_at', '') or ''), '')
            if key <= cutoff:
                rows.append(row)
    rows.sort(key=lambda x: (str(x.get('date', '') or ''), str(x.get('created_at', '') or ''), funder_event_order(x), str(x.get('receipt_no', '') or ''), str(x.get('party', '') or '')))
    return rows

def active_funder_snapshot(db, date_value='', created_at=''):
    funders = {x.get('name', ''): x for x in normalized_funders(db) if x.get('name', '')}
    running = {name: 0.0 for name in funders}
    for ev in _funder_cash_events_as_of(db, date_value, created_at):
        name = ev.get('party', '')
        if not name or name not in running:
            continue
        amount = _safe_float(ev.get('amount', 0))
        source = ev.get('source', '')
        if source in ('funder_base_capital', 'funder_capital_in'):
            running[name] = max(0.0, running.get(name, 0.0) + amount)
        elif source == 'funder_capital_out':
            running[name] = max(0.0, running.get(name, 0.0) - amount)
    rows = []
    total = 0.0
    for name, base in funders.items():
        capital = max(0.0, running.get(name, 0.0))
        if not funder_is_active_as_of(base, date_value, created_at) or capital <= 0:
            continue
        total += capital
        rows.append({
            'name': name,
            'capital': capital,
            'is_owner_capital': bool(base.get('is_owner_capital', False)),
            'active': bool(base.get('active', True)),
        })
    for row in rows:
        row['ratio'] = (row['capital'] / total) if total > 0 else 0.0
    return rows, total


def funder_is_active_as_of(funder_row, date_value='', created_at=''):
    history = list(funder_row.get('status_history', []) or [])
    if not history:
        return bool(funder_row.get('active', True))
    cutoff = (str(date_value or ''), str(created_at or ''), '')
    state = bool(history[0].get('active', True))
    for row in sorted(history, key=lambda x: (str(x.get('date', '') or ''), str(x.get('created_at', '') or ''))):
        key = (str(row.get('date', '') or ''), str(row.get('created_at', '') or ''), '')
        if key <= cutoff:
            state = bool(row.get('active', state))
        else:
            break
    return state


def _initial_sale_payment_events(db):
    events = []
    grouped = {}
    for s in db.get('sales', []):
        paid = _safe_float(s.get('paid_amount', 0))
        total = _safe_float(s.get('total', 0))
        if paid <= 0 or total <= 0:
            continue
        gid = s.get('invoice_group_id') or f"single-{ensure_invoice_no(db, s)}"
        g = grouped.setdefault(gid, {
            'event_uid': f'initpay_{gid}',
            'event_type': 'payment',
            'payment_uid': f'initpay_{gid}',
            'date': s.get('date', ''),
            'created_at': s.get('created_at', ''),
            'customer': s.get('customer', ''),
            'source': 'sales_group',
            'amount': 0.0,
            'allocations': [],
        })
        realized_sale_amount = min(total, paid)
        gross_profit = _safe_float(s.get('profit', 0)) * (realized_sale_amount / total) if total > 0 else 0.0
        g['amount'] += realized_sale_amount
        g['allocations'].append({
            'sale_id': s.get('id', ''),
            'invoice_group_id': gid,
            'invoice_no': ensure_invoice_no(db, s),
            'customer': s.get('customer', ''),
            'sale_date': s.get('date', ''),
            'item': s.get('item', ''),
            'amount': realized_sale_amount,
            'gross_profit': gross_profit,
            'line_total': total,
        })
    return sorted(events + list(grouped.values()), key=lambda x: (str(x.get('date', '') or ''), str(x.get('created_at', '') or ''), str(x.get('event_uid', '') or '')))



def _customer_payment_events(db):
    events = []
    for ev in customer_payment_allocations(db):
        p = ev.get('payment_row', {})
        allocations = []
        for alloc in ev.get('allocations', []):
            allocations.append({
                'sale_id': alloc.get('sale_id', ''),
                'invoice_group_id': alloc.get('invoice_group_id', ''),
                'invoice_no': alloc.get('invoice_no', ''),
                'customer': alloc.get('customer', ''),
                'sale_date': alloc.get('sale_date', ''),
                'item': alloc.get('item', ''),
                'amount': round(_safe_float(alloc.get('amount', 0)), 2),
                'gross_profit': round(_safe_float(alloc.get('gross_profit', 0)), 2),
                'line_total': round(_safe_float(alloc.get('line_total', 0)), 2),
            })
        if allocations:
            events.append({
                'event_uid': f"custpay_{p.get('receipt_no', '') or p.get('created_at', '')}",
                'event_type': 'payment',
                'payment_uid': f"custpay_{p.get('receipt_no', '') or p.get('created_at', '')}",
                'date': p.get('date', ''),
                'created_at': p.get('created_at', ''),
                'customer': ev.get('customer', ''),
                'source': p.get('source', ''),
                'amount': round(sum(_safe_float(a.get('amount', 0)) for a in allocations), 2),
                'allocations': allocations,
                'receipt_no': p.get('receipt_no', ''),
                'notes': p.get('notes', ''),
            })
    return events

def _profit_ledger_has_value(events, entries):
    try:
        events_total = sum(abs(_safe_float(x.get('net_profit', 0))) for x in (events or []))
        entries_total = sum(abs(_safe_float(x.get('amount', 0))) for x in (entries or []) if x.get('entry_type', '') == 'profit')
        return (events_total > 0.009) or (entries_total > 0.009)
    except Exception:
        return False


def _profit_entries_breakdown(entries):
    totals = {
        'partner_total': 0.0,
        'funder_total': 0.0,
        'owner_financing_total': 0.0,
        'external_funders_total': 0.0,
    }
    for row in entries or []:
        if row.get('entry_type', '') != 'profit':
            continue
        amount = _safe_float(row.get('amount', 0))
        if row.get('beneficiary_type', '') == 'partner':
            totals['partner_total'] += amount
        elif row.get('beneficiary_type', '') == 'funder':
            totals['funder_total'] += amount
            if str(row.get('beneficiary_name', '') or '').strip() == 'هضاب':
                totals['owner_financing_total'] += amount
            else:
                totals['external_funders_total'] += amount
    return totals



def build_profit_ledger(db):
    old_events = deepcopy(db.get('profit_events', []) or [])
    old_entries = deepcopy(db.get('profit_entries', []) or [])
    total_expenses = expenses_total(db)
    total_damaged = damaged_loss(db)

    positive_events = _initial_sale_payment_events(db) + _customer_payment_events(db)
    positive_events.sort(key=lambda x: (str(x.get('date', '') or ''), str(x.get('created_at', '') or ''), str(x.get('event_uid', '') or '')))

    customer_names = {
        str(x.get('customer', '') or '').strip()
        for x in db.get('sales', [])
    } | {
        str(x.get('customer', '') or '').strip()
        for x in db.get('returns', [])
    } | {
        str(x.get('party', '') or '').strip()
        for x in db.get('cash', [])
        if x.get('source') in ('customer_payment', 'opening_customer_payment')
    }
    customer_names = {x for x in customer_names if x}
    customer_snapshots = {name: customer_receivable_snapshot(db, name) for name in customer_names}
    return_state_map = {}
    for snap in customer_snapshots.values():
        for ret_id, state in (snap.get('return_states', {}) or {}).items():
            return_state_map[ret_id] = state

    combined_events = []
    for ev in positive_events:
        combined_events.append({
            'kind': 'payment',
            'date': str(ev.get('date', '') or ''),
            'created_at': str(ev.get('created_at', '') or ''),
            'sort_ref': str(ev.get('event_uid', '') or ''),
            'payload': ev,
        })
    for ret in db.get('returns', []):
        combined_events.append({
            'kind': 'return',
            'date': str(ret.get('date', '') or ''),
            'created_at': str(ret.get('created_at', '') or ''),
            'sort_ref': str(ret.get('id', '') or ''),
            'payload': ret,
        })
    combined_events.sort(key=lambda x: (x.get('date', '') or '', x.get('created_at', '') or '', 0 if x.get('kind') == 'payment' else 1, x.get('sort_ref', '')))

    events = []
    entries = []
    payment_fragments = []
    allocated_expenses = 0.0
    allocated_damaged = 0.0
    remaining_expenses = round(max(0.0, total_expenses), 2)
    remaining_damaged = round(max(0.0, total_damaged), 2)

    owner_pct = _safe_float(db.get('settings', {}).get('owner_profit_pct', 50.0))
    partner_pct = _safe_float(db.get('settings', {}).get('partner_profit_pct', 50.0))
    configured_pool_pct = _safe_float(db.get('settings', {}).get('funders_profit_pool_pct', 33.3333))
    funders_pool_ratio = min(1.0, max(0.0, configured_pool_pct / 100.0))
    total_pct = owner_pct + partner_pct
    owner_ratio = (owner_pct / total_pct) if total_pct > 0 else 0.5

    for wrapped in combined_events:
        if wrapped.get('kind') == 'payment':
            ev = wrapped.get('payload', {}) or {}
            gross_profit = sum(_safe_float(a.get('gross_profit', 0)) for a in ev.get('allocations', []))
            realized_sale_amount = sum(_safe_float(a.get('amount', 0)) for a in ev.get('allocations', []))
            if gross_profit <= 0 or realized_sale_amount <= 0:
                continue

            remaining_after_expenses = gross_profit
            incr_expenses = round(min(remaining_after_expenses, remaining_expenses), 2)
            remaining_after_expenses = round(remaining_after_expenses - incr_expenses, 2)
            remaining_expenses = round(max(0.0, remaining_expenses - incr_expenses), 2)
            allocated_expenses = round(allocated_expenses + incr_expenses, 2)

            incr_damaged = round(min(remaining_after_expenses, remaining_damaged), 2)
            remaining_after_expenses = round(remaining_after_expenses - incr_damaged, 2)
            remaining_damaged = round(max(0.0, remaining_damaged - incr_damaged), 2)
            allocated_damaged = round(allocated_damaged + incr_damaged, 2)

            net_profit = max(0.0, round(remaining_after_expenses, 2))
            funders_snapshot, total_capital = active_funder_snapshot(db, ev.get('date', ''), ev.get('created_at', ''))
            funders_pool = round(net_profit * funders_pool_ratio, 2)
            owner_capital_share = 0.0
            external_total = 0.0
            event_entries = []
            for f in funders_snapshot:
                amount = round(funders_pool * _safe_float(f.get('ratio', 0)), 2)
                if amount <= 0:
                    continue
                if f.get('is_owner_capital', False):
                    owner_capital_share += amount
                    event_entries.append({
                        'id': generate_id('pe'),
                        'event_uid': ev['event_uid'],
                        'entry_type': 'profit',
                        'beneficiary_type': 'funder',
                        'beneficiary_name': f.get('name', ''),
                        'amount': amount,
                        'date': ev.get('date', ''),
                        'created_at': ev.get('created_at', ''),
                        'notes': f"حصة تمويل هضاب راجعة للشراكة من تسديد {ev.get('payment_uid','')}",
                    })
                else:
                    external_total += amount
                    event_entries.append({
                        'id': generate_id('pe'),
                        'event_uid': ev['event_uid'],
                        'entry_type': 'profit',
                        'beneficiary_type': 'funder',
                        'beneficiary_name': f.get('name', ''),
                        'amount': amount,
                        'date': ev.get('date', ''),
                        'created_at': ev.get('created_at', ''),
                        'notes': f"ربح ممول من تسديد {ev.get('payment_uid','')}",
                    })
            partners_profit = round(max(0.0, net_profit - external_total), 2)
            hidab_amount = round(partners_profit * owner_ratio, 2)
            mostafa_amount = round(partners_profit - hidab_amount, 2)
            event_entries.extend([
                {'id': generate_id('pe'), 'event_uid': ev['event_uid'], 'entry_type': 'profit', 'beneficiary_type': 'partner', 'beneficiary_name': 'هضاب', 'amount': hidab_amount, 'date': ev.get('date', ''), 'created_at': ev.get('created_at', ''), 'notes': f"ربح هضاب من تسديد {ev.get('payment_uid','')}"},
                {'id': generate_id('pe'), 'event_uid': ev['event_uid'], 'entry_type': 'profit', 'beneficiary_type': 'partner', 'beneficiary_name': 'مصطفى', 'amount': mostafa_amount, 'date': ev.get('date', ''), 'created_at': ev.get('created_at', ''), 'notes': f"ربح مصطفى من تسديد {ev.get('payment_uid','')}"},
            ])
            event = {
                'event_uid': ev['event_uid'],
                'event_type': 'payment',
                'date': ev.get('date', ''),
                'created_at': ev.get('created_at', ''),
                'payment_uid': ev.get('payment_uid', ''),
                'customer': ev.get('customer', ''),
                'gross_profit': round(gross_profit, 2),
                'net_profit': net_profit,
                'realized_sale_amount': round(realized_sale_amount, 2),
                'allocated_expenses': incr_expenses,
                'allocated_damaged': incr_damaged,
                'funders_pool': round(funders_pool, 2),
                'external_funders_total': round(external_total, 2),
                'owner_capital_share': round(owner_capital_share, 2),
                'partners_profit': round(partners_profit, 2),
                'entries': event_entries,
            }
            for alloc in ev.get('allocations', []):
                gp = _safe_float(alloc.get('gross_profit', 0))
                frag_ratio = (gp / gross_profit) if gross_profit > 0 else 0.0
                frag_entries = []
                for ent in event_entries:
                    share = round(_safe_float(ent.get('amount', 0)) * frag_ratio, 2)
                    if share == 0:
                        continue
                    frag_entries.append({
                        'beneficiary_type': ent.get('beneficiary_type', ''),
                        'beneficiary_name': ent.get('beneficiary_name', ''),
                        'amount': share,
                    })
                payment_fragments.append({
                    'fragment_uid': generate_id('frag'),
                    'event_uid': ev['event_uid'],
                    'sale_id': alloc.get('sale_id', ''),
                    'invoice_no': alloc.get('invoice_no', ''),
                    'customer': alloc.get('customer', ''),
                    'item': alloc.get('item', ''),
                    'date': ev.get('date', ''),
                    'created_at': ev.get('created_at', ''),
                    'sale_amount_total': round(_safe_float(alloc.get('amount', 0)), 2),
                    'sale_amount_remaining': round(_safe_float(alloc.get('amount', 0)), 2),
                    'gross_profit_total': round(gp, 2),
                    'gross_profit_remaining': round(gp, 2),
                    'entries_total': frag_entries,
                    'entries_remaining': [dict(x) for x in frag_entries],
                })
            events.append(event)
            entries.extend(event_entries)
            continue

        ret = wrapped.get('payload', {}) or {}
        ret_state = return_state_map.get(ret.get('id', ''), {}) or {}
        own_unpaid_absorbed = round(_safe_float(ret_state.get('own_sale_applied', 0)), 2)
        remaining_sale_value = round(max(0.0, (_safe_float(ret.get('qty', 0)) * _safe_float(ret.get('unit_price', 0))) - own_unpaid_absorbed), 2)
        if remaining_sale_value <= 0:
            continue
        sale_id = ret.get('sale_id', '')
        fragments = [f for f in payment_fragments if f.get('sale_id') == sale_id and _safe_float(f.get('sale_amount_remaining', 0)) > 0]
        fragments.sort(key=lambda x: (str(x.get('date', '') or ''), str(x.get('created_at', '') or ''), str(x.get('fragment_uid', '') or '')))
        reverse_entries = []
        reversed_gross = 0.0
        for frag in fragments:
            if remaining_sale_value <= 0:
                break
            take = min(_safe_float(frag.get('sale_amount_remaining', 0)), remaining_sale_value)
            if take <= 0:
                continue
            ratio = take / _safe_float(frag.get('sale_amount_total', 1) or 1)
            frag['sale_amount_remaining'] = round(_safe_float(frag.get('sale_amount_remaining', 0)) - take, 2)
            remaining_sale_value = round(remaining_sale_value - take, 2)
            gross_piece = round(_safe_float(frag.get('gross_profit_total', 0)) * ratio, 2)
            frag['gross_profit_remaining'] = round(_safe_float(frag.get('gross_profit_remaining', 0)) - gross_piece, 2)
            reversed_gross += gross_piece
            for ent in frag.get('entries_remaining', []):
                amount_piece = round(_safe_float(ent.get('amount', 0)) * ratio, 2)
                if amount_piece == 0:
                    continue
                reverse_entries.append({
                    'id': generate_id('pe'),
                    'event_uid': f"reverse_{ret.get('id', '')}_{frag.get('fragment_uid', '')}",
                    'entry_type': 'reverse',
                    'beneficiary_type': ent.get('beneficiary_type', ''),
                    'beneficiary_name': ent.get('beneficiary_name', ''),
                    'amount': -amount_piece,
                    'date': ret.get('date', ''),
                    'created_at': ret.get('created_at', ''),
                    'notes': f"عكس ربح بسبب مرتجع فاتورة #{ret.get('invoice_no','')}",
                })
                ent['amount'] = round(_safe_float(ent.get('amount', 0)) - amount_piece, 2)
        if reverse_entries:
            event_uid = f"reverse_return_{ret.get('id', '')}"
            for ent in reverse_entries:
                ent['event_uid'] = event_uid
            events.append({
                'event_uid': event_uid,
                'event_type': 'return_reverse',
                'date': ret.get('date', ''),
                'created_at': ret.get('created_at', ''),
                'payment_uid': ret.get('id', ''),
                'customer': ret.get('customer', ''),
                'gross_profit': -round(reversed_gross, 2),
                'net_profit': round(sum(_safe_float(x.get('amount', 0)) for x in reverse_entries), 2),
                'realized_sale_amount': -round((_safe_float(ret.get('qty', 0)) * _safe_float(ret.get('unit_price', 0))), 2),
                'allocated_expenses': 0.0,
                'allocated_damaged': 0.0,
                'funders_pool': 0.0,
                'external_funders_total': 0.0,
                'owner_capital_share': 0.0,
                'partners_profit': 0.0,
                'entries': reverse_entries,
            })
            entries.extend(reverse_entries)

    events.sort(key=lambda x: (str(x.get('date', '') or ''), str(x.get('created_at', '') or ''), str(x.get('event_uid', '') or '')))
    entries.sort(key=lambda x: (str(x.get('date', '') or ''), str(x.get('created_at', '') or ''), str(x.get('event_uid', '') or ''), str(x.get('beneficiary_name', '') or '')))
    new_has_value = _profit_ledger_has_value(events, entries)
    old_has_value = _profit_ledger_has_value(old_events, old_entries)
    should_keep_old = (not new_has_value) and old_has_value

    if should_keep_old:
        db['profit_events'] = old_events
        db['profit_entries'] = old_entries
        return old_events, old_entries

    db['profit_events'] = events
    db['profit_entries'] = entries
    return events, entries

def partner_profit_entries(db, partner_name):
    _, entries = build_profit_ledger(db)
    rows = [x for x in entries if x.get('beneficiary_type') == 'partner' and x.get('beneficiary_name') == partner_name]
    rows.sort(key=lambda x: (str(x.get('date', '') or ''), str(x.get('created_at', '') or ''), str(x.get('event_uid', '') or '')))
    return rows


def funder_profit_entries(db, funder_name):
    _, entries = build_profit_ledger(db)
    rows = [x for x in entries if x.get('beneficiary_type') == 'funder' and x.get('beneficiary_name') == funder_name]
    rows.sort(key=lambda x: (str(x.get('date', '') or ''), str(x.get('created_at', '') or ''), str(x.get('event_uid', '') or '')))
    return rows

ensure_single_data_location()

def load_db():
    INVOICES_DIR.mkdir(parents=True, exist_ok=True)
    ATTACHMENTS_DIR.mkdir(parents=True, exist_ok=True)
    BACKUPS_DIR.mkdir(parents=True, exist_ok=True)
    # بدء نظيف: لا تنسخ أي data.json قديم من مجلد البرنامج
    if not PASSWORD_FILE.exists():
        legacy_pass = BASE_DIR / 'pass.txt'
        if legacy_pass.exists():
            try:
                shutil.copy2(legacy_pass, PASSWORD_FILE)
            except Exception:
                pass
    if not DB_FILE.exists():
        save_db(deepcopy(DEFAULT_DB))
    try:
        data = json.loads(DB_FILE.read_text(encoding='utf-8'))
    except Exception:
        data = deepcopy(DEFAULT_DB)
        save_db(data)
    for k, v in DEFAULT_DB.items():
        if k not in data:
            data[k] = deepcopy(v)
    for k, v in DEFAULT_DB['settings'].items():
        data['settings'].setdefault(k, v)

    data = normalize_db(data)
    try:
        save_db(data)
    except Exception:
        pass
    return data


def factory_reset_all_data():
    """إعادة ضبط المصنع: حذف كل البيانات المحلية وإرجاع النظام لأول تشغيل."""
    APP_DATA_DIR.mkdir(parents=True, exist_ok=True)
    for folder in (INVOICES_DIR, ATTACHMENTS_DIR, BACKUPS_DIR):
        try:
            if folder.exists():
                shutil.rmtree(folder, ignore_errors=True)
        except Exception:
            pass
        try:
            folder.mkdir(parents=True, exist_ok=True)
        except Exception:
            pass
    try:
        if DB_FILE.exists():
            DB_FILE.unlink()
    except Exception:
        pass
    try:
        if PASSWORD_FILE.exists():
            PASSWORD_FILE.unlink()
    except Exception:
        pass
    fresh = deepcopy(DEFAULT_DB)
    save_db(fresh)
    return fresh


def is_primary_button_text(txt):
    txt = (txt or '').strip()
    starts = ('➕', '💰', '📥', '📊', '🧮', 'تسجيل الدخول', 'حفظ')
    return txt.startswith(starts) or 'إضافة' in txt or 'تسجيل' in txt or 'إنشاء' in txt or txt == 'حفظ'


def apply_theme_to_widget(widget):
    tooltip_style = f'''
    QToolTip {{
        background-color: {CARD};
        color: {MUTED};
        border: 1px solid {BORDER};
        padding: 6px 10px;
        border-radius: 8px;
        font-size: 11px;
    }}
    QTabWidget::pane {{ border:1px solid {BORDER}; background:{CARD}; border-radius:14px; margin-top:8px; }}
    QTabBar::tab {{ background:{rgba_from_hex(CARD,0.82)}; color:{TEXT}; min-width:170px; min-height:40px; padding:10px 14px; margin:4px 4px 0 0; border-radius:12px; font-size:15px; font-weight:800; }}
    QTabBar::tab:selected {{ background:{ACCENT}; color:{TEXT_ON_ACCENT}; }}
    QTabBar::tab:hover {{ background:{ACCENT2}; color:{TEXT_ON_ACCENT}; }}
    '''
    widget.setStyleSheet(WINDOW_STYLE + INPUT_STYLE + tooltip_style)
    for table in widget.findChildren(QTableWidget):
        table.setStyleSheet(TABLE_STYLE)
        try:
            table.setAlternatingRowColors(True)
            table.verticalHeader().setDefaultSectionSize(34)
        except Exception:
            pass
    for button in widget.findChildren(QPushButton):
        button.setStyleSheet(BUTTON_STYLE if is_primary_button_text(button.text()) else SECONDARY_BUTTON)
        try:
            button.setMinimumHeight(max(button.minimumHeight(), 42))
        except Exception:
            pass
    for spin in widget.findChildren(QSpinBox):
        tune_numeric_widget(spin)
    for spin in widget.findChildren(QDoubleSpinBox):
        tune_numeric_widget(spin)
    for frame in widget.findChildren(QFrame):
        if frame.layout() is not None:
            frame.setStyleSheet(CARD_FRAME_STYLE)


def set_theme(theme_key):
    build_styles(theme_key)
    app = QApplication.instance()
    if app:
        app.setStyleSheet(f"QWidget{{background-color:{DARK}; color:{TEXT};}} QLabel{{color:{TEXT};}}")
    for widget in QApplication.topLevelWidgets() if app else []:
        try:
            if hasattr(widget, 'apply_theme'):
                widget.apply_theme()
            else:
                apply_theme_to_widget(widget)
        except Exception:
            pass




def opening_data(db):
    op = db.setdefault('opening', {
        'opening_cash': 0.0,
        'operating_cash': 0.0,
        'customers': [],
        'suppliers': [],
        'stock_items': [],
        'old_funders': [],
        'old_totals': {
            'old_sales_total': 0.0,
            'old_purchases_total': 0.0,
            'old_expenses_total': 0.0,
            'old_funders_withdrawals_total': 0.0,
            'old_hidab_withdrawals_total': 0.0,
            'old_mustafa_withdrawals_total': 0.0,
            'old_profit_clearing_total': 0.0
        },
        'start_date': '',
        'hidab_opening_balance': 0.0,
        'mustafa_opening_balance': 0.0,
        'opening_locked': False
    })
    op.setdefault('customers', [])
    op.setdefault('suppliers', [])
    op.setdefault('stock_items', [])
    op.setdefault('old_funders', [])
    op.setdefault('old_totals', {})
    for k, v in {
        'old_sales_total': 0.0,
        'old_purchases_total': 0.0,
        'old_expenses_total': 0.0,
        'old_funders_withdrawals_total': 0.0,
        'old_hidab_withdrawals_total': 0.0,
        'old_mustafa_withdrawals_total': 0.0,
        'old_profit_clearing_total': 0.0
    }.items():
        op['old_totals'].setdefault(k, v)
    op.setdefault('opening_cash', 0.0)
    op.setdefault('operating_cash', 0.0)
    op.setdefault('start_date', '')
    op.setdefault('hidab_opening_balance', 0.0)
    op.setdefault('mustafa_opening_balance', 0.0)
    op.setdefault('opening_locked', False)
    return op


def old_funder_net_amount(row):
    capital = float(row.get('capital', 0) or 0)
    withdrawals = float(row.get('withdrawals', 0) or 0)
    return max(0.0, capital - withdrawals)


def opening_old_funders_operating_balance(db):
    return sum(old_funder_net_amount(x) for x in opening_data(db).get('old_funders', []))


def opening_old_paid_funders_profit_total(db):
    return sum(float(x.get('paid_profit', 0) or 0) for x in opening_data(db).get('old_funders', []))




def calc_old_profit_clearing_from_values(old_sales_total=0.0, old_purchases_total=0.0, old_expenses_total=0.0,
                                         old_funders_withdrawals_total=0.0, old_hidab_withdrawals_total=0.0,
                                         old_mustafa_withdrawals_total=0.0, old_paid_funders_profit_total=0.0):
    sales = float(old_sales_total or 0)
    purchases = float(old_purchases_total or 0)
    expenses = float(old_expenses_total or 0)
    funders_w = float(old_funders_withdrawals_total or 0)
    hidab_w = float(old_hidab_withdrawals_total or 0)
    mustafa_w = float(old_mustafa_withdrawals_total or 0)
    paid_funders_profit = float(old_paid_funders_profit_total or 0)
    return sales - purchases - expenses - funders_w - hidab_w - mustafa_w - paid_funders_profit
def opening_old_totals_balance(db):
    op = opening_data(db)
    t = op.get('old_totals', {})
    sales = float(t.get('old_sales_total', 0) or 0)
    purchases = float(t.get('old_purchases_total', 0) or 0)
    expenses = float(t.get('old_expenses_total', 0) or 0)
    funders_w = float(t.get('old_funders_withdrawals_total', 0) or 0)
    hidab_w = float(t.get('old_hidab_withdrawals_total', 0) or 0)
    mustafa_w = float(t.get('old_mustafa_withdrawals_total', 0) or 0)
    clearing = float(t.get('old_profit_clearing_total', 0) or 0)
    paid_funders_profit = opening_old_paid_funders_profit_total(db)
    return sales - purchases - expenses - funders_w - hidab_w - mustafa_w - clearing - paid_funders_profit


def opening_cash_balance(db):
    op = opening_data(db)
    auto_operating = opening_old_funders_operating_balance(db)
    op['operating_cash'] = auto_operating
    # نعتمد الرصيد الصافي المعتمد من النظام بعد تنزيل الديون القديمة من التهيئة.
    # هذا الحقل يمثّل نقطة البداية الحقيقية للقاصة، لذلك لا نخصم الديون القديمة مرة ثانية.
    return float(op.get('opening_cash', 0) or 0)


def expenses_total(db):
    return sum(float(x.get('amount', 0) or 0) for x in db.get('expenses', []))


def opening_customer_amount(db, customer_name):
    return sum(float(x.get('amount', 0) or 0) for x in opening_data(db).get('customers', []) if x.get('name', '') == customer_name)


def opening_supplier_amount(db, supplier_name):
    return sum(float(x.get('amount', 0) or 0) for x in opening_data(db).get('suppliers', []) if x.get('name', '') == supplier_name)


def opening_customer_receipts(db, customer_name):
    return sum(
        float(x.get('amount', 0) or 0)
        for x in db.get('cash', [])
        if x.get('party') == customer_name and x.get('source') in ('opening_customer_payment',)
    )

def opening_supplier_payments(db, supplier_name):
    return sum(
        float(x.get('amount', 0) or 0)
        for x in db.get('cash', [])
        if x.get('party') == supplier_name and x.get('source') in ('opening_supplier_payment',)
    )

def opening_customer_due_remaining(db, customer_name):
    total_old = opening_customer_amount(db, customer_name)
    paid_old = opening_customer_receipts(db, customer_name)
    return max(0.0, total_old - paid_old)

def opening_supplier_due_remaining(db, supplier_name):
    total_old = opening_supplier_amount(db, supplier_name)
    paid_old = opening_supplier_payments(db, supplier_name)
    return max(0.0, total_old - paid_old)


def opening_customers_total(db):
    return sum(float(x.get('amount', 0) or 0) for x in opening_data(db).get('customers', []))

def opening_suppliers_total(db):
    return sum(float(x.get('amount', 0) or 0) for x in opening_data(db).get('suppliers', []))

def inventory_value(db):
    return sum(item_inventory_value(x) for x in db.get('items', []))

def total_receivables(db):
    return sum(float(x.get('due', 0) or 0) for x in customer_due_summary(db))


def agents_custody_rows(db):
    rows = db.setdefault('agents_custody', [])
    for row in rows:
        row.setdefault('id', generate_id('agc'))
        row.setdefault('date', '')
        row.setdefault('agent', '')
        row.setdefault('type', 'وارد مستلم')
        row.setdefault('amount', 0.0)
        row.setdefault('party', '')
        row.setdefault('notes', '')
        row.setdefault('settlement_direction', '')
        row.setdefault('created_at', now_str())
    return rows


def agent_custody_effect(row):
    typ = (row.get('type', '') or '').strip()
    amount = float(row.get('amount', 0) or 0)
    if typ == 'وارد مستلم':
        return amount
    if typ in ('تحويل', 'مصروف'):
        return -amount
    if typ in ('معالجة فرق', 'تسوية'):
        direction = (row.get('settlement_direction', '') or '').strip()
        return amount if direction == 'زيادة' else -amount
    return 0.0


def agent_custody_balance(db, agent_name=''):
    total = 0.0
    agent_name = (agent_name or '').strip()
    for row in agents_custody_rows(db):
        if agent_name and (row.get('agent', '') or '').strip() != agent_name:
            continue
        total += agent_custody_effect(row)
    return round(total, 2)


def agents_custody_summary(db):
    summary = {}
    for row in agents_custody_rows(db):
        agent = (row.get('agent', '') or '').strip() or 'بدون اسم'
        entry = summary.setdefault(agent, {'agent': agent, 'incoming': 0.0, 'transfers': 0.0, 'expenses': 0.0, 'settlements': 0.0, 'balance': 0.0})
        amount = float(row.get('amount', 0) or 0)
        typ = (row.get('type', '') or '').strip()
        if typ == 'وارد مستلم':
            entry['incoming'] += amount
        elif typ == 'تحويل':
            entry['transfers'] += amount
        elif typ == 'مصروف':
            entry['expenses'] += amount
        elif typ in ('معالجة فرق', 'تسوية'):
            entry['settlements'] += agent_custody_effect(row)
        entry['balance'] += agent_custody_effect(row)
    return sorted(summary.values(), key=lambda x: x['agent'])


def total_agents_base_custody(db):
    total = 0.0
    for row in agents_custody_rows(db):
        typ = (row.get('type', '') or '').strip()
        amount = float(row.get('amount', 0) or 0)
        if typ == 'وارد مستلم':
            total += amount
        elif typ in ('تحويل', 'مصروف'):
            total -= amount
    return round(total, 2)


def total_agents_settlements(db):
    total = 0.0
    for row in agents_custody_rows(db):
        typ = (row.get('type', '') or '').strip()
        if typ in ('معالجة فرق', 'تسوية'):
            total += agent_custody_effect(row)
    return round(total, 2)


def total_agents_custody(db):
    return round(total_agents_base_custody(db) + total_agents_settlements(db), 2)

def reconciliation_data(db):
    rec = db.setdefault('reconciliation', {'current_actual_cash': 0.0, 'current_notes': '', 'records': []})
    rec.setdefault('current_actual_cash', 0.0)
    rec.setdefault('current_notes', '')
    rec.setdefault('records', [])
    rec.setdefault('actual_cash_anchor', None)
    rec.setdefault('actual_cash_anchor_book', None)
    rec.setdefault('actual_cash_anchor_date', '')
    for row in rec['records']:
        row.setdefault('id', generate_id('rec'))
        row.setdefault('date', now_str())
        row.setdefault('actual_cash', 0.0)
        row.setdefault('book_cash', 0.0)
        row.setdefault('agents_base_custody', 0.0)
        row.setdefault('settlements_value', 0.0)
        row.setdefault('agents_custody', 0.0)
        row.setdefault('customer_dues', 0.0)
        row.setdefault('payables', 0.0)
        row.setdefault('diff_cash', 0.0)
        row.setdefault('diff_with_agents', 0.0)
        row.setdefault('notes', '')
    return rec

def reconciliation_metrics(db, actual_cash=None):
    rec = reconciliation_data(db)
    actual_input = round(float(rec.get('current_actual_cash', 0) or 0) if actual_cash is None else float(actual_cash or 0), 2)
    # القاصة الدفترية لازم تُسحب من القاصة النظرية بالنظام، مو من الكاش الفعلي.
    book_cash = round(cash_balance(db), 2)
    agents_base = round(total_agents_base_custody(db), 2)
    settlements_value = round(total_agents_settlements(db), 2)
    agents_value = round(agents_base + settlements_value, 2)
    customer_dues = round(total_customer_dues(db), 2)
    payables = round(total_payables(db), 2)
    diff_cash = round(actual_input - book_cash, 2)
    diff_with_agents = round((actual_input + agents_value) - book_cash, 2)
    return {
        'actual_cash': actual_input,
        'book_cash': book_cash,
        'agents_base_custody': agents_base,
        'settlements_value': settlements_value,
        'agents_custody': agents_value,
        'customer_dues': customer_dues,
        'payables': payables,
        'diff_cash': diff_cash,
        'diff_with_agents': diff_with_agents,
    }

def total_customer_dues(db):
    return sum(float(x.get('due', 0) or 0) for x in customer_due_summary(db))

def total_payables(db):
    return sum(float(x.get('due', 0) or 0) for x in supplier_due_summary(db))

def total_funders_capital(db):
    funders = active_funders(db)
    total = sum(funder_effective_capital(db, f) for f in funders)
    return total if total > 0 else opening_cash_balance(db)

def total_funders_paid_profit(db):
    return sum(float(x.get('amount', 0) or 0) for x in db.get('cash', []) if x.get('source') == 'funder_profit_payment')

def actual_cash_on_hand(db):
    # الكاش الفعلي = آخر كاش فعلي مثبّت بالمطابقة + صافي تغيّر القاصة بعد ذلك التثبيت.
    # بهذا الشكل، إذا حفظ المستخدم جردًا ثم صارت حركة (بيع/مصروف/وارد...)
    # يتحرك الكاش الفعلي تلقائيًا مع الحركات بدل ما يبقى جامد على آخر رقم يدوي.
    rec = reconciliation_data(db)
    current_book = float(cash_balance(db) or 0)

    anchor_actual = rec.get('actual_cash_anchor')
    anchor_book = rec.get('actual_cash_anchor_book')
    if anchor_actual is not None and anchor_book is not None:
        return round(float(anchor_actual or 0) + (current_book - float(anchor_book or 0)), 2)

    # توافق رجعي للبيانات القديمة: إذا أكو لقطة مطابقة محفوظة، نعتبر أحدث لقطة نقطة ارتكاز.
    records = rec.get('records', []) or []
    if records:
        latest = records[0]
        if latest.get('actual_cash') is not None and latest.get('book_cash') is not None:
            return round(float(latest.get('actual_cash', 0) or 0) + (current_book - float(latest.get('book_cash', 0) or 0)), 2)

    # توافق رجعي: إذا المستخدم فقط حافِظ الكاش الحالي بدون نقاط ارتكاز، نرجع له كما هو.
    actual_input = float(rec.get('current_actual_cash', 0) or 0)
    has_manual_reconciliation = bool(str(rec.get('current_notes', '') or '').strip()) or actual_input != 0
    if has_manual_reconciliation:
        return round(actual_input, 2)
    return current_book

def cash_source_total(db, sources, *, category=None, party=None):
    if isinstance(sources, str):
        sources = {sources}
    else:
        sources = set(sources)
    total = 0.0
    for row in db.get('cash', []):
        if row.get('source') not in sources:
            continue
        if category is not None and row.get('category') != category:
            continue
        if party is not None and row.get('party') != party:
            continue
        total += float(row.get('amount', 0) or 0)
    return total

def owner_profit_payment_sum(db, person_name=''):
    person_name = (person_name or '').strip()
    total = 0.0
    for row in db.get('cash', []):
        if row.get('source') != 'owner_profit_payment':
            continue
        party = (row.get('party', '') or '').strip()
        category = (row.get('category', '') or '').strip()
        if person_name and person_name not in (party, category.replace('دفع ربح ', '').strip()):
            continue
        total += float(row.get('amount', 0) or 0)
    return total


def cash_breakdown(db):
    # معادلة القاصة الدفترية المعتمدة:
    # الرصيد الصافي المعتمد من النظام
    # - السحوبات
    # - سحوبات الممولين
    # - أرباح الممولين المدفوعة
    # - أرباح الشراكة المدفوعة فعلياً
    # - المصاريف
    # - ذمم الموردين المدفوعة فعلياً
    # + تسديدات الديون
    # + وارد البيع النقدي
    # + إضافات الممولين
    #
    # ملاحظة مهمة:
    # عهدة المندوبين تُسجَّل كعنصر منفصل عن القاصة، ولا ننزّلها هنا مرة ثانية.
    # السبب أن بعض حركات العهدة أصلًا مرتبطة بقيود قبض/صرف دخلت في سجل الكاش،
    # لذلك خصمها من القاصة هنا كان يسبب خلطًا بين "القاصة" و"العهدة" ويولّد
    # رقم قاصة أقل من الحقيقي. بدل ذلك نُرجِع القاصة الدفترية كما هي، ونُظهر
    # العهدة بحقل مستقل مع قيمة "القاصة بعد تنزيل العهدة" عند الحاجة للعرض فقط.
    opening_cash = float(opening_cash_balance(db) or 0)

    sales_in = cash_source_total(db, 'sales_group')
    debt_collections = cash_source_total(db, {'customer_payment', 'opening_customer_payment'})
    funders_additions = cash_source_total(db, 'funder_capital_in')

    purchases_paid = cash_source_total(db, {'inbound', 'supplier_payment', 'opening_supplier_payment'})
    funders_withdrawals = cash_source_total(db, 'funder_capital_out')
    owner_withdrawals = cash_source_total(db, 'withdrawal')
    expenses_paid = cash_source_total(db, 'expense')
    funders_profit_paid = cash_source_total(db, 'funder_profit_payment')
    owners_profit_paid = cash_source_total(db, 'owner_profit_payment')
    manual_in = sum(float(x.get('amount', 0) or 0) for x in db.get('cash', []) if x.get('source') == 'manual' and x.get('type') == 'إيراد')
    manual_out = sum(float(x.get('amount', 0) or 0) for x in db.get('cash', []) if x.get('source') == 'manual' and x.get('type') == 'مصروف')
    agents_custody = total_agents_custody(db)

    final_cash = (
        opening_cash
        - owner_withdrawals
        - funders_withdrawals
        - expenses_paid
        - purchases_paid
        - funders_profit_paid
        - owners_profit_paid
        - manual_out
        + debt_collections
        + sales_in
        + funders_additions
        + manual_in
    )
    cash_after_custody = final_cash - agents_custody

    return {
        'opening_cash': opening_cash,
        'sales_in': sales_in,
        'debt_collections': debt_collections,
        'funders_additions': funders_additions,
        'purchases_paid': purchases_paid,
        'funders_withdrawals': funders_withdrawals,
        'owner_withdrawals': owner_withdrawals,
        'expenses_paid': expenses_paid,
        'funders_profit_paid': funders_profit_paid,
        'owners_profit_paid': owners_profit_paid,
        'agents_custody': agents_custody,
        'cash_after_custody': cash_after_custody,
        'final_cash': final_cash,
    }

def cash_balance(db):
    return float(cash_breakdown(db).get('final_cash', 0) or 0)

def cash_balance_after_custody(db):
    return float(cash_breakdown(db).get('cash_after_custody', 0) or 0)

def sales_profit(db):
    return sum(float(x.get('profit', 0)) for x in db['sales'])


def manual_expenses(db):
    return expenses_total(db)


def returns_loss(db):
    total = 0.0
    for row in db.get('returns', []):
        if 'profit_impact' in row:
            total += float(row.get('profit_impact', 0) or 0)
            continue
        sale_id = row.get('sale_id')
        qty = int(row.get('qty', 0) or 0)
        sale = next((s for s in db.get('sales', []) if s.get('id') == sale_id), None)
        if sale:
            unit_profit = float(sale.get('unit_price', 0) or 0) - float(sale.get('buy_price', 0) or 0)
            total += qty * unit_profit
    return total

def damaged_loss(db):
    return sum(float(x.get('total', 0)) for x in db.get('damaged', []))


def operating_profit(db):
    return sales_profit(db) - manual_expenses(db) - returns_loss(db) - damaged_loss(db)


def withdrawals_sum(db, person):
    return sum(float(x.get('amount', 0)) for x in db['cash'] if x.get('source') == 'withdrawal' and x.get('category') == person)


def is_legacy_opening_cash(row):
    return row.get('source') in ('opening_customer_payment', 'opening_supplier_payment')


def customer_receipts(db, customer_name):
    return sum(
        float(x.get('amount', 0) or 0)
        for x in db['cash']
        if x.get('source') == 'customer_payment' and x.get('party') == customer_name and not is_legacy_opening_cash(x)
    )


def supplier_payments(db, supplier_name):
    return sum(
        float(x.get('amount', 0) or 0)
        for x in db['cash']
        if x.get('source') == 'supplier_payment' and x.get('party') == supplier_name and not is_legacy_opening_cash(x)
    )



def customer_payment_allocations(db, customer_name=''):
    names = []
    customer_name = str(customer_name or '').strip()
    if customer_name:
        names = [customer_name]
    else:
        names = sorted({
            str(x.get('customer', '') or '').strip()
            for x in db.get('sales', [])
        } | {
            str(x.get('party', '') or '').strip()
            for x in db.get('cash', [])
            if x.get('source') in ('customer_payment', 'opening_customer_payment')
        })
        names = [x for x in names if x]

    events = []
    for name in names:
        snapshot = customer_receivable_snapshot(db, name)
        for ev in snapshot.get('payment_events', []):
            p = ev.get('payment_row', {})
            allocations = []
            for alloc in ev.get('allocations', []):
                sale = alloc.get('sale_row', {}) or {}
                total = _safe_float(sale.get('total', 0))
                gross_profit = _safe_float(sale.get('profit', 0)) * (_safe_float(alloc.get('amount', 0)) / total) if total > 0 else 0.0
                allocations.append({
                    'sale_id': alloc.get('sale_id', ''),
                    'invoice_group_id': alloc.get('invoice_group_id', ''),
                    'invoice_no': alloc.get('invoice_no', ''),
                    'customer': alloc.get('customer', name),
                    'sale_date': alloc.get('sale_date', ''),
                    'item': alloc.get('item', ''),
                    'amount': round(_safe_float(alloc.get('amount', 0)), 2),
                    'gross_profit': round(gross_profit, 2),
                    'line_total': round(total, 2),
                    'payment_row': p,
                })
            events.append({
                'payment_row': p,
                'customer': name,
                'opening_take': round(_safe_float(ev.get('opening_take', 0)), 2),
                'allocations': allocations,
            })
    events.sort(key=lambda x: (str((x.get('payment_row', {}) or {}).get('date', '') or ''), str((x.get('payment_row', {}) or {}).get('created_at', '') or ''), str((x.get('payment_row', {}) or {}).get('receipt_no', '') or '')))
    return events

def inbound_payment_allocations(db, supplier_name=''):
    inbound_by_supplier = {}
    for row in sorted(db.get('inbound', []), key=lambda x: (str(x.get('date', '') or ''), str(x.get('created_at', '') or ''), str(x.get('id', '') or ''))):
        total = _safe_float(row.get('total', 0))
        paid = _safe_float(row.get('paid_amount', 0))
        rem = max(0.0, total - paid)
        supplier = row.get('supplier', '')
        inbound_by_supplier.setdefault(supplier, []).append({'row': row, 'remaining': rem})

    opening_due = {s.get('name', ''): max(0.0, opening_supplier_amount(db, s.get('name', '')) - opening_supplier_payments(db, s.get('name', ''))) for s in db.get('suppliers', [])}
    events = []
    payments = [x for x in db.get('cash', []) if x.get('source') in ('supplier_payment', 'opening_supplier_payment')]
    payments.sort(key=lambda x: (str(x.get('date', '') or ''), str(x.get('created_at', '') or ''), str(x.get('receipt_no', '') or '')))
    for p in payments:
        supplier = p.get('party', '')
        if supplier_name and supplier != supplier_name:
            continue
        amount = _safe_float(p.get('amount', 0))
        remaining = amount
        old_due = opening_due.get(supplier, 0.0)
        old_take = min(old_due, remaining)
        if old_take > 0:
            opening_due[supplier] = max(0.0, old_due - old_take)
            remaining = round(remaining - old_take, 2)
        allocations = []
        for holder in inbound_by_supplier.get(supplier, []):
            if remaining <= 0:
                break
            row = holder['row']
            take = min(holder['remaining'], remaining)
            if take <= 0:
                continue
            holder['remaining'] = round(holder['remaining'] - take, 2)
            remaining = round(remaining - take, 2)
            allocations.append({
                'inbound_id': row.get('id', ''),
                'supplier': supplier,
                'amount': take,
                'payment_row': p,
            })
        events.append({'payment_row': p, 'supplier': supplier, 'opening_take': old_take, 'allocations': allocations})
    return events



def current_sale_due(db, sale_or_gid):
    if isinstance(sale_or_gid, str):
        gid = sale_or_gid
        line = next((s for s in db.get('sales', []) if (s.get('invoice_group_id') or f"single-{ensure_invoice_no(db, s)}") == gid), None)
    else:
        line = sale_or_gid
        gid = line.get('invoice_group_id') or f"single-{ensure_invoice_no(db, line)}"
    if not line:
        return 0.0
    snapshot = customer_receivable_snapshot(db, line.get('customer', ''))
    return round(max(0.0, float(snapshot.get('invoice_due_map', {}).get(gid, 0.0) or 0.0)), 2)

def current_inbound_due(db, inbound_or_id):
    inbound_id = inbound_or_id if isinstance(inbound_or_id, str) else inbound_or_id.get('id', '')
    row = next((x for x in db.get('inbound', []) if x.get('id', '') == inbound_id), None)
    if not row:
        return 0.0
    base_due = max(0.0, _safe_float(row.get('total', 0)) - _safe_float(row.get('paid_amount', 0)))
    later_paid = 0.0
    for ev in inbound_payment_allocations(db, row.get('supplier', '')):
        later_paid += sum(_safe_float(a.get('amount', 0)) for a in ev.get('allocations', []) if a.get('inbound_id') == inbound_id)
    return round(max(0.0, base_due - later_paid), 2)



def customer_due_summary(db):
    rows = []
    for c in db.get('customers', []):
        name = c.get('name', '')
        opening_original = opening_customer_amount(db, name)
        old_receipts = opening_customer_receipts(db, name)

        sales_rows = [s for s in db.get('sales', []) if s.get('customer') == name]
        invoice_total = sum(float(s.get('total', 0) or 0) for s in sales_rows)
        initial_paid = sum(float(s.get('paid_amount', 0) or 0) for s in sales_rows)
        later_receipts = customer_receipts(db, name)

        snapshot = customer_receivable_snapshot(db, name)
        return_credit = float(snapshot.get('carry_credit_total', 0) or 0)
        due = float(snapshot.get('final_due', 0) or 0)
        opening_due = float(snapshot.get('opening_remaining', 0) or 0)

        gross = opening_original + invoice_total
        receipts = old_receipts + later_receipts + initial_paid

        rows.append({
            'name': name,
            'credit_sales': gross,
            'receipts': receipts,
            'return_credit': round(return_credit, 2),
            'due': round(due, 2),
            'opening_due': round(opening_due, 2),
            'opening_original': opening_original,
            'invoice_total': invoice_total,
            'initial_paid': initial_paid,
            'later_receipts': later_receipts,
        })
    return rows

def supplier_due_summary(db):
    rows = []
    for s in db.get('suppliers', []):
        name = s.get('name', '')
        opening_original = opening_supplier_amount(db, name)
        old_payments = opening_supplier_payments(db, name)

        inbound_rows = [i for i in db.get('inbound', []) if i.get('supplier') == name]
        purchase_total = sum(float(i.get('total', 0) or 0) for i in inbound_rows)
        initial_paid = sum(float(i.get('paid_amount', 0) or 0) for i in inbound_rows)
        later_payments = supplier_payments(db, name)

        gross = opening_original + purchase_total
        payments = old_payments + later_payments + initial_paid
        due = max(0.0, gross - payments)

        rows.append({
            'name': name,
            'credit_purchases': gross,
            'payments': payments,
            'due': due,
            'opening_due': max(0.0, opening_original - old_payments),
            'opening_original': opening_original,
            'purchase_total': purchase_total,
            'initial_paid': initial_paid,
            'later_payments': later_payments,
        })
    return rows

def supplier_credit_rows(db):
    rows = []
    for s in db.get('suppliers', []):
        name = s.get('name', '')
        opening_amt = opening_supplier_amount(db, name)
        credit_purchases = sum(float(i.get('due_amount', 0)) for i in db.get('inbound', []) if i.get('supplier') == name) + opening_amt
        payments = supplier_payments(db, name)
        due = max(0.0, credit_purchases - payments)
        rows.append({'name': name, 'credit_purchases': credit_purchases, 'payments': payments, 'due': due, 'opening_due': opening_amt})
    return rows


def available_cash(db):
    return cash_balance(db)

def financial_position(db):
    # الوضع المالي يقيس رأس المال التشغيلي الحالي:
    # كاش فعلي + قيمة المخزون + الذمم التشغيلية الجديدة - التزامات الموردين.
    # نستثني ديون الزبائن الافتتاحية القديمة حتى لا تنزل من القاصة أو تظهر كخسارة من نقطة البداية.
    receivables = sum(
        max(0.0, float(x.get('due', 0) or 0) - float(x.get('opening_due', 0) or 0))
        for x in customer_due_summary(db)
    )
    payables = sum(x['due'] for x in supplier_due_summary(db))
    return available_cash(db) + inventory_value(db) + receivables - payables

def capital_reference(db):
    funders_capital = sum(funder_effective_capital(db, f) for f in active_funders(db))
    if funders_capital > 0:
        return funders_capital
    return opening_cash_balance(db)

def capital_status(db):
    capital = capital_reference(db)
    position = financial_position(db)
    diff = position - capital
    pct = ((diff / capital) * 100.0) if capital else 0.0
    return {'capital': capital, 'position': position, 'diff': diff, 'pct': pct, 'loss': diff < 0}


def funder_effective_capital(db, funder_row):
    name = funder_row.get('name', '')
    base = funder_capital_base(db, name)
    base += funder_total_deposit(db, name)
    base -= funder_total_withdraw_capital(db, name)
    return max(0.0, base)


def normalized_funders(db):
    """Return a merged current+opening view of funders without dropping inactive rows.
    Historical profit snapshots depend on seeing the full list, then applying activity
    as-of each event date rather than filtering by the current active flag.
    """
    current_rows = db.get('funders', []) or []
    opening_rows = opening_data(db).get('old_funders', []) or []

    merged = {}
    for row in current_rows:
        name = (row.get('name', '') or '').strip()
        if not name:
            continue
        base = dict(row)
        ensure_funder_identity(base, name)
        base['_from_opening'] = False
        merged[name] = base

    for old in opening_rows:
        name = (old.get('name', '') or '').strip()
        if not name:
            continue
        old_capital = max(0.0, _safe_float(old.get('capital', 0)) - _safe_float(old.get('withdrawals', 0)))
        opening_stamp = str(old.get('created_at', '') or old.get('date', '') or '')
        current = merged.get(name)
        if current is None:
            base = {
                'name': name,
                'capital': old_capital,
                'base_capital': old_capital,
                'phone': '',
                'notes': '',
                'active': True,
                'is_owner_capital': bool(name == 'هضاب'),
                'created_at': opening_stamp,
                'status_history': ([] if not opening_stamp else [{'active': True, 'date': opening_stamp[:10], 'created_at': opening_stamp}]),
                '_from_opening': True,
            }
            ensure_funder_identity(base, name)
            merged[name] = base
            continue

        if _safe_float(current.get('base_capital', current.get('capital', 0))) <= 0 and old_capital > 0:
            current['base_capital'] = old_capital
        if _safe_float(current.get('capital', 0)) <= 0 and old_capital > 0:
            current['capital'] = old_capital
        if opening_stamp and not str(current.get('created_at', '') or '').strip():
            current['created_at'] = opening_stamp
        history = list(current.get('status_history', []) or [])
        if not history and opening_stamp:
            history = [{'active': bool(current.get('active', True)), 'date': opening_stamp[:10], 'created_at': opening_stamp}]
        if opening_stamp and history:
            earliest = history[0]
            if not str(earliest.get('created_at', '') or '').strip():
                history[0] = {'active': bool(earliest.get('active', current.get('active', True))), 'date': opening_stamp[:10], 'created_at': opening_stamp}
        current['status_history'] = history
        current['_from_opening'] = True
        ensure_funder_identity(current, name)

    rows = list(merged.values())
    rows.sort(key=lambda x: ((x.get('name', '') != 'هضاب'), str(x.get('name', '') or '')))
    return rows

def active_funders(db):
    return [f for f in normalized_funders(db) if f.get('active', True) and funder_effective_capital(db, f) > 0]


def profit_breakdown(db):
    build_profit_ledger(db)
    settings = db.get('settings', {})
    funders = active_funders(db)
    funder_rows = []
    external_total = 0.0
    owner_capital_profit = 0.0
    partner_entries = [x for x in db.get('profit_entries', []) if x.get('beneficiary_type') == 'partner']
    funder_entries = [x for x in db.get('profit_entries', []) if x.get('beneficiary_type') == 'funder']

    configured_pool_pct = _safe_float(settings.get('funders_profit_pool_pct', 33.3333))
    for f in funders:
        name = f.get('name', '')
        cap = funder_effective_capital(db, f)
        earned = sum(_safe_float(x.get('amount', 0)) for x in funder_entries if x.get('beneficiary_name') == name)
        paid = funder_total_paid_profit(db, name)
        pending = round(earned - paid, 2)
        row = {
            'name': name,
            'capital': cap,
            'ratio': funder_ratio_pct(db, f),
            'amount': round(earned, 2),
            'paid': round(paid, 2),
            'pending': pending,
            'is_owner_capital': bool(f.get('is_owner_capital', False)),
            'type_label': 'رأس مال هضاب' if bool(f.get('is_owner_capital', False)) else 'ممّول خارجي',
        }
        funder_rows.append(row)
        if row['is_owner_capital']:
            owner_capital_profit += earned
        else:
            external_total += earned

    hidab_share = sum(_safe_float(x.get('amount', 0)) for x in partner_entries if x.get('beneficiary_name') == 'هضاب')
    mostafa_share = sum(_safe_float(x.get('amount', 0)) for x in partner_entries if x.get('beneficiary_name') == 'مصطفى')
    hidab_withdrawals = withdrawals_sum(db, 'سحوبات هضاب')
    mostafa_withdrawals = withdrawals_sum(db, 'سحوبات مصطفى')
    hidab_profit_paid = owner_profit_payment_sum(db, 'هضاب')
    mostafa_profit_paid = owner_profit_payment_sum(db, 'مصطفى')
    operating_profit_value = round(sum(_safe_float(x.get('net_profit', 0)) for x in db.get('profit_events', [])), 2)
    funders_pool = round(sum(max(0.0, _safe_float(x.get('funders_pool', 0))) for x in db.get('profit_events', []) if x.get('event_type') == 'payment'), 2)
    owner_pct = _safe_float(settings.get('owner_profit_pct', 50.0))
    partner_pct = _safe_float(settings.get('partner_profit_pct', 50.0))
    total_pct = owner_pct + partner_pct
    owner_ratio = (owner_pct / total_pct * 100.0) if total_pct > 0 else 50.0
    partner_ratio = (partner_pct / total_pct * 100.0) if total_pct > 0 else 50.0
    distributable_profit = round(hidab_share + mostafa_share, 2)
    return {
        'operating_profit': operating_profit_value,
        'funders_pool_pct': configured_pool_pct,
        'funders_pool': funders_pool,
        'funders_rows': funder_rows,
        'external_funders_total': round(external_total, 2),
        'owner_capital_profit': round(owner_capital_profit, 2),
        'owner_profit_added_back': round(owner_capital_profit, 2),
        'distributable_profit': distributable_profit,
        'hidab_share': round(hidab_share, 2),
        'mostafa_share': round(mostafa_share, 2),
        'hidab_withdrawals': round(hidab_withdrawals, 2),
        'mostafa_withdrawals': round(mostafa_withdrawals, 2),
        'hidab_profit_paid': round(hidab_profit_paid, 2),
        'mostafa_profit_paid': round(mostafa_profit_paid, 2),
        'hidab_net': round(hidab_share - hidab_withdrawals - hidab_profit_paid, 2),
        'mostafa_net': round(mostafa_share - mostafa_withdrawals - mostafa_profit_paid, 2),
        'owner_ratio_pct': owner_ratio,
        'partner_ratio_pct': partner_ratio,
    }

def profit_ui_snapshot(db):
    build_profit_ledger(db)
    op = opening_data(db)

    operating_profit = round(sum(_safe_float(x.get("net_profit", 0)) for x in db.get("profit_events", [])), 2)
    if abs(operating_profit) < 0.009:
        operating_profit = round(net_profit(db), 2)
    external_funders_profit = 0.0
    hidab_financing_profit = 0.0
    hidab_partner_profit = 0.0
    mostafa_partner_profit = 0.0

    funders = normalized_funders(db)
    funder_index = {str(f.get("name", "")).strip(): f for f in funders}
    entries = db.get("profit_entries", [])

    for row in entries:
        amount = float(row.get("amount", 0) or 0)
        btype = row.get("beneficiary_type", "")
        bname = row.get("beneficiary_name", "")

        if btype == "funder":
            funder = funder_index.get(str(bname or '').strip())
            if funder and bool(funder.get("is_owner_capital", False)):
                hidab_financing_profit += amount
            else:
                external_funders_profit += amount
        elif btype == "partner":
            if bname == "هضاب":
                hidab_partner_profit += amount
            elif bname == "مصطفى":
                mostafa_partner_profit += amount

    partnership_profit = hidab_partner_profit + mostafa_partner_profit
    hidab_withdrawals = float(withdrawals_sum(db, "سحوبات هضاب") or 0)
    mostafa_withdrawals = float(withdrawals_sum(db, "سحوبات مصطفى") or 0)
    hidab_opening_balance = float(op.get("hidab_opening_balance", 0) or 0)
    mostafa_opening_balance = float(op.get("mustafa_opening_balance", 0) or 0)

    hidab_final = hidab_partner_profit - hidab_withdrawals + hidab_opening_balance
    mostafa_final = mostafa_partner_profit - mostafa_withdrawals + mostafa_opening_balance

    return {
        "operating_profit": round(operating_profit, 2),
        "external_funders_profit": round(external_funders_profit, 2),
        "hidab_financing_profit": round(hidab_financing_profit, 2),
        "partnership_profit": round(partnership_profit, 2),
        "hidab_partner_profit": round(hidab_partner_profit, 2),
        "mostafa_partner_profit": round(mostafa_partner_profit, 2),
        "hidab_withdrawals": round(hidab_withdrawals, 2),
        "mostafa_withdrawals": round(mostafa_withdrawals, 2),
        "hidab_opening_balance": round(hidab_opening_balance, 2),
        "mostafa_opening_balance": round(mostafa_opening_balance, 2),
        "hidab_final": round(hidab_final, 2),
        "mostafa_final": round(mostafa_final, 2),
    }


def validate_profit_consistency(db):
    snap = profit_ui_snapshot(db)
    pb = profit_breakdown(db)
    op = opening_data(db)
    expected = {
        "external_funders_profit": round(float(pb.get("external_funders_total", 0) or 0), 2),
        "hidab_financing_profit": round(float(pb.get("owner_capital_profit", 0) or 0), 2),
        "partnership_profit": round(float(pb.get("distributable_profit", 0) or 0), 2),
        "hidab_final": round(float(pb.get("hidab_share", 0) or 0) - float(pb.get("hidab_withdrawals", 0) or 0) + float(op.get("hidab_opening_balance", 0) or 0), 2),
        "mostafa_final": round(float(pb.get("mostafa_share", 0) or 0) - float(pb.get("mostafa_withdrawals", 0) or 0) + float(op.get("mustafa_opening_balance", 0) or 0), 2),
    }
    actual = {
        "external_funders_profit": round(float(snap.get("external_funders_profit", 0) or 0), 2),
        "hidab_financing_profit": round(float(snap.get("hidab_financing_profit", 0) or 0), 2),
        "partnership_profit": round(float(snap.get("partnership_profit", 0) or 0), 2),
        "hidab_final": round(float(snap.get("hidab_final", 0) or 0), 2),
        "mostafa_final": round(float(snap.get("mostafa_final", 0) or 0), 2),
    }
    diffs = {}
    for key, report_value in expected.items():
        if round(report_value, 2) != round(actual[key], 2):
            diffs[key] = {"report": round(report_value, 2), "ui": round(actual[key], 2)}
    return {"ok": len(diffs) == 0, "diffs": diffs}


def unpaid_external_funders_profit(db):
    pb = profit_breakdown(db)
    return sum(float(x.get('pending', 0) or 0) for x in pb.get('funders_rows', []) if not x.get('is_owner_capital', False))


def total_commitments(db):
    st = person_profit_status(db)
    return total_payables(db) + unpaid_external_funders_profit(db) + float(st.get('owners_pending_profit', 0) or 0)


def net_cash_after_commitments(db):
    return actual_cash_on_hand(db) - total_commitments(db)

def person_profit_status(db):
    pb = profit_breakdown(db)
    op = opening_data(db)

    hidab_opening = float(op.get('hidab_opening_balance', 0) or 0)
    mostafa_opening = float(op.get('mustafa_opening_balance', 0) or 0)

    # ربح الفترة القابل للتوزيع للشراكة = ربح الشراكة المباشر + حصة هضاب التمويلية الراجعة للشراكة
    partnership_period_profit = float(pb.get('distributable_profit', 0) or 0) + float(pb.get('owner_capital_profit', 0) or 0)
    period_share_each = round(partnership_period_profit / 2.0, 2)

    hidab_share = period_share_each
    mostafa_share = period_share_each
    hidab_withdrawals = float(pb.get('hidab_withdrawals', 0) or 0)
    mostafa_withdrawals = float(pb.get('mostafa_withdrawals', 0) or 0)
    hidab_profit_paid = owner_profit_payment_sum(db, 'هضاب')
    mostafa_profit_paid = owner_profit_payment_sum(db, 'مصطفى')

    # العجز = الرصيد/العجز الافتتاحي + السحوبات/المدفوع - ربح الفترة
    hidab_total_deficit_before_profit = round(hidab_opening + hidab_withdrawals + hidab_profit_paid, 2)
    mostafa_total_deficit_before_profit = round(mostafa_opening + mostafa_withdrawals + mostafa_profit_paid, 2)

    hidab_settled = min(hidab_share, hidab_total_deficit_before_profit)
    mostafa_settled = min(mostafa_share, mostafa_total_deficit_before_profit)

    hidab_deficit = max(0.0, round(hidab_total_deficit_before_profit - hidab_share, 2))
    mostafa_deficit = max(0.0, round(mostafa_total_deficit_before_profit - mostafa_share, 2))

    hidab_surplus = max(0.0, round(hidab_share - hidab_total_deficit_before_profit, 2))
    mostafa_surplus = max(0.0, round(mostafa_share - mostafa_total_deficit_before_profit, 2))

    return {
        'hidab_balance': hidab_deficit,
        'mostafa_balance': mostafa_deficit,
        'hidab_deficit': hidab_deficit,
        'mostafa_deficit': mostafa_deficit,
        'hidab_overdraw': 0.0,
        'mostafa_overdraw': 0.0,
        'hidab_opening_balance': hidab_opening,
        'mustafa_opening_balance': mostafa_opening,
        'hidab_profit_paid': hidab_profit_paid,
        'mostafa_profit_paid': mostafa_profit_paid,
        'hidab_withdrawals': hidab_withdrawals,
        'mostafa_withdrawals': mostafa_withdrawals,
        'hidab_share': hidab_share,
        'mostafa_share': mostafa_share,
        'hidab_total_deficit_before_profit': hidab_total_deficit_before_profit,
        'mostafa_total_deficit_before_profit': mostafa_total_deficit_before_profit,
        'hidab_settled': hidab_settled,
        'mostafa_settled': mostafa_settled,
        'hidab_surplus': hidab_surplus,
        'mostafa_surplus': mostafa_surplus,
        'owners_pending_profit': hidab_deficit + mostafa_deficit,
        'total_deficit': hidab_deficit + mostafa_deficit,
    }

def funder_ratio_pct(db, funder_row):
    if not funder_row.get('active', True):
        return 0.0
    funders = active_funders(db)
    total_capital = sum(funder_effective_capital(db, f) for f in funders)
    if not total_capital:
        return 0.0
    cap = funder_effective_capital(db, funder_row)
    return (cap / total_capital) * 100.0

def net_profit(db):
    return operating_profit(db)


class QuickCalculatorPanel(QFrame):
    def __init__(self, title='حاسبة سريعة', note='للمراجعة السريعة داخل الصفحة.', compact=False, parent=None):
        super().__init__(parent)
        self.setObjectName('quickCalcPanel')
        self.setStyleSheet(CARD_FRAME_STYLE)
        self.setMinimumWidth(280 if compact else 300)
        wrap = QVBoxLayout(self)
        wrap.setContentsMargins(16, 16, 16, 16)
        wrap.setSpacing(10)
        title_lbl = QLabel(title)
        title_lbl.setAlignment(Qt.AlignRight)
        title_lbl.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        note_lbl = QLabel(note)
        note_lbl.setAlignment(Qt.AlignRight)
        note_lbl.setWordWrap(True)
        note_lbl.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        self.display = QLineEdit()
        self.display.setAlignment(Qt.AlignRight)
        self.display.setReadOnly(True)
        self.display.setPlaceholderText('0')
        self.display.setMinimumHeight(52 if compact else 54)
        self.display.setStyleSheet(f"font-size:{20 if compact else 22}px;font-weight:900;padding:10px 14px;border-radius:14px;border:1px solid {BORDER};background:{CARD};")
        wrap.addWidget(title_lbl)
        wrap.addWidget(note_lbl)
        wrap.addWidget(self.display)
        grid = QGridLayout()
        grid.setHorizontalSpacing(8)
        grid.setVerticalSpacing(8)
        buttons = [
            ('C', 0, 0), ('⌫', 0, 1), ('%', 0, 2), ('÷', 0, 3),
            ('7', 1, 0), ('8', 1, 1), ('9', 1, 2), ('×', 1, 3),
            ('4', 2, 0), ('5', 2, 1), ('6', 2, 2), ('-', 2, 3),
            ('1', 3, 0), ('2', 3, 1), ('3', 3, 2), ('+', 3, 3),
            ('0', 4, 0), ('000', 4, 1), ('.', 4, 2), ('=', 4, 3),
        ]
        for txt, row, col in buttons:
            btn = QPushButton(txt)
            btn.setMinimumHeight(42 if compact else 44)
            btn.setStyleSheet(BUTTON_STYLE if txt == '=' else SECONDARY_BUTTON)
            btn.clicked.connect(lambda _=False, value=txt: self.press(value))
            grid.addWidget(btn, row, col)
        wrap.addLayout(grid)

    def press(self, value):
        current = self.display.text().strip()
        if value == 'C':
            self.display.clear()
            return
        if value == '⌫':
            self.display.setText(current[:-1])
            return
        if value == '=':
            self.evaluate()
            return
        mapped = {'×': '*', '÷': '/'}
        token = mapped.get(value, value)
        if value == '%':
            token = '/100'
        if value in ('+', '-', '×', '÷') and (not current or current[-1] in '+-*/.'):
            return
        if value == '.' and (not current or current[-1] in '+-*/'):
            token = '0.'
        self.display.setText(current + token)

    def evaluate(self):
        expr = (self.display.text() or '').strip()
        if not expr:
            return
        if not all(ch in '0123456789+-*/(). ' for ch in expr):
            return QMessageBox.warning(self, 'تنبيه', 'العملية تحتوي رموز غير مدعومة.')
        try:
            result = eval(expr, {'__builtins__': {}}, {})
            number = float(result)
            if abs(number - int(number)) < 0.0000001:
                shown = str(int(number))
            else:
                shown = ('{:.2f}'.format(number)).rstrip('0').rstrip('.')
            self.display.setText(shown)
        except ZeroDivisionError:
            QMessageBox.warning(self, 'تنبيه', 'ما يصير القسمة على صفر.')
        except Exception:
            QMessageBox.warning(self, 'تنبيه', 'العملية غير صحيحة.')


class SummaryCard(QFrame):
    def __init__(self, title, value='', note='', accent=None):
        super().__init__()
        self.accent = accent
        self.setObjectName('summaryCard')
        self.setMinimumHeight(108)
        self.setMaximumHeight(122)
        self.setCursor(Qt.PointingHandCursor)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
        lay = QVBoxLayout(self)
        lay.setContentsMargins(18, 14, 18, 14)
        lay.setSpacing(4)
        self.title_label = QLabel(title)
        self.title_label.setAlignment(Qt.AlignRight)
        self.title_label.setWordWrap(True)
        self.title_label.setFrameShape(QFrame.NoFrame)
        self.value_label = QLabel(value)
        self.value_label.setAlignment(Qt.AlignRight)
        self.value_label.setFrameShape(QFrame.NoFrame)
        self.note_label = QLabel(note)
        self.note_label.setAlignment(Qt.AlignRight)
        self.note_label.setWordWrap(True)
        self.note_label.setFrameShape(QFrame.NoFrame)
        self.note_label.hide()
        lay.addWidget(self.title_label)
        lay.addWidget(self.value_label)
        lay.addStretch(1)
        self.shadow = QGraphicsDropShadowEffect(self)
        self.shadow.setBlurRadius(26)
        self.shadow.setOffset(0, 10)
        self.setGraphicsEffect(self.shadow)
        self.shadow_anim = QPropertyAnimation(self.shadow, b"blurRadius", self)
        self.shadow_anim.setDuration(180)
        self.shadow_anim.setEasingCurve(QEasingCurve.OutCubic)
        self.offset_anim = QPropertyAnimation(self.shadow, b"offset", self)
        self.offset_anim.setDuration(180)
        self.offset_anim.setEasingCurve(QEasingCurve.OutCubic)
        self.anim_group = QParallelAnimationGroup(self)
        self.anim_group.addAnimation(self.shadow_anim)
        self.anim_group.addAnimation(self.offset_anim)
        self.apply_theme()
    def set_value(self, value):
        self.value_label.setText(value)
    def setText(self, value):
        self.value_label.setText(value)
    def text(self):
        return self.value_label.text()
    def set_note(self, note):
        self.note_label.setText(note)
    def _animate_shadow(self, blur, y_offset):
        self.anim_group.stop()
        self.shadow_anim.setStartValue(self.shadow.blurRadius())
        self.shadow_anim.setEndValue(blur)
        self.offset_anim.setStartValue(self.shadow.offset())
        self.offset_anim.setEndValue(QPoint(0, y_offset))
        self.anim_group.start()
    def enterEvent(self, event):
        self._animate_shadow(42, 16)
        return super().enterEvent(event)
    def leaveEvent(self, event):
        self._animate_shadow(26, 10)
        return super().leaveEvent(event)
    def _sync_scroll_slider(self, minimum, maximum):
        try:
            self.scroll_slider.blockSignals(True)
            self.scroll_slider.setRange(int(minimum), int(maximum if maximum > minimum else minimum + 1))
            self.scroll_slider.setPageStep(max(40, int(self.scroll.verticalScrollBar().pageStep())))
            self.scroll_slider.setValue(int(self.scroll.verticalScrollBar().value()))
        finally:
            self.scroll_slider.blockSignals(False)

    def apply_theme(self):
        accent = self.accent or ACCENT
        self.setStyleSheet(f"QFrame#summaryCard{{background:qlineargradient(x1:0,y1:0,x2:1,y2:1, stop:0 {rgba_from_hex(CARD,0.99)}, stop:1 {rgba_from_hex(CARD,0.88)});border:none;border-radius:30px;}} QLabel{{background:transparent;border:none;outline:none;color:{TEXT};padding:0;margin:0;}}")
        self.shadow.setColor(QColor(rgba_from_hex('#000000', 0.18)))
        self.title_label.setStyleSheet(f'font-size:11px;font-weight:800;color:{MUTED};letter-spacing:0.2px;border:none;outline:none;background:transparent;padding:0;')
        self.value_label.setStyleSheet(f'font-size:22px;font-weight:900;color:{TEXT};border:none;outline:none;background:transparent;padding:0;')
        self.note_label.setStyleSheet(f'font-size:10px;font-weight:800;color:{accent};border:none;outline:none;background:transparent;padding:0;')

class DashboardActionCard(QFrame):
    clicked = Signal()
    def __init__(self, emoji, title, subtitle=''):
        super().__init__()
        self.setObjectName('dashboardActionCard')
        self.setCursor(Qt.PointingHandCursor)
        self.setMinimumHeight(176)
        self.setMaximumHeight(188)
        self.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
        box = QVBoxLayout(self)
        box.setContentsMargins(20, 18, 20, 18)
        box.setSpacing(8)
        self.badge = QLabel(emoji)
        self.badge.setAlignment(Qt.AlignCenter)
        self.badge.setFixedSize(52, 52)
        self.title_lbl = QLabel(title)
        self.title_lbl.setAlignment(Qt.AlignRight)
        self.title_lbl.setWordWrap(True)
        self.subtitle_lbl = QLabel(subtitle)
        self.subtitle_lbl.setAlignment(Qt.AlignRight)
        self.subtitle_lbl.setWordWrap(True)
        self.cta_lbl = QLabel('فتح القسم')
        self.cta_lbl.setAlignment(Qt.AlignRight)
        box.addWidget(self.badge, alignment=Qt.AlignRight)
        box.addWidget(self.title_lbl)
        box.addWidget(self.subtitle_lbl)
        box.addStretch(1)
        box.addWidget(self.cta_lbl)
        self.shadow = QGraphicsDropShadowEffect(self)
        self.shadow.setBlurRadius(28)
        self.shadow.setOffset(0, 10)
        self.setGraphicsEffect(self.shadow)
        self.shadow_anim = QPropertyAnimation(self.shadow, b"blurRadius", self)
        self.shadow_anim.setDuration(180)
        self.shadow_anim.setEasingCurve(QEasingCurve.OutCubic)
        self.offset_anim = QPropertyAnimation(self.shadow, b"offset", self)
        self.offset_anim.setDuration(180)
        self.offset_anim.setEasingCurve(QEasingCurve.OutCubic)
        self.anim_group = QParallelAnimationGroup(self)
        self.anim_group.addAnimation(self.shadow_anim)
        self.anim_group.addAnimation(self.offset_anim)
        self.apply_theme()
    def set_subtitle(self, text):
        self.subtitle_lbl.setText(text)
    def set_meta(self, text):
        self.cta_lbl.setText(text)
    def _animate_shadow(self, blur, y_offset):
        self.anim_group.stop()
        self.shadow_anim.setStartValue(self.shadow.blurRadius())
        self.shadow_anim.setEndValue(blur)
        self.offset_anim.setStartValue(self.shadow.offset())
        self.offset_anim.setEndValue(QPoint(0, y_offset))
        self.anim_group.start()
    def _sync_scroll_slider(self, minimum, maximum):
        try:
            self.scroll_slider.blockSignals(True)
            self.scroll_slider.setRange(int(minimum), int(maximum if maximum > minimum else minimum + 1))
            self.scroll_slider.setPageStep(max(40, int(self.scroll.verticalScrollBar().pageStep())))
            self.scroll_slider.setValue(int(self.scroll.verticalScrollBar().value()))
        finally:
            self.scroll_slider.blockSignals(False)

    def apply_theme(self):
        self.default_style = f"QFrame#dashboardActionCard{{background:qlineargradient(x1:0,y1:0,x2:1,y2:1, stop:0 {rgba_from_hex(CARD,0.99)}, stop:1 {rgba_from_hex(CARD,0.88)});border:none;border-radius:34px;}} QLabel{{background:transparent;border:none;color:{TEXT};}}"
        self.hover_style = f"QFrame#dashboardActionCard{{background:qlineargradient(x1:0,y1:0,x2:1,y2:1, stop:0 {rgba_from_hex(CARD,1.0)}, stop:1 {rgba_from_hex(CARD,0.94)});border:none;border-radius:34px;}} QLabel{{background:transparent;border:none;color:{TEXT};}}"
        self.setStyleSheet(self.default_style)
        self.shadow.setColor(QColor(rgba_from_hex('#000000', 0.20)))
        self.badge.setStyleSheet(f'background:{rgba_from_hex(ACCENT,0.14)};color:{TEXT};border:none;border-radius:18px;font-size:22px;font-weight:900;')
        self.title_lbl.setStyleSheet(f'font-size:20px;font-weight:900;color:{TEXT};')
        self.subtitle_lbl.setStyleSheet(f'font-size:11px;font-weight:700;color:{MUTED};line-height:1.35;')
        self.cta_lbl.setStyleSheet(f'font-size:11px;font-weight:900;color:{ACCENT};')
    def enterEvent(self, event):
        self.setStyleSheet(self.hover_style)
        self._animate_shadow(48, 18)
        return super().enterEvent(event)
    def leaveEvent(self, event):
        self.setStyleSheet(self.default_style)
        self._animate_shadow(28, 10)
        return super().leaveEvent(event)
    def mousePressEvent(self, event):
        self.clicked.emit()
        return super().mousePressEvent(event)

def mindflow_tab_style(min_width=190):
    return f"""
    QTabWidget::pane {{
        border: 1px solid {rgba_from_hex(TEXT,0.08)};
        background: {rgba_from_hex(CARD,0.78)};
        border-radius: 28px;
        top: -4px;
    }}
    QTabBar::tab {{
        min-width: {min_width}px;
        min-height: 48px;
        padding: 10px 18px;
        margin: 4px 8px 10px 0;
        border-radius: 18px;
        background: {rgba_from_hex(CARD,0.56)};
        color: {TEXT};
        font-size: 14px;
        font-weight: 800;
        border: 1px solid {rgba_from_hex(TEXT,0.08)};
    }}
    QTabBar::tab:hover {{
        background: {rgba_from_hex(CARD,0.92)};
        border: 1px solid {rgba_from_hex(ACCENT,0.26)};
    }}
    QTabBar::tab:selected {{
        background: {ACCENT};
        color: {TEXT_ON_ACCENT};
        border: 1px solid {ACCENT};
    }}
    """


def mindflow_frame():
    frame = QFrame()
    frame.setStyleSheet(CARD_FRAME_STYLE)
    return frame


def mindflow_section(title, subtitle=''):
    frame = mindflow_frame()
    layout = QVBoxLayout(frame)
    layout.setContentsMargins(20, 18, 20, 18)
    layout.setSpacing(14)
    title_lbl = QLabel(title)
    title_lbl.setAlignment(Qt.AlignRight)
    title_lbl.setStyleSheet(f'font-size:22px;font-weight:900;color:{TEXT};')
    layout.addWidget(title_lbl)
    if subtitle:
        sub_lbl = QLabel(subtitle)
        sub_lbl.setAlignment(Qt.AlignRight)
        sub_lbl.setWordWrap(True)
        sub_lbl.setStyleSheet(f'font-size:12px;font-weight:600;color:{MUTED};')
        layout.addWidget(sub_lbl)
    return frame, layout


def make_stat_card(title, note='', accent=None):
    return SummaryCard(title, '0', note, accent)


def generate_id(prefix='id'):
    return f"{prefix}_{datetime.now().strftime('%Y%m%d%H%M%S%f')}"


def new_id(prefix='id'):
    return generate_id(prefix)


def normalize_db(data):
    data.setdefault('returns', [])
    data.setdefault('damaged', [])
    data.setdefault('inventory_ledger', [])
    data.setdefault('expenses', [])
    data.setdefault('agents_custody', [])
    rec = data.setdefault('reconciliation', {'current_actual_cash': 0.0, 'current_notes': '', 'records': []})
    rec.setdefault('current_actual_cash', 0.0)
    rec.setdefault('current_notes', '')
    rec.setdefault('records', [])
    opening_data(data)
    for s in data.get('sales', []):
        s.setdefault('id', generate_id('sale'))
        ensure_invoice_no(data, s)
        total = float(s.get('total', 0) or 0)
        paid = float(s.get('paid_amount', total if s.get('payment_type', 'نقدي') == 'نقدي' else 0) or 0)
        paid = max(0.0, min(total, paid))
        s.setdefault('payment_type', 'نقدي' if paid >= total else ('آجل' if paid == 0 else 'جزئي'))
        s['paid_amount'] = paid
        s['due_amount'] = max(0.0, total - paid)
    for i in data.get('inbound', []):
        i.setdefault('id', generate_id('inb'))
        total = float(i.get('total', 0) or 0)
        paid = float(i.get('paid_amount', total if i.get('payment_type', 'نقدي') == 'نقدي' else 0) or 0)
        paid = max(0.0, min(total, paid))
        i.setdefault('payment_type', 'نقدي' if paid >= total else ('آجل' if paid == 0 else 'جزئي'))
        i['paid_amount'] = paid
        i['due_amount'] = max(0.0, total - paid)

    id_prefix_map = {
        'cash': 'cash',
        'returns': 'ret',
        'damaged': 'dmg',
        'expenses': 'exp',
        'profit_distributions': 'pd',
        'profit_events': 'pe',
        'agents_custody': 'agc',
    }
    for list_key, prefix in id_prefix_map.items():
        for row in data.get(list_key, []):
            row.setdefault('id', generate_id(prefix))

    for row in rec.get('records', []):
        row.setdefault('id', generate_id('rec'))

    for item in data.get('items', []):
        ensure_item_inventory_fields(item)
    ensure_inventory_baseline(data)
    for f in data.get('funders', []):
        f.setdefault('is_owner_capital', False)
        f.setdefault('active', True)
    for r in data.get('returns', []):
        r.setdefault('credit_amount', float(r.get('total', 0) or 0))
        r.setdefault('credit_used', 0.0)
        r.setdefault('cash_paid_out', 0.0)
        r.setdefault('status', 'متبقي')
    return data


class BaseDialog(QDialog):
    def __init__(self, title, parent=None):
        super().__init__(parent)
        self.setWindowTitle(title)
        self.resize(440, 380)
        self.apply_theme()
    def _sync_scroll_slider(self, minimum, maximum):
        try:
            self.scroll_slider.blockSignals(True)
            self.scroll_slider.setRange(int(minimum), int(maximum if maximum > minimum else minimum + 1))
            self.scroll_slider.setPageStep(max(40, int(self.scroll.verticalScrollBar().pageStep())))
            self.scroll_slider.setValue(int(self.scroll.verticalScrollBar().value()))
        finally:
            self.scroll_slider.blockSignals(False)

    def apply_theme(self):
        apply_theme_to_widget(self)


class ItemDialog(BaseDialog):
    def __init__(self, parent=None, item=None):
        super().__init__('إضافة صنف' if item is None else 'تعديل صنف', parent)
        layout = QVBoxLayout(self)
        form = QFormLayout()
        self.code = QLineEdit()
        self.name = QLineEdit()
        self.unit = QComboBox(); self.unit.addItems(['حبة','كرتون','كغم','لتر','قطعة'])
        self.qty = QSpinBox(); self.qty.setRange(0, 10_000_000)
        self.buy_price = QDoubleSpinBox(); self.buy_price.setRange(0, 1_000_000_000); self.buy_price.setDecimals(0)
        self.sell_price = QDoubleSpinBox(); self.sell_price.setRange(0, 1_000_000_000); self.sell_price.setDecimals(0)
        self.notes = QTextEdit(); self.notes.setFixedHeight(80)
        for l, w in [('كود الصنف', self.code), ('اسم الصنف', self.name), ('الوحدة', self.unit), ('الكمية الافتتاحية', self.qty), ('سعر الشراء', self.buy_price), ('سعر البيع', self.sell_price), ('ملاحظات', self.notes)]:
            form.addRow(l + ':', w)
        layout.addLayout(form)
        btns = QHBoxLayout(); s = QPushButton('حفظ'); c = QPushButton('إلغاء'); s.setStyleSheet(BUTTON_STYLE); c.setStyleSheet(SECONDARY_BUTTON)
        s.clicked.connect(self.accept); c.clicked.connect(self.reject); btns.addWidget(s); btns.addWidget(c); layout.addLayout(btns)
        if item:
            self.code.setText(item.get('code','')); self.name.setText(item.get('name','')); self.unit.setCurrentText(item.get('unit','حبة'))
            self.qty.setValue(int(item.get('qty',0))); self.buy_price.setValue(float(item.get('buy_price',0))); self.sell_price.setValue(float(item.get('sell_price',0)))
            self.notes.setPlainText(item.get('notes',''))
    def get_data(self):
        return {
            'code': self.code.text().strip(), 'name': self.name.text().strip(), 'unit': self.unit.currentText(),
            'qty': int(self.qty.value()), 'buy_price': float(self.buy_price.value()), 'sell_price': float(self.sell_price.value()),
            'notes': self.notes.toPlainText().strip()
        }


class PersonDialog(BaseDialog):
    def __init__(self, title, parent=None, data=None):
        super().__init__(title, parent)
        layout = QVBoxLayout(self); form = QFormLayout()
        self.name = QLineEdit(); self.phone = QLineEdit(); self.address = QLineEdit(); self.notes = QTextEdit(); self.notes.setFixedHeight(80)
        for l, w in [('الاسم', self.name), ('الهاتف', self.phone), ('العنوان', self.address), ('ملاحظات', self.notes)]: form.addRow(l + ':', w)
        layout.addLayout(form)
        btns = QHBoxLayout(); s = QPushButton('حفظ'); c = QPushButton('إلغاء'); s.setStyleSheet(BUTTON_STYLE); c.setStyleSheet(SECONDARY_BUTTON)
        s.clicked.connect(self.accept); c.clicked.connect(self.reject); btns.addWidget(s); btns.addWidget(c); layout.addLayout(btns)
        if data:
            self.name.setText(data.get('name','')); self.phone.setText(data.get('phone','')); self.address.setText(data.get('address','')); self.notes.setPlainText(data.get('notes',''))
    def get_data(self):
        return {'name': self.name.text().strip(), 'phone': self.phone.text().strip(), 'address': self.address.text().strip(), 'notes': self.notes.toPlainText().strip()}


class FunderDialog(BaseDialog):
    def __init__(self, parent=None, data=None):
        super().__init__('إضافة ممول' if data is None else 'تعديل ممول', parent)
        layout = QVBoxLayout(self); form = QFormLayout()
        self.name = QLineEdit(); self.capital = QDoubleSpinBox(); self.capital.setRange(0, 1_000_000_000_000); self.capital.setDecimals(0)
        self.phone = QLineEdit(); self.notes = QTextEdit(); self.notes.setFixedHeight(80); self.active = QCheckBox('مفعل'); self.owner_capital = QCheckBox('هذا رأس مال هضاب')
        self.active.setChecked(True)
        for l, w in [('اسم الممول', self.name), ('رأس المال', self.capital), ('الهاتف', self.phone), ('ملاحظات', self.notes), ('الحالة', self.active), ('النوع', self.owner_capital)]: form.addRow(l + ':', w)
        layout.addLayout(form)
        btns = QHBoxLayout(); s = QPushButton('حفظ'); c = QPushButton('إلغاء'); s.setStyleSheet(BUTTON_STYLE); c.setStyleSheet(SECONDARY_BUTTON)
        s.clicked.connect(self.accept); c.clicked.connect(self.reject); btns.addWidget(s); btns.addWidget(c); layout.addLayout(btns)
        if data:
            self.name.setText(data.get('name','')); self.capital.setValue(float(data.get('capital',0))); self.phone.setText(data.get('phone','')); self.notes.setPlainText(data.get('notes','')); self.active.setChecked(bool(data.get('active', True))); self.owner_capital.setChecked(bool(data.get('is_owner_capital', False)))
    def get_data(self):
        return {'name': self.name.text().strip(), 'capital': float(self.capital.value()), 'phone': self.phone.text().strip(), 'notes': self.notes.toPlainText().strip(), 'active': self.active.isChecked(), 'is_owner_capital': self.owner_capital.isChecked()}


def setup_section_pages(tab_widget, page_title, button_titles, note=''):
    """حوّل التبويبات إلى صفحات مستقلة: بوابة داخلية بسيطة + صفحات مستقلة مع زر رجوع."""
    if tab_widget.count() and tab_widget.tabText(0) in ('الرئيسية', 'اختيار القسم'):
        tab_widget.removeTab(0)

    home = QWidget()
    lay = QVBoxLayout(home)
    lay.setContentsMargins(8, 8, 8, 8)
    lay.setSpacing(14)

    hero = QFrame(); hero.setStyleSheet(CARD_FRAME_STYLE)
    hero_box = QVBoxLayout(hero); hero_box.setContentsMargins(18, 18, 18, 18); hero_box.setSpacing(8)
    title = QLabel(page_title)
    title.setAlignment(Qt.AlignRight)
    title.setStyleSheet('font-size:20px;font-weight:900;background:transparent;border:none;')
    hero_box.addWidget(title)
    note_lbl = QLabel(note or 'اختر الصفحة المطلوبة داخل هذا القسم. كل صفحة مستقلة وبيها رجوع واضح فقط.')
    note_lbl.setWordWrap(True)
    note_lbl.setAlignment(Qt.AlignRight)
    note_lbl.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
    hero_box.addWidget(note_lbl)
    lay.addWidget(hero)

    card = QFrame(); card.setStyleSheet(CARD_FRAME_STYLE)
    box = QVBoxLayout(card); box.setContentsMargins(18, 18, 18, 18); box.setSpacing(10)
    grid = QGridLayout(); grid.setHorizontalSpacing(10); grid.setVerticalSpacing(10)
    for i, title_txt in enumerate(button_titles):
        btn = QPushButton(title_txt)
        btn.setMinimumHeight(58)
        btn.setStyleSheet(BUTTON_STYLE)
        btn.clicked.connect(lambda _=False, idx=i+1, tabs=tab_widget: tabs.setCurrentIndex(idx))
        grid.addWidget(btn, i // 2, i % 2)
    box.addLayout(grid)
    lay.addWidget(card)
    lay.addStretch(1)

    tab_widget.insertTab(0, home, 'اختيار القسم')
    for idx in range(1, tab_widget.count()):
        page = tab_widget.widget(idx)
        page_layout = page.layout()
        if page_layout is not None:
            first_item = page_layout.itemAt(0)
            has_back = False
            if first_item is not None and first_item.layout() is not None:
                for j in range(first_item.layout().count()):
                    w = first_item.layout().itemAt(j).widget()
                    if isinstance(w, QPushButton) and 'الرجوع' in w.text():
                        has_back = True
                        break
            if not has_back:
                top = QHBoxLayout(); top.setContentsMargins(0, 0, 0, 0)
                top.addStretch(1)
                back_btn = QPushButton('↩ الرجوع')
                back_btn.setStyleSheet(SECONDARY_BUTTON)
                back_btn.setMinimumHeight(44)
                back_btn.clicked.connect(lambda _=False, tabs=tab_widget: tabs.setCurrentIndex(0))
                top.addWidget(back_btn)
                page_layout.insertLayout(0, top)
    tab_widget.tabBar().hide()
    tab_widget.setCurrentIndex(0)


class BaseWindow(QWidget):
    def __init__(self, main, title):
        super().__init__(main)
        self.main = main
        self.setWindowTitle(title)
        self.setWindowFlag(Qt.Window, True)
        self.resize(1440, 900)
        self.apply_theme()

        self.outer_layout = QVBoxLayout(self)
        self.outer_layout.setContentsMargins(22, 18, 22, 18)
        self.outer_layout.setSpacing(12)

        self.header_card = QFrame(); self.header_card.setObjectName('innerTopBar')
        top = QHBoxLayout(self.header_card); top.setContentsMargins(18, 14, 18, 14); top.setSpacing(12)
        back = QPushButton('العودة'); back.setStyleSheet(SECONDARY_BUTTON); back.setMinimumHeight(46); back.clicked.connect(self.close)
        title_box = QVBoxLayout(); title_box.setSpacing(2)
        self.title_lbl = QLabel(title); self.title_lbl.setAlignment(Qt.AlignRight); self.title_lbl.setStyleSheet(f'font-size:28px;font-weight:900;color:{TEXT};')
        self.subtitle_lbl = QLabel('واجهة داخلية واضحة، ناعمة، ومنسقة بأسلوب حديث.')
        self.subtitle_lbl.setAlignment(Qt.AlignRight); self.subtitle_lbl.setWordWrap(True); self.subtitle_lbl.setStyleSheet(f'font-size:12px;font-weight:600;color:{MUTED};')
        title_box.addWidget(self.title_lbl); title_box.addWidget(self.subtitle_lbl)
        top.addWidget(back, alignment=Qt.AlignLeft); top.addLayout(title_box, 1)
        self.outer_layout.addWidget(self.header_card)

        self.scroll = QScrollArea()
        self.scroll.setWidgetResizable(True)
        self.scroll.setFrameShape(QFrame.NoFrame)
        self.scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarAlwaysOff)
        self.scroll.setVerticalScrollBarPolicy(Qt.ScrollBarAlwaysOn)
        self.scroll.setStyleSheet('QScrollArea{background:transparent;border:none;} QScrollBar:vertical{background:rgba(255,255,255,0.04);width:14px;margin:6px 2px;border-radius:7px;} QScrollBar::handle:vertical{background:rgba(255,255,255,0.28);border-radius:7px;min-height:46px;} QScrollBar::add-line:vertical,QScrollBar::sub-line:vertical{height:0;}')
        self.body = QWidget()
        self.body.setObjectName('baseWindowBody')
        self.scroll.setWidget(self.body)

        self.scroll_row = QHBoxLayout()
        self.scroll_row.setContentsMargins(0, 0, 0, 0)
        self.scroll_row.setSpacing(10)
        self.scroll_row.addWidget(self.scroll, 1)

        self.scroll_rail = QFrame()
        self.scroll_rail.setObjectName('scrollRail')
        self.scroll_rail.setStyleSheet(f"QFrame#scrollRail{{background:{rgba_from_hex(CARD,0.72)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:20px;}}")
        rail_box = QVBoxLayout(self.scroll_rail)
        rail_box.setContentsMargins(8, 10, 8, 10)
        rail_box.setSpacing(8)
        self.scroll_up_btn = QToolButton()
        self.scroll_up_btn.setText('▲')
        self.scroll_up_btn.setStyleSheet(SECONDARY_BUTTON)
        self.scroll_up_btn.setMinimumSize(44, 44)
        self.scroll_down_btn = QToolButton()
        self.scroll_down_btn.setText('▼')
        self.scroll_down_btn.setStyleSheet(SECONDARY_BUTTON)
        self.scroll_down_btn.setMinimumSize(44, 44)
        self.scroll_slider = QSlider(Qt.Vertical)
        self.scroll_slider.setInvertedAppearance(False)
        self.scroll_slider.setStyleSheet(f"QSlider::groove:vertical{{background:{rgba_from_hex(TEXT,0.07)};width:12px;border-radius:6px;}} QSlider::handle:vertical{{background:{ACCENT};height:44px;margin:-4px;border-radius:8px;}}")
        self.scroll_slider.setMinimumHeight(180)
        rail_box.addWidget(self.scroll_up_btn, 0, Qt.AlignHCenter)
        rail_box.addWidget(self.scroll_slider, 1)
        rail_box.addWidget(self.scroll_down_btn, 0, Qt.AlignHCenter)
        self.scroll_row.addWidget(self.scroll_rail, 0, Qt.AlignTop)
        self.outer_layout.addLayout(self.scroll_row, 1)

        self.layout = QVBoxLayout(self.body)
        self.layout.setContentsMargins(6, 4, 6, 8)
        self.layout.setSpacing(14)

        bar = self.scroll.verticalScrollBar()
        self.scroll_slider.setRange(bar.minimum(), max(bar.maximum(), 0))
        self.scroll_slider.valueChanged.connect(bar.setValue)
        bar.valueChanged.connect(self.scroll_slider.setValue)
        bar.rangeChanged.connect(self._sync_scroll_slider)
        self.scroll_up_btn.clicked.connect(lambda: bar.setValue(max(bar.minimum(), bar.value() - max(220, bar.pageStep() // 2))))
        self.scroll_down_btn.clicked.connect(lambda: bar.setValue(min(bar.maximum(), bar.value() + max(220, bar.pageStep() // 2))))
        self._sync_scroll_slider(bar.minimum(), bar.maximum())
    def _sync_scroll_slider(self, minimum, maximum):
        try:
            self.scroll_slider.blockSignals(True)
            self.scroll_slider.setRange(int(minimum), int(maximum if maximum > minimum else minimum + 1))
            self.scroll_slider.setPageStep(max(40, int(self.scroll.verticalScrollBar().pageStep())))
            self.scroll_slider.setValue(int(self.scroll.verticalScrollBar().value()))
        finally:
            self.scroll_slider.blockSignals(False)

    def apply_theme(self):
        apply_theme_to_widget(self)
        try: self.header_card.setStyleSheet(f'QFrame#innerTopBar{{background-color:{rgba_from_hex(CARD,0.74)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:26px;}}')
        except Exception: pass
    def closeEvent(self, event):
        try:
            self.hide()
            self.main.show()
            self.main.raise_()
            self.main.activateWindow()
        except Exception:
            pass
        event.accept()
    @property
    def db(self):
        return self.main.db
    def save(self):
        self.main.save_all()


class ItemsWindow(BaseWindow):
    def __init__(self, main):
        super().__init__(main, '📦 إدارة الأصناف')
        self.subtitle_lbl.setText('ترتيب واضح للأصناف والمخزون بنفس روح اللوحة الرئيسية، مع بحث سريع ومؤشرات مختصرة.')

        hero = QFrame(); hero.setStyleSheet(CARD_FRAME_STYLE)
        hero_box = QVBoxLayout(hero); hero_box.setContentsMargins(20, 18, 20, 18); hero_box.setSpacing(10)
        hero_title = QLabel('لوحة الأصناف')
        hero_title.setAlignment(Qt.AlignRight)
        hero_title.setStyleSheet('font-size:20px;font-weight:900;background:transparent;border:none;')
        hero_note = QLabel('إدارة الأصناف، متابعة المخزون، ومعرفة القيمة الإجمالية بشكل مرتب وسريع.')
        hero_note.setAlignment(Qt.AlignRight); hero_note.setWordWrap(True)
        hero_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        hero_box.addWidget(hero_title); hero_box.addWidget(hero_note)
        self.layout.addWidget(hero)

        self.tabs = style_funders_tabs(QTabWidget())
        self.tabs.setTabPosition(QTabWidget.North)
        self.layout.addWidget(self.tabs, 1)

        manage_tab = QWidget(); manage_layout = QVBoxLayout(manage_tab); manage_layout.setContentsMargins(8,8,8,8); manage_layout.setSpacing(12)
        summary_tab = QWidget(); summary_layout = QVBoxLayout(summary_tab); summary_layout.setContentsMargins(8,8,8,8); summary_layout.setSpacing(12)
        table_tab = QWidget(); table_layout = QVBoxLayout(table_tab); table_layout.setContentsMargins(8,8,8,8); table_layout.setSpacing(12)

        actions_card = QFrame(); actions_card.setStyleSheet(CARD_FRAME_STYLE)
        actions_box = QVBoxLayout(actions_card); actions_box.setContentsMargins(18,18,18,18); actions_box.setSpacing(10)
        actions_title = QLabel('إدارة الأصناف')
        actions_title.setAlignment(Qt.AlignRight)
        actions_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        actions_box.addWidget(actions_title)
        self.search = QLineEdit(); self.search.setPlaceholderText('بحث بالاسم أو الكود'); self.search.textChanged.connect(self.refresh_table)
        self.search.setMinimumHeight(44)
        actions_box.addWidget(self.search)
        btn_grid = QGridLayout(); btn_grid.setHorizontalSpacing(10); btn_grid.setVerticalSpacing(10)
        specs = [('➕ إضافة', self.add_item, BUTTON_STYLE), ('✏️ تعديل', self.edit_item, SECONDARY_BUTTON), ('🗑 حذف', self.delete_item, SECONDARY_BUTTON)]
        for i,(txt, fn, style) in enumerate(specs):
            b = QPushButton(txt); b.setStyleSheet(style); b.setMinimumHeight(46); b.clicked.connect(fn); btn_grid.addWidget(b, i // 2, i % 2)
        actions_box.addLayout(btn_grid)
        manage_note = QLabel('كل جزء بالأصناف صار ضمن تبويب مستقل مثل الممولين حتى يبقى الشغل أهدأ وأوضح.')
        manage_note.setAlignment(Qt.AlignRight); manage_note.setWordWrap(True)
        manage_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        actions_box.addWidget(manage_note)
        manage_layout.addWidget(actions_card)
        manage_layout.addStretch(1)

        summary_card = QFrame(); summary_card.setStyleSheet(CARD_FRAME_STYLE)
        summary_box = QVBoxLayout(summary_card); summary_box.setContentsMargins(18,18,18,18); summary_box.setSpacing(10)
        summary_title = QLabel('ملخص الأصناف')
        summary_title.setAlignment(Qt.AlignRight)
        summary_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        summary_box.addWidget(summary_title)
        self.count_card = SummaryCard('عدد الأصناف', '0', 'عدد المواد المسجلة')
        self.qty_card = SummaryCard('إجمالي المخزون', '0', 'مجموع الكميات الحالية')
        self.value_card = SummaryCard('قيمة المخزون', '0', 'إجمالي القيمة التقديرية')
        for c in (self.count_card, self.qty_card, self.value_card):
            c.setMinimumHeight(96); c.setMaximumHeight(112); summary_box.addWidget(c)
        self.summary = QLabel(); self.summary.setAlignment(Qt.AlignRight); self.summary.setWordWrap(True)
        self.summary.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        summary_box.addWidget(self.summary)
        summary_layout.addWidget(summary_card)
        summary_layout.addStretch(1)

        table_card = QFrame(); table_card.setStyleSheet(CARD_FRAME_STYLE)
        table_box = QVBoxLayout(table_card); table_box.setContentsMargins(16,16,16,16); table_box.setSpacing(10)
        table_title = QLabel('جدول الأصناف')
        table_title.setAlignment(Qt.AlignRight)
        table_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        table_box.addWidget(table_title)
        self.table = QTableWidget(); self.table.setStyleSheet(TABLE_STYLE); self.table.setColumnCount(8)
        self.table.setHorizontalHeaderLabels(['#','الكود','اسم الصنف','الوحدة','المخزون','شراء','بيع','قيمة المخزون'])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch); self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table_box.addWidget(self.table, 1)
        table_layout.addWidget(table_card, 1)

        self.tabs.addTab(manage_tab, 'إدارة الأصناف')
        self.tabs.addTab(summary_tab, 'الملخص')
        self.tabs.addTab(table_tab, 'السجل')
        self.refresh_table()
    def filtered(self):
        q = self.search.text().strip().lower()
        if not q: return self.db['items']
        return [x for x in self.db['items'] if q in x.get('name','').lower() or q in x.get('code','').lower()]
    def load_selected_row(self, switch_tab=False):
        r = self.table.currentRow()
        rows = self.db.get('expenses', [])
        if r < 0 or r >= len(rows):
            return
        row = rows[r]
        try:
            self.date.setDate(QDate.fromString(str(row.get('date', '')), 'yyyy-MM-dd') or QDate.currentDate())
        except Exception:
            self.date.setDate(QDate.currentDate())
        self.category.setCurrentText(str(row.get('category', '') or ''))
        self.amount.setValue(float(row.get('amount', 0) or 0))
        self.notes.setPlainText(str(row.get('notes', '') or ''))
        if switch_tab:
            self.tabs.setCurrentIndex(0)

    def refresh_table(self):
        data = self.filtered(); self.table.setRowCount(len(data)); total_qty = total_val = 0
        for r, it in enumerate(data):
            ensure_item_inventory_fields(it)
            qty = int(it.get('qty',0)); val = item_inventory_value(it); total_qty += qty; total_val += val
            vals = [r+1, it.get('code',''), it.get('name',''), it.get('unit',''), qty, fmt_money(item_avg_cost(it)), fmt_money(it.get('sell_price',0)), fmt_money(val)]
            for c, v in enumerate(vals): self.table.setItem(r, c, QTableWidgetItem(str(v)))
        self.count_card.set_value(str(len(data)))
        self.qty_card.set_value(f"{total_qty:,}")
        self.value_card.set_value(fmt_money(total_val))
        self.summary.setText(f'عدد الأصناف: {len(data)} | إجمالي المخزون: {total_qty} | قيمة المخزون: {fmt_money(total_val)} د.ع')
    def real_index(self):
        row = self.table.currentRow();
        if row < 0 or row >= len(self.filtered()): return None
        target = self.filtered()[row].get('name')
        for i, x in enumerate(self.db['items']):
            if x.get('name') == target: return i
    def add_item(self):
        while True:
            d = ItemDialog(self)
            result = d.exec()
            self.raise_(); self.activateWindow()
            if not result:
                break
            item = d.get_data()
            if not item['name']:
                QMessageBox.warning(self,'تنبيه','اسم الصنف مطلوب')
                continue
            self.db['items'].append(item)
            self.save()
            self.refresh_table()
    def edit_item(self):
        i = self.real_index();
        if i is None: return QMessageBox.warning(self,'تنبيه','اختر صنف')
        d = ItemDialog(self, self.db['items'][i])
        if d.exec():
            old = self.db['items'][i]['name']; new = d.get_data(); self.db['items'][i].update(new)
            for row in self.db['sales']:
                if row.get('item') == old: row['item'] = new['name']
            for row in self.db['inbound']:
                if row.get('item') == old: row['item'] = new['name']
            self.save(); self.refresh_table()
    def delete_item(self):
        i = self.real_index();
        if i is None: return QMessageBox.warning(self,'تنبيه','اختر صنف')
        name = self.db['items'][i]['name']
        if any(x.get('item') == name for x in self.db['sales'] + self.db['inbound']):
            return QMessageBox.warning(self,'منع','الصنف مربوط بحركات')
        if QMessageBox.question(self,'تأكيد',f'حذف {name}؟') == QMessageBox.Yes:
            self.db['items'].pop(i); self.save(); self.refresh_table()


class PeopleWindow(BaseWindow):
    def __init__(self, main, key, title, icon):
        super().__init__(main, f'{icon} {title}')
        self.key = key; self.entity = title
        self.subtitle_lbl.setText(f'واجهة {title} صارت أرتب وأكثر تقسيمًا: إدارة، ملخص، وسجل منفصل حتى تبقى الصفحة هادئة وواضحة.')

        hero = QFrame(); hero.setStyleSheet(CARD_FRAME_STYLE)
        hero_box = QVBoxLayout(hero); hero_box.setContentsMargins(18,16,18,16); hero_box.setSpacing(6)
        hero_title = QLabel(f'لوحة {title}')
        hero_title.setAlignment(Qt.AlignRight)
        hero_title.setStyleSheet('font-size:20px;font-weight:900;background:transparent;border:none;')
        hero_note = QLabel('إدارة الأسماء، متابعة الأرصدة، والوصول السريع إلى كشف الحركة والتسديد ضمن تبويبات أوضح.')
        hero_note.setAlignment(Qt.AlignRight); hero_note.setWordWrap(True)
        hero_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        hero_box.addWidget(hero_title); hero_box.addWidget(hero_note)
        self.layout.addWidget(hero)

        debt_label = 'الدين علينا' if self.key == 'suppliers' else 'الدين لنا'
        self.tabs = style_funders_tabs(QTabWidget())
        self.layout.addWidget(self.tabs, 1)

        manage_tab = QWidget(); manage_layout = QVBoxLayout(manage_tab); manage_layout.setContentsMargins(8,8,8,8); manage_layout.setSpacing(12)
        actions_card = QFrame(); actions_card.setStyleSheet(CARD_FRAME_STYLE)
        actions_box = QVBoxLayout(actions_card); actions_box.setContentsMargins(16,16,16,16); actions_box.setSpacing(10)
        actions_title = QLabel('بحث وإجراءات')
        actions_title.setAlignment(Qt.AlignRight)
        actions_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        actions_box.addWidget(actions_title)
        self.search = QLineEdit(); self.search.setPlaceholderText(f'بحث عن {title}'); self.search.textChanged.connect(self.refresh_table); self.search.setMinimumHeight(44)
        actions_box.addWidget(self.search)
        btn_grid = QGridLayout(); btn_grid.setHorizontalSpacing(10); btn_grid.setVerticalSpacing(10)
        specs = [('➕ إضافة', self.add_person, BUTTON_STYLE), ('✏️ تعديل', self.edit_person, SECONDARY_BUTTON), ('🗑 حذف', self.delete_person, SECONDARY_BUTTON), ('💳 كشف/تسديد', self.open_dues, BUTTON_STYLE)]
        for i,(txt, fn, style) in enumerate(specs):
            b = QPushButton(txt); b.setStyleSheet(style); b.setMinimumHeight(46); b.clicked.connect(fn); btn_grid.addWidget(b, i // 2, i % 2)
        actions_box.addLayout(btn_grid)
        if self.key == 'customers':
            statement_card = QFrame(); statement_card.setStyleSheet(CARD_FRAME_STYLE)
            statement_box = QVBoxLayout(statement_card); statement_box.setContentsMargins(12,12,12,12); statement_box.setSpacing(8)
            statement_title = QLabel('كشف الزبون المباشر')
            statement_title.setAlignment(Qt.AlignRight)
            statement_title.setStyleSheet('font-size:14px;font-weight:900;background:transparent;border:none;')
            statement_box.addWidget(statement_title)
            self.statement_customer_cb = QComboBox(); self.statement_customer_cb.setStyleSheet(INPUT_STYLE); self.statement_customer_cb.setMinimumHeight(44)
            statement_box.addWidget(self.statement_customer_cb)
            self.statement_open_btn = QPushButton('🧾 فتح كشف الزبون'); self.statement_open_btn.setStyleSheet(BUTTON_STYLE); self.statement_open_btn.setMinimumHeight(46); self.statement_open_btn.clicked.connect(self.open_customer_statement_from_combo)
            statement_box.addWidget(self.statement_open_btn)
            statement_hint = QLabel('اختَر الزبون من القائمة المنسدلة وافتح كشفه مباشرة من هنا.')
            statement_hint.setAlignment(Qt.AlignRight); statement_hint.setWordWrap(True)
            statement_hint.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
            statement_box.addWidget(statement_hint)
            actions_box.addWidget(statement_card)
        manage_note = QLabel('الترتيب هنا مخصص للإضافة والتعديل والوصول السريع إلى كشف الحركة بدون زحمة الجدول.')
        manage_note.setAlignment(Qt.AlignRight); manage_note.setWordWrap(True)
        manage_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        actions_box.addWidget(manage_note)
        manage_layout.addWidget(actions_card)
        manage_layout.addStretch(1)

        summary_tab = QWidget(); summary_layout = QVBoxLayout(summary_tab); summary_layout.setContentsMargins(8,8,8,8); summary_layout.setSpacing(12)
        summary_card = QFrame(); summary_card.setStyleSheet(CARD_FRAME_STYLE)
        summary_box = QVBoxLayout(summary_card); summary_box.setContentsMargins(16,16,16,16); summary_box.setSpacing(10)
        summary_title = QLabel('ملخص سريع')
        summary_title.setAlignment(Qt.AlignRight)
        summary_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        summary_box.addWidget(summary_title)
        cards_grid = QGridLayout(); cards_grid.setHorizontalSpacing(10); cards_grid.setVerticalSpacing(10)
        self.count_card = SummaryCard('عدد السجلات', '0', 'عدد الأسماء الظاهرة')
        self.due_card = SummaryCard(debt_label, '0', 'إجمالي الرصيد الحالي')
        for i, c in enumerate((self.count_card, self.due_card)):
            c.setMinimumHeight(104); c.setMaximumHeight(120); cards_grid.addWidget(c, 0, i)
        summary_box.addLayout(cards_grid)
        self.summary = QLabel(); self.summary.setAlignment(Qt.AlignRight); self.summary.setWordWrap(True)
        self.summary.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        summary_box.addWidget(self.summary)
        summary_layout.addWidget(summary_card)
        summary_layout.addStretch(1)

        table_tab = QWidget(); table_layout = QVBoxLayout(table_tab); table_layout.setContentsMargins(8,8,8,8); table_layout.setSpacing(12)
        table_card = QFrame(); table_card.setStyleSheet(CARD_FRAME_STYLE)
        table_box = QVBoxLayout(table_card); table_box.setContentsMargins(16,16,16,16); table_box.setSpacing(10)
        table_title = QLabel(f'جدول {title}')
        table_title.setAlignment(Qt.AlignRight)
        table_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        table_box.addWidget(table_title)
        self.table = QTableWidget(); self.table.setStyleSheet(TABLE_STYLE); self.table.setColumnCount(6)
        self.table.setHorizontalHeaderLabels(['#','الاسم','الهاتف','العنوان',debt_label,'ملاحظات']); self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table_box.addWidget(self.table, 1)
        table_layout.addWidget(table_card, 1)

        self.tabs.addTab(manage_tab, 'الإدارة')
        self.tabs.addTab(summary_tab, 'الملخص')
        self.tabs.addTab(table_tab, 'السجل')
        self.refresh_table()
    def filtered(self):
        q = self.search.text().strip().lower()
        data = self.db[self.key]
        if not q: return data
        return [x for x in data if q in x.get('name','').lower() or q in x.get('phone','').lower()]
    def person_due(self, name):
        rows = supplier_due_summary(self.db) if self.key == 'suppliers' else customer_due_summary(self.db)
        for row in rows:
            if row['name'] == name:
                return row['due']
        return 0.0
    def refresh_table(self):
        data = self.filtered(); self.table.setRowCount(len(data)); total_due = 0
        for r, p in enumerate(data):
            due = self.person_due(p.get('name', '')); total_due += due
            vals = [r+1, p.get('name',''), p.get('phone',''), p.get('address',''), fmt_money(due), p.get('notes','')]
            for c, v in enumerate(vals): self.table.setItem(r,c,QTableWidgetItem(str(v)))
        if hasattr(self, 'count_card'): self.count_card.set_value(str(len(data)))
        if hasattr(self, 'due_card'): self.due_card.set_value(fmt_money(total_due))
        self.summary.setText(f'عدد السجلات: {len(data)} | إجمالي الديون: {fmt_money(total_due)} د.ع')
        if self.key == 'customers' and hasattr(self, 'statement_customer_cb'):
            current_name = self.statement_customer_cb.currentText().strip()
            names = [str((row or {}).get('name', '')).strip() for row in self.db.get('customers', []) if str((row or {}).get('name', '')).strip()]
            self.statement_customer_cb.blockSignals(True)
            self.statement_customer_cb.clear()
            self.statement_customer_cb.addItems(names)
            if current_name and current_name in names:
                self.statement_customer_cb.setCurrentText(current_name)
            self.statement_customer_cb.blockSignals(False)
    def real_index(self):
        row = self.table.currentRow();
        if row < 0 or row >= len(self.filtered()): return None
        target = self.filtered()[row].get('name')
        for i, x in enumerate(self.db[self.key]):
            if x.get('name') == target: return i
    def add_person(self):
        while True:
            d = PersonDialog(f'إضافة {self.entity}', self)
            result = d.exec()
            self.raise_(); self.activateWindow()
            if not result:
                break
            p = d.get_data()
            name = p.get('name','').strip()
            if not name:
                QMessageBox.warning(self,'تنبيه','الاسم مطلوب')
                continue
            if any(x.get('name','').strip() == name for x in self.db[self.key]):
                QMessageBox.warning(self,'منع','الاسم موجود مسبقًا')
                continue
            p['name'] = name
            self.db[self.key].append(p)
            self.save()
            self.refresh_table()
    def edit_person(self):
        i = self.real_index();
        if i is None: return QMessageBox.warning(self,'تنبيه','اختر سجل')
        d = PersonDialog(f'تعديل {self.entity}', self, self.db[self.key][i])
        if d.exec():
            old = self.db[self.key][i]['name']; new = d.get_data()
            new_name = new.get('name','').strip()
            if not new_name: return QMessageBox.warning(self,'تنبيه','الاسم مطلوب')
            if any(idx != i and x.get('name','').strip() == new_name for idx, x in enumerate(self.db[self.key])):
                return QMessageBox.warning(self,'منع','الاسم موجود مسبقًا')
            new['name'] = new_name
            self.db[self.key][i].update(new)
            link_key = 'customer' if self.key == 'customers' else 'supplier'
            rows = self.db['sales'] if self.key == 'customers' else self.db['inbound']
            for row in rows:
                if row.get(link_key) == old: row[link_key] = new['name']
            if self.key == 'customers':
                for row in self.db.get('returns', []):
                    if row.get('customer') == old:
                        row['customer'] = new['name']
            for row in self.db['cash']:
                if row.get('party') == old and row.get('source') in ['customer_payment', 'supplier_payment', 'opening_customer_payment', 'opening_supplier_payment', 'sales_group', 'inbound']:
                    row['party'] = new['name']
            opening_key = 'customers' if self.key == 'customers' else 'suppliers'
            for row in opening_data(self.db).get(opening_key, []):
                if row.get('name', '') == old:
                    row['name'] = new['name']
            self.save(); self.refresh_table()

    def delete_person(self):
        i = self.real_index()
        if i is None:
            return QMessageBox.warning(self, 'تنبيه', 'اختر سجل')
        name = self.db[self.key][i]['name']
        if self.person_due(name) > 0:
            return QMessageBox.warning(self, 'منع', 'لا يمكن الحذف لأن عليه/له دين قائم')

        if self.key == 'customers':
            if any(x.get('customer') == name for x in self.db.get('sales', [])):
                return QMessageBox.warning(self, 'منع', 'مربوط بمبيعات')
            if any(x.get('customer') == name for x in self.db.get('returns', [])):
                return QMessageBox.warning(self, 'منع', 'مربوط بمرتجعات محفوظة')
            if any(x.get('name', '') == name for x in opening_data(self.db).get('customers', [])):
                return QMessageBox.warning(self, 'منع', 'لا يمكن حذف زبون له رصيد افتتاحي أو تاريخ افتتاحي محفوظ')
            if any(x.get('party') == name and x.get('source') in ('customer_payment', 'opening_customer_payment') for x in self.db.get('cash', [])):
                return QMessageBox.warning(self, 'منع', 'لا يمكن حذف زبون له وصولات قبض أو تاريخ تسديد محفوظ')
        else:
            if any(x.get('supplier') == name for x in self.db.get('inbound', [])):
                return QMessageBox.warning(self, 'منع', 'مربوط بوارد')
            if any(x.get('name', '') == name for x in opening_data(self.db).get('suppliers', [])):
                return QMessageBox.warning(self, 'منع', 'لا يمكن حذف مورد له رصيد افتتاحي أو تاريخ افتتاحي محفوظ')
            if any(x.get('party') == name and x.get('source') in ('supplier_payment', 'opening_supplier_payment') for x in self.db.get('cash', [])):
                return QMessageBox.warning(self, 'منع', 'لا يمكن حذف مورد له وصل تسديد أو تاريخ مدفوعات محفوظ')

        if QMessageBox.question(self, 'تأكيد', f'حذف {name}؟') == QMessageBox.Yes:
            self.db[self.key].pop(i)
            self.save()
            self.refresh_table()
    def open_customer_statement_direct(self):
        if self.key != 'customers':
            return
        i = self.real_index()
        if i is None:
            return QMessageBox.warning(self, 'تنبيه', 'اختر الزبون أولاً')
        try:
            name = (self.db.get(self.key, [])[i] or {}).get('name', '').strip()
        except Exception:
            name = ''
        if not name:
            return QMessageBox.warning(self, 'تنبيه', 'تعذر تحديد اسم الزبون المختار')
        try:
            dlg = CustomerStatementDialog(self, self.db, name)
            dlg.exec()
        except Exception as exc:
            QMessageBox.warning(self, 'تعذر فتح الكشف', f'تعذر إنشاء أو فتح كشف الزبون: {exc}')

    def open_customer_statement_from_combo(self):
        if self.key != 'customers':
            return
        name = ''
        if hasattr(self, 'statement_customer_cb'):
            name = self.statement_customer_cb.currentText().strip()
        if not name:
            return QMessageBox.warning(self, 'تنبيه', 'اختر الزبون من القائمة أولاً')
        try:
            dlg = CustomerStatementDialog(self, self.db, name)
            dlg.exec()
        except Exception as exc:
            QMessageBox.warning(self, 'تعذر فتح الكشف', f'تعذر إنشاء أو فتح كشف الزبون: {exc}')

    def open_dues(self):
        key = 'customer_dues' if self.key == 'customers' else 'supplier_dues'
        self.main.show_win(key, DuesWindow, self.key)



class DuesWindow(BaseWindow):
    def __init__(self, main, people_key):
        self.people_key = people_key
        title = '💳 ديون الزبائن' if people_key == 'customers' else '💳 ديون الموردين'
        super().__init__(main, title)
        target_label = 'الزبائن' if people_key == 'customers' else 'الموردين'
        self.subtitle_lbl.setText(f'واجهة {target_label} صارت بنفس روح المبيعات: صفحة رئيسية واضحة، أقسام مستقلة، وحقول أكبر بدون تغيير أي معادلة.')

        self.person_cb = QComboBox()
        self.amount = tune_numeric_widget(QDoubleSpinBox()); self.amount.setRange(0,1_000_000_000); self.amount.setDecimals(0)
        self.date = QDateEdit(); fix_date_edit_widget(self.date); self.date.setDate(QDate.currentDate())
        self.notes = QLineEdit(); self.notes.setPlaceholderText('ملاحظات التسديد / القبض')
        self.method = QComboBox(); self.method.addItems(['نقد', 'تحويل']); self.method.currentTextChanged.connect(self.toggle_account)
        self.account_no = QLineEdit(); self.account_no.setPlaceholderText('رقم الحساب المحول له')
        self.refresh_combos(); self.person_cb.currentTextChanged.connect(self.refresh_details)
        self._style_dues_inputs()

        self.summary_boxes = QLabel(); self.summary_boxes.setAlignment(Qt.AlignRight); self.summary_boxes.setWordWrap(True)
        self.summary_boxes.setStyleSheet(f'font-size:13px;font-weight:700;color:{MUTED};padding:4px 2px;')
        self.details_summary_box = QLabel(); self.details_summary_box.setAlignment(Qt.AlignRight); self.details_summary_box.setWordWrap(True)
        self.details_summary_box.setStyleSheet(f'font-size:13px;font-weight:700;color:{MUTED};padding:4px 2px;')

        self.dues_stack = QStackedWidget()
        self.layout.addWidget(self.dues_stack, 1)

        self.dues_home = QWidget(); self.dues_home_layout = QVBoxLayout(self.dues_home); self.dues_home_layout.setContentsMargins(18,18,18,18); self.dues_home_layout.setSpacing(16)
        self.entry_tab = QWidget(); self.entry_tab_layout = QVBoxLayout(self.entry_tab); self.entry_tab_layout.setContentsMargins(0,0,0,0); self.entry_tab_layout.setSpacing(0)
        self.summary_tab = QWidget(); self.summary_tab_layout = QVBoxLayout(self.summary_tab); self.summary_tab_layout.setContentsMargins(0,0,0,0); self.summary_tab_layout.setSpacing(0)
        self.details_tab = QWidget(); self.details_tab_layout = QVBoxLayout(self.details_tab); self.details_tab_layout.setContentsMargins(0,0,0,0); self.details_tab_layout.setSpacing(0)
        self.payments_tab = QWidget(); self.payments_tab_layout = QVBoxLayout(self.payments_tab); self.payments_tab_layout.setContentsMargins(0,0,0,0); self.payments_tab_layout.setSpacing(0)
        self.dues_stack.addWidget(self.dues_home)
        self.dues_stack.addWidget(self.entry_tab)
        self.dues_stack.addWidget(self.summary_tab)
        self.dues_stack.addWidget(self.details_tab)
        self.dues_stack.addWidget(self.payments_tab)

        self._build_dues_home(target_label)
        self._build_dues_entry_page(target_label)

        summary_scroll, summary_content = self._make_dues_page('ملخص الحسابات', f'عرض كامل ومريح لكل {target_label} مع الأرصدة الحالية.')
        summary_content_layout = summary_content.layout()
        summary_card = QFrame()
        summary_card.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.82)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:24px;")
        summary_card_layout = QVBoxLayout(summary_card)
        summary_card_layout.setContentsMargins(14,14,14,14)
        summary_card_layout.setSpacing(10)
        self.table = QTableWidget()
        self.table.setStyleSheet(TABLE_STYLE)
        if self.people_key == 'customers':
            self.table.setColumnCount(7)
            headers = ['#','الاسم','الدين القديم','الدين الجديد','إجمالي التسديدات','رصيد المرتجعات','المتبقي النهائي']
        else:
            self.table.setColumnCount(5)
            headers = ['#','الاسم','إجمالي الحركة','الواصلات/الدفعات','المتبقي']
        self.table.setHorizontalHeaderLabels(headers)
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.itemSelectionChanged.connect(self.sync_selected_person_from_table)
        self.table.setMinimumHeight(420)
        summary_card_layout.addWidget(self.table)
        summary_content_layout.addWidget(summary_card)
        self.summary_tab_layout.addWidget(summary_scroll)

        details_scroll, details_content = self._make_dues_page('تفاصيل الحركة', 'تفصيل القديم والجديد والتسديدات والمرتجعات حسب الاسم المحدد من الحقل العلوي أو من الملخص.')
        details_content_layout = details_content.layout()
        details_info = QLabel('اختَر الاسم من الأعلى أو من جدول الملخص، وبعدها تنزل التفاصيل هنا بشكل أوضح.')
        details_info.setWordWrap(True)
        details_info.setStyleSheet(f'font-size:13px;font-weight:700;color:{MUTED};padding:0 4px;')
        details_content_layout.addWidget(details_info)
        details_toolbar = QHBoxLayout(); details_toolbar.setSpacing(10)
        details_toolbar.addWidget(self.details_summary_box, 1)
        if self.people_key == 'customers':
            self.statement_btn = QPushButton('🖨 طباعة كشف الزبون')
            self.statement_btn.setStyleSheet(BUTTON_STYLE)
            self.statement_btn.setMinimumHeight(52)
            self.statement_btn.clicked.connect(self.open_customer_statement)
            details_toolbar.addWidget(self.statement_btn, 0)
        details_content_layout.addLayout(details_toolbar)
        details_card = QFrame()
        details_card.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.82)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:24px;")
        details_card_layout = QVBoxLayout(details_card)
        details_card_layout.setContentsMargins(14,14,14,14)
        details_card_layout.setSpacing(10)
        self.details = QTableWidget(); self.details.setStyleSheet(TABLE_STYLE); self.details.setColumnCount(8)
        self.details.setHorizontalHeaderLabels(['#','القسم','المرجع','التاريخ','الصنف/البيان','الكمية','القيمة','المتبقي على المرجع'])
        self.details.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.details.setSelectionBehavior(QAbstractItemView.SelectRows); self.details.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.details.setMinimumHeight(420)
        details_card_layout.addWidget(self.details)
        details_content_layout.addWidget(details_card)
        self.details_tab_layout.addWidget(details_scroll)

        payments_scroll, payments_content = self._make_dues_page('سجل التسديدات', 'كل الحركات المرحّلة لهذا القسم تظهر هنا بسجل مريح وواضح.')
        payments_content_layout = payments_content.layout()
        payments_card = QFrame()
        payments_card.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.82)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:24px;")
        payments_card_layout = QVBoxLayout(payments_card)
        payments_card_layout.setContentsMargins(14,14,14,14)
        payments_card_layout.setSpacing(10)
        self.payments = QTableWidget(); self.payments.setStyleSheet(TABLE_STYLE); self.payments.setColumnCount(8)
        self.payments.setHorizontalHeaderLabels(['#','التاريخ','الاسم','النوع','المبلغ','الطريقة','الحساب','ملاحظات'])
        self.payments.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.payments.setSelectionBehavior(QAbstractItemView.SelectRows); self.payments.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.payments.setMinimumHeight(420)
        payments_card_layout.addWidget(self.payments)
        payments_content_layout.addWidget(payments_card)
        self.payments_tab_layout.addWidget(payments_scroll)

        self.last_receipt_path = None
        self.toggle_account()
        self._show_dues_page(0)
        self.refresh_table()

    def _style_dues_inputs(self):
        widgets = [self.person_cb, self.amount, self.date, self.method, self.account_no, self.notes]
        for w in widgets:
            try:
                w.setMinimumHeight(56)
            except Exception:
                pass
        combo_style = f"""
            QComboBox, QDateEdit {{
                background-color:{rgba_from_hex(DARK,0.78)};
                border:1px solid {rgba_from_hex(TEXT,0.10)};
                border-radius:18px;
                padding:12px 42px 12px 14px;
                color:{TEXT};
                font-size:20px;
                font-weight:800;
            }}
            QComboBox::drop-down, QDateEdit::drop-down {{
                subcontrol-origin: padding;
                subcontrol-position: left center;
                width:34px;
                border:none;
                margin-left:12px;
            }}
            QComboBox QAbstractItemView {{
                background-color:{CARD};
                color:{TEXT};
                selection-background-color:{ACCENT};
                selection-color:{TEXT_ON_ACCENT};
                border:1px solid {rgba_from_hex(TEXT,0.10)};
            }}
        """
        line_style = f"""
            QLineEdit, QTextEdit, QDoubleSpinBox, QSpinBox {{
                background-color:{rgba_from_hex(DARK,0.78)};
                border:1px solid {rgba_from_hex(TEXT,0.10)};
                border-radius:18px;
                padding:12px 14px;
                color:{TEXT};
                font-size:20px;
                font-weight:800;
            }}
        """
        for w in [self.person_cb, self.date, self.method]:
            w.setStyleSheet(combo_style)
        for w in [self.amount, self.account_no, self.notes]:
            w.setStyleSheet(line_style)

    def _build_input_card(self, label_text, widget):
        card = QFrame()
        card.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.55)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:22px;")
        layout = QVBoxLayout(card)
        layout.setContentsMargins(12,10,12,12)
        layout.setSpacing(8)
        lbl = QLabel(label_text)
        lbl.setStyleSheet(f'font-size:15px;font-weight:900;color:{TEXT};padding:0 4px;')
        layout.addWidget(lbl)
        layout.addWidget(widget)
        return card

    def _build_dues_entry_page(self, target_label):
        scroll, content = self._make_dues_page('تسجيل حركة مباشرة', f'صفحة إدخال مريحة لتسجيل تسديدات {target_label} بدون تراكب وبنفس النسخة الثابتة.')
        content_layout = content.layout()
        form_card = QFrame()
        form_card.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.82)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:24px;")
        form_layout = QVBoxLayout(form_card)
        form_layout.setContentsMargins(18,18,18,18)
        form_layout.setSpacing(14)
        title = QLabel('بيانات الحركة')
        title.setStyleSheet(f'font-size:22px;font-weight:900;color:{TEXT};')
        form_layout.addWidget(title)
        grid = QGridLayout(); grid.setHorizontalSpacing(14); grid.setVerticalSpacing(14)
        grid.setColumnStretch(0,1); grid.setColumnStretch(1,1)
        grid.addWidget(self._build_input_card('الاسم', self.person_cb), 0, 0)
        grid.addWidget(self._build_input_card('المبلغ', self.amount), 0, 1)
        grid.addWidget(self._build_input_card('التاريخ', self.date), 1, 0)
        grid.addWidget(self._build_input_card('الطريقة', self.method), 1, 1)
        form_layout.addLayout(grid)
        form_layout.addWidget(self._build_input_card('رقم الحساب', self.account_no))
        form_layout.addWidget(self._build_input_card('ملاحظات', self.notes))
        form_layout.addWidget(self.summary_boxes)
        btn_row = QHBoxLayout(); btn_row.setSpacing(10)
        back_btn = QPushButton('↩ الرئيسية')
        back_btn.setStyleSheet(SECONDARY_BUTTON)
        back_btn.clicked.connect(lambda: self._show_dues_page(0))
        receipt_btn = QPushButton('🧾 فتح آخر وصل')
        receipt_btn.setStyleSheet(SECONDARY_BUTTON)
        receipt_btn.clicked.connect(self.open_last_receipt)
        save_btn = QPushButton('💳 تسديد مباشر')
        save_btn.setStyleSheet(BUTTON_STYLE)
        save_btn.clicked.connect(self.add_payment)
        btn_row.addWidget(back_btn)
        btn_row.addStretch(1)
        btn_row.addWidget(receipt_btn)
        btn_row.addWidget(save_btn)
        form_layout.addLayout(btn_row)
        content_layout.addWidget(form_card)
        content_layout.addStretch(1)
        self.entry_tab_layout.addWidget(scroll)

    def _build_dues_home(self, target_label):
        hero = QFrame()
        hero.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.82)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:26px;")
        hero_layout = QVBoxLayout(hero)
        hero_layout.setContentsMargins(18,18,18,18)
        hero_layout.setSpacing(12)
        title_lbl = QLabel(f'{target_label} / الصفحة الرئيسية')
        title_lbl.setStyleSheet(f'font-size:28px;font-weight:900;color:{TEXT};')
        note_lbl = QLabel('اختر الصفحة المطلوبة من هنا: ملخص الحسابات، التفاصيل، أو سجل التسديدات. نفس روح المبيعات وبنفس النسخة الثابتة.')
        note_lbl.setWordWrap(True)
        note_lbl.setStyleSheet(f'font-size:13px;font-weight:700;color:{MUTED};')
        hero_layout.addWidget(title_lbl)
        hero_layout.addWidget(note_lbl)
        stat_grid = QGridLayout(); stat_grid.setHorizontalSpacing(14); stat_grid.setVerticalSpacing(14)
        self.dues_total_card = make_stat_card('إجمالي المتبقي', 'الرصيد القائم حالياً', ACCENT)
        self.dues_people_card = make_stat_card(f'عدد {target_label}', 'الأسماء الظاهرة في الملخص', '#a997ff')
        self.dues_focus_card = make_stat_card('الحالة المحددة', 'يتحدث حسب الاسم المختار', '#ff9bca')
        stat_grid.addWidget(self.dues_total_card, 0, 0, 1, 2)
        stat_grid.addWidget(self.dues_people_card, 1, 0)
        stat_grid.addWidget(self.dues_focus_card, 1, 1)
        hero_layout.addLayout(stat_grid)
        self.dues_home_layout.addWidget(hero)

        nav_card = QFrame()
        nav_card.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.72)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:24px;")
        nav_layout = QVBoxLayout(nav_card)
        nav_layout.setContentsMargins(18,18,18,18)
        nav_layout.setSpacing(14)
        nav_title = QLabel('الأقسام')
        nav_title.setStyleSheet(f'font-size:24px;font-weight:900;color:{TEXT};')
        nav_layout.addWidget(nav_title)
        btn_grid = QGridLayout(); btn_grid.setHorizontalSpacing(12); btn_grid.setVerticalSpacing(12)
        for idx, (txt, note, page_idx) in enumerate([
            ('تسجيل حركة مباشرة', 'إدخال الاسم والمبلغ والتسديد بصفحة مريحة', 1),
            ('ملخص الحسابات', 'كشف سريع لكل الأسماء والأرصدة', 2),
            ('تفاصيل الحركة', 'تفصيل الاسم المحدد بالفواتير والتسديدات', 3),
            ('سجل التسديدات', 'كل الوصلات والحركات المرحّلة', 4),
        ]):
            btn = QPushButton(txt)
            btn.setMinimumHeight(74)
            btn.setStyleSheet(f"QPushButton{{text-align:right;padding:18px 22px;border-radius:22px;background-color:{rgba_from_hex(DARK,0.78)};border:1px solid {rgba_from_hex(TEXT,0.08)};font-size:20px;font-weight:900;color:{TEXT};}} QPushButton:hover{{border:1px solid {rgba_from_hex(ACCENT,0.45)};background-color:{rgba_from_hex('#ffffff',0.05)};}}")
            btn.clicked.connect(lambda _=False, i=page_idx: self._show_dues_page(i))
            info = QLabel(note)
            info.setWordWrap(True)
            info.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};padding:0 4px;')
            wrap = QFrame(); wrap.setStyleSheet('background:transparent;border:none;')
            wrap_layout = QVBoxLayout(wrap); wrap_layout.setContentsMargins(0,0,0,0); wrap_layout.setSpacing(6)
            wrap_layout.addWidget(btn)
            wrap_layout.addWidget(info)
            btn_grid.addWidget(wrap, idx // 2, idx % 2)
        nav_layout.addLayout(btn_grid)
        self.dues_home_layout.addWidget(nav_card)
        self.dues_home_layout.addStretch(1)
        self.dues_home_layout.addStretch(1)

    def _show_dues_page(self, index):
        self.dues_stack.setCurrentIndex(index)

    def _make_dues_page(self, title, note):
        scroll = QScrollArea(); scroll.setWidgetResizable(True)
        scroll.setStyleSheet("QScrollArea{background:transparent;border:none;} QScrollArea > QWidget > QWidget{background:transparent;}")
        content = QWidget()
        layout = QVBoxLayout(content)
        layout.setContentsMargins(18,18,18,28)
        layout.setSpacing(16)
        header = QFrame()
        header.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.82)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:24px;")
        header_layout = QVBoxLayout(header)
        header_layout.setContentsMargins(16,16,16,16)
        header_layout.setSpacing(10)
        top = QHBoxLayout(); top.setSpacing(10)
        back_btn = QPushButton('↩ الرئيسية')
        back_btn.setStyleSheet(SECONDARY_BUTTON)
        back_btn.clicked.connect(lambda: self._show_dues_page(0))
        title_lbl = QLabel(title)
        title_lbl.setStyleSheet(f'font-size:28px;font-weight:900;color:{TEXT};')
        top.addWidget(back_btn, 0)
        top.addWidget(title_lbl, 1)
        header_layout.addLayout(top)
        note_lbl = QLabel(note)
        note_lbl.setWordWrap(True)
        note_lbl.setStyleSheet(f'font-size:13px;font-weight:700;color:{MUTED};')
        header_layout.addWidget(note_lbl)
        layout.addWidget(header)
        scroll.setWidget(content)
        return scroll, content

    def toggle_account(self):
        self.account_no.setVisible(self.method.currentText() == 'تحويل')

    def refresh_combos(self):
        if hasattr(self, 'person_cb'):
            current = self.person_cb.currentText()
            self.person_cb.clear(); self.person_cb.addItems([x.get('name','') for x in self.db.get(self.people_key, [])])
            idx = self.person_cb.findText(current)
            if idx >= 0:
                self.person_cb.setCurrentIndex(idx)

    def sync_selected_person_from_table(self):
        row = self.table.currentRow()
        if row < 0:
            return
        item = self.table.item(row, 1)
        if not item:
            return
        name = (item.text() or '').strip()
        if not name:
            return
        idx = self.person_cb.findText(name)
        if idx >= 0 and self.person_cb.currentIndex() != idx:
            self.person_cb.setCurrentIndex(idx)

    def due_rows(self):
        return customer_due_summary(self.db) if self.people_key == 'customers' else supplier_due_summary(self.db)

    def payment_source(self):
        return 'customer_payment' if self.people_key == 'customers' else 'supplier_payment'

    def payment_type(self):
        return 'إيراد' if self.people_key == 'customers' else 'مصروف'

    def payment_category(self):
        return 'وصل قبض دين' if self.people_key == 'customers' else 'وصل تسديد مورد'

    def opening_rows_key(self):
        return 'customers' if self.people_key == 'customers' else 'suppliers'

    def save_opening_due(self):
        name = self.person_cb.currentText().strip()
        amount = float(self.amount.value())
        if not name or amount <= 0:
            return QMessageBox.warning(self, 'تنبيه', 'اختر الاسم وأدخل مبلغ الدين القديم')
        rows = opening_data(self.db).setdefault(self.opening_rows_key(), [])
        for row in rows:
            if row.get('name', '') == name:
                row['amount'] = amount
                row['created_at'] = row.get('created_at', now_str())
                break
        else:
            rows.append({'name': name, 'amount': amount, 'created_at': now_str()})
        self.amount.setValue(0)
        self.notes.clear()
        self.account_no.clear()
        self.save()
        self.refresh_table()
        QMessageBox.information(self, 'تم', 'تم حفظ الدين القديم ضمن الرصيد الافتتاحي بدون تأثير على الأرباح.')

    def delete_opening_due(self):
        name = self.person_cb.currentText().strip()
        if not name:
            return QMessageBox.warning(self, 'تنبيه', 'اختر الاسم أولاً')
        rows = opening_data(self.db).setdefault(self.opening_rows_key(), [])
        before = len(rows)
        rows[:] = [x for x in rows if x.get('name', '') != name]
        if len(rows) == before:
            return QMessageBox.information(self, 'تنبيه', 'ماكو دين قديم محفوظ لهذا الاسم.')
        self.save()
        self.refresh_table()
        QMessageBox.information(self, 'تم', 'تم حذف الدين القديم.')


    def all_refs(self, name):
        rows = []
        if not name:
            return rows

        if self.people_key == 'customers':
            snapshot = customer_receivable_snapshot(self.db, name)
            opening_original = opening_customer_amount(self.db, name)
            opening_remaining = float(snapshot.get('opening_remaining', 0) or 0)
            if opening_original:
                rows.append({
                    'ref': 'دين قديم',
                    'date': '',
                    'item': 'رصيد سابق',
                    'qty': '',
                    'value': opening_original,
                    'due': opening_remaining,
                    'kind': 'opening',
                    'gid': 'opening',
                })

            grouped = {}
            for s in self.db.get('sales', []):
                if s.get('customer') != name:
                    continue
                gid = s.get('invoice_group_id') or f"single-{ensure_invoice_no(self.db, s)}"
                g = grouped.setdefault(gid, {
                    'ref': f"فاتورة #{ensure_invoice_no(self.db, s)}",
                    'date': s.get('date',''),
                    'item': [],
                    'qty': 0,
                    'value': 0.0,
                    'due': 0.0,
                    'kind': 'invoice',
                    'gid': gid,
                    'invoice_no': ensure_invoice_no(self.db, s),
                })
                g['item'].append(s.get('item',''))
                g['qty'] += int(s.get('qty',0) or 0)
                g['value'] += float(s.get('total',0) or 0)
            for gid, g in grouped.items():
                g['due'] = current_sale_due(self.db, gid)
            invoice_rows = []
            for g in grouped.values():
                items = [x for x in g.get('item', []) if x]
                g['item'] = ' + '.join(items[:2]) + (f" + {len(items)-2}..." if len(items) > 2 else '')
                invoice_rows.append(g)
            invoice_rows.sort(key=lambda x: (x.get('date',''), x.get('invoice_no', 0)))
            rows.extend(invoice_rows)

            return_state_map = snapshot.get('return_states', {}) or {}
            return_rows = []
            for ret in self.db.get('returns', []):
                if ret.get('customer') != name:
                    continue
                ret_state = return_state_map.get(ret.get('id',''), {})
                return_rows.append({
                    'ref': f"مرتجع فاتورة #{ret.get('invoice_no','')}",
                    'date': ret.get('date',''),
                    'item': ret.get('item',''),
                    'qty': ret.get('qty',''),
                    'value': float(ret.get('credit_amount', ret.get('total', 0)) or 0),
                    'due': round(float(ret_state.get('remaining_credit', 0) or 0), 2),
                    'kind': 'return',
                    'gid': ret.get('id',''),
                })
            rows.extend(return_rows)

            pay_rows = []
            for p in self.db.get('cash', []):
                if p.get('party') != name or p.get('source') not in ('customer_payment', 'opening_customer_payment'):
                    continue
                pay_rows.append({
                    'ref': f"قبض #{p.get('receipt_no','')}",
                    'date': p.get('date',''),
                    'item': p.get('notes','') or 'تسديد',
                    'qty': '',
                    'value': float(p.get('amount',0) or 0),
                    'due': 0.0,
                    'kind': 'payment',
                    'gid': p.get('receipt_no',''),
                })
            rows.extend(pay_rows)
            rows.sort(key=lambda x: (x.get('date','9999-99-99') or '9999-99-99', {'opening':0,'invoice':1,'return':2,'payment':3}.get(x.get('kind'), 9)))
            return rows

        opening_remaining = opening_supplier_due_remaining(self.db, name)
        opening_original = opening_supplier_amount(self.db, name)
        if opening_original:
            rows.append({
                'ref': 'دين قديم',
                'date': '',
                'item': 'رصيد سابق',
                'qty': '',
                'value': opening_original,
                'due': opening_remaining,
                'kind': 'opening',
                'gid': 'opening',
            })

        for row in self.db.get('inbound', []):
            if row.get('supplier') != name:
                continue
            rows.append({
                'ref': f"وارد {row.get('item','')}",
                'date': row.get('date',''),
                'item': row.get('item',''),
                'qty': row.get('qty',''),
                'value': float(row.get('total',0) or 0),
                'due': current_inbound_due(self.db, row),
                'kind': 'invoice',
                'gid': row.get('id',''),
            })

        for p in self.db.get('cash', []):
            if p.get('party') != name or p.get('source') not in ('supplier_payment', 'opening_supplier_payment'):
                continue
            rows.append({
                'ref': f"تسديد #{p.get('receipt_no','')}",
                'date': p.get('date',''),
                'item': p.get('notes','') or 'تسديد',
                'qty': '',
                'value': float(p.get('amount',0) or 0),
                'due': 0.0,
                'kind': 'payment',
                'gid': p.get('receipt_no',''),
            })
        rows.sort(key=lambda x: (x.get('date','9999-99-99') or '9999-99-99', {'opening':0,'invoice':1,'payment':2}.get(x.get('kind'), 9)))
        return rows
    def refresh_details(self):
        name = self.person_cb.currentText().strip()
        refs = self.all_refs(name) if name else []
        self.details.setRowCount(len(refs))
        section_labels = {
            'opening': 'الديون القديمة',
            'invoice': 'الفواتير الجديدة',
            'payment': 'التسديدات',
            'return': 'المرتجعات',
        }
        for r, row in enumerate(refs):
            vals = [
                r+1,
                section_labels.get(row.get('kind',''), row.get('kind','')),
                row.get('ref',''),
                row.get('date',''),
                row.get('item',''),
                row.get('qty',''),
                fmt_money(row.get('value',0)),
                fmt_money(row.get('due',0)),
            ]
            for c, v in enumerate(vals):
                self.details.setItem(r, c, QTableWidgetItem(str(v)))

        due = next((x for x in self.due_rows() if x['name'] == name), {'due': 0, 'credit_sales': 0, 'credit_purchases': 0, 'receipts': 0, 'payments': 0, 'return_credit':0, 'opening_due':0, 'opening_original':0, 'invoice_total':0, 'later_receipts':0, 'initial_paid':0})
        if self.people_key == 'customers':
            old_debt = float(due.get('opening_original', 0) or 0)
            new_debt = float(due.get('invoice_total', 0) or 0)
            total_payments = float(due.get('receipts', 0) or 0)
            ret_credit = float(due.get('return_credit', 0) or 0)
            final_due = float(due.get('due', 0) or 0)
            summary_text = f"الزبون: {name or '—'} | الدين القديم: {fmt_money(old_debt)} د.ع | الدين الجديد: {fmt_money(new_debt)} د.ع | إجمالي التسديدات: {fmt_money(total_payments)} د.ع | رصيد المرتجعات: {fmt_money(ret_credit)} د.ع | المتبقي النهائي: {fmt_money(final_due)} د.ع"
            self.summary_boxes.setText(summary_text)
            self.details_summary_box.setText(summary_text)
            if hasattr(self, 'dues_focus_card'):
                self.dues_focus_card.set_value(name or '—')
                self.dues_focus_card.set_note(f'المتبقي النهائي: {fmt_money(final_due)} د.ع')
        else:
            base = due.get('credit_purchases', 0)
            mov = due.get('payments', 0)
            opening_due = float(due.get('opening_due', 0) or 0)
            old_txt = f" | المتبقي من القديم: {fmt_money(opening_due)} د.ع" if opening_due else ''
            summary_text = f"إجمالي القديم + الجديد: {fmt_money(base)} د.ع | مجموع التسديدات: {fmt_money(mov)} د.ع | المتبقي على {name or '—'}: {fmt_money(due.get('due',0))} د.ع{old_txt}"
            self.summary_boxes.setText(summary_text)
            self.details_summary_box.setText(summary_text)
            if hasattr(self, 'dues_focus_card'):
                self.dues_focus_card.set_value(name or '—')
                self.dues_focus_card.set_note(f'المتبقي الحالي: {fmt_money(due.get("due",0))} د.ع')

    def refresh_table(self):
        self.refresh_combos()
        rows = self.due_rows()
        self.table.setRowCount(len(rows))
        total_due = total_base = 0
        for r, row in enumerate(rows):
            due = float(row.get('due', 0) or 0)
            total_due += due
            if self.people_key == 'customers':
                old_debt = float(row.get('opening_original', 0) or 0)
                new_debt = float(row.get('invoice_total', 0) or 0)
                payments = float(row.get('receipts', 0) or 0)
                return_credit = float(row.get('return_credit', 0) or 0)
                total_base += old_debt + new_debt
                vals = [r+1, row['name'], fmt_money(old_debt), fmt_money(new_debt), fmt_money(payments), fmt_money(return_credit), fmt_money(due)]
            else:
                base = row.get('credit_purchases', 0)
                moves = row.get('payments', 0)
                total_base += base
                vals = [r+1, row['name'], fmt_money(base), fmt_money(moves), fmt_money(due)]
            for c, v in enumerate(vals):
                self.table.setItem(r, c, QTableWidgetItem(str(v)))
        cash_rows = [x for x in self.db.get('cash', []) if x.get('source') in (self.payment_source(), 'opening_customer_payment' if self.people_key == 'customers' else 'opening_supplier_payment')]
        self.payments.setRowCount(len(cash_rows))
        for r, row in enumerate(cash_rows):
            vals = [r+1, row.get('date',''), row.get('party',''), row.get('category',''), fmt_money(row.get('amount',0)), row.get('payment_method','نقد'), row.get('account_no',''), row.get('notes','')]
            for c, v in enumerate(vals):
                self.payments.setItem(r, c, QTableWidgetItem(str(v)))
        if hasattr(self, 'dues_total_card'):
            self.dues_total_card.set_value(f'{fmt_money(total_due)} د.ع')
            self.dues_total_card.set_note(f'إجمالي الحركة: {fmt_money(total_base)} د.ع')
        if hasattr(self, 'dues_people_card'):
            label = 'زبون' if self.people_key == 'customers' else 'مورد'
            self.dues_people_card.set_value(f'{len(rows)} {label}')
            self.dues_people_card.set_note(f'عدد التسديدات المسجلة: {len(cash_rows)}')
        self.refresh_details()

    def add_payment(self):
        name = self.person_cb.currentText().strip()
        amount = float(self.amount.value())
        if not name or amount <= 0:
            return QMessageBox.warning(self,'تنبيه','أكمل البيانات')

        due = 0.0
        for row in self.due_rows():
            if row['name'] == name:
                due = float(row.get('due', 0) or 0)
                break
        if amount > due:
            return QMessageBox.warning(self,'تنبيه', f'المبلغ أكبر من المتبقي ({fmt_money(due)})')

        rec_no = next_receipt_no(self.db, 'next_due_receipt_no')
        pay_date = self.date.date().toString('yyyy-MM-dd')
        method = self.method.currentText()
        account_no = self.account_no.text().strip()
        notes = self.notes.text().strip()
        title = 'وصل قبض' if self.people_key == 'customers' else 'وصل تسديد'

        source = self.payment_source()
        category = self.payment_category()
        save_source = source

        # إذا عليه دين قديم فقط وما عنده حركات جديدة، نخليه على مصدر الافتتاحي حتى يبقى واضح بالكشف
        if self.people_key == 'customers':
            has_new = any(s.get('customer') == name for s in self.db.get('sales', []))
            if opening_customer_due_remaining(self.db, name) > 0 and not has_new:
                save_source = 'opening_customer_payment'
                category = 'تسديد دين قديم'
        else:
            has_new = any(i.get('supplier') == name for i in self.db.get('inbound', []))
            if opening_supplier_due_remaining(self.db, name) > 0 and not has_new:
                save_source = 'opening_supplier_payment'
                category = 'تسديد دين قديم'

        self.db['cash'].append({
            'date': pay_date,
            'type': self.payment_type(),
            'category': category,
            'party': name,
            'amount': amount,
            'notes': notes,
            'source': save_source,
            'payment_method': method,
            'account_no': account_no,
            'receipt_no': rec_no,
            'created_at': now_str()
        })

        extra = [('نوع العملية', 'تسديد مباشر تلقائي'), ('آلية التوزيع', 'الأقدم أولاً')]
        self.last_receipt_path = write_receipt_file(self.db, save_source, rec_no, title, pay_date, name, amount, method, account_no, notes, extra)

        self.amount.setValue(0)
        self.notes.clear()
        self.account_no.clear()
        self.save()
        self.refresh_table()

    def open_last_receipt(self):
        if self.last_receipt_path and Path(self.last_receipt_path).exists():
            import os
            os.startfile(str(self.last_receipt_path))
        else:
            QMessageBox.information(self, 'تنبيه', 'ماكو وصل منشأ بعد.')

    def open_customer_statement(self):
        if self.people_key != 'customers':
            return
        name = self.person_cb.currentText().strip()
        if not name:
            return QMessageBox.warning(self, 'تنبيه', 'اختر الزبون أولاً')
        try:
            path = write_customer_statement_file(self.db, name)
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(path.resolve())))
        except Exception as exc:
            QMessageBox.warning(self, 'تعذر فتح الكشف', f'تعذر إنشاء أو فتح كشف الزبون: {exc}')

    def open_customer_statement_direct(self):
        if self.key != 'customers':
            return
        i = self.real_index()
        if i is None:
            return QMessageBox.warning(self, 'تنبيه', 'اختر الزبون أولاً')
        try:
            name = (self.db.get(self.key, [])[i] or {}).get('name', '').strip()
        except Exception:
            name = ''
        if not name:
            return QMessageBox.warning(self, 'تنبيه', 'تعذر تحديد اسم الزبون المختار')
        try:
            dlg = CustomerStatementDialog(self, self.db, name)
            dlg.exec()
        except Exception as exc:
            QMessageBox.warning(self, 'تعذر فتح الكشف', f'تعذر إنشاء أو فتح كشف الزبون: {exc}')


class FundersWindow(BaseWindow):
    def __init__(self, main):
        super().__init__(main, '🏦 الممولين')
        self.subtitle_lbl.setText('واجهة الممولين صارت مقسمة بهدوء: كل جزء داخل تبويب مستقل وواضح بدون زحمة.')

        hero = QFrame(); hero.setStyleSheet(CARD_FRAME_STYLE)
        hero_box = QVBoxLayout(hero); hero_box.setContentsMargins(20,18,20,18); hero_box.setSpacing(10)
        hero_title = QLabel('لوحة الممولين')
        hero_title.setAlignment(Qt.AlignRight)
        hero_title.setStyleSheet('font-size:20px;font-weight:900;background:transparent;border:none;')
        hero_note = QLabel('ملخص سريع للممولين الفعّالين، نسبهم الحالية، وكشف حركات كل ممول داخل تبويبات واضحة.')
        hero_note.setAlignment(Qt.AlignRight); hero_note.setWordWrap(True)
        hero_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        hero_box.addWidget(hero_title); hero_box.addWidget(hero_note)
        self.layout.addWidget(hero)

        self.tabs = style_funders_tabs(QTabWidget())
        self.funders_actions_tab = QWidget(); self.funders_actions_layout = QVBoxLayout(self.funders_actions_tab); self.funders_actions_layout.setContentsMargins(10,10,10,10); self.funders_actions_layout.setSpacing(12)
        self.funders_summary_tab = QWidget(); self.funders_summary_tab_layout = QVBoxLayout(self.funders_summary_tab); self.funders_summary_tab_layout.setContentsMargins(10,10,10,10); self.funders_summary_tab_layout.setSpacing(12)
        self.funders_table_tab = QWidget(); self.funders_table_layout = QVBoxLayout(self.funders_table_tab); self.funders_table_layout.setContentsMargins(10,10,10,10); self.funders_table_layout.setSpacing(12)
        self.funders_movements_tab = QWidget(); self.funders_movements_tab_layout = QVBoxLayout(self.funders_movements_tab); self.funders_movements_tab_layout.setContentsMargins(10,10,10,10); self.funders_movements_tab_layout.setSpacing(12)
        self.tabs.addTab(self.funders_actions_tab, 'إجراءات الممولين')
        self.tabs.addTab(self.funders_summary_tab, 'ملخص التمويل')
        self.tabs.addTab(self.funders_table_tab, 'جدول الممولين')
        self.tabs.addTab(self.funders_movements_tab, 'كشف الحركات')
        self.layout.addWidget(self.tabs, 1)

        actions_card = QFrame(); actions_card.setStyleSheet(CARD_FRAME_STYLE)
        actions_box = QVBoxLayout(actions_card); actions_box.setContentsMargins(18,18,18,18); actions_box.setSpacing(10)
        actions_title = QLabel('إجراءات الممولين')
        actions_title.setAlignment(Qt.AlignRight)
        actions_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        actions_box.addWidget(actions_title)
        specs = [
            ('➕ إضافة ممول', self.add_funder, BUTTON_STYLE), ('✏️ تعديل', self.edit_funder, SECONDARY_BUTTON), ('🗑 حذف', self.delete_funder, SECONDARY_BUTTON),
            ('🔁 تفعيل/إيقاف', self.toggle_active, SECONDARY_BUTTON), ('💵 إضافة رصيد', self.deposit_capital, BUTTON_STYLE), ('📤 سحب رصيد', self.withdraw_capital, SECONDARY_BUTTON),
            ('✅ تم تسديد الممول', self.pay_profit, BUTTON_STYLE), ('🧾 فتح آخر وصل', self.open_last_receipt, SECONDARY_BUTTON)
        ]
        grid = QGridLayout(); grid.setHorizontalSpacing(10); grid.setVerticalSpacing(10)
        for i,(txt, fn, style) in enumerate(specs):
            b = QPushButton(txt); b.setStyleSheet(style); b.setMinimumHeight(46); b.clicked.connect(fn); grid.addWidget(b, i // 2, i % 2)
        actions_box.addLayout(grid)
        self.funders_actions_layout.addWidget(actions_card)
        self.funders_actions_layout.addStretch(1)

        metrics_card = QFrame(); metrics_card.setStyleSheet(CARD_FRAME_STYLE)
        metrics_box = QVBoxLayout(metrics_card); metrics_box.setContentsMargins(18,18,18,18); metrics_box.setSpacing(10)
        metrics_title = QLabel('ملخص التمويل')
        metrics_title.setAlignment(Qt.AlignRight)
        metrics_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        metrics_box.addWidget(metrics_title)
        self.funder_count_card = SummaryCard('عدد الممولين', '0', 'عدد الممولين الظاهرين')
        self.active_capital_card = SummaryCard('إجمالي التمويل الفعال', '0', 'رأس المال الفعّال الحالي')
        self.pending_profit_card = SummaryCard('المستحقات الحالية', '0', 'إجمالي أرباح الممولين غير المسددة')
        metrics_grid = QGridLayout(); metrics_grid.setHorizontalSpacing(10); metrics_grid.setVerticalSpacing(10)
        for i, c in enumerate((self.funder_count_card, self.active_capital_card, self.pending_profit_card)):
            c.setMinimumHeight(100); c.setMaximumHeight(116); metrics_grid.addWidget(c, i // 2, i % 2)
        metrics_box.addLayout(metrics_grid)
        self.funders_summary_tab_layout.addWidget(metrics_card)

        table_card = QFrame(); table_card.setStyleSheet(CARD_FRAME_STYLE)
        table_box = QVBoxLayout(table_card); table_box.setContentsMargins(16,16,16,16); table_box.setSpacing(10)
        table_title = QLabel('جدول الممولين')
        table_title.setAlignment(Qt.AlignRight)
        table_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        table_box.addWidget(table_title)
        self.table = QTableWidget(); self.table.setStyleSheet(TABLE_STYLE); self.table.setColumnCount(11)
        self.table.setHorizontalHeaderLabels(['#','اسم الممول','الرصيد الفعال','النسبة الحالية','حصة الثلث','المسدّد كلياً','المتبقي له','النوع','الهاتف','الحالة','ملاحظات'])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.itemSelectionChanged.connect(self.refresh_movements)
        table_box.addWidget(self.table, 1)
        self.summary = QLabel(); self.summary.setAlignment(Qt.AlignRight); self.summary.setWordWrap(True)
        self.summary.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        table_box.addWidget(self.summary)
        self.funders_table_layout.addWidget(table_card, 1)

        details_card = QFrame(); details_card.setStyleSheet(CARD_FRAME_STYLE)
        details_box = QVBoxLayout(details_card); details_box.setContentsMargins(16,16,16,16); details_box.setSpacing(10)
        details_title = QLabel('التفصيل والسجل')
        details_title.setAlignment(Qt.AlignRight)
        details_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        details_box.addWidget(details_title)
        self.details_label = QLabel(); self.details_label.setAlignment(Qt.AlignRight); self.details_label.setWordWrap(True)
        self.details_label.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        details_box.addWidget(self.details_label)
        self.movements = QTableWidget(); self.movements.setStyleSheet(TABLE_STYLE); self.movements.setColumnCount(11)
        self.movements.setHorizontalHeaderLabels(['#','التاريخ','الممول','الحركة','المبلغ','الرصيد بعد الحركة','النسبة بعد الحركة','إجمالي ربحه','المسدّد','المتبقي','ملاحظات'])
        self.movements.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.movements.setSelectionBehavior(QAbstractItemView.SelectRows); self.movements.setEditTriggers(QAbstractItemView.NoEditTriggers)
        details_box.addWidget(self.movements, 1)
        self.funders_movements_tab_layout.addWidget(details_card, 1)

        self.last_receipt_path = None
        self.refresh_table()

    def funders_list(self):
        rows = [ensure_funder_identity(dict(x), x.get('name', '')) for x in normalized_funders(self.db)]
        return rows if rows else [ensure_funder_identity(dict(x), x.get('name', '')) for x in self.db.get('funders', [])]

    def selected_funder(self):
        r = self.table.currentRow()
        funders = self.funders_list()
        if r < 0 or r >= len(funders):
            QMessageBox.warning(self, 'تنبيه', 'اختر ممول')
            return None
        return funders[r]

    def refresh_table(self):
        build_profit_ledger(self.db)
        data = self.funders_list(); self.table.setRowCount(len(data)); active_total = pending_total = 0.0
        pb = profit_breakdown(self.db)
        pb_map = {x.get('name',''): x for x in pb.get('funders_rows', [])}
        for r, x in enumerate(data):
            cap = funder_effective_capital(self.db, x)
            if x.get('active', True):
                active_total += cap
            p = pb_map.get(x.get('name',''), {})
            pending = float(p.get('pending', 0) or 0)
            pending_total += pending
            vals = [
                r+1,
                x.get('name',''),
                fmt_money(cap),
                fmt_pct(funder_ratio_pct(self.db, x)),
                fmt_money(p.get('amount', 0)),
                fmt_money(p.get('paid', 0)),
                fmt_money(pending),
                'رأس مال هضاب' if x.get('is_owner_capital', False) else 'ممّول خارجي',
                x.get('phone',''),
                'مفعل' if x.get('active', True) else 'موقوف',
                x.get('notes','')
            ]
            for c, v in enumerate(vals):
                self.table.setItem(r, c, QTableWidgetItem(str(v)))
        self.funder_count_card.set_value(str(len(data)))
        self.active_capital_card.set_value(fmt_money(active_total))
        self.pending_profit_card.set_value(fmt_money(pending_total))
        self.summary.setText(
            f'عدد الممولين: {len(data)} | إجمالي التمويل الفعال: {fmt_money(active_total)} د.ع | '
            f'إجمالي مستحقاتهم الحالية: {fmt_money(pending_total)} د.ع | '
            f'النسب تتحدث تلقائياً حسب الرصيد الفعال بعد كل إضافة أو سحب.'
        )
        if data:
            row = self.table.currentRow()
            if row < 0:
                self.table.selectRow(0)
            else:
                self.refresh_movements()
        else:
            self.movements.setRowCount(0)
            self.details_label.setText('لا يوجد ممولون حالياً.')

    def refresh_movements(self):
        funder = None
        r = self.table.currentRow()
        funders = self.funders_list()
        if 0 <= r < len(funders):
            funder = funders[r]
        rows = funder_movement_rows(self.db, funder.get('name','') if funder else None)
        self.movements.setRowCount(len(rows))
        for i, row in enumerate(rows):
            vals = [
                i+1,
                row.get('date','') or '—',
                row.get('party',''),
                row.get('movement',''),
                fmt_money(row.get('amount',0)),
                fmt_money(row.get('capital_after',0)),
                fmt_pct(float(row.get('ratio_after',0) or 0)),
                fmt_money(row.get('earned_profit',0)),
                fmt_money(row.get('paid_profit',0)),
                fmt_money(row.get('pending_profit',0)),
                row.get('notes','')
            ]
            for c, v in enumerate(vals):
                self.movements.setItem(i, c, QTableWidgetItem(str(v)))
        if funder:
            pb_row = next((x for x in profit_breakdown(self.db).get('funders_rows', []) if x.get('name','') == funder.get('name','')), {})
            self.details_label.setText(
                f"كشف تفصيلي للممول: {funder.get('name','')} | الرصيد الأساسي: {fmt_money(funder_capital_base(self.db, funder.get('name','')))} د.ع | "
                f"الرصيد الفعال الحالي: {fmt_money(funder_effective_capital(self.db, funder))} د.ع | "
                f"نسبته الحالية: {fmt_pct(funder_ratio_pct(self.db, funder))} | "
                f"إجمالي ربحه: {fmt_money(pb_row.get('amount',0))} د.ع | "
                f"المسدّد له: {fmt_money(pb_row.get('paid',0))} د.ع | "
                f"المتبقي له: {fmt_money(pb_row.get('pending',0))} د.ع"
            )
        else:
            self.details_label.setText('اختر ممولاً لعرض كشفه التفصيلي.')

    def add_funder(self):
        while True:
            d = FunderDialog(self)
            result = d.exec()
            self.raise_(); self.activateWindow()
            if not result:
                break
            x = d.get_data(); name = x.get('name','').strip()
            if not name:
                QMessageBox.warning(self,'تنبيه','اسم الممول مطلوب')
                continue
            if any(f.get('name','').strip() == name for f in self.db.get('funders', [])):
                QMessageBox.warning(self,'منع','اسم الممول موجود مسبقًا')
                continue
            x['name'] = name
            stamp = now_str()
            x.setdefault('created_at', stamp)
            x['base_capital'] = max(0.0, _safe_float(x.get('capital', 0)))
            x['status_history'] = [{'active': bool(x.get('active', True)), 'date': stamp[:10], 'created_at': stamp}]
            self.db.setdefault('funders', []).append(x)
            self.save(); self.refresh_table()

    def edit_funder(self):
        selected = self.selected_funder()
        if not selected:
            return
        f = resolve_funder_record(self.db, selected, create_missing=True)
        i = self.db['funders'].index(f)
        old_name = f.get('name','').strip()
        dialog_seed = dict(f)
        dialog_seed['capital'] = round(funder_effective_capital(self.db, f), 2)
        d = FunderDialog(self, dialog_seed)
        if d.exec():
            data = d.get_data(); name = data.get('name','').strip()
            if not name:
                return QMessageBox.warning(self,'تنبيه','اسم الممول مطلوب')
            if any(idx != i and fx.get('name','').strip() == name for idx, fx in enumerate(self.db.get('funders', []))):
                return QMessageBox.warning(self,'منع','اسم الممول موجود مسبقًا')
            stamp = now_str()
            desired_capital = max(0.0, _safe_float(data.get('capital', 0)))
            current_effective = round(funder_effective_capital(self.db, f), 2)
            capital_delta = round(desired_capital - current_effective, 2)
            prev_active = bool(f.get('active', True))
            new_active = bool(data.get('active', True))
            data['name'] = name
            f['name'] = name
            f['phone'] = data.get('phone','')
            f['notes'] = data.get('notes','')
            f['is_owner_capital'] = bool(data.get('is_owner_capital', False))
            f['capital'] = desired_capital
            f.setdefault('base_capital', max(0.0, _safe_float(f.get('capital', 0))))
            history = list(f.get('status_history', []) or [])
            if not history:
                base_stamp = str(f.get('created_at', '') or stamp)
                history = [{'active': prev_active, 'date': base_stamp[:10], 'created_at': base_stamp}]
            if prev_active != new_active:
                history.append({'active': new_active, 'date': stamp[:10], 'created_at': stamp})
            f['active'] = new_active
            f['status_history'] = sorted(history, key=lambda x: (str(x.get('date', '') or ''), str(x.get('created_at', '') or '')))
            if abs(capital_delta) > 0.009:
                self.db.setdefault('cash', []).append({
                    'date': stamp[:10], 'type': 'إيراد' if capital_delta > 0 else 'مصروف', 'category': 'تمويل', 'party': old_name,
                    'amount': abs(capital_delta), 'notes': f"تسوية تعديل رصيد الممول {old_name}",
                    'source': 'funder_capital_in' if capital_delta > 0 else 'funder_capital_out', 'created_at': stamp,
                    'payment_method': 'تعديل', 'account_no': '', 'receipt_no': generate_receipt_no(self.db, stamp[:10]),
                })
            if old_name != name:
                rename_funder_references(self.db, old_name, name)
            self.save(); self.refresh_table()

    def delete_funder(self):
        selected = self.selected_funder()
        if not selected: return
        f = resolve_funder_record(self.db, selected, create_missing=False) or selected
        name = f.get('name','')
        has_movements = any(x.get('party','') == name and x.get('source') in ('funder_capital_in','funder_capital_out','funder_profit_payment') for x in self.db.get('cash', []))
        has_profit_rows = any(x.get('beneficiary_type') == 'funder' and x.get('beneficiary_name','') == name for x in self.db.get('profit_entries', []))
        has_opening_capital = opening_old_funder_capital(self.db, name) > 0
        has_effective_capital = funder_effective_capital(self.db, f) > 0
        if has_movements or has_profit_rows or has_opening_capital or has_effective_capital:
            return QMessageBox.warning(self,'منع','لا يمكن حذف ممول لديه رأس مال أو حركات أو أرباح محفوظة. صفِّ رصيده أولاً.')
        if QMessageBox.question(self,'تأكيد',f"حذف {name}؟") == QMessageBox.Yes:
            self.db['funders'].remove(f); self.save(); self.refresh_table()

    def toggle_active(self):
        selected = self.selected_funder()
        if not selected:
            return
        f = resolve_funder_record(self.db, selected, create_missing=True)
        prev_state = bool(f.get('active', True))
        new_state = not prev_state
        history = list(f.get('status_history', []) or [])
        stamp = now_str()
        if not history:
            base_stamp = str(f.get('created_at', '') or '') or stamp
            history.append({'active': prev_state, 'date': base_stamp[:10], 'created_at': base_stamp})
        f['active'] = new_state
        history.append({'active': new_state, 'date': stamp[:10], 'created_at': stamp})
        f['status_history'] = history
        self.save(); self.refresh_table()

    def amount_dialog(self, title, label='المبلغ'):
        d = QDialog(self)
        d.setWindowTitle(title)
        d.resize(560, 320)
        d.setMinimumSize(560, 320)
        apply_theme_to_widget(d)

        layout = QVBoxLayout(d)
        layout.setContentsMargins(18, 18, 18, 18)
        layout.setSpacing(14)

        form = QFormLayout()
        form.setLabelAlignment(Qt.AlignRight)
        form.setFormAlignment(Qt.AlignTop)
        form.setHorizontalSpacing(18)
        form.setVerticalSpacing(12)

        amount = QDoubleSpinBox()
        amount.setRange(0, 1_000_000_000_000)
        amount.setDecimals(0)

        date = QDateEdit()
        fix_date_edit_widget(date)
        date.setDate(QDate.currentDate())

        method = QComboBox()
        method.addItems(['نقد', 'تحويل'])

        account = QLineEdit()
        account.setPlaceholderText('رقم الحساب')

        notes = QLineEdit()
        notes.setPlaceholderText('ملاحظات')

        form.addRow(label + ':', amount)
        form.addRow('التاريخ:', date)
        form.addRow('الطريقة:', method)
        form.addRow('رقم الحساب:', account)
        form.addRow('ملاحظات:', notes)
        layout.addLayout(form)

        btns = QHBoxLayout()
        s = QPushButton('حفظ')
        c = QPushButton('إلغاء')
        s.clicked.connect(d.accept)
        c.clicked.connect(d.reject)
        btns.addWidget(s)
        btns.addWidget(c)
        layout.addLayout(btns)

        return d, amount, method, account, date, notes

    def funder_amount_dialog(self, title, label='المبلغ'):
        d = QDialog(self)
        d.setWindowTitle(title)
        d.resize(560, 340)
        d.setMinimumSize(560, 340)
        apply_theme_to_widget(d)

        layout = QVBoxLayout(d)
        layout.setContentsMargins(18, 18, 18, 18)
        layout.setSpacing(14)

        form = QFormLayout()
        form.setLabelAlignment(Qt.AlignRight)
        form.setFormAlignment(Qt.AlignTop)
        form.setHorizontalSpacing(18)
        form.setVerticalSpacing(12)

        amount = QDoubleSpinBox()
        amount.setRange(0, 1_000_000_000_000)
        amount.setDecimals(0)

        date = QDateEdit()
        fix_date_edit_widget(date)
        date.setDate(QDate.currentDate())
        date.setToolTip('اختر تاريخ العملية الفعلي للممول')

        method = QComboBox()
        method.addItems(['نقد', 'تحويل'])

        account = QLineEdit()
        account.setPlaceholderText('رقم الحساب')

        notes = QLineEdit()
        notes.setPlaceholderText('ملاحظات')

        form.addRow(label + ':', amount)
        form.addRow('التاريخ:', date)
        form.addRow('الطريقة:', method)
        form.addRow('رقم الحساب:', account)
        form.addRow('ملاحظات:', notes)
        layout.addLayout(form)

        btns = QHBoxLayout()
        s = QPushButton('حفظ')
        c = QPushButton('إلغاء')
        s.clicked.connect(d.accept)
        c.clicked.connect(d.reject)
        btns.addWidget(s)
        btns.addWidget(c)
        layout.addLayout(btns)

        return d, amount, method, account, date, notes


    def deposit_capital(self):
        f = self.selected_funder()
        if not f: return
        d, amount, method, account, date, notes = self.funder_amount_dialog('إضافة رصيد للممول')
        if d.exec():
            val = float(amount.value())
            if val <= 0: return QMessageBox.warning(self,'تنبيه','أدخل مبلغ صحيح')
            rec_no = next_receipt_no(self.db, 'next_funder_receive_no')
            self.db['cash'].append({'date': date.date().toString('yyyy-MM-dd'),'type':'إيراد','category':'استلام تمويل','party':f.get('name',''),'amount':val,'notes':notes.text().strip(),'source':'funder_capital_in','payment_method':method.currentText(),'account_no':account.text().strip(),'receipt_no':rec_no,'created_at':now_str()})
            self.last_receipt_path = write_receipt_file(self.db, 'funder_receive', rec_no, 'وصل استلام ممول', date.date().toString('yyyy-MM-dd'), f.get('name',''), val, method.currentText(), account.text().strip(), notes.text().strip(), [('نوع الحركة', 'إضافة رصيد ممول'), ('النسبة بعد الإضافة', fmt_pct(funder_current_ratio_pct(self.db, f.get('name',''))))])
            self.save(); self.refresh_table()

    def withdraw_capital(self):
        f = self.selected_funder()
        if not f: return
        d, amount, method, account, date, notes = self.funder_amount_dialog('سحب رصيد ممول')
        if d.exec():
            val = float(amount.value())
            if val <= 0: return QMessageBox.warning(self,'تنبيه','أدخل مبلغ صحيح')
            if val > funder_effective_capital(self.db, f): return QMessageBox.warning(self,'تنبيه','السحب أكبر من الرصيد الفعال للممول')
            rec_no = next_receipt_no(self.db, 'next_funder_pay_no')
            self.db['cash'].append({'date': date.date().toString('yyyy-MM-dd'),'type':'مصروف','category':'سحب رأس مال ممول','party':f.get('name',''),'amount':val,'notes':notes.text().strip(),'source':'funder_capital_out','payment_method':method.currentText(),'account_no':account.text().strip(),'receipt_no':rec_no,'created_at':now_str()})
            self.last_receipt_path = write_receipt_file(self.db, 'funder_pay', rec_no, 'وصل تسديد ممول', date.date().toString('yyyy-MM-dd'), f.get('name',''), val, method.currentText(), account.text().strip(), notes.text().strip(), [('نوع الحركة', 'سحب من رصيد الممول'), ('النسبة بعد السحب', fmt_pct(funder_current_ratio_pct(self.db, f.get('name',''))))])
            self.save(); self.refresh_table()

    def pay_profit(self):
        f = self.selected_funder()
        if not f: return
        if bool(f.get('is_owner_capital', False)):
            return QMessageBox.information(self, 'تنبيه', 'هذا رأس مال هضاب، وحصته من التمويل تُضاف تلقائياً إلى أرباح هضاب ومصطفى ولا تُصرف هنا كممول مستقل.')
        pb = {x.get('name',''): x for x in profit_breakdown(self.db).get('funders_rows', [])}
        pending = float(pb.get(f.get('name',''), {}).get('pending', 0) or 0)
        d, amount, method, account, date, notes = self.funder_amount_dialog('تسديد مستحق ممول')
        amount.setValue(pending)
        if d.exec():
            val = float(amount.value())
            if val <= 0: return QMessageBox.warning(self,'تنبيه','أدخل مبلغ صحيح')
            if val > pending: return QMessageBox.warning(self,'تنبيه', f'المبلغ أكبر من المستحق ({fmt_money(pending)})')
            rec_no = next_receipt_no(self.db, 'next_funder_pay_profit_no')
            self.db['cash'].append({'date': date.date().toString('yyyy-MM-dd'),'type':'مصروف','category':'تسديد أرباح ممول','party':f.get('name',''),'amount':val,'notes':notes.text().strip(),'source':'funder_profit_payment','payment_method':method.currentText(),'account_no':account.text().strip(),'receipt_no':rec_no,'created_at':now_str()})
            self.last_receipt_path = write_receipt_file(self.db, 'funder_profit_payment', rec_no, 'وصل تسديد أرباح ممول', date.date().toString('yyyy-MM-dd'), f.get('name',''), val, method.currentText(), account.text().strip(), notes.text().strip(), [('نوع الحركة', 'تسديد أرباح ممول')])
            self.save(); self.refresh_table()

    def open_last_receipt(self):
        if self.last_receipt_path and Path(self.last_receipt_path).exists():
            import os
            os.startfile(str(self.last_receipt_path))
        else:
            QMessageBox.information(self, 'تنبيه', 'ماكو وصل منشأ بعد.')



class WarehouseWindow(BaseWindow):
    def __init__(self, main):
        super().__init__(main, '🏬 المخزن')
        self.subtitle_lbl.setText('تم فصل المخزن إلى تبويبين واضحين: ملخص المخزون ثم تفاصيل حركة الصنف، بدون تغيير أي معادلة.')

        self.tabs = style_funders_tabs(QTabWidget())
        self.layout.addWidget(self.tabs, 1)

        self.summary_tab = QWidget()
        self.summary_tab_layout = QVBoxLayout(self.summary_tab)
        self.summary_tab_layout.setContentsMargins(12, 12, 12, 12)
        self.summary_tab_layout.setSpacing(12)

        self.details_tab = QWidget()
        self.details_tab_layout = QVBoxLayout(self.details_tab)
        self.details_tab_layout.setContentsMargins(12, 12, 12, 12)
        self.details_tab_layout.setSpacing(12)

        self.tabs.addTab(self.summary_tab, 'ملخص المخزون')
        self.tabs.addTab(self.details_tab, 'تفاصيل حركة الصنف')

        top = QHBoxLayout()
        self.search = QLineEdit()
        self.search.setPlaceholderText('بحث بالمخزن')
        self.search.textChanged.connect(self.refresh_table)
        self.only_low = QCheckBox('فقط الناقص')
        self.only_low.stateChanged.connect(self.refresh_table)
        top.addWidget(self.search, 1)
        top.addWidget(self.only_low)
        self.summary_tab_layout.addLayout(top)

        self.table = QTableWidget()
        self.table.setStyleSheet(TABLE_STYLE)
        self.table.setColumnCount(9)
        self.table.setHorizontalHeaderLabels(['#','الصنف','الوحدة','المخزون الحالي','حد التنبيه','متوسط الكلفة','البيع','قيمة المخزون','الكود'])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.table.itemSelectionChanged.connect(self.refresh_details)
        self.summary_tab_layout.addWidget(self.table, 1)

        self.summary = QLabel()
        self.summary.setAlignment(Qt.AlignRight)
        self.summary.setWordWrap(True)
        self.summary_tab_layout.addWidget(self.summary)

        hint = QLabel('اختر الصنف من الملخص، وبعدها افتح تبويب تفاصيل الحركة لمشاهدة حركة الداخل والخارج والرصيد.')
        hint.setAlignment(Qt.AlignRight)
        hint.setWordWrap(True)
        hint.setStyleSheet(f'font-size:13px;font-weight:700;color:{MUTED};')
        self.details_tab_layout.addWidget(hint)

        self.details_title = QLabel('تفاصيل حركة الصنف')
        self.details_title.setAlignment(Qt.AlignRight)
        self.details_title.setStyleSheet(f'font-size:22px;font-weight:900;color:{TEXT};')
        self.details_tab_layout.addWidget(self.details_title)

        self.detail_summary = QLabel()
        self.detail_summary.setAlignment(Qt.AlignRight)
        self.detail_summary.setWordWrap(True)
        self.detail_summary.setStyleSheet(f'font-size:13px;font-weight:700;color:{MUTED};padding:4px 2px;')
        self.details_tab_layout.addWidget(self.detail_summary)

        self.details = QTableWidget()
        self.details.setStyleSheet(TABLE_STYLE)
        self.details.setColumnCount(9)
        self.details.setHorizontalHeaderLabels(['#','التاريخ','الحركة','داخل','خارج','الرصيد','كلفة الوحدة','المرجع','ملاحظات'])
        self.details.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.details.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.details.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.details_tab_layout.addWidget(self.details, 1)

        self.refresh_table()

    def filtered_rows(self):
        q = self.search.text().strip().lower()
        low = int(self.db['settings'].get('low_stock_threshold', 5))
        rows = []
        for it in self.db['items']:
            ensure_item_inventory_fields(it)
            if q and q not in it.get('name', '').lower() and q not in it.get('code', '').lower():
                continue
            if self.only_low.isChecked() and int(it.get('qty', 0)) > low:
                continue
            rows.append(it)
        return rows

    def refresh_table(self):
        rows = self.filtered_rows()
        low = int(self.db['settings'].get('low_stock_threshold', 5))
        self.table.setRowCount(len(rows))
        total_qty = 0
        total_val = 0.0
        for r, it in enumerate(rows):
            qty = int(it.get('qty', 0))
            total_qty += qty
            val = item_inventory_value(it)
            total_val += val
            vals = [r+1, it.get('name',''), it.get('unit',''), qty, low, fmt_money(item_avg_cost(it)), fmt_money(it.get('sell_price',0)), fmt_money(val), it.get('code','')]
            for c, v in enumerate(vals):
                self.table.setItem(r, c, QTableWidgetItem(str(v)))
        self.summary.setText(f'أصناف المخزن: {len(rows)} | إجمالي الكمية: {total_qty} | قيمة المخزون: {fmt_money(total_val)} د.ع')
        if rows:
            row = self.table.currentRow()
            if row < 0 or row >= len(rows):
                self.table.selectRow(0)
        else:
            self.details.setRowCount(0)
            self.detail_summary.setText('')
        self.refresh_details()

    def selected_item(self):
        row = self.table.currentRow()
        rows = self.filtered_rows()
        if row < 0 or row >= len(rows):
            return None
        return rows[row]

    def refresh_details(self):
        item = self.selected_item()
        if not item:
            self.details.setRowCount(0)
            self.detail_summary.setText('')
            return
        rows = inventory_movements_for_item(self.db, item.get('name', ''))
        self.details.setRowCount(len(rows))
        balance = 0
        total_in = 0
        total_out = 0
        type_labels = {
            'opening': 'افتتاحي', 'inbound': 'وارد', 'sale': 'بيع', 'return': 'مرتجع صالح', 'damage': 'تالف', 'adjustment': 'تعديل'
        }
        for r, row in enumerate(rows):
            qty_in = int(row.get('qty_in', 0) or 0)
            qty_out = int(row.get('qty_out', 0) or 0)
            balance += qty_in - qty_out
            total_in += qty_in
            total_out += qty_out
            ref = row.get('reference_id', '')
            vals = [
                r+1,
                row.get('date', ''),
                type_labels.get(row.get('movement_type', ''), row.get('movement_type', '')),
                qty_in,
                qty_out,
                balance,
                fmt_money(row.get('unit_cost', 0)),
                ref,
                row.get('notes', '')
            ]
            for c, v in enumerate(vals):
                self.details.setItem(r, c, QTableWidgetItem(str(v)))
        self.detail_summary.setText(
            f"الصنف: {item.get('name','')} | الداخل: {total_in} | الخارج: {total_out} | الرصيد الحالي: {int(item.get('qty',0))} | قيمة الصنف: {fmt_money(item_inventory_value(item))} د.ع"
        )

class InboundWindow(BaseWindow):
    def __init__(self, main):
        super().__init__(main, '📥 الوارد')
        self.subtitle_lbl.setText('واجهة الوارد صارت بنفس منطق الممولين: كل شغلة داخل تبويب مستقل وواضح بدون تكدس.')

        hero = QFrame(); hero.setStyleSheet(CARD_FRAME_STYLE)
        hero_box = QVBoxLayout(hero); hero_box.setContentsMargins(20,18,20,18); hero_box.setSpacing(10)
        hero_title = QLabel('لوحة الوارد')
        hero_title.setAlignment(Qt.AlignRight)
        hero_title.setStyleSheet('font-size:20px;font-weight:900;background:transparent;border:none;')
        hero_note = QLabel('قسّمنا الوارد مثل صفحة الممولين: إدارة الوارد، ملخص الوارد، سجل الحركات، وذمم الموردين. كل جزء بتبويب مستقل وواضح.')
        hero_note.setAlignment(Qt.AlignRight); hero_note.setWordWrap(True)
        hero_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        hero_box.addWidget(hero_title); hero_box.addWidget(hero_note)
        self.layout.addWidget(hero)

        self.tabs = style_funders_tabs(QTabWidget())
        self.layout.addWidget(self.tabs, 1)

        actions_tab = QWidget(); actions_layout = QVBoxLayout(actions_tab); actions_layout.setContentsMargins(10,10,10,10); actions_layout.setSpacing(12)
        form_tab = QWidget(); form_layout = QVBoxLayout(form_tab); form_layout.setContentsMargins(8,8,8,8); form_layout.setSpacing(12)
        top = QHBoxLayout(); top.setSpacing(12)

        self.inbound_calc = QuickCalculatorPanel('حاسبة الوارد', 'ثابتة أعلى اليمين داخل سجل الوارد حتى تبقى قريبة من التفاصيل والحركات المحفوظة.', compact=True)
        input_card = QFrame(); input_card.setStyleSheet(CARD_FRAME_STYLE)
        input_box = QVBoxLayout(input_card); input_box.setContentsMargins(18,18,18,18); input_box.setSpacing(10)
        input_title = QLabel('تسجيل حركة وارد')
        input_title.setAlignment(Qt.AlignRight)
        input_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        input_box.addWidget(input_title)
        form = QFormLayout(); form.setLabelAlignment(Qt.AlignRight); form.setFormAlignment(Qt.AlignRight); form.setHorizontalSpacing(14); form.setVerticalSpacing(12)
        self.date = QDateEdit(); fix_date_edit_widget(self.date); self.date.setDate(QDate.currentDate())
        self.item = QComboBox(); self.supplier = QComboBox(); self.refresh_combos()
        self.qty = tune_numeric_widget(QSpinBox()); self.qty.setRange(1,10_000_000)
        self.cost = tune_numeric_widget(QDoubleSpinBox()); self.cost.setRange(0,1_000_000_000); self.cost.setDecimals(0)
        self.payment = QComboBox(); self.payment.addItems(['نقدي','آجل','جزئي']); self.payment.currentTextChanged.connect(self.sync_paid_limits)
        self.paid = tune_numeric_widget(QDoubleSpinBox()); self.paid.setRange(0,1_000_000_000); self.paid.setDecimals(0)
        for w in (self.date, self.item, self.supplier, self.payment):
            w.setMinimumHeight(46)
            w.setStyleSheet(f'font-size:15px;font-weight:700;padding:8px 10px;border:1px solid {BORDER};border-radius:12px;background:{CARD};')
        for w in (self.qty, self.cost, self.paid):
            w.setMinimumHeight(48)
            w.setStyleSheet(f'font-size:16px;font-weight:800;padding:8px 10px;border:1px solid {BORDER};border-radius:12px;background:{CARD};')
        self.notes = QTextEdit(); self.notes.setFixedHeight(120); self.notes.setPlaceholderText('ملاحظات الحركة'); self.notes.setStyleSheet(f'font-size:15px;font-weight:700;padding:10px;border:1px solid {BORDER};border-radius:12px;background:{CARD};')
        form.addRow('التاريخ:', self.date); form.addRow('الصنف:', self.item); form.addRow('المورد:', self.supplier); form.addRow('الكمية:', self.qty)
        form.addRow('سعر الوحدة:', self.cost); form.addRow('طريقة الدفع:', self.payment); form.addRow('المدفوع:', self.paid)
        input_box.addLayout(form)
        notes_lbl = QLabel('الملاحظات')
        notes_lbl.setAlignment(Qt.AlignRight)
        notes_lbl.setStyleSheet(f'font-size:12px;font-weight:800;color:{MUTED};background:transparent;border:none;')
        input_box.addWidget(notes_lbl); input_box.addWidget(self.notes)
        top.addWidget(input_card, 3)

        actions_card = QFrame(); actions_card.setStyleSheet(CARD_FRAME_STYLE)
        actions_box = QVBoxLayout(actions_card); actions_box.setContentsMargins(18,18,18,18); actions_box.setSpacing(10)
        actions_title = QLabel('إجراءات ومرفقات')
        actions_title.setAlignment(Qt.AlignRight)
        actions_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        actions_box.addWidget(actions_title)
        self.attachment_path = ''
        attach = QPushButton('📎 إرفاق فاتورة'); view_att = QPushButton('👁️ عرض المرفق'); edit_btn = QPushButton('✏️ تعديل الحركة'); s = QPushButton('💾 تسجيل الوارد'); d = QPushButton('🗑 حذف الحركة')
        for b,style in [(attach, SECONDARY_BUTTON),(view_att, SECONDARY_BUTTON),(edit_btn, SECONDARY_BUTTON),(s, BUTTON_STYLE),(d, SECONDARY_BUTTON)]: b.setStyleSheet(style); b.setMinimumHeight(46)
        attach.clicked.connect(self.select_attachment); view_att.clicked.connect(self.view_attachment); edit_btn.clicked.connect(self.edit_row); s.clicked.connect(self.add_row); d.clicked.connect(self.delete_row)
        actions_box.addWidget(attach); actions_box.addWidget(view_att); actions_box.addWidget(edit_btn); actions_box.addWidget(s); actions_box.addWidget(d)
        action_note = QLabel('المدفوع يسجَّل تلقائياً بالصندوق، والباقي يبقى ضمن ذمم الموردين.')
        action_note.setWordWrap(True); action_note.setAlignment(Qt.AlignRight)
        action_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        actions_box.addWidget(action_note)
        top.addWidget(actions_card, 2)
        top.setStretch(0, 1)
        top.setStretch(1, 3)
        top.setStretch(2, 2)
        form_layout.addLayout(top)

        self.summary_card = QFrame(); self.summary_card.setStyleSheet(CARD_FRAME_STYLE)
        summary_box = QVBoxLayout(self.summary_card); summary_box.setContentsMargins(16,14,16,14); summary_box.setSpacing(6)
        summary_title = QLabel('ملخص سريع')
        summary_title.setAlignment(Qt.AlignRight)
        summary_title.setStyleSheet('font-size:15px;font-weight:900;background:transparent;border:none;')
        self.summary = QLabel(); self.summary.setAlignment(Qt.AlignRight); self.summary.setWordWrap(True)
        self.summary.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        summary_box.addWidget(summary_title); summary_box.addWidget(self.summary)
        form_layout.addWidget(self.summary_card)
        form_layout.addStretch(1)

        actions_wrap = QFrame(); actions_wrap.setStyleSheet(CARD_FRAME_STYLE)
        actions_wrap_box = QVBoxLayout(actions_wrap); actions_wrap_box.setContentsMargins(18,18,18,18); actions_wrap_box.setSpacing(10)
        actions_tab_title = QLabel('إدارة الوارد')
        actions_tab_title.setAlignment(Qt.AlignRight)
        actions_tab_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        actions_tab_note = QLabel('هذا التبويب مخصص فقط لإدخال حركة الوارد ومرفقاتها والإجراءات السريعة، بنفس منطق أول تبويب في صفحة الممولين.')
        actions_tab_note.setAlignment(Qt.AlignRight); actions_tab_note.setWordWrap(True)
        actions_tab_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        actions_wrap_box.addWidget(actions_tab_title)
        actions_wrap_box.addWidget(actions_tab_note)
        actions_top_row = QGridLayout()
        actions_top_row.setHorizontalSpacing(12)
        actions_top_row.setVerticalSpacing(12)
        actions_stack = QWidget()
        actions_stack_box = QVBoxLayout(actions_stack)
        actions_stack_box.setContentsMargins(0, 0, 0, 0)
        actions_stack_box.setSpacing(12)
        actions_stack_box.addWidget(input_card)
        actions_stack_box.addWidget(actions_card)
        actions_stack_box.addWidget(self.summary_card)
        actions_top_row.addWidget(actions_stack, 0, 0)
        actions_top_row.setColumnStretch(0, 1)
        actions_wrap_box.addLayout(actions_top_row)
        actions_layout.addWidget(actions_wrap, 1)

        metrics_tab = QWidget(); metrics_layout = QVBoxLayout(metrics_tab); metrics_layout.setContentsMargins(8,8,8,8); metrics_layout.setSpacing(12)
        metrics_wrap = QFrame(); metrics_wrap.setStyleSheet(CARD_FRAME_STYLE)
        metrics_box = QVBoxLayout(metrics_wrap); metrics_box.setContentsMargins(16,16,16,16); metrics_box.setSpacing(12)
        metrics_title = QLabel('مؤشرات الوارد')
        metrics_title.setAlignment(Qt.AlignRight)
        metrics_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        metrics_box.addWidget(metrics_title)
        grid = QGridLayout(); grid.setHorizontalSpacing(10); grid.setVerticalSpacing(10)
        self.inbound_count_card = SummaryCard('عدد الحركات', '0', 'إجمالي قيود الوارد')
        self.inbound_total_card = SummaryCard('إجمالي الوارد', '0', 'القيمة الكلية للوارد')
        self.inbound_due_card = SummaryCard('المتبقي للموردين', '0', 'الذمم المفتوحة من الوارد')
        self.inbound_paid_card = SummaryCard('المدفوع فعلياً', '0', 'ما تم تسجيله بالصندوق')
        for i,c in enumerate((self.inbound_count_card, self.inbound_total_card, self.inbound_due_card, self.inbound_paid_card)):
            c.setMinimumHeight(96); c.setMaximumHeight(112); grid.addWidget(c, i // 2, i % 2)
        metrics_box.addLayout(grid)
        metrics_layout.addWidget(metrics_wrap)
        metrics_layout.addStretch(1)

        table_tab = QWidget(); table_layout = QVBoxLayout(table_tab); table_layout.setContentsMargins(8,8,8,8); table_layout.setSpacing(12)
        table_wrap = QFrame(); table_wrap.setStyleSheet(CARD_FRAME_STYLE)
        table_box = QVBoxLayout(table_wrap); table_box.setContentsMargins(16,16,16,16); table_box.setSpacing(10)
        table_top = QGridLayout(); table_top.setHorizontalSpacing(12); table_top.setVerticalSpacing(12)
        table_title = QLabel('سجل الوارد')
        table_title.setAlignment(Qt.AlignRight)
        table_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        table_note = QLabel('رتبنا الحاسبة بهذا التبويب حتى تبقى قريبة من تفاصيل الحركات المحفوظة بدل صفحة الإدخال.')
        table_note.setAlignment(Qt.AlignRight); table_note.setWordWrap(True)
        table_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        table_title_box = QVBoxLayout(); table_title_box.setContentsMargins(0,0,0,0); table_title_box.setSpacing(4)
        table_title_box.addWidget(table_title)
        table_title_box.addWidget(table_note)
        table_title_wrap = QWidget(); table_title_wrap.setLayout(table_title_box)
        table_top.addWidget(table_title_wrap, 0, 0)
        table_top.addWidget(self.inbound_calc, 0, 1, 1, 1, Qt.AlignTop | Qt.AlignRight)
        table_top.setColumnStretch(0, 1)
        table_box.addLayout(table_top)
        self.table = QTableWidget(); self.table.setStyleSheet(TABLE_STYLE); self.table.setColumnCount(11); self.table.setHorizontalHeaderLabels(['#','التاريخ','الصنف','المورد','الكمية','سعر الوحدة','المجموع','الدفع','المدفوع','الباقي','المرفق'])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch); self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table_box.addWidget(self.table, 1)
        table_layout.addWidget(table_wrap, 1)

        dues_tab = QWidget(); dues_layout = QVBoxLayout(dues_tab); dues_layout.setContentsMargins(8,8,8,8); dues_layout.setSpacing(12)
        dues_wrap = QFrame(); dues_wrap.setStyleSheet(CARD_FRAME_STYLE)
        dues_box = QVBoxLayout(dues_wrap); dues_box.setContentsMargins(16,16,16,16); dues_box.setSpacing(10)
        dues_title = QLabel('ذمم الموردين من الوارد')
        dues_title.setAlignment(Qt.AlignRight)
        dues_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        dues_box.addWidget(dues_title)
        self.dues_label = QLabel(); self.dues_label.setAlignment(Qt.AlignRight); self.dues_label.setWordWrap(True)
        self.dues_label.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        dues_box.addWidget(self.dues_label)
        self.supplier_dues_table = QTableWidget(); self.supplier_dues_table.setStyleSheet(TABLE_STYLE); self.supplier_dues_table.setColumnCount(5)
        self.supplier_dues_table.setHorizontalHeaderLabels(['#','المورد','عدد الحركات المفتوحة','إجمالي الوارد','المتبقي عليه'])
        self.supplier_dues_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.supplier_dues_table.setSelectionBehavior(QAbstractItemView.SelectRows); self.supplier_dues_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        dues_box.addWidget(self.supplier_dues_table, 1)
        dues_layout.addWidget(dues_wrap, 1)

        self.tabs.addTab(metrics_tab, 'ملخص الوارد')
        self.tabs.addTab(actions_tab, 'إدارة الوارد')
        self.tabs.addTab(table_tab, 'سجل الوارد')
        self.tabs.addTab(dues_tab, 'ذمم الموردين')
        self.tabs.setCurrentIndex(0)
        self.sync_paid_limits(); self.refresh_table()
    def select_attachment(self):
        path, _ = QFileDialog.getOpenFileName(self, 'اختيار فاتورة المورد', str(APP_DATA_DIR), 'Images/PDF (*.png *.jpg *.jpeg *.webp *.bmp *.pdf)')
        if path:
            self.attachment_path = path
            QMessageBox.information(self, 'تم', 'تم اختيار المرفق وسيُحفظ مع حركة الوارد.')

    def view_attachment(self):
        r = self.table.currentRow()
        if r < 0:
            return QMessageBox.warning(self, 'تنبيه', 'اختر حركة أولاً')
        rec = self.db['inbound'][r]
        rel = rec.get('attachment', '')
        if not rel:
            return QMessageBox.information(self, 'تنبيه', 'لا يوجد مرفق لهذه الحركة')
        full = resolve_app_file(rel)
        if not full.exists():
            return QMessageBox.warning(self, 'خطأ', 'الملف المرفق غير موجود')
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(full.resolve())))

    def refresh_combos(self):
        if hasattr(self,'item'): self.item.clear(); self.item.addItems([x.get('name','') for x in self.db['items']])
        if hasattr(self,'supplier'): self.supplier.clear(); self.supplier.addItems([x.get('name','') for x in self.db['suppliers']])
    def sync_paid_limits(self):
        total = int(self.qty.value()) * float(self.cost.value())
        self.paid.setMaximum(max(total, 1_000_000_000))
        mode = self.payment.currentText()
        if mode == 'نقدي': self.paid.setValue(total)
        elif mode == 'آجل': self.paid.setValue(0)
        elif self.paid.value() > total: self.paid.setValue(total)
    def refresh_table(self):
        self.refresh_combos(); data = self.db['inbound']; self.table.setRowCount(len(data)); total = due = paid_total = 0
        for r,row in enumerate(data):
            current_due = current_inbound_due(self.db, row)
            total += float(row.get('total',0)); due += current_due; paid_total += float(row.get('total',0) or 0) - current_due
            vals = [r+1,row.get('date',''),row.get('item',''),row.get('supplier',''),row.get('qty',0),fmt_money(row.get('unit_cost',0)),fmt_money(row.get('total',0)),row.get('payment_type',''),fmt_money(float(row.get('total',0) or 0) - current_due),fmt_money(current_due), 'نعم' if row.get('attachment') else '—']
            for c,v in enumerate(vals): self.table.setItem(r,c,QTableWidgetItem(str(v)))
        self.inbound_count_card.set_value(str(len(data)))
        self.inbound_total_card.set_value(fmt_money(total))
        self.inbound_due_card.set_value(fmt_money(due))
        self.inbound_paid_card.set_value(fmt_money(paid_total))
        self.summary.setText(f'حركات الوارد: {len(data)} | إجمالي الوارد: {fmt_money(total)} د.ع | المتبقي للموردين: {fmt_money(due)} د.ع')
        self.refresh_supplier_dues()

    def refresh_supplier_dues(self):
        rows = {}
        for rec in self.db.get('inbound', []):
            supplier = str(rec.get('supplier', '') or '').strip() or '—'
            bucket = rows.setdefault(supplier, {'count': 0, 'total': 0.0, 'due': 0.0})
            bucket['total'] += float(rec.get('total', 0) or 0)
            due_amt = current_inbound_due(self.db, rec)
            if due_amt > 0:
                bucket['count'] += 1
                bucket['due'] += due_amt
        ordered = [(name, vals) for name, vals in rows.items() if vals['total'] or vals['due']]
        ordered.sort(key=lambda x: x[1]['due'], reverse=True)
        self.supplier_dues_table.setRowCount(len(ordered))
        total_due = 0.0
        for r, (name, vals) in enumerate(ordered):
            total_due += vals['due']
            line = [r + 1, name, vals['count'], fmt_money(vals['total']), fmt_money(vals['due'])]
            for c, v in enumerate(line):
                self.supplier_dues_table.setItem(r, c, QTableWidgetItem(str(v)))
        self.dues_label.setText(
            f'عدد الموردين المرتبطين بالوارد: {len(ordered)} | إجمالي الذمم المفتوحة: {fmt_money(total_due)} د.ع | '
            f'هذا التبويب مخصص فقط لتفصيل المتبقي لكل مورد بشكل مستقل مثل نسق صفحة الممولين.'
        )

    def add_row(self):
        if not self.db['items']: return QMessageBox.warning(self,'تنبيه','أضف صنف أولاً')
        if not self.db['suppliers']: return QMessageBox.warning(self,'تنبيه','أضف مورد أولاً')
        item_name = self.item.currentText().strip(); supplier = self.supplier.currentText().strip(); qty = int(self.qty.value()); cost = float(self.cost.value())
        if not item_name or not supplier or qty <= 0: return QMessageBox.warning(self,'تنبيه','أكمل البيانات')
        if cost <= 0: return QMessageBox.warning(self,'تنبيه','سعر الشراء يجب أن يكون أكبر من صفر')
        total = qty * cost
        if total <= 0: return QMessageBox.warning(self,'تنبيه','إجمالي الوارد يجب أن يكون أكبر من صفر')
        paid = min(total, max(0.0, float(self.paid.value())))
        mode = self.payment.currentText()
        if mode == 'نقدي': paid = total
        elif mode == 'آجل': paid = 0.0
        elif paid <= 0 or paid >= total: return QMessageBox.warning(self,'تنبيه','في الدفع الجزئي أدخل مدفوع بين صفر والمجموع')
        rec_id = generate_id('inb')
        attachment_rel = ''
        if self.attachment_path:
            ATTACHMENTS_DIR.mkdir(exist_ok=True)
            src_path = Path(self.attachment_path)
            if src_path.exists():
                target = ATTACHMENTS_DIR / f'{rec_id}{src_path.suffix.lower()}'
                shutil.copy2(src_path, target)
                attachment_rel = app_relative_path(target)
        self.db['inbound'].append({'id': rec_id, 'date': self.date.date().toString('yyyy-MM-dd'),'item': item_name,'supplier': supplier,'qty': qty,'unit_cost': cost,'total': total,'payment_type': mode,'paid_amount': paid,'due_amount': total-paid,'notes': self.notes.toPlainText().strip(),'attachment': attachment_rel,'created_at': now_str()})
        for it in self.db['items']:
            if it.get('name') == item_name:
                inventory_add_stock(it, qty, cost)
                add_inventory_movement(self.db, it, 'inbound', qty_in=qty, unit_cost=cost, reference_type='inbound', reference_id=rec_id, date=self.date.date().toString('yyyy-MM-dd'), notes=f"وارد من {supplier}", movement_uid=f"inbound_{rec_id}_{it.get('name','')}")
                break
        if paid > 0:
            self.db['cash'].append({'date': self.date.date().toString('yyyy-MM-dd'),'type':'مصروف','category':'وارد بضاعة','party':supplier,'amount':paid,'notes':f'شراء {qty} من {item_name}','source':'inbound','ref_id': rec_id, 'created_at': now_str()})
        self.qty.setValue(1); self.cost.setValue(0); self.paid.setValue(0); self.notes.clear(); self.attachment_path = ''; self.save(); self.refresh_table()
    def edit_row(self):
        r = self.table.currentRow()
        if r < 0:
            return QMessageBox.warning(self, 'تنبيه', 'اختر حركة وارد من السجل أولاً')
        rows = self.db.get('inbound', [])
        if r >= len(rows):
            return QMessageBox.warning(self, 'تنبيه', 'تعذر العثور على حركة الوارد المحددة')
        row = rows[r]
        if current_inbound_due(self.db, row) < max(0.0, float(row.get('total', 0) or 0) - float(row.get('paid_amount', 0) or 0)):
            return QMessageBox.warning(self, 'منع', 'لا يمكن تعديل حركة الوارد بعد وجود تسديدات لاحقة مرتبطة بها.')
        item_name = self.item.currentText().strip()
        supplier = self.supplier.currentText().strip()
        qty = int(self.qty.value())
        cost = float(self.cost.value())
        if not item_name or not supplier or qty <= 0:
            return QMessageBox.warning(self, 'تنبيه', 'أكمل بيانات الوارد المعدّل')
        if cost <= 0:
            return QMessageBox.warning(self, 'تنبيه', 'سعر الشراء يجب أن يكون أكبر من صفر')
        total = qty * cost
        if total <= 0:
            return QMessageBox.warning(self, 'تنبيه', 'إجمالي الوارد يجب أن يكون أكبر من صفر')
        paid = min(total, max(0.0, float(self.paid.value())))
        mode = self.payment.currentText()
        if mode == 'نقدي':
            paid = total
        elif mode == 'آجل':
            paid = 0.0
        elif paid <= 0 or paid >= total:
            return QMessageBox.warning(self, 'تنبيه', 'في الدفع الجزئي أدخل مدفوع بين صفر والمجموع')
        old_item = next((it for it in self.db['items'] if it.get('name') == row.get('item')), None)
        if old_item:
            if int(row.get('qty', 0) or 0) > int(old_item.get('qty', 0) or 0):
                return QMessageBox.warning(self, 'منع', 'لا يمكن تعديل الوارد لأن جزءاً من الكمية تم تصريفه من المخزن')
            inventory_remove_stock(old_item, int(row.get('qty', 0) or 0), float(row.get('unit_cost', 0) or 0))
            remove_inventory_movement(self.db, f"inbound_{row.get('id')}_{old_item.get('name','')}")
        new_item = next((it for it in self.db['items'] if it.get('name') == item_name), None)
        if not new_item:
            return QMessageBox.warning(self, 'تنبيه', 'الصنف المحدد غير موجود')
        inventory_add_stock(new_item, qty, cost)
        add_inventory_movement(self.db, new_item, 'inbound', qty_in=qty, unit_cost=cost, reference_type='inbound', reference_id=row.get('id'), date=self.date.date().toString('yyyy-MM-dd'), notes=f"وارد من {supplier}", movement_uid=f"inbound_{row.get('id')}_{new_item.get('name','')}")
        row.update({
            'date': self.date.date().toString('yyyy-MM-dd'),
            'item': item_name,
            'supplier': supplier,
            'qty': qty,
            'unit_cost': cost,
            'total': total,
            'payment_type': mode,
            'paid_amount': paid,
            'due_amount': round(total - paid, 2),
            'notes': self.notes.toPlainText().strip(),
            'updated_at': now_str(),
        })
        for cash_row in self.db.get('cash', []):
            if cash_row.get('source') == 'inbound' and cash_row.get('ref_id') == row.get('id'):
                if paid > 0:
                    cash_row.update({'date': row['date'], 'type': 'مصروف', 'category': 'وارد بضاعة', 'party': supplier, 'amount': paid, 'notes': f'شراء {qty} من {item_name}', 'updated_at': now_str()})
                else:
                    self.db['cash'].remove(cash_row)
                break
        else:
            if paid > 0:
                self.db['cash'].append({'date': row['date'], 'type': 'مصروف', 'category': 'وارد بضاعة', 'party': supplier, 'amount': paid, 'notes': f'شراء {qty} من {item_name}', 'source': 'inbound', 'ref_id': row.get('id'), 'created_at': now_str()})
        self.save(); self.refresh_table(); self.main.refresh_dashboard()
        QMessageBox.information(self, 'تم', 'تم تعديل حركة الوارد بنجاح.')

    def delete_row(self):
        r = self.table.currentRow();
        if r < 0: return QMessageBox.warning(self,'تنبيه','اختر حركة')
        rec = self.db['inbound'][r]
        linked_followup = 0.0
        for ev in inbound_payment_allocations(self.db, rec.get('supplier', '')):
            linked_followup += sum(_safe_float(a.get('amount', 0)) for a in ev.get('allocations', []) if a.get('inbound_id') == rec.get('id'))
        if linked_followup > 0:
            return QMessageBox.warning(self, 'منع', 'لا يمكن حذف حركة الوارد لأن عليها تسديدات لاحقة مرتبطة. احذف أو عدّل التسديدات أولاً.')
        if QMessageBox.question(self,'تأكيد','حذف الحركة وعكس أثرها؟') != QMessageBox.Yes: return
        for it in self.db['items']:
            if it.get('name') == rec.get('item'):
                if int(rec.get('qty',0)) > int(it.get('qty',0)):
                    return QMessageBox.warning(self,'منع','لا يمكن حذف الوارد لأن جزءاً من الكمية تم تصريفه من المخزن')
                inventory_remove_stock(it, int(rec.get('qty',0)), float(rec.get('unit_cost',0) or 0))
                remove_inventory_movement(self.db, f"inbound_{rec.get('id')}_{it.get('name','')}")
                break
        for i in range(len(self.db['cash'])-1, -1, -1):
            c = self.db['cash'][i]
            if c.get('ref_id') == rec.get('id') and c.get('source') == 'inbound':
                self.db['cash'].pop(i)
        self.db['inbound'].pop(r); self.save(); self.refresh_table()

class SalesWindow(BaseWindow):
    def __init__(self, main):
        super().__init__(main, '💰 المبيعات')
        self.subtitle_lbl.setText('واجهة بيع واسعة ومريحة: كل تبويب صفحة كاملة تقريباً، حقول أكبر، ورجوع سريع للرئيسية بدون تغيير أي معادلة.')
        self.cart = []
        self.editing_group_id = None
        self.editing_invoice_no = None

        self.sales_stack = QStackedWidget()
        self.layout.addWidget(self.sales_stack, 1)

        self.sales_home = QWidget(); self.sales_home_layout = QVBoxLayout(self.sales_home); self.sales_home_layout.setContentsMargins(18,18,18,18); self.sales_home_layout.setSpacing(16)
        self.invoice_entry_tab = QWidget(); self.invoice_entry_layout = QVBoxLayout(self.invoice_entry_tab); self.invoice_entry_layout.setContentsMargins(18,18,18,18); self.invoice_entry_layout.setSpacing(14)
        self.invoices_tab = QWidget(); self.invoices_tab_layout = QVBoxLayout(self.invoices_tab); self.invoices_tab_layout.setContentsMargins(10,10,10,10); self.invoices_tab_layout.setSpacing(12)

        self.sales_stack.addWidget(self.sales_home)
        self.sales_stack.addWidget(self.invoice_entry_tab)
        self.sales_stack.addWidget(self.invoices_tab)

        self._build_sales_home()

        self.date = QDateEdit(); fix_date_edit_widget(self.date); self.date.setDate(QDate.currentDate())
        self.item = QComboBox(); self.customer = QComboBox(); self.refresh_combos(); self.item.currentTextChanged.connect(self.sync_price)
        self.qty = tune_numeric_widget(QSpinBox()); self.qty.setRange(1,10_000_000)
        self.price = QDoubleSpinBox(); self.price.setRange(0,1_000_000_000); self.price.setDecimals(0)
        self.payment = QComboBox(); self.payment.addItems(['نقدي','آجل','جزئي']); self.payment.currentTextChanged.connect(self.sync_paid_limits)
        self.paid = tune_numeric_widget(QDoubleSpinBox()); self.paid.setRange(0,1_000_000_000); self.paid.setDecimals(0)
        self.notes = QLineEdit(); self.notes.setPlaceholderText('ملاحظات الفاتورة')
        self.qty.valueChanged.connect(self.sync_paid_limits); self.price.valueChanged.connect(self.sync_paid_limits)
        self.sync_price()
        self._enlarge_sales_widgets()

        entry_scroll = QScrollArea(); entry_scroll.setWidgetResizable(True)
        entry_scroll.setStyleSheet("QScrollArea{background:transparent;border:none;} QScrollArea > QWidget > QWidget{background:transparent;}")
        entry_container = QWidget()
        entry_main = QVBoxLayout(entry_container)
        entry_main.setContentsMargins(0, 0, 0, 28)
        entry_main.setSpacing(16)

        entry_header = self._build_sales_page_card('تسجيل الفاتورة', 'رتبنا الإدخال بنفس روح صفحة إدخال المبيعات: بيانات الفاتورة أولاً، ثم الصنف، ثم السلة، ثم إجراءات الحفظ بشكل أنظف.')
        entry_main.addWidget(entry_header)

        top_row = QGridLayout()
        top_row.setHorizontalSpacing(12)
        top_row.setVerticalSpacing(12)
        self.sales_calc = QuickCalculatorPanel('حاسبة المبيعات', 'ثابتة أعلى اليمين داخل الفواتير المحفوظة حتى تبقى قريبة من المراجعة والتدقيق.', compact=True)

        form_card = QFrame()
        form_card.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.82)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:24px;")
        form_layout = QVBoxLayout(form_card)
        form_layout.setContentsMargins(20, 20, 20, 20)
        form_layout.setSpacing(14)
        form_title = QLabel('بيانات الفاتورة')
        form_title.setAlignment(Qt.AlignRight)
        form_title.setStyleSheet(f'font-size:20px;font-weight:900;color:{TEXT};')
        form_layout.addWidget(form_title)
        top_grid = QGridLayout(); top_grid.setHorizontalSpacing(14); top_grid.setVerticalSpacing(14)
        top_grid.setColumnStretch(1, 1); top_grid.setColumnStretch(3, 1)
        top_grid.addWidget(QLabel('التاريخ'), 0, 0)
        top_grid.addWidget(self.date, 0, 1)
        top_grid.addWidget(QLabel('الزبون'), 0, 2)
        top_grid.addWidget(self.customer, 0, 3)
        top_grid.addWidget(QLabel('نوع الدفع'), 1, 0)
        top_grid.addWidget(self.payment, 1, 1)
        top_grid.addWidget(QLabel('المقبوض'), 1, 2)
        top_grid.addWidget(self.paid, 1, 3)
        form_layout.addLayout(top_grid)
        notes_label = QLabel('ملاحظات الفاتورة')
        notes_label.setAlignment(Qt.AlignRight)
        notes_label.setStyleSheet(f'font-size:14px;font-weight:800;color:{MUTED};')
        form_layout.addWidget(notes_label)
        self.notes.setMinimumHeight(64)
        form_layout.addWidget(self.notes)
        top_row.addWidget(form_card, 0, 0)
        top_row.setColumnStretch(0, 1)
        entry_main.addLayout(top_row)

        item_card = QFrame()
        item_card.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.82)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:24px;")
        item_layout = QVBoxLayout(item_card)
        item_layout.setContentsMargins(20, 20, 20, 20)
        item_layout.setSpacing(14)
        item_title = QLabel('إضافة صنف إلى الفاتورة')
        item_title.setAlignment(Qt.AlignRight)
        item_title.setStyleSheet(f'font-size:20px;font-weight:900;color:{TEXT};')
        item_layout.addWidget(item_title)
        item_grid = QGridLayout(); item_grid.setHorizontalSpacing(14); item_grid.setVerticalSpacing(14)
        item_grid.setColumnStretch(1, 1); item_grid.setColumnStretch(3, 1)
        item_grid.addWidget(QLabel('الصنف'), 0, 0)
        item_grid.addWidget(self.item, 0, 1, 1, 3)
        item_grid.addWidget(QLabel('الكمية'), 1, 0)
        item_grid.addWidget(self.qty, 1, 1)
        item_grid.addWidget(QLabel('سعر الوحدة'), 1, 2)
        item_grid.addWidget(self.price, 1, 3)
        item_layout.addLayout(item_grid)
        item_buttons = QHBoxLayout(); item_buttons.setSpacing(12)
        add_line = QPushButton('➕ إضافة للفاتورة'); add_line.setStyleSheet(BUTTON_STYLE); add_line.clicked.connect(self.add_to_cart); add_line.setMinimumHeight(56)
        remove_line = QPushButton('➖ حذف السطر المحدد'); remove_line.setStyleSheet(SECONDARY_BUTTON); remove_line.clicked.connect(self.remove_cart_line); remove_line.setMinimumHeight(56)
        item_buttons.addWidget(add_line)
        item_buttons.addWidget(remove_line)
        item_buttons.addStretch(1)
        item_layout.addLayout(item_buttons)
        entry_main.addWidget(item_card)

        cart_card = QFrame()
        cart_card.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.82)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:24px;")
        cart_layout = QVBoxLayout(cart_card)
        cart_layout.setContentsMargins(20, 20, 20, 20)
        cart_layout.setSpacing(14)
        cart_title = QLabel('سلة الفاتورة الحالية')
        cart_title.setAlignment(Qt.AlignRight)
        cart_title.setStyleSheet(f'font-size:20px;font-weight:900;color:{TEXT};')
        cart_layout.addWidget(cart_title)
        self.cart_table = QTableWidget(); self.cart_table.setStyleSheet(TABLE_STYLE); self.cart_table.setColumnCount(5)
        self.cart_table.setHorizontalHeaderLabels(['#','الصنف','الكمية','سعر الوحدة','المجموع'])
        self.cart_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.cart_table.setSelectionBehavior(QAbstractItemView.SelectRows); self.cart_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.cart_table.itemSelectionChanged.connect(self.load_selected_cart_line)
        self.cart_table.setMinimumHeight(300)
        cart_layout.addWidget(self.cart_table)
        self.cart_summary = QLabel(); self.cart_summary.setAlignment(Qt.AlignRight)
        self.cart_summary.setStyleSheet(f'font-size:15px;font-weight:800;color:{MUTED};padding:6px 0;')
        cart_layout.addWidget(self.cart_summary)
        entry_main.addWidget(cart_card)

        actions_card = QFrame()
        actions_card.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.82)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:24px;")
        actions_layout = QVBoxLayout(actions_card)
        actions_layout.setContentsMargins(20, 20, 20, 20)
        actions_layout.setSpacing(14)
        actions_title = QLabel('عمليات الفاتورة')
        actions_title.setAlignment(Qt.AlignRight)
        actions_title.setStyleSheet(f'font-size:20px;font-weight:900;color:{TEXT};')
        actions_layout.addWidget(actions_title)
        actions_note = QLabel('هذا التبويب مخصص لتسجيل الفاتورة الحالية فقط. التعديل والحذف نقلناهن إلى تبويب الفواتير المحفوظة حتى يبقى الإدخال أنظف وأوضح.')
        actions_note.setWordWrap(True)
        actions_note.setAlignment(Qt.AlignRight)
        actions_note.setStyleSheet(f'font-size:14px;font-weight:700;color:{MUTED};')
        actions_layout.addWidget(actions_note)
        action_buttons = QGridLayout(); action_buttons.setHorizontalSpacing(12); action_buttons.setVerticalSpacing(12)
        save_btn = QPushButton('💾 تسجيل الفاتورة'); save_btn.setStyleSheet(BUTTON_STYLE); save_btn.setMinimumHeight(58); save_btn.clicked.connect(self.finalize_sale)
        open_btn = QPushButton('🧾 فتح الفاتورة'); open_btn.setStyleSheet(SECONDARY_BUTTON); open_btn.setMinimumHeight(58); open_btn.clicked.connect(self.open_invoice)
        print_btn = QPushButton('🖨️ فتح للطباعة'); print_btn.setStyleSheet(SECONDARY_BUTTON); print_btn.setMinimumHeight(58); print_btn.clicked.connect(self.print_invoice)
        clear_btn = QPushButton('🧹 تفريغ الفاتورة'); clear_btn.setStyleSheet(SECONDARY_BUTTON); clear_btn.setMinimumHeight(58); clear_btn.clicked.connect(self.clear_cart)
        action_buttons.addWidget(save_btn, 0, 0, 1, 2)
        action_buttons.addWidget(open_btn, 1, 0)
        action_buttons.addWidget(print_btn, 1, 1)
        action_buttons.addWidget(clear_btn, 2, 0, 1, 2)
        actions_layout.addLayout(action_buttons)
        entry_main.addWidget(actions_card)
        entry_main.addSpacing(32)

        entry_scroll.setWidget(entry_container)
        entry_scroll.setViewportMargins(0, 0, 0, 12)
        self.invoice_entry_layout.addWidget(entry_scroll, 1)


        invoices_card = self._build_sales_page_card('أرشيف الفواتير', 'واجهة الأرشيف صارت منفصلة ومريحة حتى تختار الفاتورة بسهولة.')
        invoices_layout = invoices_card.layout()
        invoices_top = QGridLayout(); invoices_top.setHorizontalSpacing(12); invoices_top.setVerticalSpacing(12)
        invoices_title_box = QVBoxLayout(); invoices_title_box.setContentsMargins(0,0,0,0); invoices_title_box.setSpacing(4)
        invoices_title = QLabel('الفواتير المحفوظة')
        invoices_title.setAlignment(Qt.AlignRight)
        invoices_title.setStyleSheet(f'font-size:20px;font-weight:900;color:{TEXT};')
        invoices_note = QLabel('نقلنا الحاسبة لهذا التبويب حتى تبقى قريبة من الأرشيف والتفاصيل بدل صفحة الإدخال.')
        invoices_note.setAlignment(Qt.AlignRight); invoices_note.setWordWrap(True)
        invoices_note.setStyleSheet(f'font-size:13px;font-weight:700;color:{MUTED};')
        invoices_title_box.addWidget(invoices_title)
        invoices_title_box.addWidget(invoices_note)
        invoices_title_wrap = QWidget(); invoices_title_wrap.setLayout(invoices_title_box)
        invoices_top.addWidget(invoices_title_wrap, 0, 0)
        invoices_top.addWidget(self.sales_calc, 0, 1, 1, 1, Qt.AlignTop | Qt.AlignRight)
        invoices_top.setColumnStretch(0, 1)
        invoices_layout.addLayout(invoices_top)
        self.table = QTableWidget(); self.table.setStyleSheet(TABLE_STYLE); self.table.setColumnCount(11)
        self.table.setHorizontalHeaderLabels(['#','رقم الفاتورة','التاريخ','الأصناف','الزبون','الكمية','المجموع','الدفع','المقبوض','الباقي','الأسطر'])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch); self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.table.setMinimumHeight(500)
        invoices_layout.addWidget(self.table, 1)

        archive_actions = QHBoxLayout(); archive_actions.setSpacing(12)
        archive_actions.addStretch(1)
        open_saved_btn = QPushButton('🧾 فتح الفاتورة')
        open_saved_btn.setStyleSheet(BUTTON_STYLE)
        open_saved_btn.setMinimumHeight(52)
        open_saved_btn.clicked.connect(self.open_invoice)
        archive_actions.addWidget(open_saved_btn)
        print_saved_btn = QPushButton('🖨️ فتح للطباعة')
        print_saved_btn.setStyleSheet(SECONDARY_BUTTON)
        print_saved_btn.setMinimumHeight(52)
        print_saved_btn.clicked.connect(self.print_invoice)
        archive_actions.addWidget(print_saved_btn)
        edit_saved_btn = QPushButton('✏️ تعديل الفاتورة')
        edit_saved_btn.setStyleSheet(SECONDARY_BUTTON)
        edit_saved_btn.setMinimumHeight(52)
        edit_saved_btn.clicked.connect(self.edit_invoice)
        archive_actions.addWidget(edit_saved_btn)
        delete_saved_btn = QPushButton('🗑 حذف الفاتورة')
        delete_saved_btn.setStyleSheet(SECONDARY_BUTTON)
        delete_saved_btn.setMinimumHeight(52)
        delete_saved_btn.clicked.connect(self.delete_row)
        archive_actions.addWidget(delete_saved_btn)
        invoices_layout.addLayout(archive_actions)

        self.summary = QLabel(); self.summary.setAlignment(Qt.AlignRight); self.summary.setStyleSheet(f'font-size:15px;font-weight:800;color:{MUTED};padding:6px 0;')
        invoices_layout.addWidget(self.summary)
        self.invoices_tab_layout.addWidget(invoices_card, 1)

        self.sync_paid_limits(); self.refresh_cart_table(); self.refresh_table(); self.back_to_sales_home()

    def _back_to_main(self):
        self.close()

    def open_sales_page(self, page):
        self.sales_stack.setCurrentWidget(page)

    def back_to_sales_home(self):
        self.sales_stack.setCurrentWidget(self.sales_home)

    def _build_sales_home(self):
        hero = QFrame()
        hero.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.82)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:28px;")
        hero_layout = QVBoxLayout(hero)
        hero_layout.setContentsMargins(26, 24, 26, 24)
        hero_layout.setSpacing(10)
        title = QLabel('المبيعات')
        title.setAlignment(Qt.AlignRight)
        title.setStyleSheet(f'font-size:28px;font-weight:900;color:{TEXT};')
        subtitle = QLabel('اختَر الصفحة المطلوبة داخل المبيعات. كل شاشة مستقلة وبيها رجوع واضح.')
        subtitle.setWordWrap(True)
        subtitle.setAlignment(Qt.AlignRight)
        subtitle.setStyleSheet(f'font-size:14px;font-weight:700;color:{MUTED};')
        hero_layout.addWidget(title)
        hero_layout.addWidget(subtitle)
        self.sales_home_layout.addWidget(hero)

        grid_wrap = QFrame()
        grid_wrap.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.60)};border:1px solid {rgba_from_hex(TEXT,0.06)};border-radius:28px;")
        grid = QGridLayout(grid_wrap)
        grid.setContentsMargins(18, 18, 18, 18)
        grid.setHorizontalSpacing(14)
        grid.setVerticalSpacing(14)

        actions = [
            ('💾 تسجيل فاتورة', 'الاسم والصنف وكل الإدخالات بصفحة وحدة مع أزرار الحفظ.', self.invoice_entry_tab),
            ('🗂 الفواتير المحفوظة', 'افتح أرشيف الفواتير المحفوظة بصفحة مستقلة.', self.invoices_tab),
        ]

        for i, (label, desc, page) in enumerate(actions):
            btn = QPushButton(f'{label}\n{desc}')
            btn.setCursor(Qt.PointingHandCursor)
            btn.setMinimumHeight(108)
            btn.setStyleSheet(f"QPushButton{{text-align:right;padding:18px 22px;border-radius:24px;background-color:{rgba_from_hex(CARD,0.96)};border:1px solid {rgba_from_hex(TEXT,0.08)};font-size:18px;font-weight:800;color:{TEXT};}}QPushButton:hover{{border:1px solid {rgba_from_hex(ACCENT,0.45)};background-color:{rgba_from_hex('#ffffff',0.98)};}}")
            btn.clicked.connect(lambda _=False, target=page: self.open_sales_page(target))
            grid.addWidget(btn, i // 2, i % 2)

        self.sales_home_layout.addWidget(grid_wrap, 1)

        footer = QHBoxLayout()
        footer.addStretch(1)
        main_back = QPushButton('↩ الرجوع للرئيسية العامة')
        main_back.setMinimumHeight(50)
        main_back.setStyleSheet(SECONDARY_BUTTON)
        main_back.clicked.connect(self._back_to_main)
        footer.addWidget(main_back)
        self.sales_home_layout.addLayout(footer)

    def _build_sales_page_card(self, title, subtitle):
        card = QFrame()
        card.setStyleSheet(f"background-color:{rgba_from_hex(CARD,0.82)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:28px;")
        lay = QVBoxLayout(card)
        lay.setContentsMargins(24, 22, 24, 22)
        lay.setSpacing(16)
        top = QHBoxLayout(); top.setSpacing(12)
        back_btn = QPushButton('↩ الرئيسية')
        back_btn.setStyleSheet(SECONDARY_BUTTON)
        back_btn.setMinimumHeight(48)
        back_btn.clicked.connect(self.back_to_sales_home)
        title_box = QVBoxLayout(); title_box.setSpacing(4)
        title_lbl = QLabel(title); title_lbl.setAlignment(Qt.AlignRight); title_lbl.setStyleSheet(f'font-size:24px;font-weight:900;color:{TEXT};')
        sub_lbl = QLabel(subtitle); sub_lbl.setAlignment(Qt.AlignRight); sub_lbl.setWordWrap(True); sub_lbl.setStyleSheet(f'font-size:13px;font-weight:700;color:{MUTED};')
        title_box.addWidget(title_lbl); title_box.addWidget(sub_lbl)
        top.addWidget(back_btn, alignment=Qt.AlignLeft)
        top.addLayout(title_box, 1)
        lay.addLayout(top)
        return card

    def _enlarge_sales_widgets(self):
        widgets = [self.date, self.item, self.customer, self.qty, self.price, self.payment, self.paid, self.notes]
        style = f"""
        QLineEdit, QComboBox, QDateEdit, QSpinBox, QDoubleSpinBox {{
            min-height: 54px;
            padding: 10px 16px;
            padding-right: 14px;
            font-size: 18px;
            font-weight: 700;
            border-radius: 16px;
        }}
        QComboBox::drop-down, QDateEdit::drop-down {{
            width: 34px;
            subcontrol-origin: padding;
            subcontrol-position: center left;
            border: none;
            margin-left: 8px;
        }}
        QComboBox QAbstractItemView {{ font-size: 16px; }}
        """
        for w in widgets:
            try:
                w.setMinimumHeight(54)
                w.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
                if isinstance(w, QLineEdit):
                    w.setClearButtonEnabled(True)
                if isinstance(w, (QComboBox, QDateEdit)):
                    w.setLayoutDirection(Qt.LeftToRight)
                    try:
                        w.lineEdit().setAlignment(Qt.AlignRight)
                    except Exception:
                        pass
                w.setStyleSheet(style)
            except Exception:
                pass

    def select_attachment(self):
        path, _ = QFileDialog.getOpenFileName(self, 'اختيار فاتورة المورد', str(APP_DATA_DIR), 'Images/PDF (*.png *.jpg *.jpeg *.webp *.bmp *.pdf)')
        if path:
            self.attachment_path = path
            QMessageBox.information(self, 'تم', 'تم اختيار المرفق وسيُحفظ مع حركة الوارد.')

    def view_attachment(self):
        r = self.table.currentRow()
        if r < 0:
            return QMessageBox.warning(self, 'تنبيه', 'اختر حركة أولاً')
        rec = self.db['inbound'][r]
        rel = rec.get('attachment', '')
        if not rel:
            return QMessageBox.information(self, 'تنبيه', 'لا يوجد مرفق لهذه الحركة')
        full = resolve_app_file(rel)
        if not full.exists():
            return QMessageBox.warning(self, 'خطأ', 'الملف المرفق غير موجود')
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(full.resolve())))

    def refresh_combos(self):
        if hasattr(self,'item'):
            self.item.clear(); self.item.addItems([x.get('name','') for x in self.db['items']])
        if hasattr(self,'customer'):
            self.customer.clear(); self.customer.addItems([x.get('name','') for x in self.db['customers']])

    def sync_price(self):
        name = self.item.currentText().strip()
        for it in self.db['items']:
            if it.get('name') == name:
                self.price.setValue(float(it.get('sell_price',0))); self.sync_paid_limits(); return
        self.price.setValue(0); self.sync_paid_limits()

    def cart_total(self):
        return sum(float(x.get('total',0)) for x in self.cart)

    def sync_paid_limits(self):
        total = self.cart_total()
        self.paid.setMaximum(max(total, 1_000_000_000))
        mode = self.payment.currentText()
        if mode == 'نقدي':
            self.paid.setValue(total)
        elif mode == 'آجل':
            self.paid.setValue(0)
        elif self.paid.value() > total:
            self.paid.setValue(total)

    def get_item(self, name):
        for it in self.db['items']:
            if it.get('name') == name: return it

    def reserved_qty(self, item_name):
        return sum(int(x.get('qty',0)) for x in self.cart if x.get('item') == item_name)


    def add_to_cart(self):
        if not self.db['items']:
            return QMessageBox.warning(self,'تنبيه','أضف صنف أولاً')
        if not self.db['customers']:
            return QMessageBox.warning(self,'تنبيه','أضف زبون أولاً')
        item_name = self.item.currentText().strip()
        qty = int(self.qty.value())
        unit_price = float(self.price.value())
        item = self.get_item(item_name)
        if not item:
            return QMessageBox.warning(self,'تنبيه','الصنف غير موجود')
        if qty <= 0:
            return QMessageBox.warning(self,'تنبيه','الكمية يجب أن تكون أكبر من صفر')
        if unit_price <= 0:
            return QMessageBox.warning(self,'تنبيه','سعر البيع يجب أن يكون أكبر من صفر')
        stock = int(item.get('qty',0))
        available = stock - self.reserved_qty(item_name)
        if qty > available:
            return QMessageBox.warning(self,'تنبيه', f'الكمية أكبر من المتاح بعد الحجز ({available})')
        avg_cost = item_avg_cost(item)
        line_total = qty * unit_price
        if line_total <= 0:
            return QMessageBox.warning(self,'تنبيه','إجمالي السطر يجب أن يكون أكبر من صفر')
        line = {'item': item_name, 'qty': qty, 'unit_price': unit_price, 'buy_price': avg_cost, 'total': line_total}
        self.cart.append(line)
        self.qty.setValue(1)
        self.refresh_cart_table()
        self.open_sales_page(self.invoice_entry_tab)
    def remove_cart_line(self):
        r = self.cart_table.currentRow()
        if r < 0 or r >= len(self.cart):
            return QMessageBox.warning(self,'تنبيه','اختر سطر من الفاتورة')
        self.cart.pop(r)
        self.refresh_cart_table()

    def load_selected_cart_line(self):
        r = self.cart_table.currentRow()
        if r < 0 or r >= len(self.cart):
            return
        line = self.cart[r]
        idx = self.item.findText(line.get('item',''))
        if idx >= 0:
            self.item.setCurrentIndex(idx)
        self.qty.setValue(int(line.get('qty',0) or 0))
        self.price.setValue(float(line.get('unit_price',0) or 0))

    def clear_cart(self):
        self.cart = []
        self.notes.clear()
        self.editing_group_id = None
        self.editing_invoice_no = None
        self.refresh_cart_table()

    def edit_invoice(self):
        sale = self.selected_sale()
        if not sale:
            return
        lines = get_invoice_sales(self.db, sale)
        sale_ids = {x.get('id') for x in lines}
        gid = sale.get('invoice_group_id') or f"single-{ensure_invoice_no(self.db, sale)}"
        linked_followup = 0.0
        for ev in customer_payment_allocations(self.db, sale.get('customer', '')):
            linked_followup += sum(_safe_float(a.get('amount', 0)) for a in ev.get('allocations', []) if a.get('sale_id') in sale_ids or a.get('invoice_group_id') == gid)
        if linked_followup > 0:
            return QMessageBox.warning(self, 'منع', 'لا يمكن تعديل الفاتورة بعد وجود تسديدات لاحقة مرتبطة بها. عدّل أو احذف التسديدات أولاً.')
        if any(x.get('sale_id') in sale_ids for x in self.db.get('returns', [])):
            return QMessageBox.warning(self, 'منع', 'لا يمكن تعديل الفاتورة بعد وجود مرتجعات مرتبطة بها. احذف المرتجعات أولاً.')
        self.cart = []
        for line in lines:
            self.cart.append({
                'item': line.get('item',''),
                'qty': int(line.get('qty',0) or 0),
                'unit_price': float(line.get('unit_price',0) or 0),
                'buy_price': float(line.get('buy_price',0) or 0),
                'total': float(line.get('total',0) or 0),
                'profit': 0.0,
            })
        first = lines[0]
        qd = QDate.fromString(first.get('date',''), 'yyyy-MM-dd')
        if qd.isValid():
            self.date.setDate(qd)
        cust_idx = self.customer.findText(first.get('customer',''))
        if cust_idx >= 0:
            self.customer.setCurrentIndex(cust_idx)
        self.notes.setText(first.get('notes','') or '')
        total = sum(float(x.get('total',0) or 0) for x in lines)
        paid = sum(float(x.get('paid_amount',0) or 0) for x in lines)
        if paid >= total:
            mode = 'نقدي'
        elif paid <= 0:
            mode = 'آجل'
        else:
            mode = 'جزئي'
        self.payment.setCurrentText(mode)
        self.paid.setValue(paid)
        self.editing_group_id = first.get('invoice_group_id') or f"single-{ensure_invoice_no(self.db, first)}"
        self.editing_invoice_no = ensure_invoice_no(self.db, first)
        self.refresh_cart_table()
        if self.cart:
            self.cart_table.selectRow(0)
            self.load_selected_cart_line()
        self.open_sales_page(self.invoice_entry_tab)
        QMessageBox.information(self, 'تم', f'تم تحميل الفاتورة #{self.editing_invoice_no} كاملة للتعديل: الزبون، التاريخ، نوع الدفع، المقبوض، الملاحظات، وكل الأسطر. بعد التعديل اضغط تسجيل الفاتورة لحفظ التحديث.')

    def refresh_cart_table(self):
        self.cart_table.setRowCount(len(self.cart))
        total = 0
        for r, row in enumerate(self.cart):
            total += float(row.get('total',0))
            vals = [r+1, row.get('item',''), row.get('qty',0), fmt_money(row.get('unit_price',0)), fmt_money(row.get('total',0))]
            for c, v in enumerate(vals): self.cart_table.setItem(r,c,QTableWidgetItem(str(v)))
        self.cart_summary.setText(f'إجمالي الفاتورة الحالية: {fmt_money(total)} د.ع | عدد الأسطر: {len(self.cart)}')
        self.sync_paid_limits()

    def invoice_groups(self):
        groups = []
        seen = set()
        for s in self.db['sales']:
            key = s.get('invoice_group_id') or f"single-{ensure_invoice_no(self.db, s)}"
            if key in seen:
                continue
            seen.add(key)
            lines = get_invoice_sales(self.db, s)
            groups.append({'sale': s, 'lines': lines})
        return groups

    def refresh_table(self):
        self.refresh_combos(); groups = self.invoice_groups(); self.table.setRowCount(len(groups)); total = due = 0
        for r, grp in enumerate(groups):
            sale = grp['sale']; lines = grp['lines']
            g_total = sum(float(x.get('total',0)) for x in lines)
            g_due = current_sale_due(self.db, sale)
            g_qty = sum(int(x.get('qty',0)) for x in lines)
            items_label = ' + '.join([x.get('item','') for x in lines[:3]])
            if len(lines) > 3:
                items_label += f' + {len(lines)-3}...'
            total += g_total; due += g_due
            ensure_invoice_no(self.db, sale)
            vals = [r+1, sale.get('invoice_no',''), sale.get('date',''), items_label, sale.get('customer',''), g_qty, fmt_money(g_total), sale.get('payment_type',''), fmt_money(g_total - g_due), fmt_money(g_due), len(lines)]
            for c,v in enumerate(vals): self.table.setItem(r,c,QTableWidgetItem(str(v)))
        self.summary.setText(f'فواتير البيع: {len(groups)} | إجمالي المبيعات: {fmt_money(total)} د.ع | ديون الزبائن: {fmt_money(due)} د.ع')

    def selected_sale(self):
        r = self.table.currentRow()
        groups = self.invoice_groups()
        if r < 0 or r >= len(groups):
            QMessageBox.warning(self,'تنبيه','اختر فاتورة من الجدول')
            return None
        return groups[r]['sale']


    def open_invoice(self):
        sale = self.selected_sale()
        if not sale:
            return
        show_invoice_dialog(self, self.db, sale, print_after=False)

    def print_invoice(self):
        sale = self.selected_sale()
        if not sale:
            return
        show_invoice_dialog(self, self.db, sale, print_after=True)

    def finalize_sale(self):
        if not self.cart:
            return QMessageBox.warning(self,'تنبيه','أضف صنف واحد على الأقل للفاتورة')
        customer = self.customer.currentText().strip()
        if not customer:
            return QMessageBox.warning(self,'تنبيه','اختر زبون')
        total = self.cart_total()
        if total <= 0 or any(float(x.get('unit_price', 0) or 0) <= 0 or float(x.get('total', 0) or 0) <= 0 for x in self.cart):
            return QMessageBox.warning(self,'تنبيه','لا يمكن حفظ فاتورة بسعر بيع صفر أو بإجمالي صفر')
        paid = min(total, max(0.0, float(self.paid.value())))
        mode = self.payment.currentText()
        if mode == 'نقدي':
            paid = total
        elif mode == 'آجل':
            paid = 0.0
        elif mode == 'جزئي' and paid <= 0:
            mode = 'آجل'
            paid = 0.0
        elif mode == 'جزئي' and paid >= total:
            mode = 'نقدي'
            paid = total

        editing_lines = []
        old_group_id = self.editing_group_id
        if old_group_id:
            editing_lines = [x for x in self.db['sales'] if (x.get('invoice_group_id') or f"single-{ensure_invoice_no(self.db, x)}") == old_group_id]
        invoice_no = self.editing_invoice_no if self.editing_invoice_no else next_invoice_no(self.db)
        group_id = old_group_id if old_group_id else generate_id('invoice')
        per_total_paid = paid / total if total else 0
        old_qty_map = {}
        for old in editing_lines:
            old_qty_map[old.get('item','')] = old_qty_map.get(old.get('item',''), 0) + int(old.get('qty',0) or 0)
        for line in self.cart:
            item = self.get_item(line['item'])
            if not item:
                return QMessageBox.warning(self,'تنبيه', f"الصنف غير موجود: {line['item']}")
            stock = int(item.get('qty',0)) + int(old_qty_map.get(line['item'], 0))
            if int(line['qty']) > stock:
                return QMessageBox.warning(self,'تنبيه', f"المخزون لا يكفي للصنف {line['item']} (المتاح {stock})")
        if editing_lines:
            for rec in editing_lines:
                for it in self.db['items']:
                    if it.get('name') == rec.get('item'):
                        inventory_add_stock(it, int(rec.get('qty',0)), float(rec.get('buy_price',0) or 0))
                        remove_inventory_movement(self.db, f"sale_{rec.get('id')}_{it.get('name','')}")
                        break
            self.db['sales'] = [x for x in self.db['sales'] if ((x.get('invoice_group_id') or f"single-{ensure_invoice_no(self.db, x)}") != old_group_id)]
            for i in range(len(self.db['cash'])-1, -1, -1):
                c = self.db['cash'][i]
                if c.get('source') == 'sales_group' and c.get('ref_id') == old_group_id:
                    self.db['cash'].pop(i)
        cash_total = 0
        for idx, line in enumerate(self.cart):
            line_total = float(line['total'])
            line_paid = round(line_total * per_total_paid, 2) if mode == 'جزئي' else (line_total if mode == 'نقدي' else 0.0)
            if idx == len(self.cart)-1 and mode == 'جزئي':
                line_paid = round(paid - cash_total, 2)
            cash_total += line_paid
            sale_id = generate_id('sale')
            item = self.get_item(line['item'])
            avg_cost = item_avg_cost(item)
            self.db['sales'].append({
                'id': sale_id, 'invoice_group_id': group_id, 'invoice_no': invoice_no,
                'date': self.date.date().toString('yyyy-MM-dd'), 'item': line['item'], 'customer': customer,
                'qty': int(line['qty']), 'unit_price': float(line['unit_price']), 'total': line_total,
                'profit': round((float(line['unit_price']) - avg_cost) * int(line['qty']), 2), 'buy_price': float(avg_cost),
                'payment_type': mode, 'paid_amount': line_paid, 'due_amount': round(line_total-line_paid,2),
                'notes': self.notes.text().strip(), 'created_at': now_str()
            })
            inventory_remove_stock(item, int(line['qty']), avg_cost)
            add_inventory_movement(self.db, item, 'sale', qty_out=int(line['qty']), unit_cost=avg_cost, sale_price=float(line['unit_price']), reference_type='sale', reference_id=sale_id, date=self.date.date().toString('yyyy-MM-dd'), notes=f"فاتورة بيع #{invoice_no}", movement_uid=f"sale_{sale_id}_{item.get('name','')}")
        if cash_total > 0:
            self.db['cash'].append({
                'date': self.date.date().toString('yyyy-MM-dd'),'type':'إيراد','category':'مبيعات','party':customer,
                'amount':cash_total,'notes':f'فاتورة بيع #{invoice_no}','source':'sales_group','ref_id': group_id,'created_at': now_str()
            })
        self.clear_cart()
        self.payment.setCurrentText('نقدي')
        self.paid.setValue(0)
        self.save()
        self.refresh_table()

    def delete_row(self):
        sale = self.selected_sale()
        if not sale: return
        linked_followup = 0.0
        sale_ids = {x.get('id') for x in get_invoice_sales(self.db, sale)}
        gid = sale.get('invoice_group_id') or f"single-{ensure_invoice_no(self.db, sale)}"
        for ev in customer_payment_allocations(self.db, sale.get('customer', '')):
            linked_followup += sum(_safe_float(a.get('amount', 0)) for a in ev.get('allocations', []) if a.get('sale_id') in sale_ids or a.get('invoice_group_id') == gid)
        if linked_followup > 0:
            return QMessageBox.warning(self,'منع','لا يمكن حذف الفاتورة لأن عليها تسديدات لاحقة مرتبطة. احذف أو عدّل التسديدات أولاً.')
        if any(x.get('sale_id') in sale_ids for x in self.db.get('returns', [])):
            return QMessageBox.warning(self,'منع','لا يمكن حذف الفاتورة لأن عليها مرتجعات محفوظة. احذف المرتجعات أولاً.')
        if QMessageBox.question(self,'تأكيد','حذف الفاتورة وعكس أثرها؟') != QMessageBox.Yes: return
        lines = get_invoice_sales(self.db, sale)
        refs = {x.get('id') for x in lines}
        group_id = sale.get('invoice_group_id')
        for rec in lines:
            for it in self.db['items']:
                if it.get('name') == rec.get('item'):
                    inventory_add_stock(it, int(rec.get('qty',0)), float(rec.get('buy_price',0) or 0))
                    remove_inventory_movement(self.db, f"sale_{rec.get('id')}_{it.get('name','')}")
        for i in range(len(self.db['cash'])-1, -1, -1):
            c = self.db['cash'][i]
            if group_id and c.get('ref_id') == group_id and c.get('source') == 'sales_group':
                self.db['cash'].pop(i)
            elif c.get('ref_id') in refs and c.get('source') == 'sales':
                self.db['cash'].pop(i)
        self.db['sales'] = [x for x in self.db['sales'] if x.get('id') not in refs]
        self.save(); self.refresh_table()

class ReturnWindow(BaseWindow):
    def __init__(self, main):
        super().__init__(main, '↩️ المرتجعات')
        self.subtitle_lbl.setText('واجهة المرتجعات صارت أوضح: كل جزء داخل تبويب مستقل حتى يبقى الشغل هادئ ومرتب.')

        hero = QFrame(); hero.setStyleSheet(CARD_FRAME_STYLE)
        hero_box = QVBoxLayout(hero); hero_box.setContentsMargins(18,16,18,16); hero_box.setSpacing(6)
        hero_title = QLabel('لوحة المرتجعات')
        hero_title.setAlignment(Qt.AlignRight)
        hero_title.setStyleSheet('font-size:20px;font-weight:900;background:transparent;border:none;')
        hero_note = QLabel('سجّل المرتجع من تبويب مستقل، راقب الملخص في تبويب منفصل، وراجع السجل بدون تزاحم.')
        hero_note.setAlignment(Qt.AlignRight); hero_note.setWordWrap(True)
        hero_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        hero_box.addWidget(hero_title); hero_box.addWidget(hero_note)
        self.layout.addWidget(hero)

        self.tabs = style_funders_tabs(QTabWidget())
        self.layout.addWidget(self.tabs, 1)

        entry_tab = QWidget(); entry_layout = QVBoxLayout(entry_tab); entry_layout.setContentsMargins(8,8,8,8); entry_layout.setSpacing(12)
        top = QVBoxLayout(); top.setSpacing(12)
        form_card = QFrame(); form_card.setStyleSheet(CARD_FRAME_STYLE)
        form_box = QVBoxLayout(form_card); form_box.setContentsMargins(16,16,16,16); form_box.setSpacing(10)
        form_title = QLabel('تسجيل مرتجع')
        form_title.setAlignment(Qt.AlignRight); form_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        form_box.addWidget(form_title)
        self.date = QDateEdit(); fix_date_edit_widget(self.date); self.date.setDate(QDate.currentDate())
        self.sale_cb = QComboBox(); self.refresh_combos()
        self.qty = tune_numeric_widget(QSpinBox()); self.qty.setRange(1,10_000_000)
        self.saleable = QCheckBox('صالح للبيع ويعود للمخزن'); self.saleable.setChecked(True)
        self.notes = QLineEdit(); self.notes.setPlaceholderText('ملاحظات')
        frm = QFormLayout(); frm.setLabelAlignment(Qt.AlignRight); frm.setFormAlignment(Qt.AlignTop); frm.setHorizontalSpacing(12); frm.setVerticalSpacing(10)
        frm.addRow('التاريخ:', self.date); frm.addRow('فاتورة البيع:', self.sale_cb); frm.addRow('الكمية المرتجعة:', self.qty)
        form_box.addLayout(frm); form_box.addWidget(self.saleable); form_box.addWidget(self.notes)
        top.addWidget(form_card)

        actions_card = QFrame(); actions_card.setStyleSheet(CARD_FRAME_STYLE)
        actions_box = QVBoxLayout(actions_card); actions_box.setContentsMargins(16,16,16,16); actions_box.setSpacing(10)
        actions_title = QLabel('إجراءات وملخص')
        actions_title.setAlignment(Qt.AlignRight); actions_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        actions_box.addWidget(actions_title)
        add = QPushButton('💾 تسجيل المرتجع'); delete = QPushButton('🗑 حذف المرتجع')
        add.setStyleSheet(BUTTON_STYLE); delete.setStyleSheet(SECONDARY_BUTTON); add.clicked.connect(self.add_row); delete.clicked.connect(self.delete_row)
        add.setMinimumHeight(46); delete.setMinimumHeight(46)
        actions_box.addWidget(add); actions_box.addWidget(delete)
        top.addWidget(actions_card)
        entry_layout.addLayout(top)
        self.tabs.addTab(entry_tab, 'تسجيل المرتجع')

        metrics_tab = QWidget(); metrics_layout = QVBoxLayout(metrics_tab); metrics_layout.setContentsMargins(8,8,8,8); metrics_layout.setSpacing(12)
        metrics_card = QFrame(); metrics_card.setStyleSheet(CARD_FRAME_STYLE)
        metrics_box = QVBoxLayout(metrics_card); metrics_box.setContentsMargins(16,16,16,16); metrics_box.setSpacing(10)
        metrics_title = QLabel('ملخص المرتجعات')
        metrics_title.setAlignment(Qt.AlignRight); metrics_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        metrics_box.addWidget(metrics_title)
        metrics_grid = QGridLayout(); metrics_grid.setHorizontalSpacing(10); metrics_grid.setVerticalSpacing(10)
        self.return_count_card = SummaryCard('عدد حركات المرتجعات', '0', 'عدد القيود الحالية')
        self.return_total_card = SummaryCard('إجمالي قيمة المرتجعات', '0', 'إجمالي القيمة المسجلة')
        for i, c in enumerate((self.return_count_card, self.return_total_card)):
            c.setMinimumHeight(100); c.setMaximumHeight(116); metrics_grid.addWidget(c, 0, i)
        metrics_box.addLayout(metrics_grid)
        self.summary = QLabel(); self.summary.setAlignment(Qt.AlignRight); self.summary.setWordWrap(True); self.summary.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        metrics_box.addWidget(self.summary)
        metrics_layout.addWidget(metrics_card)
        metrics_layout.addStretch(1)
        self.tabs.addTab(metrics_tab, 'الملخص')

        table_tab = QWidget(); table_layout = QVBoxLayout(table_tab); table_layout.setContentsMargins(8,8,8,8); table_layout.setSpacing(12)
        table_card = QFrame(); table_card.setStyleSheet(CARD_FRAME_STYLE)
        table_box = QVBoxLayout(table_card); table_box.setContentsMargins(16,16,16,16); table_box.setSpacing(10)
        table_title = QLabel('سجل المرتجعات')
        table_title.setAlignment(Qt.AlignRight); table_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        table_box.addWidget(table_title)
        self.table = QTableWidget(); self.table.setStyleSheet(TABLE_STYLE); self.table.setColumnCount(9)
        self.table.setHorizontalHeaderLabels(['#','التاريخ','الفاتورة','الصنف','الكمية','سعر البيع','المجموع','صالح للمخزن','ملاحظات'])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch); self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table_box.addWidget(self.table, 1)
        table_layout.addWidget(table_card, 1)
        self.tabs.addTab(table_tab, 'سجل المرتجعات')
        self.refresh_table()
    def select_attachment(self):
        path, _ = QFileDialog.getOpenFileName(self, 'اختيار فاتورة المورد', str(APP_DATA_DIR), 'Images/PDF (*.png *.jpg *.jpeg *.webp *.bmp *.pdf)')
        if path:
            self.attachment_path = path
            QMessageBox.information(self, 'تم', 'تم اختيار المرفق وسيُحفظ مع حركة الوارد.')

    def view_attachment(self):
        r = self.table.currentRow()
        if r < 0:
            return QMessageBox.warning(self, 'تنبيه', 'اختر حركة أولاً')
        rec = self.db['inbound'][r]
        rel = rec.get('attachment', '')
        if not rel:
            return QMessageBox.information(self, 'تنبيه', 'لا يوجد مرفق لهذه الحركة')
        full = resolve_app_file(rel)
        if not full.exists():
            return QMessageBox.warning(self, 'خطأ', 'الملف المرفق غير موجود')
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(full.resolve())))

    def refresh_combos(self):
        if not hasattr(self, 'sale_cb'): return
        self.sale_cb.clear()
        for s in self.db['sales']:
            label = f"#{s.get('invoice_no','')} | {s.get('customer','')} | {s.get('item','')} | باع {s.get('qty',0)}"
            self.sale_cb.addItem(label, s.get('id'))
    def sale_by_id(self, sid):
        for s in self.db['sales']:
            if s.get('id') == sid: return s
    def returned_qty(self, sid):
        return sum(int(x.get('qty',0)) for x in self.db.get('returns', []) if x.get('sale_id') == sid)
    def refresh_table(self):
        self.refresh_combos(); data = self.db.get('returns', []); self.table.setRowCount(len(data)); total = 0
        for r,row in enumerate(data):
            total += float(row.get('total',0))
            vals = [r+1,row.get('date',''),row.get('invoice_no',''),row.get('item',''),row.get('qty',0),fmt_money(row.get('unit_price',0)),fmt_money(row.get('total',0)), 'نعم' if row.get('saleable', True) else 'لا', row.get('notes','')]
            for c,v in enumerate(vals): self.table.setItem(r,c,QTableWidgetItem(str(v)))
        if hasattr(self, 'return_count_card'): self.return_count_card.set_value(str(len(data)))
        if hasattr(self, 'return_total_card'): self.return_total_card.set_value(fmt_money(total))
        if hasattr(self, 'return_metrics_count_card'): self.return_metrics_count_card.set_value(str(len(data)))
        if hasattr(self, 'return_metrics_total_card'): self.return_metrics_total_card.set_value(fmt_money(total))
        self.summary.setText(f'عدد المرتجعات: {len(data)} | قيمة المرتجعات: {fmt_money(total)} د.ع')
    def add_row(self):
        sid = self.sale_cb.currentData()
        sale = self.sale_by_id(sid)
        if not sale: return QMessageBox.warning(self,'تنبيه','اختر فاتورة بيع')
        qty = int(self.qty.value())
        available = int(sale.get('qty',0)) - self.returned_qty(sid)
        if qty > available: return QMessageBox.warning(self,'تنبيه', f'الكمية المرتجعة أكبر من المتاح. المتاح: {available}')
        item_name = sale.get('item','')
        saleable = self.saleable.isChecked()
        unit_cost = float(sale.get('buy_price',0) or 0)
        total = qty * float(sale.get('unit_price',0))
        rec = {'id': generate_id('return'), 'date': self.date.date().toString('yyyy-MM-dd'), 'sale_id': sid, 'invoice_no': sale.get('invoice_no',''), 'item': item_name, 'customer': sale.get('customer',''), 'qty': qty, 'unit_price': float(sale.get('unit_price',0)), 'unit_cost': unit_cost, 'saleable': saleable, 'total': total, 'credit_amount': total, 'credit_used': 0.0, 'cash_paid_out': 0.0, 'notes': self.notes.text().strip(), 'created_at': now_str()}
        if saleable:
            for it in self.db['items']:
                if it.get('name') == item_name:
                    inventory_add_stock(it, qty, unit_cost)
                    add_inventory_movement(self.db, it, 'return', qty_in=qty, unit_cost=unit_cost, reference_type='return', reference_id=rec.get('id'), date=self.date.date().toString('yyyy-MM-dd'), notes=f"مرتجع فاتورة #{sale.get('invoice_no','')}", movement_uid=f"return_{rec.get('id')}_{it.get('name','')}")
                    break
        self.db.setdefault('returns', []).append(rec)
        self.qty.setValue(1); self.saleable.setChecked(True); self.notes.clear(); self.save(); self.refresh_table()
    def delete_row(self):
        r = self.table.currentRow()
        if r < 0: return QMessageBox.warning(self,'تنبيه','اختر مرتجع')
        rec = self.db['returns'][r]
        if QMessageBox.question(self,'تأكيد','حذف المرتجع وعكس أثره؟') != QMessageBox.Yes: return
        if rec.get('saleable', True):
            for it in self.db['items']:
                if it.get('name') == rec.get('item'):
                    if int(rec.get('qty',0)) > int(it.get('qty',0)):
                        return QMessageBox.warning(self,'منع','لا يمكن حذف المرتجع لأن الكمية المرتجعة صُرفت من المخزن')
                    inventory_remove_stock(it, int(rec.get('qty',0)), float(rec.get('unit_cost',0) or 0))
                    remove_inventory_movement(self.db, f"return_{rec.get('id')}_{it.get('name','')}")
                    break
        for i in range(len(self.db['cash'])-1, -1, -1):
            c = self.db['cash'][i]
            if c.get('source') == 'return_auto' and c.get('ref_id') == rec.get('id'):
                self.db['cash'].pop(i)
        self.db['returns'].pop(r); self.save(); self.refresh_table()


class DamagedWindow(BaseWindow):
    def __init__(self, main):
        super().__init__(main, '⚠️ التالف')
        self.subtitle_lbl.setText('قسم التالف صار أوضح: كل جزء داخل تبويب مستقل حتى تبقى الصفحة مرتبة وسهلة.')

        hero = QFrame(); hero.setStyleSheet(CARD_FRAME_STYLE)
        hero_box = QVBoxLayout(hero); hero_box.setContentsMargins(18,16,18,16); hero_box.setSpacing(6)
        hero_title = QLabel('لوحة التالف')
        hero_title.setAlignment(Qt.AlignRight)
        hero_title.setStyleSheet('font-size:20px;font-weight:900;background:transparent;border:none;')
        hero_note = QLabel('سجّل التالف من تبويب مستقل، راقب الملخص في تبويب واضح، وراجع السجل بدون تكدس.')
        hero_note.setAlignment(Qt.AlignRight); hero_note.setWordWrap(True)
        hero_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        hero_box.addWidget(hero_title); hero_box.addWidget(hero_note)
        self.layout.addWidget(hero)

        self.tabs = style_funders_tabs(QTabWidget()); self.tabs.setDocumentMode(True)
        self.tabs.setStyleSheet(
            f"QTabWidget::pane {{border:1px solid {BORDER}; background:{CARD}; border-radius:14px; margin-top:8px;}}"
            f"QTabBar::tab {{min-width:170px; min-height:42px; padding:8px 14px; margin:4px 6px; border-radius:10px; background:{DARK}; color:{TEXT}; font-weight:800;}}"
            f"QTabBar::tab:selected {{background:{ACCENT2}; color:{TEXT_ON_ACCENT};}}"
        )
        self.layout.addWidget(self.tabs, 1)

        entry_tab = QWidget(); entry_layout = QVBoxLayout(entry_tab); entry_layout.setContentsMargins(8,8,8,8); entry_layout.setSpacing(12)
        top = QVBoxLayout(); top.setSpacing(12)
        form_card = QFrame(); form_card.setStyleSheet(CARD_FRAME_STYLE)
        form_box = QVBoxLayout(form_card); form_box.setContentsMargins(16,16,16,16); form_box.setSpacing(10)
        form_title = QLabel('تسجيل تالف')
        form_title.setAlignment(Qt.AlignRight); form_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        form_box.addWidget(form_title)
        self.date = QDateEdit(); fix_date_edit_widget(self.date); self.date.setDate(QDate.currentDate())
        self.item = QComboBox(); self.refresh_combos()
        self.qty = tune_numeric_widget(QSpinBox()); self.qty.setRange(1,10_000_000)
        self.notes = QLineEdit(); self.notes.setPlaceholderText('سبب التلف أو ملاحظات')
        frm = QFormLayout(); frm.setLabelAlignment(Qt.AlignRight); frm.setFormAlignment(Qt.AlignTop); frm.setHorizontalSpacing(12); frm.setVerticalSpacing(10)
        frm.addRow('التاريخ:', self.date); frm.addRow('الصنف:', self.item); frm.addRow('الكمية التالفة:', self.qty)
        form_box.addLayout(frm); form_box.addWidget(self.notes)
        top.addWidget(form_card)

        actions_card = QFrame(); actions_card.setStyleSheet(CARD_FRAME_STYLE)
        actions_box = QVBoxLayout(actions_card); actions_box.setContentsMargins(16,16,16,16); actions_box.setSpacing(10)
        actions_title = QLabel('إجراءات وملخص')
        actions_title.setAlignment(Qt.AlignRight); actions_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        actions_box.addWidget(actions_title)
        add = QPushButton('💾 تسجيل التالف'); delete = QPushButton('🗑 حذف الحركة')
        add.setStyleSheet(BUTTON_STYLE); delete.setStyleSheet(SECONDARY_BUTTON); add.clicked.connect(self.add_row); delete.clicked.connect(self.delete_row)
        add.setMinimumHeight(46); delete.setMinimumHeight(46)
        actions_box.addWidget(add); actions_box.addWidget(delete)
        self.damage_count_card = SummaryCard('حركات التالف', '0', 'عدد القيود الحالية')
        self.damage_total_card = SummaryCard('إجمالي الخسارة', '0', 'قيمة التالف المسجل')
        for c in (self.damage_count_card, self.damage_total_card):
            c.setMinimumHeight(96); c.setMaximumHeight(112); actions_box.addWidget(c)
        top.addWidget(actions_card)
        entry_layout.addLayout(top)
        self.tabs.addTab(entry_tab, 'تسجيل التالف')

        metrics_tab = QWidget(); metrics_layout = QVBoxLayout(metrics_tab); metrics_layout.setContentsMargins(8,8,8,8); metrics_layout.setSpacing(12)
        metrics_card = QFrame(); metrics_card.setStyleSheet(CARD_FRAME_STYLE)
        metrics_box = QVBoxLayout(metrics_card); metrics_box.setContentsMargins(16,16,16,16); metrics_box.setSpacing(10)
        metrics_title = QLabel('ملخص التالف')
        metrics_title.setAlignment(Qt.AlignRight); metrics_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        metrics_box.addWidget(metrics_title)
        metrics_grid = QGridLayout(); metrics_grid.setHorizontalSpacing(10); metrics_grid.setVerticalSpacing(10)
        self.damage_metrics_count_card = SummaryCard('عدد حركات التالف', '0', 'عدد القيود الحالية')
        self.damage_metrics_total_card = SummaryCard('إجمالي الخسارة', '0', 'إجمالي القيمة المسجلة')
        for i, c in enumerate((self.damage_metrics_count_card, self.damage_metrics_total_card)):
            c.setMinimumHeight(100); c.setMaximumHeight(116); metrics_grid.addWidget(c, 0, i)
        metrics_box.addLayout(metrics_grid)
        self.summary = QLabel(); self.summary.setAlignment(Qt.AlignRight); self.summary.setWordWrap(True); self.summary.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        metrics_box.addWidget(self.summary)
        metrics_layout.addWidget(metrics_card)
        metrics_layout.addStretch(1)
        self.tabs.addTab(metrics_tab, 'الملخص')

        table_tab = QWidget(); table_layout = QVBoxLayout(table_tab); table_layout.setContentsMargins(8,8,8,8); table_layout.setSpacing(12)
        table_card = QFrame(); table_card.setStyleSheet(CARD_FRAME_STYLE)
        table_box = QVBoxLayout(table_card); table_box.setContentsMargins(16,16,16,16); table_box.setSpacing(10)
        table_title = QLabel('سجل التالف')
        table_title.setAlignment(Qt.AlignRight); table_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        table_box.addWidget(table_title)
        self.table = QTableWidget(); self.table.setStyleSheet(TABLE_STYLE); self.table.setColumnCount(7)
        self.table.setHorizontalHeaderLabels(['#','التاريخ','الصنف','الكمية','سعر الشراء','الخسارة','ملاحظات'])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch); self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table_box.addWidget(self.table, 1)
        table_layout.addWidget(table_card, 1)
        self.tabs.addTab(table_tab, 'سجل التالف')
        self.refresh_table()
    def select_attachment(self):
        path, _ = QFileDialog.getOpenFileName(self, 'اختيار فاتورة المورد', str(APP_DATA_DIR), 'Images/PDF (*.png *.jpg *.jpeg *.webp *.bmp *.pdf)')
        if path:
            self.attachment_path = path
            QMessageBox.information(self, 'تم', 'تم اختيار المرفق وسيُحفظ مع حركة الوارد.')

    def view_attachment(self):
        r = self.table.currentRow()
        if r < 0:
            return QMessageBox.warning(self, 'تنبيه', 'اختر حركة أولاً')
        rec = self.db['inbound'][r]
        rel = rec.get('attachment', '')
        if not rel:
            return QMessageBox.information(self, 'تنبيه', 'لا يوجد مرفق لهذه الحركة')
        full = resolve_app_file(rel)
        if not full.exists():
            return QMessageBox.warning(self, 'خطأ', 'الملف المرفق غير موجود')
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(full.resolve())))

    def refresh_combos(self):
        if hasattr(self,'item'): self.item.clear(); self.item.addItems([x.get('name','') for x in self.db['items']])
    def refresh_table(self):
        self.refresh_combos(); data = self.db.get('damaged', []); self.table.setRowCount(len(data)); total = 0
        for r,row in enumerate(data):
            total += float(row.get('total',0))
            vals = [r+1,row.get('date',''),row.get('item',''),row.get('qty',0),fmt_money(row.get('unit_cost',0)),fmt_money(row.get('total',0)),row.get('notes','')]
            for c,v in enumerate(vals): self.table.setItem(r,c,QTableWidgetItem(str(v)))
        if hasattr(self, 'damage_count_card'): self.damage_count_card.set_value(str(len(data)))
        if hasattr(self, 'damage_total_card'): self.damage_total_card.set_value(fmt_money(total))
        self.summary.setText(f'حركات التالف: {len(data)} | إجمالي الخسارة: {fmt_money(total)} د.ع')
    def add_row(self):
        item_name = self.item.currentText().strip(); qty = int(self.qty.value())
        target = None
        for it in self.db['items']:
            if it.get('name') == item_name: target = it; break
        if not target: return QMessageBox.warning(self,'تنبيه','اختر صنف')
        if qty > int(target.get('qty',0)): return QMessageBox.warning(self,'تنبيه','الكمية التالفة أكبر من المخزون')
        unit_cost = item_avg_cost(target); total = qty * unit_cost
        inventory_remove_stock(target, qty, unit_cost)
        rec_id = generate_id('damaged')
        add_inventory_movement(self.db, target, 'damage', qty_out=qty, unit_cost=unit_cost, reference_type='damaged', reference_id=rec_id, date=self.date.date().toString('yyyy-MM-dd'), notes=self.notes.text().strip(), movement_uid=f"damage_{rec_id}_{target.get('name','')}")
        self.db.setdefault('damaged', []).append({'id': rec_id, 'date': self.date.date().toString('yyyy-MM-dd'), 'item': item_name, 'qty': qty, 'unit_cost': unit_cost, 'total': total, 'notes': self.notes.text().strip(), 'created_at': now_str()})
        self.qty.setValue(1); self.notes.clear(); self.save(); self.refresh_table()
    def delete_row(self):
        r = self.table.currentRow()
        if r < 0: return QMessageBox.warning(self,'تنبيه','اختر حركة')
        rec = self.db['damaged'][r]
        if QMessageBox.question(self,'تأكيد','حذف الحركة وعكس أثرها؟') != QMessageBox.Yes: return
        for it in self.db['items']:
            if it.get('name') == rec.get('item'):
                inventory_add_stock(it, int(rec.get('qty',0)), float(rec.get('unit_cost',0) or 0))
                remove_inventory_movement(self.db, f"damage_{rec.get('id')}_{it.get('name','')}")
        self.db['damaged'].pop(r); self.save(); self.refresh_table()


class CashWindow(BaseWindow):
    def __init__(self, main):
        super().__init__(main, '🧾 الصندوق')
        self.subtitle_lbl.setText('متابعة القاصة الدفترية وحركات الإيراد والمصروف بنفس روح اللوحة الرئيسية، لكن بشكل مقسّم وواضح حتى ما يطير أي قسم من الصفحة.')

        hero = QFrame(); hero.setObjectName('cashHero')
        hero.setStyleSheet(CARD_FRAME_STYLE)
        hero_box = QVBoxLayout(hero)
        hero_box.setContentsMargins(20, 18, 20, 18)
        hero_box.setSpacing(10)
        hero_title = QLabel('لوحة الصندوق')
        hero_title.setAlignment(Qt.AlignRight)
        hero_title.setStyleSheet('font-size:20px;font-weight:900;background:transparent;border:none;')
        hero_note = QLabel('قسمنا شاشة الصندوق إلى تبويبات: إدخال وإجراءات، مؤشرات، وسجل الحركات حتى تبقى الصفحة مرتبة ومفهومة.')
        hero_note.setAlignment(Qt.AlignRight)
        hero_note.setWordWrap(True)
        hero_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        hero_box.addWidget(hero_title)
        hero_box.addWidget(hero_note)
        self.layout.addWidget(hero)

        self.tabs = style_funders_tabs(QTabWidget())
        self.layout.addWidget(self.tabs, 1)

        setup_tab = QWidget(); setup_layout = QVBoxLayout(setup_tab); setup_layout.setContentsMargins(8, 8, 8, 8); setup_layout.setSpacing(12)
        top = QVBoxLayout(); top.setSpacing(12)

        input_card = QFrame(); input_card.setObjectName('cashInputCard'); input_card.setStyleSheet(CARD_FRAME_STYLE)
        input_wrap = QVBoxLayout(input_card)
        input_wrap.setContentsMargins(18, 18, 18, 18)
        input_wrap.setSpacing(10)
        input_title = QLabel('إضافة حركة جديدة')
        input_title.setAlignment(Qt.AlignRight)
        input_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        input_wrap.addWidget(input_title)
        form = QFormLayout()
        form.setLabelAlignment(Qt.AlignRight)
        form.setFormAlignment(Qt.AlignRight)
        form.setHorizontalSpacing(12)
        form.setVerticalSpacing(10)
        self.date = QDateEdit(); fix_date_edit_widget(self.date); self.date.setDate(QDate.currentDate())
        self.type_cb = QComboBox(); self.type_cb.addItems(['إيراد','مصروف'])
        self.category = QComboBox(); self.update_categories()
        self.type_cb.currentTextChanged.connect(self.update_categories)
        self.party = QLineEdit(); self.party.setPlaceholderText('الجهة')
        self.amount = tune_numeric_widget(QDoubleSpinBox()); self.amount.setRange(0,1_000_000_000); self.amount.setDecimals(0)
        self.notes = QTextEdit(); self.notes.setFixedHeight(96); self.notes.setPlaceholderText('ملاحظات الحركة أو سببها'); self.notes.setStyleSheet('font-size:14px;')
        form.addRow('التاريخ:', self.date)
        form.addRow('النوع:', self.type_cb)
        form.addRow('التصنيف:', self.category)
        form.addRow('الجهة:', self.party)
        form.addRow('المبلغ:', self.amount)
        input_wrap.addLayout(form)
        notes_lbl = QLabel('ملاحظات')
        notes_lbl.setAlignment(Qt.AlignRight)
        notes_lbl.setStyleSheet(f'font-size:12px;font-weight:800;color:{MUTED};background:transparent;border:none;')
        input_wrap.addWidget(notes_lbl)
        input_wrap.addWidget(self.notes)
        top.addWidget(input_card)

        actions_card = QFrame(); actions_card.setObjectName('cashActionsCard'); actions_card.setStyleSheet(CARD_FRAME_STYLE)
        actions_wrap = QVBoxLayout(actions_card)
        actions_wrap.setContentsMargins(18, 18, 18, 18)
        actions_wrap.setSpacing(10)
        actions_title = QLabel('إجراءات سريعة')
        actions_title.setAlignment(Qt.AlignRight)
        actions_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        actions_wrap.addWidget(actions_title)
        self.add_btn = QPushButton('➕ إضافة حركة')
        self.delete_btn = QPushButton('🗑 حذف الحركة المحددة')
        self.add_btn.setStyleSheet(BUTTON_STYLE)
        self.delete_btn.setStyleSheet(SECONDARY_BUTTON)
        self.add_btn.setMinimumHeight(46)
        self.delete_btn.setMinimumHeight(46)
        self.add_btn.clicked.connect(self.add_row)
        self.delete_btn.clicked.connect(self.delete_row)
        actions_wrap.addWidget(self.add_btn)
        actions_wrap.addWidget(self.delete_btn)
        action_note = QLabel('الحركات المترحلة تلقائيًا من المبيعات أو الوارد تُحذف من مصدرها الأصلي، مو من هنا.')
        action_note.setWordWrap(True)
        action_note.setAlignment(Qt.AlignRight)
        action_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        actions_wrap.addWidget(action_note)
        top.addWidget(actions_card)
        setup_layout.addLayout(top)

        self.summary_card = QFrame(); self.summary_card.setObjectName('cashSummaryCard'); self.summary_card.setStyleSheet(CARD_FRAME_STYLE)
        summary_box = QVBoxLayout(self.summary_card)
        summary_box.setContentsMargins(16, 14, 16, 14)
        summary_box.setSpacing(6)
        summary_title = QLabel('ملخص سريع')
        summary_title.setAlignment(Qt.AlignRight)
        summary_title.setStyleSheet('font-size:15px;font-weight:900;background:transparent;border:none;')
        self.summary = QLabel(); self.summary.setWordWrap(True); self.summary.setAlignment(Qt.AlignRight)
        self.summary.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        summary_box.addWidget(summary_title)
        summary_box.addWidget(self.summary)
        setup_layout.addWidget(self.summary_card)
        setup_layout.addStretch(1)

        metrics_tab = QWidget(); metrics_layout = QVBoxLayout(metrics_tab); metrics_layout.setContentsMargins(8,8,8,8); metrics_layout.setSpacing(12)
        metrics_wrap = QFrame(); metrics_wrap.setObjectName('cashMetricsWrap'); metrics_wrap.setStyleSheet(CARD_FRAME_STYLE)
        metrics_box = QVBoxLayout(metrics_wrap)
        metrics_box.setContentsMargins(16, 16, 16, 16)
        metrics_box.setSpacing(12)
        metrics_title = QLabel('مؤشرات الصندوق')
        metrics_title.setAlignment(Qt.AlignRight)
        metrics_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        metrics_box.addWidget(metrics_title)
        self.metrics = QGridLayout()
        self.metrics.setHorizontalSpacing(10)
        self.metrics.setVerticalSpacing(10)
        self.balance_card = SummaryCard('رصيد الصندوق الحالي', '0', 'القاصة الدفترية الحالية')
        self.in_card = SummaryCard('إجمالي الإيراد', '0', 'كل ما دخل للصندوق')
        self.out_card = SummaryCard('إجمالي المصروف', '0', 'كل ما خرج من الصندوق')
        self.hidab_card = SummaryCard('سحوبات هضاب', '0', 'السحوبات المسجلة على هضاب')
        self.mustafa_card = SummaryCard('سحوبات مصطفى', '0', 'السحوبات المسجلة على مصطفى')
        self.hidab_profit_card = SummaryCard('أرباح هضاب المدفوعة', '0', 'المدفوع من أرباح هضاب')
        self.mustafa_profit_card = SummaryCard('أرباح مصطفى المدفوعة', '0', 'المدفوع من أرباح مصطفى')
        self.moves_card = SummaryCard('عدد الحركات', '0', 'إجمالي قيود الصندوق الحالية')
        cards = [self.balance_card, self.in_card, self.out_card, self.hidab_card, self.mustafa_card, self.hidab_profit_card, self.mustafa_profit_card, self.moves_card]
        for i, card in enumerate(cards):
            card.setMinimumHeight(96); card.setMaximumHeight(112)
            self.metrics.addWidget(card, i // 2, i % 2)
        self.metrics.setColumnStretch(0, 1)
        self.metrics.setColumnStretch(1, 1)
        metrics_box.addLayout(self.metrics)
        metrics_layout.addWidget(metrics_wrap)
        metrics_layout.addStretch(1)

        table_tab = QWidget(); table_layout = QVBoxLayout(table_tab); table_layout.setContentsMargins(8,8,8,8); table_layout.setSpacing(12)
        table_wrap = QFrame(); table_wrap.setObjectName('cashTableWrap'); table_wrap.setStyleSheet(CARD_FRAME_STYLE)
        table_box = QVBoxLayout(table_wrap)
        table_box.setContentsMargins(16, 16, 16, 16)
        table_box.setSpacing(10)
        table_title = QLabel('سجل حركات الصندوق')
        table_title.setAlignment(Qt.AlignRight)
        table_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        table_box.addWidget(table_title)
        self.table = QTableWidget(); self.table.setStyleSheet(TABLE_STYLE); self.table.setColumnCount(7); self.table.setHorizontalHeaderLabels(['#','التاريخ','النوع','التصنيف','الجهة','المبلغ','ملاحظات']); self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch); self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table_box.addWidget(self.table)
        table_layout.addWidget(table_wrap, 1)

        self.tabs.addTab(setup_tab, 'الإدخال والإجراءات')
        self.tabs.addTab(metrics_tab, 'المؤشرات')
        self.tabs.addTab(table_tab, 'سجل الحركات')
        self.refresh_table()

    def update_categories(self):
        cur = self.category.currentText()
        self.category.clear()
        if self.type_cb.currentText() == 'إيراد':
            self.category.addItems(['إيراد يدوي','مبيعات'])
        else:
            self.category.addItems(['مصروف يدوي','وارد بضاعة','مرتجعات','سحوبات هضاب','سحوبات مصطفى','دفع ربح هضاب','دفع ربح مصطفى'])
        i = self.category.findText(cur)
        if i >= 0:
            self.category.setCurrentIndex(i)

    def refresh_table(self):
        data = self.db['cash']; self.table.setRowCount(len(data)); total_in = total_out = 0
        for r,row in enumerate(data):
            amount = float(row.get('amount',0))
            if row.get('type') == 'إيراد': total_in += amount
            else: total_out += amount
            vals = [r+1,row.get('date',''),row.get('type',''),row.get('category',''),row.get('party',''),fmt_money(amount),row.get('notes','')]
            for c,v in enumerate(vals): self.table.setItem(r,c,QTableWidgetItem(str(v)))
        balance = float(cash_balance(self.db) or 0)
        hidab_w = withdrawals_sum(self.db, 'سحوبات هضاب')
        mustafa_w = withdrawals_sum(self.db, 'سحوبات مصطفى')
        hidab_profit = owner_profit_payment_sum(self.db, 'هضاب')
        mustafa_profit = owner_profit_payment_sum(self.db, 'مصطفى')
        self.balance_card.set_value(f'{fmt_money(balance)} د.ع')
        self.in_card.set_value(f'{fmt_money(total_in)} د.ع')
        self.out_card.set_value(f'{fmt_money(total_out)} د.ع')
        self.hidab_card.set_value(f'{fmt_money(hidab_w)} د.ع')
        self.mustafa_card.set_value(f'{fmt_money(mustafa_w)} د.ع')
        self.hidab_profit_card.set_value(f'{fmt_money(hidab_profit)} د.ع')
        self.mustafa_profit_card.set_value(f'{fmt_money(mustafa_profit)} د.ع')
        self.moves_card.set_value(str(len(data)))
        self.summary.setText(
            f'رصيد الصندوق الحالي = {fmt_money(balance)} د.ع. إجمالي الإيراد {fmt_money(total_in)} د.ع مقابل إجمالي المصروف {fmt_money(total_out)} د.ع. '
            f'سحوبات هضاب {fmt_money(hidab_w)} د.ع، سحوبات مصطفى {fmt_money(mustafa_w)} د.ع، وأرباح الشركاء المدفوعة: هضاب {fmt_money(hidab_profit)} د.ع ومصطفى {fmt_money(mustafa_profit)} د.ع.'
        )

    def add_row(self):
        amount = float(self.amount.value())
        if amount <= 0: return QMessageBox.warning(self,'تنبيه','أدخل مبلغ صحيح')
        typ = self.type_cb.currentText(); category = self.category.currentText().strip()
        source = 'manual'
        party = self.party.text().strip()
        if category in ['سحوبات هضاب','سحوبات مصطفى']:
            source = 'withdrawal'
            if not party:
                party = category.replace('سحوبات ', '').strip()
        elif category in ['دفع ربح هضاب','دفع ربح مصطفى']:
            source = 'owner_profit_payment'
            if not party:
                party = category.replace('دفع ربح ', '').strip()
        rec = {'date': self.date.date().toString('yyyy-MM-dd'),'type':typ,'category':category,'party':party,'amount':amount,'notes':self.notes.toPlainText().strip(),'source':source,'created_at':now_str()}
        self.db['cash'].append(rec)
        self.amount.setValue(0); self.party.clear(); self.notes.clear(); self.save(); self.refresh_table(); self.main.refresh_dashboard()

    def delete_row(self):
        r = self.table.currentRow()
        if r < 0: return QMessageBox.warning(self,'تنبيه','اختر حركة')
        rec = self.db['cash'][r]
        if rec.get('source') in ['sales', 'sales_group', 'inbound', 'return_auto']:
            return QMessageBox.warning(self,'تنبيه','هذه حركة مترحلة تلقائيًا من شاشة أخرى. احذفها من مصدرها.')
        if QMessageBox.question(self,'تأكيد','حذف الحركة؟') != QMessageBox.Yes: return
        self.db['cash'].pop(r); self.save(); self.refresh_table(); self.main.refresh_dashboard()


class AgentCustodyDialog(BaseDialog):
    def __init__(self, parent=None, row=None, agents=None):
        super().__init__('إضافة حركة عهدة' if row is None else 'تعديل حركة عهدة', parent)
        layout = QVBoxLayout(self)
        form = QFormLayout()
        self.date = QDateEdit(); fix_date_edit_widget(self.date); self.date.setDate(QDate.currentDate())
        self.agent = QComboBox(); self.agent.setEditable(True)
        known_agents = [x for x in (agents or []) if x]
        if 'كرار' not in known_agents:
            known_agents.insert(0, 'كرار')
        self.agent.addItems(known_agents)
        self.type_cb = QComboBox(); self.type_cb.addItems(['وارد مستلم','تحويل','مصروف','معالجة فرق'])
        self.settlement_direction = QComboBox(); self.settlement_direction.addItems(['نقصان','زيادة'])
        self.amount = QDoubleSpinBox(); self.amount.setRange(0, 1_000_000_000); self.amount.setDecimals(0)
        self.party = QLineEdit(); self.party.setPlaceholderText('الجهة / المطعم / المرسل إليه')
        self.notes = QTextEdit(); self.notes.setFixedHeight(90)
        form.addRow('التاريخ:', self.date)
        form.addRow('المندوب:', self.agent)
        form.addRow('نوع الحركة:', self.type_cb)
        form.addRow('اتجاه المعالجة:', self.settlement_direction)
        form.addRow('المبلغ:', self.amount)
        form.addRow('الجهة:', self.party)
        form.addRow('ملاحظات:', self.notes)
        layout.addLayout(form)
        btns = QHBoxLayout(); save_btn = QPushButton('حفظ'); cancel_btn = QPushButton('إلغاء')
        save_btn.setStyleSheet(BUTTON_STYLE); cancel_btn.setStyleSheet(SECONDARY_BUTTON)
        save_btn.clicked.connect(self.accept); cancel_btn.clicked.connect(self.reject)
        btns.addWidget(save_btn); btns.addWidget(cancel_btn); layout.addLayout(btns)
        self.type_cb.currentTextChanged.connect(self.update_state)
        if row:
            date_text = row.get('date', '') or QDate.currentDate().toString('yyyy-MM-dd')
            qd = QDate.fromString(date_text, 'yyyy-MM-dd')
            if qd.isValid():
                self.date.setDate(qd)
            self.agent.setCurrentText(row.get('agent', ''))
            old_type = row.get('type', 'وارد مستلم')
            self.type_cb.setCurrentText('معالجة فرق' if old_type == 'تسوية' else old_type)
            self.settlement_direction.setCurrentText(row.get('settlement_direction', 'نقصان') or 'نقصان')
            self.amount.setValue(float(row.get('amount', 0) or 0))
            self.party.setText(row.get('party', ''))
            self.notes.setPlainText(row.get('notes', ''))
        self.update_state()

    def update_state(self):
        is_settlement = self.type_cb.currentText().strip() == 'معالجة فرق'
        self.settlement_direction.setEnabled(is_settlement)
        self.settlement_direction.setVisible(is_settlement)

    def get_data(self):
        return {
            'date': self.date.date().toString('yyyy-MM-dd'),
            'agent': self.agent.currentText().strip(),
            'type': self.type_cb.currentText().strip(),
            'settlement_direction': self.settlement_direction.currentText().strip() if self.type_cb.currentText().strip() == 'معالجة فرق' else '',
            'amount': float(self.amount.value()),
            'party': self.party.text().strip(),
            'notes': self.notes.toPlainText().strip(),
        }


class AgentsCustodyWindow(BaseWindow):
    def __init__(self, main):
        super().__init__(main, '🧾 عهدة المندوبين')
        self.subtitle_lbl.setText('واجهة منظمة بنفس روح اللوحة الرئيسية لمتابعة عهد المندوبين، التسويات، وصافي كل مندوب بشكل أوضح.')

        hero = QFrame(); hero.setStyleSheet(CARD_FRAME_STYLE)
        hero_box = QVBoxLayout(hero)
        hero_box.setContentsMargins(18, 16, 18, 16)
        hero_box.setSpacing(6)
        hero_title = QLabel('متابعة العهدة والتسوية')
        hero_title.setAlignment(Qt.AlignRight)
        hero_title.setStyleSheet('font-size:20px;font-weight:900;background:transparent;border:none;')
        hero_note = QLabel('الوارد المستلم يزيد العهدة، والتحويل والمصروف ينزلانها، ومعالجة الفرق تعدّل الصافي بدون ترحيل تلقائي للصندوق.')
        hero_note.setAlignment(Qt.AlignRight)
        hero_note.setWordWrap(True)
        hero_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        hero_box.addWidget(hero_title)
        hero_box.addWidget(hero_note)
        self.layout.addWidget(hero)

        metrics_wrap = QFrame(); metrics_wrap.setStyleSheet(CARD_FRAME_STYLE)
        metrics_box = QVBoxLayout(metrics_wrap)
        metrics_box.setContentsMargins(16, 16, 16, 16)
        metrics_box.setSpacing(12)
        metrics_title = QLabel('مؤشرات العهدة')
        metrics_title.setAlignment(Qt.AlignRight)
        metrics_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        metrics_box.addWidget(metrics_title)
        self.metrics = QGridLayout()
        self.metrics.setHorizontalSpacing(10)
        self.metrics.setVerticalSpacing(10)
        self.lbl_incoming = SummaryCard('الوارد المستلم', '0', 'الحركات التي تزيد العهدة')
        self.lbl_transfers = SummaryCard('التحويلات', '0', 'المبالغ التي تنزل من العهدة')
        self.lbl_expenses = SummaryCard('المصاريف والتسويات', '0', 'المصاريف ومعالجة الفرق')
        self.lbl_balance = SummaryCard('صافي العهدة', '0', 'المتبقي النهائي على المندوبين')
        cards = [self.lbl_incoming, self.lbl_transfers, self.lbl_expenses, self.lbl_balance]
        for i, w in enumerate(cards):
            w.setMinimumHeight(106)
            w.setMaximumHeight(120)
            self.metrics.addWidget(w, i // 2, i % 2)
        self.metrics.setColumnStretch(0, 1)
        self.metrics.setColumnStretch(1, 1)
        metrics_box.addLayout(self.metrics)
        self.layout.addWidget(metrics_wrap)

        top_wrap = QFrame(); top_wrap.setStyleSheet(CARD_FRAME_STYLE)
        top_outer = QVBoxLayout(top_wrap)
        top_outer.setContentsMargins(16, 16, 16, 16)
        top_outer.setSpacing(12)
        top_title = QLabel('بحث وإجراءات سريعة')
        top_title.setAlignment(Qt.AlignRight)
        top_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        top_outer.addWidget(top_title)
        top = QVBoxLayout()
        top.setSpacing(12)

        controls_card = QFrame(); controls_card.setStyleSheet(CARD_FRAME_STYLE)
        controls_box = QVBoxLayout(controls_card)
        controls_box.setContentsMargins(14, 14, 14, 14)
        controls_box.setSpacing(10)
        controls_title = QLabel('بحث وتصفية')
        controls_title.setAlignment(Qt.AlignRight)
        controls_title.setStyleSheet('font-size:15px;font-weight:900;background:transparent;border:none;')
        controls_box.addWidget(controls_title)
        self.search = QLineEdit(); self.search.setPlaceholderText('بحث بالمندوب أو الجهة أو الملاحظات'); self.search.textChanged.connect(self.refresh_table)
        controls_box.addWidget(self.search)
        self.agent_filter = QComboBox(); self.agent_filter.setEditable(False)
        self.agent_filter.currentTextChanged.connect(self.refresh_table)
        filter_lbl = QLabel('تصفية المندوب')
        filter_lbl.setAlignment(Qt.AlignRight)
        filter_lbl.setStyleSheet(f'font-size:12px;font-weight:800;color:{MUTED};background:transparent;border:none;')
        controls_box.addWidget(filter_lbl)
        controls_box.addWidget(self.agent_filter)

        actions_card = QFrame(); actions_card.setStyleSheet(CARD_FRAME_STYLE)
        actions_box = QGridLayout(actions_card)
        actions_box.setContentsMargins(14, 14, 14, 14)
        actions_box.setHorizontalSpacing(10)
        actions_box.setVerticalSpacing(10)
        actions_title = QLabel('إجراءات سريعة')
        actions_title.setAlignment(Qt.AlignRight)
        actions_title.setStyleSheet('font-size:15px;font-weight:900;background:transparent;border:none;')
        actions_box.addWidget(actions_title, 0, 0, 1, 2)
        for i, (txt, fn, style) in enumerate([
            ('➕ إضافة حركة', self.add_row, BUTTON_STYLE),
            ('✏️ تعديل', self.edit_row, SECONDARY_BUTTON),
            ('🗑 حذف', self.delete_row, SECONDARY_BUTTON),
            ('📄 تقرير المندوب', self.export_agent_report, SECONDARY_BUTTON),
        ]):
            b = QPushButton(txt); b.setStyleSheet(style); b.setMinimumHeight(40); b.clicked.connect(fn)
            actions_box.addWidget(b, 1 + i // 2, i % 2)

        top.addWidget(controls_card)
        top.addWidget(actions_card)
        top_outer.addLayout(top)
        self.layout.addWidget(top_wrap)
        table_wrap = QFrame(); table_wrap.setStyleSheet(CARD_FRAME_STYLE)
        table_box = QVBoxLayout(table_wrap)
        table_box.setContentsMargins(16, 16, 16, 16)
        table_box.setSpacing(10)
        table_title = QLabel('سجل حركات العهدة')
        table_title.setAlignment(Qt.AlignRight)
        table_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        table_box.addWidget(table_title)
        self.table = QTableWidget(); self.table.setStyleSheet(TABLE_STYLE); self.table.setColumnCount(9)
        self.table.setHorizontalHeaderLabels(['#','التاريخ','المندوب','نوع الحركة','اتجاه المعالجة','الجهة','المبلغ','الأثر على العهدة','ملاحظات'])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table_box.addWidget(self.table)
        self.layout.addWidget(table_wrap)

        summary_card = QFrame(); summary_card.setStyleSheet(CARD_FRAME_STYLE)
        summary_box = QVBoxLayout(summary_card)
        summary_box.setContentsMargins(16, 14, 16, 14)
        summary_box.setSpacing(6)
        summary_title = QLabel('شرح سريع')
        summary_title.setAlignment(Qt.AlignRight)
        summary_title.setStyleSheet('font-size:15px;font-weight:900;background:transparent;border:none;')
        self.summary = QLabel(); self.summary.setAlignment(Qt.AlignRight); self.summary.setWordWrap(True)
        self.summary.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        summary_box.addWidget(summary_title)
        summary_box.addWidget(self.summary)
        self.layout.addWidget(summary_card)

        self.refresh_agents()
        self.refresh_table()
    def refresh_agents(self):
        current = self.agent_filter.currentText()
        self.agent_filter.blockSignals(True)
        self.agent_filter.clear()
        self.agent_filter.addItem('الكل')
        agents = sorted({(x.get('agent', '') or '').strip() for x in agents_custody_rows(self.db) if (x.get('agent', '') or '').strip()})
        if 'كرار' not in agents:
            agents.insert(0, 'كرار')
        self.agent_filter.addItems(agents)
        idx = self.agent_filter.findText(current)
        self.agent_filter.setCurrentIndex(idx if idx >= 0 else 0)
        self.agent_filter.blockSignals(False)

    def filtered_rows(self):
        rows = list(agents_custody_rows(self.db))
        agent = self.agent_filter.currentText().strip()
        if agent and agent != 'الكل':
            rows = [x for x in rows if (x.get('agent', '') or '').strip() == agent]
        q = self.search.text().strip().lower()
        if q:
            rows = [x for x in rows if q in (x.get('agent', '') or '').lower() or q in (x.get('party', '') or '').lower() or q in (x.get('notes', '') or '').lower()]
        return rows

    def add_row(self):
        d = AgentCustodyDialog(self, agents=[x['agent'] for x in agents_custody_summary(self.db)])
        if d.exec():
            row = d.get_data()
            if not row['agent']:
                return QMessageBox.warning(self, 'تنبيه', 'اسم المندوب مطلوب')
            if row['amount'] <= 0:
                return QMessageBox.warning(self, 'تنبيه', 'أدخل مبلغًا صحيحًا')
            row['id'] = generate_id('agc')
            row['created_at'] = now_str()
            self.db.setdefault('agents_custody', []).append(row)
            self.save(); self.refresh_agents(); self.refresh_table(); self.main.refresh_dashboard()

    def edit_row(self):
        row = self.selected_row()
        if not row:
            return
        d = AgentCustodyDialog(self, row=row, agents=[x['agent'] for x in agents_custody_summary(self.db)])
        if d.exec():
            new = d.get_data()
            if not new['agent']:
                return QMessageBox.warning(self, 'تنبيه', 'اسم المندوب مطلوب')
            if new['amount'] <= 0:
                return QMessageBox.warning(self, 'تنبيه', 'أدخل مبلغًا صحيحًا')
            row.update(new)
            self.save(); self.refresh_agents(); self.refresh_table(); self.main.refresh_dashboard()

    def delete_row(self):
        row = self.selected_row()
        if not row:
            return
        if QMessageBox.question(self, 'تأكيد', 'حذف حركة العهدة المحددة؟') != QMessageBox.Yes:
            return
        rows = agents_custody_rows(self.db)
        for i, x in enumerate(rows):
            if x.get('id') == row.get('id'):
                rows.pop(i)
                break
        self.save(); self.refresh_agents(); self.refresh_table(); self.main.refresh_dashboard()

    def current_agent_name(self):
        agent = self.agent_filter.currentText().strip()
        if agent and agent != 'الكل':
            return agent
        row = self.selected_row(silent=True)
        if row and (row.get('agent', '') or '').strip():
            return (row.get('agent', '') or '').strip()
        return ''

    def selected_row(self, silent=False):
        r = self.table.currentRow()
        rows = self.filtered_rows()
        if r < 0 or r >= len(rows):
            if not silent:
                QMessageBox.warning(self, 'تنبيه', 'اختر حركة أولاً')
            return None
        target_id = rows[r].get('id')
        for row in agents_custody_rows(self.db):
            if row.get('id') == target_id:
                return row
        return None

    def build_agent_report_text(self, agent_name):
        agent_name = (agent_name or '').strip()
        rows = [x for x in agents_custody_rows(self.db) if (x.get('agent', '') or '').strip() == agent_name]
        rows.sort(key=lambda x: ((x.get('date', '') or ''), (x.get('created_at', '') or '')))
        incoming = transfers = expenses = settlements = balance = 0.0
        lines = []
        lines.append(f'كشف عهدة المندوب: {agent_name}')
        lines.append('=' * 70)
        lines.append('')
        for idx, row in enumerate(rows, 1):
            amount = float(row.get('amount', 0) or 0)
            effect = agent_custody_effect(row)
            typ = (row.get('type', '') or '').strip()
            if typ == 'وارد مستلم':
                incoming += amount
            elif typ == 'تحويل':
                transfers += amount
            elif typ == 'مصروف':
                expenses += amount
            elif typ == 'معالجة فرق':
                settlements += amount if (row.get('settlement_direction', '') or '').strip() == 'لصالح العهدة' else -amount
            balance += effect
            lines.extend([
                f'{idx}) التاريخ: {row.get("date", "")}',
                f'   نوع الحركة: {typ}',
                f'   اتجاه المعالجة: {row.get("settlement_direction", "")}',
                f'   الجهة: {row.get("party", "")}',
                f'   المبلغ: {fmt_money(amount)} د.ع',
                f'   الأثر على العهدة: {fmt_money(effect)} د.ع',
                f'   الملاحظات: {row.get("notes", "")}',
                '-' * 70,
            ])
        lines.extend([
            '',
            'الملخص',
            '=' * 70,
            f'إجمالي الوارد المستلم: {fmt_money(incoming)} د.ع',
            f'إجمالي التحويلات: {fmt_money(transfers)} د.ع',
            f'إجمالي المصاريف: {fmt_money(expenses)} د.ع',
            f'صافي التسويات: {fmt_money(settlements)} د.ع',
            f'صافي عهدة المندوب: {fmt_money(balance)} د.ع',
            f'عدد الحركات: {len(rows)}',
            '',
            f'تاريخ إصدار الكشف: {now_str()}',
        ])
        return '\n'.join(lines)

    def export_agent_report(self):
        agent_name = self.current_agent_name()
        if not agent_name:
            return QMessageBox.warning(self, 'تنبيه', 'اختر مندوبًا من التصفية أو حدد حركة له أولاً')
        rows = [x for x in agents_custody_rows(self.db) if (x.get('agent', '') or '').strip() == agent_name]
        if not rows:
            return QMessageBox.warning(self, 'تنبيه', 'لا توجد حركات لهذا المندوب')
        dlg = AgentCustodyReportPreviewDialog(self.db, agent_name, self)
        dlg.exec()

    def refresh_table(self):
        rows = self.filtered_rows()
        self.table.setRowCount(len(rows))
        incoming = transfers = expenses = balance = 0.0
        for r, row in enumerate(rows):
            amount = float(row.get('amount', 0) or 0)
            effect = agent_custody_effect(row)
            typ = (row.get('type', '') or '').strip()
            if typ == 'وارد مستلم':
                incoming += amount
            elif typ == 'تحويل':
                transfers += amount
            elif typ == 'مصروف':
                expenses += amount
            balance += effect
            vals = [
                r + 1,
                row.get('date', ''),
                row.get('agent', ''),
                typ,
                row.get('settlement_direction', ''),
                row.get('party', ''),
                fmt_money(amount),
                fmt_money(effect),
                row.get('notes', ''),
            ]
            for c, v in enumerate(vals):
                self.table.setItem(r, c, QTableWidgetItem(str(v)))
        self.lbl_incoming.set_value(f'{fmt_money(incoming)} د.ع')
        self.lbl_transfers.set_value(f'{fmt_money(transfers)} د.ع')
        self.lbl_expenses.set_value(f'{fmt_money(expenses)} د.ع')
        self.lbl_balance.set_value(f'{fmt_money(balance)} د.ع')
        overall = total_agents_custody(self.db)
        current_agent = self.agent_filter.currentText().strip() or 'الكل'
        self.summary.setText(
            f'العرض الحالي: {current_agent} | عدد الحركات: {len(rows)} | إجمالي عهد المندوبين حالياً: {fmt_money(overall)} د.ع | '
            f'المبلغ الموجب يعني المتبقي على المندوب، والسالب يعني صرف/تحويل زائد يحتاج مراجعة. '
            f'الوارد المستلم يزيد العهدة، أما التحويل والمصروف والتسويات فتخفضها أو تعدّلها حسب الحركة.'
        )

class ReconciliationWindow(BaseWindow):
    def __init__(self, main):
        super().__init__(main, '🧮 صفحة المطابقة')
        self.subtitle_lbl.setText('تصفية سريعة وواضحة بين الكاش الفعلي بيدك، القاصة الدفترية، عهدة المندوبين، والفرق النهائي.')

        hero = QFrame(); hero.setObjectName('reconHero')
        hero.setStyleSheet(CARD_FRAME_STYLE)
        hero_box = QVBoxLayout(hero)
        hero_box.setContentsMargins(20, 18, 20, 18)
        hero_box.setSpacing(10)
        hero_title = QLabel('لوحة المطابقة')
        hero_title.setAlignment(Qt.AlignRight)
        hero_title.setStyleSheet('font-size:20px;font-weight:900;background:transparent;border:none;')
        hero_note = QLabel('نفس روح اللوحة الرئيسية لكن مخصصة لضبط القاصة والعهدة والفرق النهائي بسرعة ووضوح.')
        hero_note.setAlignment(Qt.AlignRight)
        hero_note.setWordWrap(True)
        hero_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        hero_box.addWidget(hero_title)
        hero_box.addWidget(hero_note)
        self.layout.addWidget(hero)

        self.recon_calc = QuickCalculatorPanel('حاسبة سريعة', 'ثابتة أعلى اليمين داخل صفحة المطابقة للمراجعة السريعة وقت التدقيق.', compact=True)
        recon_calc_row = QHBoxLayout()
        recon_calc_row.setContentsMargins(0, 0, 0, 0)
        recon_calc_row.setSpacing(12)
        recon_calc_row.addStretch(1)
        recon_calc_row.addWidget(self.recon_calc, 0, Qt.AlignTop | Qt.AlignRight)
        self.layout.addLayout(recon_calc_row)

        metrics_wrap = QFrame(); metrics_wrap.setObjectName('reconMetricsWrap'); metrics_wrap.setStyleSheet(CARD_FRAME_STYLE)
        metrics_box = QVBoxLayout(metrics_wrap)
        metrics_box.setContentsMargins(16, 16, 16, 16)
        metrics_box.setSpacing(12)
        metrics_title = QLabel('مؤشرات المطابقة')
        metrics_title.setAlignment(Qt.AlignRight)
        metrics_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        metrics_box.addWidget(metrics_title)
        self.metrics = QGridLayout()
        self.metrics.setHorizontalSpacing(10)
        self.metrics.setVerticalSpacing(10)
        self.metric_labels = {}
        self.metric_cards = {}
        cards = [
            ('book_cash', 'القاصة الدفترية', 'الموجود نظرياً حسب النظام'),
            ('actual_cash', 'الكاش الفعلي', 'الموجود بيدك حالياً'),
            ('agents_base_custody', 'عهدة المندوبين الأساسية', 'الوارد/التحويل/المصروف قبل التسويات'),
            ('settlements_value', 'التسوية / فرق المندوبين', 'الفرق الذي وضعته في خانة التسوية أو معالجة الفرق'),
            ('agents_custody', 'عهدة المندوبين النهائية', 'العهدة بعد إضافة/طرح التسويات'),
            ('diff_cash', 'الفرق النقدي', 'الكاش الفعلي ناقص/زايد عن القاصة'),
            ('diff_with_agents', 'الفرق بعد العهدة والتسوية', 'الكاش الفعلي + عهدة المندوبين النهائية - القاصة'),
            ('customer_dues', 'ديون الزبائن', 'للعلم والمتابعة فقط'),
            ('payables', 'ذمم الموردين', 'التزامات مفتوحة'),
        ]
        for i, (key, title, note) in enumerate(cards):
            card = SummaryCard(title, '0', note)
            card.setMinimumHeight(96)
            card.setMaximumHeight(112)
            self.metric_labels[key] = card.value_label
            self.metric_cards[key] = card
            self.metrics.addWidget(card, i // 2, i % 2)
        self.metrics.setColumnStretch(0, 1)
        self.metrics.setColumnStretch(1, 1)
        metrics_box.addLayout(self.metrics)
        self.layout.addWidget(metrics_wrap)

        top = QGridLayout()
        top.setHorizontalSpacing(12)
        top.setVerticalSpacing(12)

        input_card = QFrame(); input_card.setObjectName('reconInputCard'); input_card.setStyleSheet(CARD_FRAME_STYLE)
        input_wrap = QVBoxLayout(input_card)
        input_wrap.setContentsMargins(18, 18, 18, 18)
        input_wrap.setSpacing(10)
        input_title = QLabel('المدخلات الحالية')
        input_title.setAlignment(Qt.AlignRight)
        input_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        input_wrap.addWidget(input_title)
        form = QFormLayout()
        form.setLabelAlignment(Qt.AlignRight)
        form.setFormAlignment(Qt.AlignRight)
        form.setHorizontalSpacing(12)
        form.setVerticalSpacing(10)
        self.actual_cash = QDoubleSpinBox(); self.actual_cash.setRange(-1_000_000_000, 1_000_000_000); self.actual_cash.setDecimals(0); self.actual_cash.valueChanged.connect(self.refresh_view)
        form.addRow('الكاش الفعلي بيدك:', self.actual_cash)
        input_wrap.addLayout(form)
        notes_lbl = QLabel('ملاحظات المطابقة')
        notes_lbl.setAlignment(Qt.AlignRight)
        notes_lbl.setStyleSheet(f'font-size:12px;font-weight:800;color:{MUTED};background:transparent;border:none;')
        input_wrap.addWidget(notes_lbl)
        self.notes = QTextEdit(); self.notes.setFixedHeight(86); self.notes.setPlaceholderText('اكتب سبب الفرق أو ملاحظتك الحالية')
        input_wrap.addWidget(self.notes)
        actions_card = QFrame(); actions_card.setObjectName('reconActionsCard'); actions_card.setStyleSheet(CARD_FRAME_STYLE)
        actions_wrap = QVBoxLayout(actions_card)
        actions_wrap.setContentsMargins(18, 18, 18, 18)
        actions_wrap.setSpacing(10)
        actions_title = QLabel('إجراءات سريعة')
        actions_title.setAlignment(Qt.AlignRight)
        actions_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        actions_wrap.addWidget(actions_title)
        for txt, fn, style in [
            ('💾 حفظ القيمة الحالية', self.save_current_state, SECONDARY_BUTTON),
            ('📌 حفظ لقطة مطابقة', self.save_snapshot, BUTTON_STYLE),
            ('🗑 حذف اللقطة المحددة', self.delete_snapshot, SECONDARY_BUTTON),
        ]:
            b = QPushButton(txt); b.setStyleSheet(style); b.setMinimumHeight(46); b.clicked.connect(fn); actions_wrap.addWidget(b)
        actions_wrap.addStretch()

        left_cards = QVBoxLayout()
        left_cards.setContentsMargins(0, 0, 0, 0)
        left_cards.setSpacing(12)
        left_cards.addWidget(input_card)
        left_cards.addWidget(actions_card)
        top.addLayout(left_cards, 0, 0)
        top.setColumnStretch(0, 1)

        self.layout.addLayout(top)

        self.summary_card = QFrame(); self.summary_card.setObjectName('reconSummaryCard'); self.summary_card.setStyleSheet(CARD_FRAME_STYLE)
        summary_box = QVBoxLayout(self.summary_card)
        summary_box.setContentsMargins(16, 14, 16, 14)
        summary_box.setSpacing(6)
        summary_title = QLabel('شرح سريع')
        summary_title.setAlignment(Qt.AlignRight)
        summary_title.setStyleSheet('font-size:15px;font-weight:900;background:transparent;border:none;')
        self.summary = QLabel(); self.summary.setWordWrap(True); self.summary.setAlignment(Qt.AlignRight)
        self.summary.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        summary_box.addWidget(summary_title)
        summary_box.addWidget(self.summary)
        self.layout.addWidget(self.summary_card)

        table_wrap = QFrame(); table_wrap.setObjectName('reconTableWrap'); table_wrap.setStyleSheet(CARD_FRAME_STYLE)
        table_box = QVBoxLayout(table_wrap)
        table_box.setContentsMargins(16, 16, 16, 16)
        table_box.setSpacing(10)
        table_title = QLabel('سجل لقطات المطابقة')
        table_title.setAlignment(Qt.AlignRight)
        table_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        table_box.addWidget(table_title)
        self.table = QTableWidget(); self.table.setStyleSheet(TABLE_STYLE); self.table.setColumnCount(10)
        self.table.setHorizontalHeaderLabels(['#','التاريخ','الكاش الفعلي','القاصة الدفترية','العهدة الأساسية','التسوية','عهدة المندوبين النهائية','الفرق النقدي','الفرق بعد العهدة والتسوية','ملاحظات'])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table_box.addWidget(self.table)
        self.layout.addWidget(table_wrap)

        self._loading = False
        self.refresh_view()
        self.refresh_table()

    def _state(self):
        return reconciliation_data(self.db)

    def calc_press(self, value):
        current = self.calc_display.text().strip()
        if value == 'C':
            self.calc_display.clear()
            return
        if value == '⌫':
            self.calc_display.setText(current[:-1])
            return
        if value == '=':
            self.calc_evaluate()
            return
        mapped = {'×': '*', '÷': '/'}
        token = mapped.get(value, value)
        if value == '%':
            token = '/100'
        if value in ('+', '-', '×', '÷') and (not current or current[-1] in '+-*/.'):
            return
        if value == '.' and (not current or current[-1] in '+-*/'):
            token = '0.'
        self.calc_display.setText(current + token)

    def calc_evaluate(self):
        expr = (self.calc_display.text() or '').strip()
        if not expr:
            return
        if not all(ch in '0123456789+-*/(). ' for ch in expr):
            return QMessageBox.warning(self, 'تنبيه', 'العملية تحتوي رموز غير مدعومة.')
        try:
            result = eval(expr, {'__builtins__': {}}, {})
            number = float(result)
            if abs(number - int(number)) < 0.0000001:
                shown = str(int(number))
            else:
                shown = ('{:.2f}'.format(number)).rstrip('0').rstrip('.')
            self.calc_display.setText(shown)
        except ZeroDivisionError:
            QMessageBox.warning(self, 'تنبيه', 'ما يصير القسمة على صفر.')
        except Exception:
            QMessageBox.warning(self, 'تنبيه', 'العملية غير صحيحة.')

    def save_current_state(self):
        rec = self._state()
        rec['current_actual_cash'] = float(self.actual_cash.value())
        rec['current_notes'] = self.notes.toPlainText().strip()
        rec['actual_cash_anchor'] = float(self.actual_cash.value())
        rec['actual_cash_anchor_book'] = float(cash_balance(self.db) or 0)
        rec['actual_cash_anchor_date'] = now_str()
        self.save(); self.main.refresh_dashboard(); self.refresh_view()
        QMessageBox.information(self, 'تم', 'تم حفظ القيمة الحالية وملاحظات المطابقة.')

    def save_snapshot(self):
        rec = self._state()
        metrics = reconciliation_metrics(self.db, self.actual_cash.value())
        row = {
            'id': generate_id('rec'),
            'date': now_str(),
            'actual_cash': metrics['actual_cash'],
            'book_cash': metrics['book_cash'],
            'agents_base_custody': metrics['agents_base_custody'],
            'settlements_value': metrics['settlements_value'],
            'agents_custody': metrics['agents_custody'],
            'customer_dues': metrics['customer_dues'],
            'payables': metrics['payables'],
            'diff_cash': metrics['diff_cash'],
            'diff_with_agents': metrics['diff_with_agents'],
            'notes': self.notes.toPlainText().strip(),
        }
        rec['current_actual_cash'] = row['actual_cash']
        rec['current_notes'] = row['notes']
        rec['actual_cash_anchor'] = row['actual_cash']
        rec['actual_cash_anchor_book'] = row['book_cash']
        rec['actual_cash_anchor_date'] = row['date']
        rec.setdefault('records', []).insert(0, row)
        self.save(); self.refresh_table(); self.refresh_view(); self.main.refresh_dashboard()
        QMessageBox.information(self, 'تم', 'تم حفظ لقطة المطابقة بنجاح.')

    def delete_snapshot(self):
        r = self.table.currentRow()
        if r < 0:
            return QMessageBox.warning(self, 'تنبيه', 'حدد لقطة مطابقة للحذف')
        rows = self._state().get('records', [])
        if r >= len(rows):
            return
        if QMessageBox.question(self, 'تأكيد', 'حذف لقطة المطابقة المحددة؟') != QMessageBox.Yes:
            return
        rows.pop(r)
        rec = self._state()
        if rows:
            latest = rows[0]
            rec['actual_cash_anchor'] = float(latest.get('actual_cash', 0) or 0)
            rec['actual_cash_anchor_book'] = float(latest.get('book_cash', 0) or 0)
            rec['actual_cash_anchor_date'] = latest.get('date', '')
            rec['current_actual_cash'] = float(latest.get('actual_cash', 0) or 0)
            rec['current_notes'] = latest.get('notes', '') or ''
        else:
            rec['actual_cash_anchor'] = None
            rec['actual_cash_anchor_book'] = None
            rec['actual_cash_anchor_date'] = ''
            rec['current_actual_cash'] = 0.0
            rec['current_notes'] = ''
        self.save(); self.refresh_table(); self.refresh_view(); self.main.refresh_dashboard()

    def refresh_view(self):
        rec = self._state()
        sender = self.sender()
        if sender is None:
            self._loading = True
            self.actual_cash.setValue(float(rec.get('current_actual_cash', 0) or 0))
            self.notes.setPlainText(rec.get('current_notes', '') or '')
            self._loading = False
        metrics = reconciliation_metrics(self.db, self.actual_cash.value())
        for key, lbl in self.metric_labels.items():
            value = metrics.get(key, 0)
            if key in ('diff_cash', 'diff_with_agents'):
                color = '#22c55e' if abs(value) < 0.01 else ('#f59e0b' if value > 0 else '#ef4444')
                lbl.setStyleSheet(f'font-size:22px;font-weight:900;color:{color};')
            else:
                lbl.setStyleSheet('font-size:22px;font-weight:900;')
            lbl.setText(f'{fmt_money(value)} د.ع')
        self.summary.setText('الفرق النقدي = الكاش الفعلي - القاصة الدفترية. عهدة المندوبين النهائية = العهدة الأساسية + التسوية/معالجة الفرق. والفرق بعد العهدة والتسوية = (الكاش الفعلي + عهدة المندوبين النهائية) - القاصة الدفترية. إذا صار الناتج قريب للصفر فالوضع متوازن أكثر.')
        if sender is not None and not self._loading:
            rec['current_actual_cash'] = float(self.actual_cash.value())

    def refresh_table(self):
        rows = list(self._state().get('records', []))
        self.table.setRowCount(len(rows))
        for r, row in enumerate(rows):
            vals = [
                str(r+1), row.get('date', ''), fmt_money(row.get('actual_cash', 0)), fmt_money(row.get('book_cash', 0)),
                fmt_money(row.get('agents_base_custody', 0)), fmt_money(row.get('settlements_value', 0)), fmt_money(row.get('agents_custody', 0)),
                fmt_money(row.get('diff_cash', 0)), fmt_money(row.get('diff_with_agents', 0)), row.get('notes', '')
            ]
            for c, v in enumerate(vals):
                it = QTableWidgetItem(v)
                if c in (7, 8):
                    num = float(row.get('diff_cash', 0) if c == 7 else row.get('diff_with_agents', 0))
                    if abs(num) < 0.01:
                        it.setForeground(QColor('#22c55e'))
                    elif num < 0:
                        it.setForeground(QColor('#ef4444'))
                    else:
                        it.setForeground(QColor('#f59e0b'))
                self.table.setItem(r, c, it)

class ExpensesWindow(BaseWindow):
    def __init__(self, main):
        super().__init__(main, '🧾 المصاريف')
        self.subtitle_lbl.setText('قسمنا شاشة المصاريف حتى تبقى أوضح، وكبرنا الكتابة وحقول الإدخال حتى تصير القراءة والإملاء أسهل.')

        hero = QFrame(); hero.setStyleSheet(CARD_FRAME_STYLE)
        hero_box = QVBoxLayout(hero)
        hero_box.setContentsMargins(20, 18, 20, 18)
        hero_box.setSpacing(8)
        title = QLabel('لوحة المصاريف')
        title.setAlignment(Qt.AlignRight)
        title.setStyleSheet('font-size:20px;font-weight:900;background:transparent;border:none;')
        note = QLabel('أدخل المصروف من تبويب مستقل، وراجع السجل والملخص من تبويب ثاني حتى ما تبقى الشاشة مزحومة.')
        note.setAlignment(Qt.AlignRight); note.setWordWrap(True)
        note.setStyleSheet(f'font-size:13px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        hero_box.addWidget(title); hero_box.addWidget(note)
        self.layout.addWidget(hero)

        self.tabs = style_funders_tabs(QTabWidget())
        self.layout.addWidget(self.tabs, 1)

        add_tab = QWidget(); add_layout = QVBoxLayout(add_tab); add_layout.setContentsMargins(8,8,8,8); add_layout.setSpacing(12)
        form_card = QFrame(); form_card.setStyleSheet(CARD_FRAME_STYLE)
        form_wrap = QVBoxLayout(form_card)
        form_wrap.setContentsMargins(18,18,18,18)
        form_wrap.setSpacing(10)
        form_title = QLabel('إضافة مصروف جديد')
        form_title.setAlignment(Qt.AlignRight)
        form_title.setStyleSheet('font-size:17px;font-weight:900;background:transparent;border:none;')
        form_wrap.addWidget(form_title)
        form = QFormLayout(); form.setLabelAlignment(Qt.AlignRight); form.setFormAlignment(Qt.AlignRight); form.setHorizontalSpacing(12); form.setVerticalSpacing(12)
        self.date = QDateEdit(); fix_date_edit_widget(self.date); self.date.setDate(QDate.currentDate())
        self.category = QComboBox(); self.category.setEditable(True); self.category.addItems(['إيجار','رواتب','نقل','كهرباء','صيانة','اتصالات','مصاريف أخرى'])
        self.category.setStyleSheet('font-size:14px;')
        self.amount = tune_numeric_widget(QDoubleSpinBox()); self.amount.setRange(0, 1_000_000_000); self.amount.setDecimals(0)
        self.notes = QTextEdit(); self.notes.setFixedHeight(110); self.notes.setPlaceholderText('اكتب الملاحظات أو سبب المصروف بشكل أوضح'); self.notes.setStyleSheet('font-size:15px; font-weight:600;')
        form.addRow('التاريخ:', self.date)
        form.addRow('النوع:', self.category)
        form.addRow('المبلغ:', self.amount)
        form_wrap.addLayout(form)
        notes_lbl = QLabel('الملاحظات / الإملاء')
        notes_lbl.setAlignment(Qt.AlignRight)
        notes_lbl.setStyleSheet(f'font-size:13px;font-weight:800;color:{MUTED};background:transparent;border:none;')
        form_wrap.addWidget(notes_lbl)
        form_wrap.addWidget(self.notes)
        add_layout.addWidget(form_card)

        actions_card = QFrame(); actions_card.setStyleSheet(CARD_FRAME_STYLE)
        actions_wrap = QVBoxLayout(actions_card)
        actions_wrap.setContentsMargins(18,18,18,18)
        actions_wrap.setSpacing(10)
        actions_title = QLabel('إجراءات المصاريف')
        actions_title.setAlignment(Qt.AlignRight)
        actions_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        actions_wrap.addWidget(actions_title)
        add = QPushButton('➕ إضافة مصروف'); edit = QPushButton('✏️ تعديل المصروف'); delete = QPushButton('🗑 حذف المصروف')
        add.setStyleSheet(BUTTON_STYLE); edit.setStyleSheet(SECONDARY_BUTTON); delete.setStyleSheet(SECONDARY_BUTTON)
        add.setMinimumHeight(48); edit.setMinimumHeight(48); delete.setMinimumHeight(48)
        add.clicked.connect(self.add_row); edit.clicked.connect(self.edit_row); delete.clicked.connect(self.delete_row)
        actions_wrap.addWidget(add); actions_wrap.addWidget(edit); actions_wrap.addWidget(delete)
        action_note = QLabel('كل مصروف يُرحّل تلقائيًا إلى الصندوق وينزل من الربح، لذلك الحذف يعكس أثره أيضًا.')
        action_note.setWordWrap(True); action_note.setAlignment(Qt.AlignRight)
        action_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        actions_wrap.addWidget(action_note)
        add_layout.addWidget(actions_card)
        add_layout.addStretch(1)

        list_tab = QWidget(); list_layout = QVBoxLayout(list_tab); list_layout.setContentsMargins(8,8,8,8); list_layout.setSpacing(12)
        summary_card = QFrame(); summary_card.setStyleSheet(CARD_FRAME_STYLE)
        summary_wrap = QVBoxLayout(summary_card)
        summary_wrap.setContentsMargins(16,14,16,14)
        summary_wrap.setSpacing(6)
        sum_title = QLabel('ملخص المصاريف')
        sum_title.setAlignment(Qt.AlignRight)
        sum_title.setStyleSheet('font-size:15px;font-weight:900;background:transparent;border:none;')
        self.summary = QLabel(); self.summary.setAlignment(Qt.AlignRight); self.summary.setWordWrap(True)
        self.summary.setStyleSheet(f'font-size:13px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        summary_wrap.addWidget(sum_title); summary_wrap.addWidget(self.summary)
        list_layout.addWidget(summary_card)
        table_card = QFrame(); table_card.setStyleSheet(CARD_FRAME_STYLE)
        table_wrap = QVBoxLayout(table_card)
        table_wrap.setContentsMargins(16,16,16,16)
        table_wrap.setSpacing(10)
        table_title = QLabel('سجل المصاريف')
        table_title.setAlignment(Qt.AlignRight)
        table_title.setStyleSheet('font-size:16px;font-weight:900;background:transparent;border:none;')
        table_wrap.addWidget(table_title)
        self.table = QTableWidget(); self.table.setStyleSheet(TABLE_STYLE); self.table.setColumnCount(6)
        self.table.setHorizontalHeaderLabels(['#','التاريخ','النوع','المبلغ','ملاحظات','مرحل للصندوق'])
        self.table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.table.setSelectionBehavior(QAbstractItemView.SelectRows); self.table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table_wrap.addWidget(self.table)
        list_layout.addWidget(table_card, 1)

        self.tabs.addTab(add_tab, 'إضافة مصروف')
        self.tabs.addTab(list_tab, 'سجل المصاريف')
        self.table.itemSelectionChanged.connect(self.load_selected_row)
        self.table.itemDoubleClicked.connect(lambda *_: self.load_selected_row(switch_tab=True))
        self.refresh_table()

    def refresh_table(self):
        rows = self.db.get('expenses', [])
        self.table.setRowCount(len(rows))
        total = 0.0
        for r, row in enumerate(rows):
            amount = float(row.get('amount', 0) or 0); total += amount
            vals = [r+1, row.get('date',''), row.get('category',''), fmt_money(amount), row.get('notes',''), 'نعم' if row.get('cash_synced', True) else 'لا']
            for c, v in enumerate(vals): self.table.setItem(r, c, QTableWidgetItem(str(v)))
        self.summary.setText(f'عدد المصاريف: {len(rows)} | مجموع المصاريف: {fmt_money(total)} د.ع | هذه المصاريف تنزل من الربح والصندوق تلقائيًا')

    def add_row(self):
        amount = float(self.amount.value())
        category = self.category.currentText().strip()
        if amount <= 0 or not category:
            return QMessageBox.warning(self, 'تنبيه', 'أكمل بيانات المصروف')
        exp_id = generate_id('exp')
        row = {'id': exp_id, 'date': self.date.date().toString('yyyy-MM-dd'), 'category': category, 'amount': amount, 'notes': self.notes.toPlainText().strip(), 'created_at': now_str(), 'cash_synced': True}
        self.db.setdefault('expenses', []).append(row)
        self.db['cash'].append({'date': row['date'], 'type': 'مصروف', 'category': category, 'party': 'مصروف تشغيلي', 'amount': amount, 'notes': row['notes'], 'source': 'expense', 'ref_id': exp_id, 'created_at': now_str()})
        self.amount.setValue(0); self.notes.clear(); self.save(); self.refresh_table()

    def load_selected_row(self, switch_tab=False):
        r = self.table.currentRow()
        rows = self.db.get('expenses', [])
        if r < 0 or r >= len(rows):
            return
        row = rows[r]
        qd = QDate.fromString(str(row.get('date', '') or ''), 'yyyy-MM-dd')
        self.date.setDate(qd if qd.isValid() else QDate.currentDate())
        self.category.setCurrentText(str(row.get('category', '') or ''))
        self.amount.setValue(float(row.get('amount', 0) or 0))
        self.notes.setPlainText(str(row.get('notes', '') or ''))
        if switch_tab:
            self.tabs.setCurrentIndex(0)

    def edit_row(self):
        r = self.table.currentRow()
        rows = self.db.get('expenses', [])
        if r < 0 or r >= len(rows):
            return QMessageBox.warning(self, 'تنبيه', 'اختر مصروفًا من السجل أولًا')
        amount = float(self.amount.value())
        category = self.category.currentText().strip()
        if amount <= 0 or not category:
            return QMessageBox.warning(self, 'تنبيه', 'أكمل بيانات المصروف')
        row = rows[r]
        row['date'] = self.date.date().toString('yyyy-MM-dd')
        row['category'] = category
        row['amount'] = amount
        row['notes'] = self.notes.toPlainText().strip()
        row['updated_at'] = now_str()
        ref_id = row.get('id')
        for cash_row in self.db.get('cash', []):
            if cash_row.get('source') == 'expense' and cash_row.get('ref_id') == ref_id:
                cash_row['date'] = row['date']
                cash_row['category'] = category
                cash_row['amount'] = amount
                cash_row['notes'] = row['notes']
                cash_row['updated_at'] = now_str()
        self.save(); self.refresh_table()
        QMessageBox.information(self, 'تم', 'تم تعديل المصروف بنجاح')

    def delete_row(self):
        r = self.table.currentRow()
        if r < 0:
            return QMessageBox.warning(self, 'تنبيه', 'اختر مصروفًا')
        row = self.db.get('expenses', [])[r]
        if QMessageBox.question(self, 'تأكيد', 'حذف المصروف وعكس أثره على الصندوق والربح؟') != QMessageBox.Yes:
            return
        for i in range(len(self.db['cash'])-1, -1, -1):
            c = self.db['cash'][i]
            if c.get('source') == 'expense' and c.get('ref_id') == row.get('id'):
                self.db['cash'].pop(i)
        self.db['expenses'].pop(r)
        self.save(); self.refresh_table()


class OpeningBalancesWindow(BaseWindow):
    def __init__(self, main):
        super().__init__(main, '🧩 التهيئة الافتتاحية')
        op = opening_data(self.db)
        self.subtitle_lbl.setText('واجهة التهيئة الافتتاحية صارت هادئة ومقسمة مثل الموردين: أساسيات واضحة، ممولون قدامى، وديون قديمة بسجل منظم.')

        hero = QFrame(); hero.setStyleSheet(CARD_FRAME_STYLE)
        hero_box = QVBoxLayout(hero); hero_box.setContentsMargins(18,16,18,16); hero_box.setSpacing(6)
        hero_title = QLabel('تهيئة البداية')
        hero_title.setAlignment(Qt.AlignRight)
        hero_title.setStyleSheet('font-size:20px;font-weight:900;background:transparent;border:none;')
        hero_note = QLabel('أدخل الرصيد الافتتاحي اليدوي، ووزع الشغل على تبويبات منفصلة للممولين القدامى والديون القديمة حتى تبقى الصفحة مرتبة.')
        hero_note.setWordWrap(True); hero_note.setAlignment(Qt.AlignRight)
        hero_note.setStyleSheet(f'font-size:12px;font-weight:700;color:{MUTED};background:transparent;border:none;')
        hero_box.addWidget(hero_title); hero_box.addWidget(hero_note)
        self.layout.addWidget(hero)

        self.tabs = QTabWidget()
        self.tabs.setTabPosition(QTabWidget.North)
        self.tabs.setDocumentMode(True)
        self.tabs.setStyleSheet(
            f"QTabWidget::pane {{border:1px solid {BORDER};background:{CARD};border-radius:14px; margin-top:8px;}}"
            f"QTabBar::tab {{min-width:190px;min-height:44px;margin:4px 6px;padding:8px 14px;border-radius:10px;background:{DARK};color:{TEXT};font-weight:800;}}"
            f"QTabBar::tab:selected {{background:{ACCENT};color:{TEXT_ON_ACCENT};}}"
            f"QTabBar::tab:hover {{background:{ACCENT2};color:{TEXT_ON_ACCENT};}}"
        )
        self.layout.addWidget(self.tabs, 1)

        def make_page(title, note=''):
            page = QWidget()
            page_layout = QVBoxLayout(page)
            page_layout.setContentsMargins(14, 14, 14, 14)
            page_layout.setSpacing(12)
            card = QFrame(); card.setStyleSheet(CARD_FRAME_STYLE)
            card_layout = QVBoxLayout(card)
            card_layout.setContentsMargins(16, 16, 16, 16)
            card_layout.setSpacing(10)
            title_lbl = QLabel(title)
            title_lbl.setAlignment(Qt.AlignRight)
            title_lbl.setStyleSheet('font-size:18px;font-weight:900;background:transparent;border:none;')
            card_layout.addWidget(title_lbl)
            if note:
                note_lbl = QLabel(note)
                note_lbl.setWordWrap(True)
                note_lbl.setAlignment(Qt.AlignRight)
                note_lbl.setStyleSheet(f'color:{MUTED};font-size:12px;font-weight:700;background:transparent;border:none;')
                card_layout.addWidget(note_lbl)
            page_layout.addWidget(card)
            page_layout.addStretch(1)
            return page, page_layout, card_layout

        basics_page, _, basics_l = make_page('الأرصدة الأساسية', 'الرصيد التشغيلي الجديد يُحسب تلقائيًا من صافي الممولين القدامى بعد تنزيل السحوبات فقط، ويظهر هنا للمتابعة.')
        cash_form = QFormLayout()
        cash_form.setLabelAlignment(Qt.AlignRight)
        cash_form.setFormAlignment(Qt.AlignTop)
        cash_form.setHorizontalSpacing(20)
        cash_form.setVerticalSpacing(12)

        self.start_date = QDateEdit(); self.start_date.setCalendarPopup(True)
        if op.get('start_date'):
            self.start_date.setDate(QDate.fromString(op.get('start_date'), 'yyyy-MM-dd'))
        else:
            self.start_date.setDate(QDate.currentDate())

        self.opening_cash = tune_numeric_widget(QDoubleSpinBox()); self.opening_cash.setRange(0, 1_000_000_000); self.opening_cash.setDecimals(0)
        self.opening_cash.setValue(float(op.get('opening_cash', 0) or 0))

        self.operating_cash = tune_numeric_widget(QDoubleSpinBox()); self.operating_cash.setRange(0, 1_000_000_000); self.operating_cash.setDecimals(0)
        self.operating_cash.setReadOnly(True); self.operating_cash.setButtonSymbols(QDoubleSpinBox.NoButtons)
        self.operating_cash.setToolTip('يُحتسب تلقائيًا من صافي الممولين القدامى بعد تنزيل سحوباتهم')

        self.hidab_opening_balance = tune_numeric_widget(QDoubleSpinBox()); self.hidab_opening_balance.setRange(-1_000_000_000, 1_000_000_000); self.hidab_opening_balance.setDecimals(0)
        self.hidab_opening_balance.setValue(float(op.get('hidab_opening_balance', 0) or 0))
        self.mustafa_opening_balance = tune_numeric_widget(QDoubleSpinBox()); self.mustafa_opening_balance.setRange(-1_000_000_000, 1_000_000_000); self.mustafa_opening_balance.setDecimals(0)
        self.mustafa_opening_balance.setValue(float(op.get('mustafa_opening_balance', 0) or 0))
        self.lock_state_lbl = QLabel()
        self.lock_state_lbl.setAlignment(Qt.AlignRight)
        self.lock_state_lbl.setStyleSheet(f'color:{MUTED};font-size:12px;')

        cash_form.addRow('تاريخ بداية التشغيل', self.start_date)
        cash_form.addRow('الرصيد الافتتاحي اليدوي', self.opening_cash)
        cash_form.addRow('الرصيد التشغيلي الجديد (تلقائي)', self.operating_cash)
        cash_form.addRow('رصيد/عجز هضاب الافتتاحي', self.hidab_opening_balance)
        cash_form.addRow('رصيد/عجز مصطفى الافتتاحي', self.mustafa_opening_balance)
        basics_l.addLayout(cash_form)

        save_btn = QPushButton('💾 حفظ التهيئة'); save_btn.clicked.connect(self.save_cash)
        save_btn.setStyleSheet(BUTTON_STYLE)
        lock_btn = QPushButton('🔒 اعتماد / فتح التهيئة'); lock_btn.clicked.connect(self.toggle_opening_lock)
        lock_btn.setStyleSheet(SECONDARY_BUTTON)
        actions = QHBoxLayout()
        actions.addWidget(save_btn)
        actions.addWidget(lock_btn)
        actions.addStretch(1)
        basics_l.addLayout(actions)

        self.cash_summary = QLabel(); self.cash_summary.setAlignment(Qt.AlignRight); self.cash_summary.setWordWrap(True)
        basics_l.addWidget(self.cash_summary)
        basics_l.addWidget(self.lock_state_lbl)
        self.tabs.addTab(basics_page, 'الأرصدة الأساسية')

        fund_page, _, fund_l = make_page('الممولون القدامى', 'المطلوب هنا فقط: رأس المال القديم والسحوبات القديمة. صافي الممول القديم = رأس المال - السحوبات، وهو الأساس في احتساب الرصيد التشغيلي الجديد.')
        fund_form = QFormLayout()
        fund_form.setLabelAlignment(Qt.AlignRight)
        fund_form.setHorizontalSpacing(20)
        fund_form.setVerticalSpacing(12)
        self.old_funder = QComboBox(); self.old_funder.setEditable(False)
        self.old_funder_capital = QDoubleSpinBox(); self.old_funder_capital.setRange(0, 1_000_000_000); self.old_funder_capital.setDecimals(0)
        self.old_funder_withdrawals = QDoubleSpinBox(); self.old_funder_withdrawals.setRange(0, 1_000_000_000); self.old_funder_withdrawals.setDecimals(0)
        old_funders_tabs = style_funders_tabs(QTabWidget()); old_funders_tabs.setDocumentMode(True)
        old_funders_tabs.setStyleSheet(self.tabs.styleSheet())
        fund_entry = QWidget(); fund_entry_l = QVBoxLayout(fund_entry); fund_entry_l.setContentsMargins(4,4,4,4); fund_entry_l.setSpacing(12)
        fund_form.addRow('الممول', self.old_funder)
        fund_form.addRow('رأس المال القديم', self.old_funder_capital)
        fund_form.addRow('السحوبات القديمة', self.old_funder_withdrawals)
        fund_entry_l.addLayout(fund_form)

        fund_actions = QHBoxLayout()
        add_old_funder = QPushButton('➕ حفظ / تحديث ممول قديم'); add_old_funder.clicked.connect(self.add_old_funder)
        add_old_funder.setStyleSheet(BUTTON_STYLE)
        fill_old_funder = QPushButton('↩ تحميل السطر المحدد'); fill_old_funder.clicked.connect(self.load_selected_old_funder)
        del_old_funder = QPushButton('🗑 حذف ممول قديم'); del_old_funder.clicked.connect(self.delete_old_funder)
        for b in (add_old_funder, fill_old_funder, del_old_funder):
            fund_actions.addWidget(b)
        fund_actions.addStretch(1)
        fund_entry_l.addLayout(fund_actions)
        fund_entry_l.addStretch(1)

        fund_table_tab = QWidget(); fund_table_l = QVBoxLayout(fund_table_tab); fund_table_l.setContentsMargins(4,4,4,4); fund_table_l.setSpacing(12)
        self.old_funders_table = QTableWidget(); self.old_funders_table.setColumnCount(5)
        self.old_funders_table.setHorizontalHeaderLabels(['#','الممول','رأس المال','السحوبات','الصافي المتبقي'])
        self.old_funders_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.old_funders_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.old_funders_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.old_funders_table.setMinimumHeight(260)
        fund_table_l.addWidget(self.old_funders_table)
        self.old_funders_summary = QLabel(); self.old_funders_summary.setAlignment(Qt.AlignRight); self.old_funders_summary.setWordWrap(True)
        fund_table_l.addWidget(self.old_funders_summary)
        old_funders_tabs.addTab(fund_entry, 'الإدخال')
        old_funders_tabs.addTab(fund_table_tab, 'السجل')
        fund_l.addWidget(old_funders_tabs)
        self.tabs.addTab(fund_page, 'الممولون القدامى')

        cust_page, _, cust_l = make_page('ديون الزبائن القديمة', 'أضف أسماء الزبائن من صفحة الزبائن أولاً، ثم اختر الاسم هنا لحفظ الدين الافتتاحي القديم فقط.')
        cust_form = QFormLayout()
        cust_form.setLabelAlignment(Qt.AlignRight)
        cust_form.setHorizontalSpacing(20)
        cust_form.setVerticalSpacing(12)
        self.old_customer = QComboBox(); self.old_customer.setEditable(False)
        self.old_customer_amount = QDoubleSpinBox(); self.old_customer_amount.setRange(0, 1_000_000_000); self.old_customer_amount.setDecimals(0)
        old_customers_tabs = style_funders_tabs(QTabWidget()); old_customers_tabs.setDocumentMode(True)
        old_customers_tabs.setStyleSheet(self.tabs.styleSheet())
        cust_entry = QWidget(); cust_entry_l = QVBoxLayout(cust_entry); cust_entry_l.setContentsMargins(4,4,4,4); cust_entry_l.setSpacing(12)
        cust_form.addRow('الزبون', self.old_customer)
        cust_form.addRow('المبلغ', self.old_customer_amount)
        cust_entry_l.addLayout(cust_form)
        cust_actions = QHBoxLayout()
        add_old_customer = QPushButton('➕ حفظ / تحديث'); add_old_customer.clicked.connect(self.add_old_customer_due); add_old_customer.setStyleSheet(BUTTON_STYLE)
        del_old_customer = QPushButton('🗑 حذف'); del_old_customer.clicked.connect(self.delete_old_customer_due)
        cust_actions.addWidget(add_old_customer); cust_actions.addWidget(del_old_customer); cust_actions.addStretch(1)
        cust_entry_l.addLayout(cust_actions)
        cust_entry_l.addStretch(1)
        cust_table_tab = QWidget(); cust_table_l = QVBoxLayout(cust_table_tab); cust_table_l.setContentsMargins(4,4,4,4); cust_table_l.setSpacing(12)
        self.old_customers_table = QTableWidget(); self.old_customers_table.setColumnCount(3)
        self.old_customers_table.setHorizontalHeaderLabels(['#','الزبون','الدين القديم'])
        self.old_customers_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.old_customers_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.old_customers_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.old_customers_table.setMinimumHeight(280)
        cust_table_l.addWidget(self.old_customers_table)
        old_customers_tabs.addTab(cust_entry, 'الإدخال')
        old_customers_tabs.addTab(cust_table_tab, 'السجل')
        cust_l.addWidget(old_customers_tabs)
        self.tabs.addTab(cust_page, 'ديون الزبائن القديمة')

        self.refresh_combos()
        self.refresh_tables()

    def _opening_inputs(self):
        return [
            self.start_date, self.opening_cash, self.hidab_opening_balance, self.mustafa_opening_balance,
            self.old_customer, self.old_customer_amount,
            self.old_funder, self.old_funder_capital, self.old_funder_withdrawals,
        ]

    def apply_opening_lock_state(self):
        locked = bool(opening_data(self.db).get('opening_locked', False))
        for w in self._opening_inputs():
            try:
                w.setEnabled(not locked)
            except Exception:
                pass
        self.lock_state_lbl.setText('حالة التهيئة: معتمدة ومقفلة' if locked else 'حالة التهيئة: مفتوحة للتعديل')

    def toggle_opening_lock(self):
        op = opening_data(self.db)
        op['opening_locked'] = not bool(op.get('opening_locked', False))
        self.save()
        self.apply_opening_lock_state()
        QMessageBox.information(self, 'تم', 'تم تحديث حالة قفل التهيئة الافتتاحية.')

    def refresh_combos(self):
        if hasattr(self, 'old_customer'):
            self.old_customer.clear()
            self.old_customer.addItems([x.get('name','') for x in self.db.get('customers', []) if x.get('name','')])
        if hasattr(self, 'old_funder'):
            self.old_funder.clear()
            self.old_funder.addItems([x.get('name','') for x in self.db.get('funders', []) if x.get('name','')])

    def refresh_tables(self):
        op = opening_data(self.db)
        self.refresh_combos()

        auto_operating = opening_old_funders_operating_balance(self.db)
        self.operating_cash.blockSignals(True)
        self.operating_cash.setValue(auto_operating)
        self.operating_cash.blockSignals(False)

        cust_rows = op.get('customers', [])
        self.old_customers_table.setRowCount(len(cust_rows))
        for r, row in enumerate(cust_rows):
            vals = [r + 1, row.get('name',''), fmt_money(row.get('amount', 0))]
            for c, v in enumerate(vals):
                self.old_customers_table.setItem(r, c, QTableWidgetItem(str(v)))

        fund_rows = op.get('old_funders', [])
        self.old_funders_table.setRowCount(len(fund_rows))
        for r, row in enumerate(fund_rows):
            vals = [
                r + 1,
                row.get('name',''),
                fmt_money(row.get('capital', 0)),
                fmt_money(row.get('withdrawals', 0)),
                fmt_money(old_funder_net_amount(row)),
            ]
            for c, v in enumerate(vals):
                self.old_funders_table.setItem(r, c, QTableWidgetItem(str(v)))

        total_opening = float(op.get('opening_cash', 0) or 0)
        total_customer_dues = sum(float(x.get('amount', 0) or 0) for x in cust_rows)
        self.cash_summary.setText(
            f'الرصيد الافتتاحي اليدوي: {fmt_money(float(op.get("opening_cash", 0) or 0))} د.ع | '
            f'الرصيد التشغيلي الجديد من صافي الممولين القدامى: {fmt_money(auto_operating)} د.ع | '
            f'رصيد/عجز هضاب الافتتاحي: {fmt_money(float(op.get("hidab_opening_balance", 0) or 0))} د.ع | '
            f'رصيد/عجز مصطفى الافتتاحي: {fmt_money(float(op.get("mustafa_opening_balance", 0) or 0))} د.ع | '
            f'إجمالي بداية الصندوق: {fmt_money(total_opening)} د.ع'
        )
        self.old_funders_summary.setText(
            f'عدد الممولين القدامى: {len(fund_rows)} | '
            f'صافي المبالغ المتبقية للممولين: {fmt_money(auto_operating)} د.ع | '
            f'إجمالي ديون الزبائن القديمة المحفوظة: {fmt_money(total_customer_dues)} د.ع'
        )
        self.apply_opening_lock_state()

    def save_cash(self):
        op = opening_data(self.db)
        if bool(op.get('opening_locked', False)):
            return QMessageBox.warning(self, 'تنبيه', 'التهيئة الافتتاحية مقفلة. افتحها أولاً إذا تريد التعديل.')
        op['opening_cash'] = float(self.opening_cash.value())
        op['operating_cash'] = opening_old_funders_operating_balance(self.db)
        op['start_date'] = self.start_date.date().toString('yyyy-MM-dd')
        op['hidab_opening_balance'] = float(self.hidab_opening_balance.value())
        op['mustafa_opening_balance'] = float(self.mustafa_opening_balance.value())
        self.save()
        self.refresh_tables()
        QMessageBox.information(self, 'تم', 'تم حفظ التهيئة الافتتاحية بنجاح.')

    def add_old_customer_due(self):
        if bool(opening_data(self.db).get('opening_locked', False)):
            return QMessageBox.warning(self, 'تنبيه', 'التهيئة الافتتاحية مقفلة. افتحها أولاً إذا تريد التعديل.')
        name = self.old_customer.currentText().strip(); amount = float(self.old_customer_amount.value())
        if not name or amount <= 0:
            return QMessageBox.warning(self, 'تنبيه', 'اختر الزبون وأدخل المبلغ')
        rows = opening_data(self.db).setdefault('customers', [])
        for row in rows:
            if row.get('name','') == name:
                row['amount'] = amount
                break
        else:
            rows.append({'name': name, 'amount': amount, 'created_at': now_str()})
        self.old_customer_amount.setValue(0)
        self.save(); self.refresh_tables()

    def delete_old_customer_due(self):
        if bool(opening_data(self.db).get('opening_locked', False)):
            return QMessageBox.warning(self, 'تنبيه', 'التهيئة الافتتاحية مقفلة. افتحها أولاً إذا تريد التعديل.')
        r = self.old_customers_table.currentRow()
        if r < 0: return QMessageBox.warning(self, 'تنبيه', 'اختر سطرًا')
        opening_data(self.db).get('customers', []).pop(r)
        self.save(); self.refresh_tables()

    def add_old_funder(self):
        if bool(opening_data(self.db).get('opening_locked', False)):
            return QMessageBox.warning(self, 'تنبيه', 'التهيئة الافتتاحية مقفلة. افتحها أولاً إذا تريد التعديل.')
        name = self.old_funder.currentText().strip()
        if not name:
            return QMessageBox.warning(self, 'تنبيه', 'أضف الممول من صفحة الممولين أولاً ثم اختره هنا')
        capital = float(self.old_funder_capital.value())
        withdrawals = float(self.old_funder_withdrawals.value())
        if capital <= 0 and withdrawals <= 0:
            return QMessageBox.warning(self, 'تنبيه', 'أدخل بيانات الممول القديم')
        rows = opening_data(self.db).setdefault('old_funders', [])
        for row in rows:
            if row.get('name','') == name:
                row.update({'capital': capital, 'withdrawals': withdrawals, 'updated_at': now_str()})
                break
        else:
            rows.append({'name': name, 'capital': capital, 'withdrawals': withdrawals, 'created_at': now_str()})
        self.old_funder_capital.setValue(0)
        self.old_funder_withdrawals.setValue(0)
        self.save(); self.refresh_tables()

    def load_selected_old_funder(self):
        r = self.old_funders_table.currentRow()
        rows = opening_data(self.db).get('old_funders', [])
        if r < 0 or r >= len(rows):
            return QMessageBox.warning(self, 'تنبيه', 'اختر ممولًا قديمًا')
        row = rows[r]
        idx = self.old_funder.findText(row.get('name',''))
        if idx >= 0:
            self.old_funder.setCurrentIndex(idx)
        self.old_funder_capital.setValue(float(row.get('capital', 0) or 0))
        self.old_funder_withdrawals.setValue(float(row.get('withdrawals', 0) or 0))

    def delete_old_funder(self):
        if bool(opening_data(self.db).get('opening_locked', False)):
            return QMessageBox.warning(self, 'تنبيه', 'التهيئة الافتتاحية مقفلة. افتحها أولاً إذا تريد التعديل.')
        r = self.old_funders_table.currentRow()
        rows = opening_data(self.db).get('old_funders', [])
        if r < 0 or r >= len(rows):
            return QMessageBox.warning(self, 'تنبيه', 'اختر سطرًا')
        rows.pop(r)
        self.save(); self.refresh_tables()



class WithdrawalsWindow(BaseWindow):
    def __init__(self, main):
        super().__init__(main, '💸 السحوبات')
        self.subtitle_lbl.setText('واجهة رئيسية للسحوبات بعناوين واضحة، وكل قسم يفتح صفحته الخاصة للإضافة والعرض والحذف بنفس روح النظام.')

        self.person_pages = {}
        self.stack = QStackedWidget()
        self.layout.addWidget(self.stack, 1)

        self.home_page = self._build_home_page()
        self.stack.addWidget(self.home_page)

        self.hidab_page, self.hidab_table, self.hidab_summary = self._build_person_page('هضاب')
        self.mustafa_page, self.mustafa_table, self.mustafa_summary = self._build_person_page('مصطفى')
        self.stack.addWidget(self.hidab_page)
        self.stack.addWidget(self.mustafa_page)

        self.person_pages.setdefault('هضاب', {})
        self.person_pages['هضاب'].update({'index': 1, 'table': self.hidab_table, 'summary': self.hidab_summary})
        self.person_pages.setdefault('مصطفى', {})
        self.person_pages['مصطفى'].update({'index': 2, 'table': self.mustafa_table, 'summary': self.mustafa_summary})

        self.refresh_table()
        self.show_home()

    def _build_home_page(self):
        page = QWidget()
        lay = QVBoxLayout(page)
        lay.setContentsMargins(0, 0, 0, 0)
        lay.setSpacing(18)

        title = QLabel('السحوبات')
        title.setAlignment(Qt.AlignRight)
        title.setStyleSheet(f'font-size:26px;font-weight:900;color:{TEXT};padding:2px 6px;')
        lay.addWidget(title)

        subtitle = QLabel('اختر الصفحة المطلوبة داخل هذا القسم. كل صفحة مستقلة وبيها رجوع واضح فقط.')
        subtitle.setAlignment(Qt.AlignRight)
        subtitle.setWordWrap(True)
        subtitle.setStyleSheet(f'font-size:14px;font-weight:700;color:{MUTED};padding:0 6px;')
        lay.addWidget(subtitle)

        grid = QGridLayout()
        grid.setHorizontalSpacing(16)
        grid.setVerticalSpacing(16)

        hidab_btn = self._create_home_card('سحب هضاب', 'إضافة سحوبات هضاب ومراجعة سجلها', 'هضاب')
        mustafa_btn = self._create_home_card('سحب مصطفى', 'إضافة سحوبات مصطفى ومراجعة سجلها', 'مصطفى')
        grid.addWidget(hidab_btn, 0, 0)
        grid.addWidget(mustafa_btn, 0, 1)
        lay.addLayout(grid)

        self.home_summary = QLabel()
        self.home_summary.setAlignment(Qt.AlignRight)
        self.home_summary.setWordWrap(True)
        self.home_summary.setStyleSheet(
            f"background:{rgba_from_hex(CARD,0.90)}; border:1px solid {rgba_from_hex(TEXT,0.08)}; border-radius:20px; padding:14px 16px; font-size:14px; font-weight:800; color:{TEXT};"
        )
        lay.addWidget(self.home_summary)
        lay.addStretch(1)
        return page

    def _create_home_card(self, title, subtitle, person_name):
        btn = QPushButton()
        btn.setCursor(Qt.PointingHandCursor)
        btn.setMinimumHeight(180)
        btn.setStyleSheet(
            f"QPushButton{{text-align:right; background:qlineargradient(x1:0,y1:0,x2:1,y2:1, stop:0 {rgba_from_hex(CARD,0.98)}, stop:1 {rgba_from_hex(DARK,0.94)});"
            f"border:1px solid {rgba_from_hex(TEXT,0.08)}; border-radius:26px; padding:18px; color:{TEXT};}}"
            f"QPushButton:hover{{border:1px solid {rgba_from_hex(ACCENT,0.34)}; background:qlineargradient(x1:0,y1:0,x2:1,y2:1, stop:0 {rgba_from_hex(CARD,1.0)}, stop:1 {rgba_from_hex(ACCENT,0.18)});}}"
            f"QPushButton:pressed{{padding-top:20px; padding-bottom:16px;}}"
        )
        btn.setText(f'{title}\n{subtitle}')
        btn.clicked.connect(lambda _=False, p=person_name: self.show_person_page(p))
        return btn

    def _build_person_page(self, person_name):
        page = QWidget()
        root = QVBoxLayout(page)
        root.setContentsMargins(0, 0, 0, 0)
        root.setSpacing(14)

        top = QHBoxLayout()
        top.setSpacing(10)
        back_btn = QPushButton('↩ رجوع')
        back_btn.setMinimumHeight(44)
        back_btn.setStyleSheet(SECONDARY_BUTTON)
        back_btn.clicked.connect(self.show_home)
        top.addWidget(back_btn, 0)
        top.addStretch(1)

        title = QLabel(f'سحب {person_name}')
        title.setAlignment(Qt.AlignRight)
        title.setStyleSheet(f'font-size:24px;font-weight:900;color:{TEXT};padding:2px 6px;')
        top.addWidget(title, 0)
        root.addLayout(top)

        form_card = QFrame()
        form_card.setStyleSheet(
            f"QFrame{{background:{rgba_from_hex(CARD,0.90)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:24px;}}"
        )
        form_wrap = QVBoxLayout(form_card)
        form_wrap.setContentsMargins(18, 16, 18, 16)
        form_wrap.setSpacing(12)

        grid = QGridLayout()
        grid.setHorizontalSpacing(12)
        grid.setVerticalSpacing(10)

        date = QDateEdit(); fix_date_edit_widget(date); date.setDate(QDate.currentDate()); date.setMinimumHeight(48)
        amount = QDoubleSpinBox(); amount.setRange(0, 1_000_000_000); amount.setDecimals(0); amount.setMinimumHeight(48)
        notes = QLineEdit(); notes.setPlaceholderText(f'ملاحظات سحب {person_name}'); notes.setMinimumHeight(48)
        add_btn = QPushButton('➕ تسجيل سحب'); add_btn.setMinimumHeight(48); add_btn.setStyleSheet(BUTTON_STYLE)
        delete_btn = QPushButton('🗑 حذف السحب المحدد'); delete_btn.setMinimumHeight(48); delete_btn.setStyleSheet(SECONDARY_BUTTON)

        controls = [('التاريخ', date, 0, 0), ('المبلغ', amount, 0, 1), ('الملاحظات', notes, 1, 0)]
        for txt, w, r, c in controls:
            label = QLabel(txt)
            label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
            label.setStyleSheet(f'font-size:13px;font-weight:800;color:{MUTED};padding:0 4px;')
            box = QVBoxLayout(); box.setSpacing(6); box.addWidget(label); box.addWidget(w)
            if w is notes:
                grid.addLayout(box, r, c, 1, 2)
            else:
                grid.addLayout(box, r, c)

        btns = QHBoxLayout(); btns.setSpacing(10)
        btns.addWidget(delete_btn)
        btns.addWidget(add_btn)
        grid.addLayout(btns, 1, 2)
        form_wrap.addLayout(grid)
        root.addWidget(form_card)

        table = QTableWidget(); table.setColumnCount(5)
        table.setHorizontalHeaderLabels(['#','التاريخ','المبلغ','الملاحظات','أثره على الرصيد'])
        table.setAlternatingRowColors(True)
        table.verticalHeader().setVisible(False)
        table.horizontalHeader().setStretchLastSection(True)
        table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeToContents)
        table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeToContents)
        table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeToContents)
        table.horizontalHeader().setSectionResizeMode(3, QHeaderView.Stretch)
        table.horizontalHeader().setSectionResizeMode(4, QHeaderView.Stretch)
        table.horizontalHeader().setMinimumHeight(44)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        table.setStyleSheet(
            f"QTableWidget{{background:{rgba_from_hex(CARD,0.98)}; color:{TEXT}; border:none; border-radius:18px; gridline-color:{rgba_from_hex(TEXT,0.06)}; font-size:15px; alternate-background-color:{rgba_from_hex(DARK,0.88)}; selection-background-color:{rgba_from_hex(ACCENT,0.22)};}}"
            f"QHeaderView::section{{background:{rgba_from_hex(ACCENT,0.96)}; color:{TEXT_ON_ACCENT}; padding:10px 12px; border:none; font-weight:900; font-size:14px;}}"
        )
        table.setShowGrid(False)
        table.setWordWrap(True)

        table_card = QFrame(); table_card.setStyleSheet(f"QFrame{{background:{rgba_from_hex(CARD,0.92)}; border:1px solid {rgba_from_hex(TEXT,0.08)}; border-radius:22px;}}")
        card_lay = QVBoxLayout(table_card); card_lay.setContentsMargins(14, 14, 14, 14); card_lay.setSpacing(12)
        card_lay.addWidget(table, 1)
        root.addWidget(table_card, 1)

        summary = QLabel(); summary.setAlignment(Qt.AlignRight); summary.setWordWrap(True)
        summary.setStyleSheet(f"background:{rgba_from_hex(DARK,0.55)}; border:1px solid {rgba_from_hex(TEXT,0.06)}; border-radius:16px; padding:12px 14px; font-size:14px; font-weight:800; color:{TEXT};")
        root.addWidget(summary)

        self.person_pages[person_name] = {'date': date, 'amount': amount, 'notes': notes, 'table': table, 'summary': summary}
        add_btn.clicked.connect(lambda _=False, p=person_name: self.add_row(p))
        delete_btn.clicked.connect(lambda _=False, p=person_name: self.delete_selected(p))
        return page, table, summary

    def show_home(self):
        self.stack.setCurrentIndex(0)

    def show_person_page(self, person_name):
        self.stack.setCurrentIndex(self.person_pages[person_name]['index'])

    def _person_rows(self, person_label):
        return [x for x in self.db.get('cash', []) if x.get('source') == 'withdrawal' and x.get('category') == person_label]

    def _fill_person_table(self, table, summary_label, person_label, deficit):
        rows = self._person_rows(person_label)
        table.setRowCount(len(rows))
        impact_txt = 'ضمن الربح'
        if deficit > 0:
            impact_txt = f"ولّد/زاد عجز {person_label.replace('سحوبات ', '').strip()}"
        for r, row in enumerate(rows):
            vals = [r+1, row.get('date',''), fmt_money(row.get('amount', 0)), row.get('notes',''), impact_txt]
            for c, v in enumerate(vals):
                item = QTableWidgetItem(str(v))
                item.setTextAlignment(Qt.AlignCenter if c in (0,1,2) else Qt.AlignRight | Qt.AlignVCenter)
                table.setItem(r, c, item)
            table.setRowHeight(r, 42)
        summary_label.setText(
            f"إجمالي {person_label}: {fmt_money(withdrawals_sum(self.db, person_label))} د.ع | "
            f"عدد الحركات: {len(rows)} | "
            f"العجز الحالي: {fmt_money(deficit)} د.ع"
        )

    def refresh_table(self):
        st = person_profit_status(self.db)
        self._fill_person_table(self.hidab_table, self.hidab_summary, 'سحوبات هضاب', st['hidab_deficit'])
        self._fill_person_table(self.mustafa_table, self.mustafa_summary, 'سحوبات مصطفى', st['mostafa_deficit'])
        self.home_summary.setText(
            f"سحوبات هضاب: {fmt_money(withdrawals_sum(self.db, 'سحوبات هضاب'))} د.ع | "
            f"سحوبات مصطفى: {fmt_money(withdrawals_sum(self.db, 'سحوبات مصطفى'))} د.ع | "
            f"عجز هضاب: {fmt_money(st['hidab_deficit'])} د.ع | "
            f"عجز مصطفى: {fmt_money(st['mostafa_deficit'])} د.ع | "
            f"إجمالي العجز: {fmt_money(st['total_deficit'])} د.ع"
        )

    def add_row(self, person_name):
        widgets = self.person_pages[person_name]
        amount = float(widgets['amount'].value())
        if amount <= 0:
            return QMessageBox.warning(self, 'تنبيه', 'أدخل مبلغ صحيح')
        person_label = 'سحوبات هضاب' if person_name == 'هضاب' else 'سحوبات مصطفى'
        self.db['cash'].append({
            'date': widgets['date'].date().toString('yyyy-MM-dd'),
            'type': 'مصروف',
            'category': person_label,
            'party': person_name,
            'amount': amount,
            'notes': widgets['notes'].text().strip(),
            'source': 'withdrawal',
            'created_at': now_str()
        })
        widgets['amount'].setValue(0)
        widgets['notes'].clear()
        self.save()
        self.refresh_table()
        self.show_person_page(person_name)

    def delete_selected(self, person_name):
        person_label = 'سحوبات هضاب' if person_name == 'هضاب' else 'سحوبات مصطفى'
        table = self.person_pages[person_name]['table']
        rows = self._person_rows(person_label)
        empty_msg = f'اختر سحبًا من قائمة {person_name}'
        view_row = table.currentRow()
        if view_row < 0 or view_row >= len(rows):
            return QMessageBox.warning(self, 'تنبيه', empty_msg)
        target = rows[view_row]
        for i in range(len(self.db['cash'])-1, -1, -1):
            row = self.db['cash'][i]
            if row is target:
                self.db['cash'].pop(i)
                break
        self.save()
        self.refresh_table()

class ProfitWindow(BaseWindow):
    def __init__(self, main):
        super().__init__(main, '📊 الأرباح اللحظية')
        self.subtitle_lbl.setText('تم تقسيم الأرباح إلى تبويبات واضحة، مع فصل الربح الكلي المحقق عن الربح القابل للتوزيع حتى ما يصير لَبس بالأرقام.')

        self.tabs = QTabWidget()
        self.tabs.setDocumentMode(True)
        self.tabs.setStyleSheet(
            f"QTabWidget::pane {{border:1px solid {BORDER}; background:{CARD}; border-radius:14px; margin-top:8px;}}"
            f"QTabBar::tab {{min-width:170px; min-height:42px; padding:10px 14px; margin:4px 6px; border-radius:12px; background:{DARK}; color:{TEXT}; font-weight:800; font-size:14px;}}"
            f"QTabBar::tab:selected {{background:{ACCENT}; color:{TEXT_ON_ACCENT};}}"
        )
        self.layout.addWidget(self.tabs, 1)

        self.overview_tab = QWidget()
        self.overview_layout = QVBoxLayout(self.overview_tab)
        self.overview_layout.setContentsMargins(12, 12, 12, 12)
        self.overview_layout.setSpacing(12)
        self.tabs.addTab(self.overview_tab, 'نظرة سريعة')

        self.overview_cards_wrap = QFrame(); self.overview_cards_wrap.setStyleSheet(CARD_FRAME_STYLE)
        self.overview_cards_layout = QGridLayout(self.overview_cards_wrap)
        self.overview_cards_layout.setContentsMargins(12, 12, 12, 12)
        self.overview_cards_layout.setHorizontalSpacing(10)
        self.overview_cards_layout.setVerticalSpacing(10)
        self.card_total_profit = SummaryCard('إجمالي الربح المحقق', '0', 'إجمالي صافي الربح المحقق من العمليات قبل تفصيل الحصص')
        self.card_external_profit = SummaryCard('ربح الممولين الخارجيين', '0', 'الأرباح المخصصة للممولين الخارجيين')
        self.card_financing_profit = SummaryCard('تمويل هضاب الراجع للشراكة', '0', 'يرجع للشراكة ثم يدخل بالتقسيم')
        self.card_partnership_profit = SummaryCard('أرباح الشراكة قبل تسديد العجز', '0', 'الربح القابل للتقسيم بين هضاب ومصطفى')
        for i, card in enumerate([self.card_total_profit, self.card_external_profit, self.card_financing_profit, self.card_partnership_profit]):
            self.overview_cards_layout.addWidget(card, i // 2, i % 2)
        self.overview_layout.addWidget(self.overview_cards_wrap)

        self.top_summary = QLabel()
        self.top_summary.setAlignment(Qt.AlignRight)
        self.top_summary.setWordWrap(True)
        self.top_summary.setStyleSheet(f'font-size:14px;font-weight:800;color:{TEXT};')
        self.overview_layout.addWidget(self.top_summary)

        self.summary_tab = QWidget()
        self.summary_layout = QVBoxLayout(self.summary_tab)
        self.summary_layout.setContentsMargins(12, 12, 12, 12)
        self.summary_layout.setSpacing(12)
        self.tabs.addTab(self.summary_tab, 'التسوية والملخص')

        self.profit_calc = QuickCalculatorPanel('حاسبة الأرباح', 'موجودة هنا للمراجعة السريعة داخل ملخص الأرباح.', compact=True)
        profit_calc_row = QHBoxLayout()
        profit_calc_row.setContentsMargins(0, 0, 0, 0)
        profit_calc_row.setSpacing(12)
        profit_calc_row.addStretch(1)
        profit_calc_row.addWidget(self.profit_calc, 0, Qt.AlignTop | Qt.AlignRight)
        self.summary_layout.addLayout(profit_calc_row)

        self.summary_table = QTableWidget()
        self.summary_table.setColumnCount(3)
        self.summary_table.setHorizontalHeaderLabels(['#', 'البند', 'القيمة'])
        self.summary_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.summary_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.summary_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.summary_table.setStyleSheet(TABLE_STYLE)
        self.summary_layout.addWidget(self.summary_table)

        self.details_table = QTableWidget()
        self.details_table.setColumnCount(6)
        self.details_table.setHorizontalHeaderLabels(['#', 'الشخص/الجهة', 'نوع الربح', 'ربح الفترة', 'المسدّد لسد العجز', 'العجز المتبقي'])
        self.details_table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        self.details_table.setSelectionBehavior(QAbstractItemView.SelectRows)
        self.details_table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        self.details_table.setStyleSheet(TABLE_STYLE)
        self.summary_layout.addWidget(self.details_table)

        self.bottom_summary = QLabel()
        self.bottom_summary.setAlignment(Qt.AlignRight)
        self.bottom_summary.setWordWrap(True)
        self.bottom_summary.setStyleSheet(f'font-size:13px;font-weight:700;color:{MUTED};')
        self.summary_layout.addWidget(self.bottom_summary)

        self.external_tab, self.external_formula, self.external_table, self.external_total = self._build_profit_tab('ربح الممولين الخارجيين')
        self.tabs.addTab(self.external_tab, 'ربح الممولين')

        self.financing_tab, self.financing_formula, self.financing_table, self.financing_total = self._build_profit_tab('ربح تمويل هضاب')
        self.tabs.addTab(self.financing_tab, 'تمويل هضاب')

        self.hidab_tab, self.hidab_formula, self.hidab_table, self.hidab_total = self._build_profit_tab('ربح هضاب')
        self.tabs.addTab(self.hidab_tab, 'ربح هضاب')

        self.mostafa_tab, self.mostafa_formula, self.mostafa_table, self.mostafa_total = self._build_profit_tab('ربح مصطفى')
        self.tabs.addTab(self.mostafa_tab, 'ربح مصطفى')

        self.refresh_table()

    def _build_profit_tab(self, title):
        tab = QWidget()
        layout = QVBoxLayout(tab)
        layout.setContentsMargins(12, 12, 12, 12)
        layout.setSpacing(12)

        header = QLabel(title)
        header.setAlignment(Qt.AlignRight)
        header.setStyleSheet(f'font-size:24px;font-weight:900;color:{TEXT};')
        layout.addWidget(header)

        formula = QTextEdit()
        formula.setReadOnly(True)
        formula.setMinimumHeight(150)
        formula.setStyleSheet(
            f"QTextEdit{{background:{rgba_from_hex(CARD,0.82)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:18px;padding:10px;color:{TEXT};font-size:15px;font-weight:700;}}"
        )
        layout.addWidget(formula)

        table = QTableWidget()
        table.setStyleSheet(TABLE_STYLE)
        table.horizontalHeader().setSectionResizeMode(QHeaderView.Stretch)
        table.setSelectionBehavior(QAbstractItemView.SelectRows)
        table.setEditTriggers(QAbstractItemView.NoEditTriggers)
        layout.addWidget(table, 1)

        total = QLabel()
        total.setAlignment(Qt.AlignRight)
        total.setWordWrap(True)
        total.setStyleSheet(f'font-size:13px;font-weight:700;color:{MUTED};')
        layout.addWidget(total)
        return tab, formula, table, total

    def refresh_table(self):
        snap = profit_ui_snapshot(self.db)
        check = validate_profit_consistency(self.db)
        pb = profit_breakdown(self.db)
        person = person_profit_status(self.db)

        total_profit_before_distribution = round(float(snap.get('operating_profit', pb.get('operating_profit', 0)) or 0), 2)
        distributed_total = round(float(snap.get('external_funders_profit', 0) or 0) + float(snap.get('hidab_financing_profit', 0) or 0) + float(snap.get('partnership_profit', 0) or 0), 2)
        partnership_before_deficit = round(float(snap.get('partnership_profit', 0) or 0) + float(snap.get('hidab_financing_profit', 0) or 0), 2)
        period_share_each = round(partnership_before_deficit / 2.0, 2)

        hidab_paid = round(float(person.get('hidab_withdrawals', 0) or 0) + float(person.get('hidab_profit_paid', 0) or 0), 2)
        mostafa_paid = round(float(person.get('mostafa_withdrawals', 0) or 0) + float(person.get('mostafa_profit_paid', 0) or 0), 2)
        hidab_opening = round(float(snap['hidab_opening_balance'] or 0), 2)
        mostafa_opening = round(float(snap['mostafa_opening_balance'] or 0), 2)
        hidab_before_profit = round(float(person.get('hidab_total_deficit_before_profit', 0) or 0), 2)
        mostafa_before_profit = round(float(person.get('mostafa_total_deficit_before_profit', 0) or 0), 2)
        hidab_settled = round(float(person.get('hidab_settled', 0) or 0), 2)
        mostafa_settled = round(float(person.get('mostafa_settled', 0) or 0), 2)
        hidab_final = round(float(person.get('hidab_deficit', 0) or 0), 2)
        mostafa_final = round(float(person.get('mostafa_deficit', 0) or 0), 2)
        hidab_surplus = round(float(person.get('hidab_surplus', 0) or 0), 2)
        mostafa_surplus = round(float(person.get('mostafa_surplus', 0) or 0), 2)

        self.card_total_profit.setText(f"{fmt_money(total_profit_before_distribution)} د.ع")
        self.card_external_profit.setText(f"{fmt_money(snap['external_funders_profit'])} د.ع")
        self.card_financing_profit.setText(f"{fmt_money(snap['hidab_financing_profit'])} د.ع")
        self.card_partnership_profit.setText(f"{fmt_money(partnership_before_deficit)} د.ع")

        summary_rows = [
            ('إجمالي الربح المحقق', total_profit_before_distribution),
            ('إجمالي الأرباح الموزعة فعليًا', distributed_total),
            ('ربح الممولين الخارجيين', snap['external_funders_profit']),
            ('حصة هضاب من التمويل الراجعة للشراكة', snap['hidab_financing_profit']),
            ('أرباح الشراكة الصافية قبل تسديد العجز', partnership_before_deficit),
            ('حصة هضاب من ربح الفترة', period_share_each),
            ('حصة مصطفى من ربح الفترة', period_share_each),
            ('العجز الافتتاحي - هضاب', hidab_opening),
            ('العجز الافتتاحي - مصطفى', mostafa_opening),
            ('إجمالي السحب/المدفوع قبل ربح الفترة - هضاب', hidab_before_profit),
            ('إجمالي السحب/المدفوع قبل ربح الفترة - مصطفى', mostafa_before_profit),
            ('المسدّد من ربح الفترة لسد العجز - هضاب', hidab_settled),
            ('المسدّد من ربح الفترة لسد العجز - مصطفى', mostafa_settled),
            ('العجز المتبقي - هضاب', hidab_final),
            ('العجز المتبقي - مصطفى', mostafa_final),
            ('الفائض القابل للسحب - هضاب', hidab_surplus),
            ('الفائض القابل للسحب - مصطفى', mostafa_surplus),
        ]
        self.summary_table.setRowCount(len(summary_rows))
        for r, (label, amount) in enumerate(summary_rows):
            vals = [r + 1, label, f"{fmt_money(amount)} د.ع"]
            for c, v in enumerate(vals):
                self.summary_table.setItem(r, c, QTableWidgetItem(str(v)))

        details_rows = [
            ('الممولون الخارجيون', 'ربح ممولين خارجيين', snap['external_funders_profit'], 0, 0),
            ('هضاب / التمويل', 'حصة تمويل راجعة للشراكة', snap['hidab_financing_profit'], 0, 0),
            ('هضاب', 'ربح الفترة قبل التسوية', period_share_each, hidab_settled, hidab_final),
            ('مصطفى', 'ربح الفترة قبل التسوية', period_share_each, mostafa_settled, mostafa_final),
        ]
        self.details_table.setRowCount(len(details_rows))
        for r, (name, ptype, earned, paid, netv) in enumerate(details_rows):
            vals = [r + 1, name, ptype, f"{fmt_money(earned)} د.ع", f"{fmt_money(paid)} د.ع", f"{fmt_money(netv)} د.ع"]
            for c, v in enumerate(vals):
                self.details_table.setItem(r, c, QTableWidgetItem(str(v)))

        status_txt = 'مطابقة 100%' if check.get('ok') else f"يوجد فرق: {check.get('diffs', {})}"
        self.top_summary.setText(
            f"حالة التطابق: {status_txt} | ربح الممولين الخارجيين: {fmt_money(snap['external_funders_profit'])} د.ع | "
            f"حصة هضاب من التمويل: {fmt_money(snap['hidab_financing_profit'])} د.ع | "
            f"أرباح الشراكة قبل تسديد العجز: {fmt_money(partnership_before_deficit)} د.ع"
        )
        self.bottom_summary.setText(
            f"هضاب — عجز افتتاحي: {fmt_money(hidab_opening)} د.ع | سحوبات/مدفوع: {fmt_money(hidab_paid)} د.ع | المسدّد من ربح الفترة: {fmt_money(hidab_settled)} د.ع | العجز المتبقي: {fmt_money(hidab_final)} د.ع\n"
            f"مصطفى — عجز افتتاحي: {fmt_money(mostafa_opening)} د.ع | سحوبات/مدفوع: {fmt_money(mostafa_paid)} د.ع | المسدّد من ربح الفترة: {fmt_money(mostafa_settled)} د.ع | العجز المتبقي: {fmt_money(mostafa_final)} د.ع"
        )

        external_rows = [x for x in pb.get('funders_rows', []) if not x.get('is_owner_capital', False)]
        self.external_table.setColumnCount(6)
        self.external_table.setHorizontalHeaderLabels(['#', 'الممول', 'الرصيد الفعال', 'النسبة', 'ربحه الحالي', 'المتبقي له'])
        self.external_table.setRowCount(len(external_rows))
        for r, row in enumerate(external_rows):
            vals = [r+1, row.get('name',''), fmt_money(row.get('capital',0)), fmt_pct(float(row.get('ratio',0) or 0)), fmt_money(row.get('amount',0)), fmt_money(row.get('pending',0))]
            for c, v in enumerate(vals):
                self.external_table.setItem(r, c, QTableWidgetItem(str(v)))
        self.external_formula.setPlainText(
            "معادلة ربح الممولين الخارجيين:\n"
            "ربح الممولين الخارجيين = مجموع قيود profit_entries للمستفيدين من نوع funder بعد استثناء رأس مال هضاب\n\n"
            f"التطبيق الحالي:\n{fmt_money(snap['external_funders_profit'])} = مجموع أرباح {len(external_rows)} ممول/ممولين خارجيين من دفتر الأرباح."
        )
        self.external_total.setText(
            f"الإجمالي الحالي: {fmt_money(snap['external_funders_profit'])} د.ع | غير المسدّد لهم: {fmt_money(sum(float(x.get('pending',0) or 0) for x in external_rows))} د.ع"
        )

        financing_rows = [x for x in pb.get('funders_rows', []) if x.get('is_owner_capital', False)]
        self.financing_table.setColumnCount(6)
        self.financing_table.setHorizontalHeaderLabels(['#', 'الاسم', 'الرصيد الفعال', 'النسبة', 'ربح التمويل', 'المسدّد/المتبقي'])
        self.financing_table.setRowCount(len(financing_rows))
        for r, row in enumerate(financing_rows):
            vals = [r+1, row.get('name',''), fmt_money(row.get('capital',0)), fmt_pct(float(row.get('ratio',0) or 0)), fmt_money(row.get('amount',0)), f"{fmt_money(row.get('paid',0))} / {fmt_money(row.get('pending',0))}"]
            for c, v in enumerate(vals):
                self.financing_table.setItem(r, c, QTableWidgetItem(str(v)))
        self.financing_formula.setPlainText(
            "معادلة ربح تمويل هضاب:\n"
            "ربح تمويل هضاب = مجموع قيود profit_entries للمستفيد funder المرتبط بممول معلَّم كرأس مال هضاب\n"
            "ثم هذه الحصة ترجع للشراكة وتدخل ضمن الربح القابل للتقسيم بين هضاب ومصطفى.\n\n"
            f"التطبيق الحالي:\n{fmt_money(snap['hidab_financing_profit'])} = مجموع أرباح رأس مال هضاب التمويلية المسجلة بدفتر الأرباح."
        )
        self.financing_total.setText(f"الإجمالي الحالي: {fmt_money(snap['hidab_financing_profit'])} د.ع")

        self.hidab_table.setColumnCount(2)
        self.hidab_table.setHorizontalHeaderLabels(['البند', 'القيمة'])
        hidab_rows = [
            ('حصة هضاب من ربح الفترة', fmt_money(period_share_each)),
            ('العجز الافتتاحي', fmt_money(hidab_opening)),
            ('السحوبات + الأرباح المدفوعة', fmt_money(hidab_paid)),
            ('إجمالي العجز قبل الربح', fmt_money(hidab_before_profit)),
            ('المسدّد من ربح الفترة', fmt_money(hidab_settled)),
            ('العجز المتبقي', fmt_money(hidab_final)),
            ('الفائض القابل للسحب', fmt_money(hidab_surplus)),
        ]
        self.hidab_table.setRowCount(len(hidab_rows))
        for r, (label, value) in enumerate(hidab_rows):
            self.hidab_table.setItem(r, 0, QTableWidgetItem(label))
            self.hidab_table.setItem(r, 1, QTableWidgetItem(f'{value} د.ع'))
        self.hidab_formula.setPlainText(
            "معادلة ربح هضاب:\n"
            "ربح هضاب للفترة = (ربح الشراكة المباشر + ربح تمويل هضاب الراجع للشراكة) ÷ 2\n"
            "إجمالي العجز قبل الربح = العجز الافتتاحي + سحوبات هضاب + الأرباح المدفوعة له\n"
            "المسدّد من ربح الفترة = أصغر قيمة بين حصة الربح وإجمالي العجز\n"
            "العجز المتبقي = أكبر قيمة بين (إجمالي العجز - حصة الربح) و 0\n\n"
            f"التطبيق الحالي:\n({fmt_money(snap['partnership_profit'])} + {fmt_money(snap['hidab_financing_profit'])}) ÷ 2 = {fmt_money(period_share_each)} د.ع\n"
            f"{fmt_money(hidab_opening)} + {fmt_money(hidab_paid)} = {fmt_money(hidab_before_profit)} د.ع\n"
            f"المسدّد = min({fmt_money(period_share_each)}, {fmt_money(hidab_before_profit)}) = {fmt_money(hidab_settled)} د.ع"
        )
        self.hidab_total.setText(f"العجز المتبقي: {fmt_money(hidab_final)} د.ع | الفائض القابل للسحب: {fmt_money(hidab_surplus)} د.ع")

        self.mostafa_table.setColumnCount(2)
        self.mostafa_table.setHorizontalHeaderLabels(['البند', 'القيمة'])
        mostafa_rows = [
            ('حصة مصطفى من ربح الفترة', fmt_money(period_share_each)),
            ('العجز الافتتاحي', fmt_money(mostafa_opening)),
            ('السحوبات + الأرباح المدفوعة', fmt_money(mostafa_paid)),
            ('إجمالي العجز قبل الربح', fmt_money(mostafa_before_profit)),
            ('المسدّد من ربح الفترة', fmt_money(mostafa_settled)),
            ('العجز المتبقي', fmt_money(mostafa_final)),
            ('الفائض القابل للسحب', fmt_money(mostafa_surplus)),
        ]
        self.mostafa_table.setRowCount(len(mostafa_rows))
        for r, (label, value) in enumerate(mostafa_rows):
            self.mostafa_table.setItem(r, 0, QTableWidgetItem(label))
            self.mostafa_table.setItem(r, 1, QTableWidgetItem(f'{value} د.ع'))
        self.mostafa_formula.setPlainText(
            "معادلة ربح مصطفى:\n"
            "ربح مصطفى للفترة = (ربح الشراكة المباشر + ربح تمويل هضاب الراجع للشراكة) ÷ 2\n"
            "إجمالي العجز قبل الربح = العجز الافتتاحي + سحوبات مصطفى + الأرباح المدفوعة له\n"
            "المسدّد من ربح الفترة = أصغر قيمة بين حصة الربح وإجمالي العجز\n"
            "العجز المتبقي = أكبر قيمة بين (إجمالي العجز - حصة الربح) و 0\n\n"
            f"التطبيق الحالي:\n({fmt_money(snap['partnership_profit'])} + {fmt_money(snap['hidab_financing_profit'])}) ÷ 2 = {fmt_money(period_share_each)} د.ع\n"
            f"{fmt_money(mostafa_opening)} + {fmt_money(mostafa_paid)} = {fmt_money(mostafa_before_profit)} د.ع\n"
            f"المسدّد = min({fmt_money(period_share_each)}, {fmt_money(mostafa_before_profit)}) = {fmt_money(mostafa_settled)} د.ع"
        )
        self.mostafa_total.setText(f"العجز المتبقي: {fmt_money(mostafa_final)} د.ع | الفائض القابل للسحب: {fmt_money(mostafa_surplus)} د.ع")


class ToastLabel(QLabel):
    def __init__(self, parent=None):
        super().__init__(parent)
        self.setAlignment(Qt.AlignCenter)
        self.setWordWrap(True)
        self.setMinimumHeight(48)
        self.setMargin(12)
        self.hide()


class ToastNotification(QFrame):
    def __init__(self, parent, title, message, level='info', timeout_ms=3200):
        super().__init__(parent)
        self.setObjectName('toastNotification')
        self.setAttribute(Qt.WA_ShowWithoutActivating, True)
        self.setWindowFlags(Qt.SubWindow | Qt.FramelessWindowHint)
        self.timeout_ms = timeout_ms
        color_map = {
            'info': ('#0f6fb5', '#ffffff'),
            'success': ('#0f9d73', '#ffffff'),
            'warning': ('#b7791f', '#ffffff'),
            'danger': ('#991b1b', '#ffffff'),
        }
        bg, fg = color_map.get(level, color_map['info'])
        self.setStyleSheet(
            f"QFrame#toastNotification{{background:{bg};color:{fg};border-radius:16px;border:1px solid rgba(255,255,255,0.18);}}"
            f"QLabel{{background:transparent;color:{fg};}}"
        )
        lay = QVBoxLayout(self)
        lay.setContentsMargins(16, 14, 16, 14)
        lay.setSpacing(4)
        t = QLabel(title)
        t.setAlignment(Qt.AlignRight)
        t.setStyleSheet('font-size:15px;font-weight:900;')
        b = QLabel(message)
        b.setAlignment(Qt.AlignRight | Qt.AlignTop)
        b.setWordWrap(True)
        b.setStyleSheet('font-size:12px;font-weight:600;')
        lay.addWidget(t)
        lay.addWidget(b)
        self.adjustSize()
        self.timer = QTimer(self)
        self.timer.setSingleShot(True)
        self.timer.timeout.connect(self.close)

    def showEvent(self, event):
        super().showEvent(event)
        self.timer.start(self.timeout_ms)


class MainWindow(QWidget):
    def __init__(self):
        super().__init__()
        self.db = load_db(); self.windows = {}
        self.sidebar_buttons = []
        self.login_window = None
        self.autosave_timer = None
        self.alert_timer = None
        self.alert_state = {}
        self.last_autosave_status = ''
        self.last_manual_save_status = ''
        self.last_export_status = ''
        self.last_alert_scan = ''
        self.last_toast = None
        self.last_notification_signature = ''
        self.autosave_counter = 0
        set_theme(self.db.get('settings', {}).get('current_theme', 'dark_lux'))
        self.setWindowTitle('مخزن النخبة'); self.resize(1560, 980)
        apply_branding(self)
        self.build_ui()
        self.apply_theme()
        self.refresh_dashboard()
        self.setup_autosave()
        self.setup_alerts_monitor()

    def paintEvent(self, event):
        paint_app_background(self, event)

    def build_ui(self):
        root = QHBoxLayout(self)
        root.setContentsMargins(14, 14, 14, 14)
        root.setSpacing(14)

        sidebar = QFrame()
        sidebar.setFixedWidth(310)
        self.sidebar = sidebar
        side_layout = QVBoxLayout(sidebar)
        side_layout.setContentsMargins(16, 16, 16, 16)
        side_layout.setSpacing(10)

        brand_logo = make_logo_label(92)
        side_layout.addWidget(brand_logo)

        brand = QLabel('مخزن النخبة')
        brand.setAlignment(Qt.AlignCenter)
        brand.setStyleSheet('font-size:26px;font-weight:900;letter-spacing:0.5px;')
        brand_sub = QLabel('القائمة الرئيسية')
        brand_sub.setAlignment(Qt.AlignCenter)
        brand_sub.setStyleSheet(f'font-size:11px;color:{MUTED};font-weight:700;')
        side_layout.addWidget(brand)
        side_layout.addWidget(brand_sub)

        self.quick_status = QLabel('جاهز للعمل')
        self.quick_status.setAlignment(Qt.AlignCenter)
        self.quick_status.setWordWrap(True)
        self.quick_status.setMinimumHeight(50)
        side_layout.addWidget(self.quick_status)

        self.toast = ToastLabel(self)
        side_layout.addWidget(self.toast)

        self.cash_chip = QLabel('الكاش الفعلي: 0 د.ع')
        self.cash_chip.setAlignment(Qt.AlignCenter)
        self.cash_chip.setWordWrap(True)
        side_layout.addWidget(self.cash_chip)

        sidebar_scroll = QScrollArea()
        sidebar_scroll.setWidgetResizable(True)
        sidebar_scroll.setFrameShape(QFrame.NoFrame)
        side_inner = QWidget()
        self.sidebar_buttons_layout = QVBoxLayout(side_inner)
        self.sidebar_buttons_layout.setContentsMargins(0, 0, 0, 0)
        self.sidebar_buttons_layout.setSpacing(8)
        sidebar_scroll.setWidget(side_inner)
        side_layout.addWidget(sidebar_scroll, 1)

        sidebar_specs = [
            ('📦  الأصناف', self.open_items),
            ('🏬  المخزن', self.open_warehouse),
            ('👥  الزبائن', self.open_customers),
            ('🚚  الموردين', self.open_suppliers),
            ('🏦  الممولين', self.open_funders),
            ('📥  الوارد', self.open_inbound),
            ('💰  المبيعات', self.open_sales),
            ('↩️  المرتجعات', self.open_returns),
            ('⚠️  التالف', self.open_damaged),
            ('🧾  الصندوق', self.open_cash),
            ('🧮  المطابقة', self.open_reconciliation),
            ('💸  السحوبات', self.open_withdrawals),
            ('🧾  المصاريف', self.open_expenses),
            ('🧾  عهدة المندوبين', self.open_agents_custody),
            ('💳  ديون الزبائن', self.open_customer_dues),
            ('💳  ديون الموردين', self.open_supplier_dues),
            ('📊  الأرباح اللحظية', self.open_profit),
            ('🟢  الرصيد الافتتاحي', self.open_opening),
        ]
        for title, fn in sidebar_specs:
            b = QPushButton(title)
            b.setCursor(Qt.PointingHandCursor)
            b.setMinimumHeight(54)
            b.clicked.connect(fn)
            self.sidebar_buttons.append(b)
            self.sidebar_buttons_layout.addWidget(b)

        utility_specs = [
            ('💾  حفظ الآن', self.save_all),
            ('🗂️  نسخة احتياطية', self.create_manual_backup),
            ('🩺  فحص سلامة النظام', self.show_health_check_dialog),
            ('🧪  التدقيق المحاسبي النهائي', self.show_accounting_audit_dialog),
            ('🧱  تقرير الاستقرار النهائي', self.show_stability_report_dialog),
            ('📂  فتح مجلد البيانات', self.open_app_data_folder),
            ('⚙️  إعدادات النظام', self.show_system_settings_dialog),
            ('🔑  تغيير كلمة السر', self.show_change_password_dialog),
            ('📜  سجل العمليات', self.show_operations_log_dialog),
            ('📧  تصدير للجيميل', self.export_to_gmail),
            ('🔓  تسجيل الخروج', self.logout),
        ]
        for title, fn in utility_specs:
            b = QPushButton(title)
            b.setCursor(Qt.PointingHandCursor)
            b.setMinimumHeight(54)
            b.clicked.connect(fn)
            self.sidebar_buttons.append(b)
            self.sidebar_buttons_layout.addWidget(b)

        self.sidebar_buttons_layout.addStretch(1)

        side_note = QLabel('')
        side_note.setWordWrap(True)
        side_note.setAlignment(Qt.AlignCenter)
        side_note.setStyleSheet(f'font-size:12px;color:{MUTED};')
        side_layout.addWidget(side_note)

        content_scroll = QScrollArea()
        content_scroll.setWidgetResizable(True)
        content_scroll.setFrameShape(QFrame.NoFrame)
        content_scroll.setStyleSheet("QScrollArea{background:transparent;border:none;} QScrollArea > QWidget > QWidget{background:transparent;}")
        content_host = QWidget()
        content_wrap = QVBoxLayout(content_host)
        content_wrap.setContentsMargins(0, 0, 0, 0)
        content_wrap.setSpacing(10)
        content_scroll.setWidget(content_host)
        root.addWidget(content_scroll, 1)
        root.addWidget(sidebar)

        hero = QFrame(); hero.setStyleSheet(CARD_FRAME_STYLE)
        hero_layout = QHBoxLayout(hero)
        hero_layout.setContentsMargins(16, 12, 16, 12)
        hero_layout.setSpacing(10)
        hero_text = QVBoxLayout()
        header = QLabel('لوحة التحكم')
        header.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        header.setStyleSheet('font-size:28px;font-weight:900;')
        sub = QLabel('واجهة مختصرة وواضحة: بطاقات متساوية ومسارات سريعة بنفس روح Mind Flow.')
        sub.setWordWrap(True)
        sub.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        sub.setStyleSheet(f'font-size:11px;color:{MUTED};font-weight:700;')
        hero_chip = QLabel('العرض الرئيسي')
        hero_chip.setAlignment(Qt.AlignCenter)
        hero_chip.setStyleSheet(f'background:{rgba_from_hex(ACCENT,0.16)};color:{TEXT};border:none;border-radius:16px;padding:6px 12px;font-size:11px;font-weight:900;')
        hero_text.addWidget(hero_chip, alignment=Qt.AlignRight)
        hero_text.addWidget(header)
        hero_text.addWidget(sub)
        hero_layout.addLayout(hero_text, 1)
        tools_box = QVBoxLayout(); tools_box.setSpacing(8)
        theme_label = QLabel('الثيم')
        theme_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        theme_label.setStyleSheet(f'font-size:12px;color:{MUTED};font-weight:800;')
        self.theme_combo = QComboBox()
        self.theme_combo.setMinimumHeight(38)
        for key, meta in THEMES.items(): self.theme_combo.addItem(meta['name'], key)
        current_theme = self.db.get('settings', {}).get('current_theme', 'dark_lux')
        idx = self.theme_combo.findData(current_theme)
        if idx >= 0: self.theme_combo.setCurrentIndex(idx)
        apply_btn = QPushButton('🎨 تطبيق الثيم'); apply_btn.clicked.connect(self.apply_selected_theme); self.apply_theme_btn = apply_btn; apply_btn.setMinimumHeight(38)
        tools_box.addWidget(theme_label)
        tools_box.addWidget(self.theme_combo)
        tools_box.addWidget(apply_btn)
        hero_layout.addLayout(tools_box)
        content_wrap.addWidget(hero)

        dashboard_wrap = QFrame(); dashboard_wrap.setObjectName('dashboardWrap'); dashboard_wrap.setStyleSheet(CARD_FRAME_STYLE)
        dashboard_outer = QVBoxLayout(dashboard_wrap)
        dashboard_outer.setContentsMargins(16, 14, 16, 14)
        dashboard_outer.setSpacing(10)
        dash_head = QHBoxLayout(); dash_head.setSpacing(10)
        dash_titles = QVBoxLayout(); dash_titles.setSpacing(4)
        dash_title = QLabel('الواجهة الرئيسية')
        dash_title.setAlignment(Qt.AlignRight); dash_title.setStyleSheet('font-size:24px;font-weight:900;')
        dash_titles.addWidget(dash_title)
        dash_head.addLayout(dash_titles, 1)
        self.focus_right = QLabel('تشغيل مستقر'); self.focus_right.setAlignment(Qt.AlignCenter); self.focus_right.setMinimumWidth(150)
        self.focus_right.setStyleSheet(f'background:{ACCENT};color:{TEXT_ON_ACCENT};border-radius:14px;padding:9px 12px;font-size:11px;font-weight:900;')
        dash_head.addWidget(self.focus_right)
        dashboard_outer.addLayout(dash_head)

        self.hero_stats_layout = QGridLayout(); self.hero_stats_layout.setHorizontalSpacing(10); self.hero_stats_layout.setVerticalSpacing(10)
        self.hero_cash_card = SummaryCard('مبلغ القاصة', '0 د.ع', 'القيمة النظرية حسب الحركة', ACCENT)
        self.hero_dues_card = SummaryCard('ديون الزبائن', '0 د.ع', 'التحصيلات القادمة', '#a997ff')
        self.hero_alert_card = SummaryCard('حالة التشغيل', 'جاهز', 'أي تنبيه مهم يطلع هنا', '#ff9bca')
        self.hero_payables_card = SummaryCard('ذمم الموردين', '0 د.ع', 'الالتزامات المفتوحة', ACCENT2)
        for i, card in enumerate([self.hero_cash_card, self.hero_dues_card, self.hero_alert_card, self.hero_payables_card]):
            self.hero_stats_layout.addWidget(card, 0, i)
            self.hero_stats_layout.setColumnStretch(i, 1)
        dashboard_outer.addLayout(self.hero_stats_layout)

        actions_title = QLabel('المسارات الرئيسية'); actions_title.setAlignment(Qt.AlignRight); actions_title.setStyleSheet('font-size:16px;font-weight:900;')
        dashboard_outer.addWidget(actions_title)
        self.dashboard_grid = QGridLayout(); self.dashboard_grid.setContentsMargins(0, 0, 0, 0); self.dashboard_grid.setHorizontalSpacing(10); self.dashboard_grid.setVerticalSpacing(10)
        self.dashboard_cards = []
        self.card_sales = DashboardActionCard('💰', 'المبيعات', 'بيع، فواتير، وتحويلات الزبائن من نفس المسار.')
        self.card_items = DashboardActionCard('📦', 'الأصناف', 'المخزون، الأسعار، والكميات بشكل مرتب وواضح.')
        self.card_profit = DashboardActionCard('📊', 'الأرباح', 'الربح، التوزيع، وحصص الشراكة بواجهة مختصرة.')
        self.card_cash = DashboardActionCard('🧾', 'الصندوق', 'القاصة، الكاش الفعلي، والعجز والحركة اليومية.')
        for card, fn in [(self.card_sales,self.open_sales),(self.card_items,self.open_items),(self.card_profit,self.open_profit),(self.card_cash,self.open_cash)]:
            card.clicked.connect(fn); self.dashboard_cards.append(card); card.setMinimumHeight(176); card.setMaximumHeight(188)
        self.dashboard_grid.addWidget(self.card_sales, 0, 0)
        self.dashboard_grid.addWidget(self.card_items, 0, 1)
        self.dashboard_grid.addWidget(self.card_profit, 1, 0)
        self.dashboard_grid.addWidget(self.card_cash, 1, 1)
        self.dashboard_grid.setColumnStretch(0, 1); self.dashboard_grid.setColumnStretch(1, 1)
        self.dashboard_grid.setRowStretch(0, 1); self.dashboard_grid.setRowStretch(1, 1)
        dashboard_outer.addLayout(self.dashboard_grid)

        self.focus_strip = QFrame(); self.focus_strip.setObjectName('focusStrip'); self.focus_strip.setStyleSheet(CARD_FRAME_STYLE)
        focus_layout = QHBoxLayout(self.focus_strip); focus_layout.setContentsMargins(12, 10, 12, 10); focus_layout.setSpacing(10)
        self.focus_left = QLabel('')
        self.focus_left.setAlignment(Qt.AlignRight); self.focus_left.setWordWrap(True); self.focus_left.setStyleSheet(f'font-size:10px;font-weight:800;color:{TEXT};')
        focus_layout.addWidget(self.focus_left, 1)
        self.focus_strip.hide()
        dashboard_outer.addWidget(self.focus_strip)
        content_wrap.addWidget(dashboard_wrap)

        lower_row = QHBoxLayout(); lower_row.setSpacing(10)

        stats_wrap = QFrame(); stats_wrap.setStyleSheet(CARD_FRAME_STYLE)
        stats_outer = QVBoxLayout(stats_wrap); stats_outer.setContentsMargins(12, 12, 12, 12); stats_outer.setSpacing(8)
        stats_title = QLabel('مؤشرات التشغيل'); stats_title.setAlignment(Qt.AlignRight); stats_title.setStyleSheet('font-size:16px;font-weight:900;')
        stats_outer.addWidget(stats_title)
        self.stat_grid = QGridLayout(); self.stat_grid.setHorizontalSpacing(8); self.stat_grid.setVerticalSpacing(8)
        self.stat_labels = {}
        for i,(k,label,note) in enumerate([('cash_total','مبلغ القاصة', 'القيمة النظرية حسب الحركة'),('agents_custody','عهدة المندوبين', 'المبالغ المتبقية على المندوبين'),('total_deficit','مبلغ العجز', 'إذا موجود يحتاج متابعة'),('customer_dues','ديون الزبائن', 'التحصيلات القادمة'),('payables','ذمم الموردين', 'الالتزامات المفتوحة')]):
            card = SummaryCard(label, '0 د.ع', note, [ACCENT, ACCENT2, '#a997ff', '#ff9bca'][i % 4]); self.stat_labels[k] = card; self.stat_grid.addWidget(card, i // 2, i % 2)
        stats_outer.addLayout(self.stat_grid)
        lower_row.addWidget(stats_wrap, 1)

        content_wrap.addLayout(lower_row)

        alerts_wrap = QFrame(); alerts_wrap.setObjectName('alertsStrip'); alerts_wrap.setStyleSheet(CARD_FRAME_STYLE)
        alerts_wrap.setFixedHeight(72)
        alerts_layout = QHBoxLayout(alerts_wrap); alerts_layout.setContentsMargins(12, 10, 12, 10); alerts_layout.setSpacing(8)
        self.alerts_title = QLabel('التنبيهات'); self.alerts_title.setAlignment(Qt.AlignRight | Qt.AlignVCenter); self.alerts_title.setStyleSheet('font-size:14px;font-weight:900;')
        self.alert_badge = QLabel('0'); self.alert_badge.setAlignment(Qt.AlignCenter); self.alert_badge.setFixedSize(40, 32)
        self.notifications_badge = QLabel('0'); self.notifications_badge.setAlignment(Qt.AlignCenter); self.notifications_badge.setFixedSize(34, 30)
        self.notifications_btn = QPushButton('🔔'); self.notifications_btn.clicked.connect(self.show_notifications_dialog); self.notifications_btn.setFixedHeight(34); self.notifications_btn.setFixedWidth(46)
        self.notifications_sound_btn = QPushButton('🔊'); self.notifications_sound_btn.clicked.connect(self.toggle_notification_sound); self.notifications_sound_btn.setFixedHeight(34); self.notifications_sound_btn.setFixedWidth(46)
        self.notifications_log_btn = QPushButton('📜'); self.notifications_log_btn.clicked.connect(self.show_alert_log_dialog); self.notifications_log_btn.setFixedHeight(34); self.notifications_log_btn.setFixedWidth(46)
        strip_right = QHBoxLayout(); strip_right.setSpacing(8)
        strip_right.addWidget(self.notifications_log_btn)
        strip_right.addWidget(self.notifications_sound_btn)
        strip_right.addWidget(self.notifications_btn)
        strip_right.addWidget(self.notifications_badge)
        strip_right.addWidget(self.alert_badge)
        strip_right.addWidget(self.alerts_title)
        self.alerts_strip_text = QLabel('الوضع ممتاز حالياً، لا توجد تنبيهات حرجة.')
        self.alerts_strip_text.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        self.alerts_strip_text.setStyleSheet(f'font-size:11px;font-weight:800;color:{TEXT};padding:0 10px;')
        self.alerts_strip_text.setMinimumHeight(42)
        self.alerts_strip_text.setWordWrap(False)
        self.alerts_strip_text.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Preferred)
        alerts_layout.addLayout(strip_right)
        alerts_layout.addWidget(self.alerts_strip_text, 1)
        self.alerts_hint = QLabel(''); self.alerts_hint.hide()
        self.alerts_list = QListWidget(); self.alerts_list.hide()
        content_wrap.addWidget(alerts_wrap)


    def _sync_scroll_slider(self, minimum, maximum):
        try:
            self.scroll_slider.blockSignals(True)
            self.scroll_slider.setRange(int(minimum), int(maximum if maximum > minimum else minimum + 1))
            self.scroll_slider.setPageStep(max(40, int(self.scroll.verticalScrollBar().pageStep())))
            self.scroll_slider.setValue(int(self.scroll.verticalScrollBar().value()))
        finally:
            self.scroll_slider.blockSignals(False)

    def apply_theme(self):
        apply_theme_to_widget(self)
        try:
            self.sidebar.setStyleSheet(f'QFrame{{background:{CARD};border:1px solid {BORDER};border-radius:22px;}}')
            self.theme_combo.setStyleSheet(INPUT_STYLE)
            self.apply_theme_btn.setStyleSheet(BUTTON_STYLE)
            self.quick_status.setStyleSheet(f'background:{ACCENT};color:{TEXT_ON_ACCENT};border-radius:14px;padding:10px;font-weight:bold;')
            self.toast.setStyleSheet(f'background:{CARD};color:{TEXT};border:1px solid {BORDER};border-radius:14px;padding:10px 12px;font-weight:700;')
            self.cash_chip.setStyleSheet(f'background:{DARK};color:{TEXT};border:1px solid {BORDER};border-radius:12px;padding:10px;font-weight:700;')
            self.alert_badge.setStyleSheet(f'background:{ACCENT};color:{TEXT_ON_ACCENT};border-radius:11px;padding:6px 12px;font-weight:900;')
            self.notifications_badge.setStyleSheet(f'background:{BORDER};color:{TEXT};border-radius:11px;padding:6px 10px;font-weight:900;')
            self.alerts_list.setStyleSheet(TABLE_STYLE + f" QListWidget::item{{padding:10px;border-bottom:1px solid {BORDER};}} QListWidget::item:selected{{background:{ACCENT2};color:{TEXT_ON_ACCENT};}}")
            self.alerts_strip_text.setStyleSheet(f'background:{DARK};color:{TEXT};border:none;border-radius:16px;padding:0 12px;font-size:11px;font-weight:800;')
            sidebar_btn_style = (
                f"QPushButton {{"
                f"background-color: {CARD};"
                f"color: {TEXT};"
                f"text-align: right;"
                f"padding: 12px 14px;"
                f"border-radius: 12px;"
                f"border: 1px solid {BORDER};"
                f"font-size: 14px;"
                f"font-weight: 700;"
                f"}}"
                f"QPushButton:hover {{"
                f"background-color: {ACCENT2};"
                f"color: {TEXT_ON_ACCENT};"
                f"border: 1px solid {ACCENT2};"
                f"}}"
            )
            dash_btn_style = BUTTON_STYLE + 'QPushButton[dashboardCard="true"]{text-align:center; line-height:1.3; font-size:16px; padding:14px; border-radius:18px;}'
            for b in self.sidebar_buttons:
                b.setStyleSheet(sidebar_btn_style)
            for b in self.findChildren(QPushButton):
                if b.property('dashboardCard'):
                    b.setStyleSheet(dash_btn_style)
        except Exception:
            pass

    def show_toast(self, title, message, level='info', timeout_ms=3200):
        try:
            if self.last_toast is not None:
                try:
                    self.last_toast.close()
                except Exception:
                    pass
            toast = ToastNotification(self, title, message, level=level, timeout_ms=timeout_ms)
            max_width = min(430, max(260, self.width() - 80))
            toast.setMaximumWidth(max_width)
            toast.adjustSize()
            x = max(16, self.width() - toast.width() - 18)
            y = 18
            toast.move(QPoint(x, y))
            toast.show()
            toast.raise_()
            self.last_toast = toast
        except Exception:
            pass

    def collect_notifications(self):
        notes = []
        payables_value = total_payables(self.db)
        cash_total_value = cash_balance(self.db)
        cash_value = actual_cash_on_hand(self.db)
        st = person_profit_status(self.db)
        deficit_value = _safe_float(st.get('total_deficit', 0))
        commitments_value = total_commitments(self.db)

        if deficit_value > 0:
            notes.append({'level': 'danger', 'title': 'عجز بالقاصة', 'message': f'يوجد عجز بقيمة {fmt_money(deficit_value)} د.ع ويحتاج متابعة.'})
        if payables_value > 0:
            notes.append({'level': 'warning', 'title': 'ذمم موردين', 'message': f'ذمم الموردين الحالية {fmt_money(payables_value)} د.ع.'})
        if commitments_value > 0:
            notes.append({'level': 'info', 'title': 'التزامات قائمة', 'message': f'إجمالي الالتزامات الحالية {fmt_money(commitments_value)} د.ع.'})
        if cash_value < 0 or cash_total_value < 0:
            notes.append({'level': 'danger', 'title': 'رصيد كاش غير طبيعي', 'message': 'أكو قيمة سالبة بالكاش أو القاصة، راجع الحركات الأخيرة.'})

        low_stock = []
        threshold = _safe_int(self.db.get('settings', {}).get('low_stock_threshold', 5))
        for item in self.db.get('items', []):
            qty = _safe_int(item.get('qty', 0))
            if qty <= threshold:
                low_stock.append(f"{item.get('name', '')} ({qty})")
        if low_stock:
            preview = '، '.join(low_stock[:4])
            extra = '' if len(low_stock) <= 4 else f' +{len(low_stock)-4}'
            notes.append({'level': 'warning', 'title': 'مخزون منخفض', 'message': f'أصناف قربت تخلص: {preview}{extra}'})

        customer_due = total_customer_dues(self.db)
        if customer_due > 0:
            notes.append({'level': 'info', 'title': 'ديون زبائن', 'message': f'إجمالي الديون المسجلة على الزبائن {fmt_money(customer_due)} د.ع.'})

        return notes

    def update_notifications_ui(self, notes=None):
        notes = notes if notes is not None else self.collect_notifications()
        count = len(notes)
        if hasattr(self, 'notifications_badge'):
            self.notifications_badge.setText(str(count))
            if count == 0:
                self.notifications_badge.setStyleSheet(f'background:{BORDER};color:{TEXT};border-radius:17px;font-weight:900;padding:6px 8px;')
                self.notifications_btn.setText('🔔 التنبيهات')
            else:
                top = max(notes, key=lambda n: notification_severity(n.get('level', 'info')))
                top_level = top.get('level', 'info')
                badge_bg, badge_fg, _ = level_colors(top_level)
                self.notifications_badge.setStyleSheet(f'background:{badge_bg};color:{badge_fg};border-radius:17px;font-weight:900;padding:6px 8px;')
                self.notifications_btn.setText(f'🔔 التنبيهات ({count})')
            self.sync_notification_sound_button()
        return notes

    def show_notifications_dialog(self):
        notes = self.update_notifications_ui()
        if not notes:
            QMessageBox.information(self, 'التنبيهات', 'ماكو تنبيهات حالياً، الوضع مستقر.')
            return
        lines = []
        for i, note in enumerate(sorted(notes, key=lambda n: notification_severity(n.get('level', 'info')), reverse=True), 1):
            icon = {'danger': '⛔', 'warning': '⚠️', 'success': '✅', 'info': 'ℹ️'}.get(note.get('level', 'info'), '•')
            lines.append(f"{i}. {icon} {note.get('title', '')}\n{note.get('message', '')}")
        QMessageBox.information(self, 'مركز التنبيهات', '\n\n'.join(lines))

    def periodic_notification_check(self):
        notes = self.update_notifications_ui()
        signature = ' | '.join(f"{n.get('level','')}::{n.get('title','')}::{n.get('message','')}" for n in notes[:8])
        if signature and signature != self.last_notification_signature:
            self.last_notification_signature = signature
            top = max(notes, key=lambda n: notification_severity(n.get('level', 'info')))
            self.show_toast(f"{top.get('title', 'تنبيه')} — {top.get('message', '')}", level=top.get('level', 'info'), timeout=4200, play_sound=True)

    def resizeEvent(self, event):
        super().resizeEvent(event)
        try:
            if self.last_toast is not None and self.last_toast.isVisible():
                x = max(16, self.width() - self.last_toast.width() - 18)
                self.last_toast.move(QPoint(x, 18))
        except Exception:
            pass

    def apply_selected_theme(self):
        theme_key = self.theme_combo.currentData()
        self.db.setdefault('settings', {})['current_theme'] = theme_key
        save_db(self.db)
        set_theme(theme_key)
        self.apply_theme()
        self.refresh_dashboard()
        self.show_toast('تم تطبيق الثيم وحفظه.', 'success')

    def notification_settings(self):
        return self.db.setdefault('settings', {})

    def sync_notification_sound_button(self):
        enabled = bool(self.notification_settings().get('notification_sound_enabled', True))
        if hasattr(self, 'notifications_sound_btn'):
            self.notifications_sound_btn.setText('🔊 الصوت: تشغيل' if enabled else '🔇 الصوت: إيقاف')

    def toggle_notification_sound(self):
        settings = self.notification_settings()
        enabled = not bool(settings.get('notification_sound_enabled', True))
        settings['notification_sound_enabled'] = enabled
        save_db(self.db)
        self.sync_notification_sound_button()
        self.show_toast('تم تشغيل صوت التنبيهات.' if enabled else 'تم إيقاف صوت التنبيهات.', 'info', 2400, play_sound=False)

    def play_notification_sound(self, level='info'):
        settings = self.notification_settings()
        if not settings.get('notification_sound_enabled', True):
            return
        if settings.get('critical_only_sound', False) and level not in ('warning', 'danger'):
            return
        try:
            QApplication.beep()
            if level == 'danger':
                QTimer.singleShot(170, QApplication.beep)
        except Exception:
            pass

    def append_notification_log(self, title, message, level='info', source='system'):
        self.db.setdefault('notifications_log', [])
        entry = {
            'title': str(title or ''),
            'message': str(message or ''),
            'level': str(level or 'info'),
            'source': str(source or 'system'),
            'created_at': now_str(),
        }
        self.db['notifications_log'].append(entry)
        self.db['notifications_log'] = self.db['notifications_log'][-300:]

    def append_operation_log(self, action, details='', level='info'):
        self.db.setdefault('operations_log', [])
        entry = {
            'timestamp': now_str(),
            'action': str(action or '').strip() or 'عملية',
            'details': str(details or '').strip(),
            'level': level or 'info',
        }
        self.db['operations_log'].append(entry)
        self.db['operations_log'] = self.db['operations_log'][-500:]

    def create_manual_backup(self):
        try:
            path = create_backup_file(self.db, 'manual_backup', self.db.get('settings', {}).get('backup_keep_files', 20))
            self.append_operation_log('نسخة احتياطية يدوية', f'تم إنشاء النسخة: {path.name}', 'success')
            self.append_notification_log('نسخة احتياطية', f'تم إنشاء نسخة احتياطية يدوية داخل {path.parent}.', 'success', source='manual_backup')
            self.show_toast('تم إنشاء نسخة احتياطية جديدة.', 'success', 2600, play_sound=False)
            QMessageBox.information(self, 'تم', f'تم إنشاء النسخة الاحتياطية بنجاح:\n{path}')
        except Exception as e:
            self.append_operation_log('فشل النسخة الاحتياطية', str(e), 'danger')
            QMessageBox.warning(self, 'خطأ', f'تعذر إنشاء النسخة الاحتياطية:\n{e}')

    def show_operations_log_dialog(self):
        dialog = QDialog(self)
        dialog.setWindowTitle('سجل العمليات')
        dialog.resize(860, 520)
        lay = QVBoxLayout(dialog)
        table = QTableWidget()
        rows = list(reversed(self.db.get('operations_log', [])))
        table.setColumnCount(4)
        table.setHorizontalHeaderLabels(['الوقت', 'العملية', 'التفاصيل', 'المستوى'])
        table.setRowCount(len(rows))
        for r, row in enumerate(rows):
            vals = [row.get('timestamp', ''), row.get('action', ''), row.get('details', ''), row.get('level', '')]
            for c, val in enumerate(vals):
                item = QTableWidgetItem(str(val))
                item.setTextAlignment(Qt.AlignCenter if c != 2 else (Qt.AlignRight | Qt.AlignVCenter))
                table.setItem(r, c, item)
        table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeToContents)
        table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeToContents)
        table.horizontalHeader().setSectionResizeMode(2, QHeaderView.Stretch)
        table.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeToContents)
        lay.addWidget(table)
        btn_row = QHBoxLayout()
        clear_btn = QPushButton('🗑️ تفريغ السجل')
        close_btn = QPushButton('إغلاق')
        btn_row.addWidget(clear_btn)
        btn_row.addStretch(1)
        btn_row.addWidget(close_btn)
        lay.addLayout(btn_row)
        clear_btn.clicked.connect(lambda: (self.db.__setitem__('operations_log', []), save_db(self.db), dialog.accept()))
        close_btn.clicked.connect(dialog.accept)
        self.apply_theme()
        dialog.exec()

    def show_system_settings_dialog(self):
        settings = self.db.setdefault('settings', {})
        dialog = QDialog(self)
        dialog.setWindowTitle('إعدادات النظام')
        dialog.resize(560, 560)
        lay = QVBoxLayout(dialog)
        form = QFormLayout()

        autosave_spin = QSpinBox(); autosave_spin.setRange(15, 3600); autosave_spin.setValue(int(settings.get('autosave_interval_sec', 60) or 60))
        low_stock_spin = QSpinBox(); low_stock_spin.setRange(0, 999999); low_stock_spin.setValue(int(settings.get('low_stock_threshold', 5) or 5))
        keep_spin = QSpinBox(); keep_spin.setRange(3, 200); keep_spin.setValue(int(settings.get('backup_keep_files', 20) or 20))
        autosave_backup_spin = QSpinBox(); autosave_backup_spin.setRange(1, 100); autosave_backup_spin.setValue(int(settings.get('autosave_backup_every', 5) or 5))

        sound_only_critical = QCheckBox('الصوت فقط للتنبيهات التحذيرية والخطرة')
        sound_only_critical.setChecked(bool(settings.get('critical_only_sound', False)))
        toast_enabled = QCheckBox('تفعيل Toast داخل الواجهة')
        toast_enabled.setChecked(bool(settings.get('toast_notifications_enabled', True)))
        backup_manual_save = QCheckBox('إنشاء نسخة احتياطية مع الحفظ اليدوي')
        backup_manual_save.setChecked(bool(settings.get('create_backup_on_manual_save', True)))
        backup_export = QCheckBox('إنشاء نسخة احتياطية مع التصدير')
        backup_export.setChecked(bool(settings.get('create_backup_on_export', True)))
        confirm_logout = QCheckBox('تأكيد قبل تسجيل الخروج')
        confirm_logout.setChecked(bool(settings.get('confirm_before_logout', True)))
        confirm_export = QCheckBox('تأكيد قبل التصدير')
        confirm_export.setChecked(bool(settings.get('confirm_before_export', False)))
        confirm_exit = QCheckBox('تأكيد قبل إغلاق البرنامج')
        confirm_exit.setChecked(bool(settings.get('confirm_before_app_exit', True)))

        form.addRow('فاصل الحفظ التلقائي (ثانية)', autosave_spin)
        form.addRow('حد المخزون المنخفض', low_stock_spin)
        form.addRow('عدد النسخ الاحتياطية المحفوظة', keep_spin)
        form.addRow('نسخة احتياطية كل X حفظ تلقائي', autosave_backup_spin)
        form.addRow('', toast_enabled)
        form.addRow('', sound_only_critical)
        form.addRow('', backup_manual_save)
        form.addRow('', backup_export)
        form.addRow('', confirm_logout)
        form.addRow('', confirm_export)
        form.addRow('', confirm_exit)
        lay.addLayout(form)

        note = QLabel('هاي الإعدادات تضبط التنبيهات والنسخ الاحتياطية والتأكيدات الذكية بدون ما تغيّر شغلك الأساسي. زر إعادة ضبط المصنع يمسح كل البيانات والإعدادات ويرجع النظام لأول تشغيل.')
        note.setWordWrap(True)
        lay.addWidget(note)

        danger_note = QLabel('تحذير: إعادة ضبط المصنع تمسح الداتا، التخصيصات، القاموس، المرفقات، الفواتير المحلية، وكلمة المرور الحالية.')
        danger_note.setWordWrap(True)
        danger_note.setStyleSheet(f'color:{ACCENT}; font-weight:800;')
        lay.addWidget(danger_note)

        btn_row = QHBoxLayout()
        factory_btn = QPushButton('إعادة ضبط المصنع')
        factory_btn.setStyleSheet(f"QPushButton{{background:#7f1d1d;color:#fff;border:1px solid #ef4444;border-radius:12px;padding:10px 18px;font-weight:900;}} QPushButton:hover{{background:#991b1b;}}")
        save_btn = QPushButton('حفظ الإعدادات')
        close_btn = QPushButton('إلغاء')
        btn_row.addWidget(factory_btn)
        btn_row.addStretch(1)
        btn_row.addWidget(save_btn)
        btn_row.addWidget(close_btn)
        lay.addLayout(btn_row)

        def commit():
            settings['autosave_interval_sec'] = int(autosave_spin.value())
            settings['low_stock_threshold'] = int(low_stock_spin.value())
            settings['backup_keep_files'] = int(keep_spin.value())
            settings['autosave_backup_every'] = int(autosave_backup_spin.value())
            settings['critical_only_sound'] = bool(sound_only_critical.isChecked())
            settings['toast_notifications_enabled'] = bool(toast_enabled.isChecked())
            settings['create_backup_on_manual_save'] = bool(backup_manual_save.isChecked())
            settings['create_backup_on_export'] = bool(backup_export.isChecked())
            settings['confirm_before_logout'] = bool(confirm_logout.isChecked())
            settings['confirm_before_export'] = bool(confirm_export.isChecked())
            settings['confirm_before_app_exit'] = bool(confirm_exit.isChecked())
            save_db(self.db)
            self.append_operation_log('تحديث إعدادات النظام', 'تم حفظ إعدادات التنبيهات والنسخ الاحتياطية والتأكيدات.', 'success')
            if getattr(self, 'autosave_timer', None) is not None:
                self.autosave_timer.setInterval(max(15000, int(settings['autosave_interval_sec']) * 1000))
            self.refresh_dashboard()
            self.show_toast('تم حفظ إعدادات النظام.', 'success', 2400, play_sound=False)
            dialog.accept()

        def do_factory_reset():
            first = QMessageBox.question(
                dialog,
                'تأكيد خطير',
                'هذي العملية راح تمسح كل بيانات البرنامج وترجعه لأول تشغيل. هل تريد المتابعة؟',
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.No,
            )
            if first != QMessageBox.Yes:
                return
            second = QMessageBox.question(
                dialog,
                'تأكيد نهائي',
                'تأكيد نهائي: سيتم حذف البيانات، الإعدادات، القاموس، المرفقات، الفواتير المحلية، وكلمة المرور. هل أنت متأكد 100%؟',
                QMessageBox.Yes | QMessageBox.No,
                QMessageBox.No,
            )
            if second != QMessageBox.Yes:
                return
            try:
                self.db = factory_reset_all_data()
                if getattr(self, 'windows', None) is not None:
                    self.windows.clear()
                if getattr(self, 'autosave_timer', None) is not None:
                    self.autosave_counter = 0
                    self.autosave_timer.setInterval(max(15000, int(self.db.get('settings', {}).get('autosave_interval_sec', 60)) * 1000))
                self.refresh_dashboard()
                self.append_operation_log('إعادة ضبط المصنع', 'تم تصفير كل بيانات النظام وإرجاعه لأول تشغيل.', 'warning')
                try:
                    save_db(self.db)
                except Exception:
                    pass
                QMessageBox.information(dialog, 'تمت العملية', 'تمت إعادة ضبط المصنع بنجاح. سيعود النظام الآن ببيانات نظيفة.')
                dialog.accept()
            except Exception as e:
                QMessageBox.warning(dialog, 'خطأ', f'تعذر تنفيذ إعادة ضبط المصنع\n{e}')

        save_btn.clicked.connect(commit)
        close_btn.clicked.connect(dialog.reject)
        factory_btn.clicked.connect(do_factory_reset)
        self.apply_theme()
        dialog.exec()

    def show_alert_log_dialog(self):
        dialog = QDialog(self)
        dialog.setWindowTitle('سجل التنبيهات')
        dialog.resize(760, 520)
        apply_theme_to_widget(dialog)
        lay = QVBoxLayout(dialog)
        lay.setContentsMargins(14, 14, 14, 14)
        lay.setSpacing(10)
        info = QLabel('هذا سجل التنبيهات المحفوظ داخل النظام. الأحدث بالأعلى.')
        info.setWordWrap(True)
        info.setAlignment(Qt.AlignRight)
        lay.addWidget(info)
        lst = QListWidget()
        lst.setStyleSheet(TABLE_STYLE + f" QListWidget::item{{padding:10px;border-bottom:1px solid {BORDER};}}")
        rows = list(reversed(self.db.get('notifications_log', [])))
        if not rows:
            lst.addItem('ماكو سجل تنبيهات محفوظ حالياً.')
        for row in rows:
            icon = {'danger': '⛔', 'warning': '⚠️', 'success': '✅', 'info': 'ℹ️'}.get(row.get('level', 'info'), '•')
            item = QListWidgetItem(f"{icon} [{row.get('created_at', '')}] {row.get('title', '')}\n{row.get('message', '')}")
            bg, fg, _ = level_colors(row.get('level', 'info'))
            item.setBackground(bg)
            item.setForeground(fg)
            item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            lst.addItem(item)
        lay.addWidget(lst, 1)
        btns = QHBoxLayout()
        clear_btn = QPushButton('🗑️ مسح السجل')
        close_btn = QPushButton('إغلاق')
        clear_btn.clicked.connect(lambda: (self.db.__setitem__('notifications_log', []), save_db(self.db), dialog.accept()))
        close_btn.clicked.connect(dialog.accept)
        btns.addWidget(clear_btn)
        btns.addWidget(close_btn)
        lay.addLayout(btns)
        apply_theme_to_widget(dialog)
        dialog.exec()

    def show_toast(self, message, level='info', timeout=3200, play_sound=False):
        try:
            if not self.notification_settings().get('toast_notifications_enabled', True):
                if play_sound:
                    self.play_notification_sound(level)
                return
            bg, fg, border = level_colors(level)
            self.toast.setText(message)
            self.toast.setStyleSheet(f'background:{bg};color:{fg};border:1px solid {border};border-radius:14px;padding:10px 12px;font-weight:800;')
            self.toast.show()
            if play_sound:
                self.play_notification_sound(level)
            QTimer.singleShot(timeout, self.toast.hide)
        except Exception:
            pass

    def collect_alerts(self):
        alerts = []
        threshold = _safe_int(self.db.get('settings', {}).get('low_stock_threshold', 5))
        low_stock = []
        for item in self.db.get('items', []):
            qty = _safe_int(item.get('qty', 0))
            if qty <= threshold:
                low_stock.append((item.get('name', ''), qty))
        low_stock.sort(key=lambda x: (x[1], x[0]))
        if low_stock:
            names = '، '.join(f"{name} ({qty})" for name, qty in low_stock[:4])
            extra = '' if len(low_stock) <= 4 else f' +{len(low_stock)-4} صنف'
            alerts.append({'level': 'warning', 'key': f'low_stock_{len(low_stock)}', 'text': f'مخزون منخفض: {names}{extra}'})

        payables_value = total_payables(self.db)
        if payables_value > 0:
            alerts.append({'level': 'warning', 'key': f'payables_{int(payables_value)}', 'text': f'ذمم الموردين الحالية: {fmt_money(payables_value)} د.ع'})

        st = person_profit_status(self.db)
        deficit_value = _safe_float(st.get('total_deficit', 0))
        if deficit_value > 0:
            alerts.append({'level': 'danger', 'key': f'deficit_{int(deficit_value)}', 'text': f'يوجد عجز فعلي بقيمة {fmt_money(deficit_value)} د.ع ويحتاج معالجة.'})

        opening_locked = bool(opening_data(self.db).get('opening_locked', False))
        if opening_locked:
            alerts.append({'level': 'info', 'key': 'opening_locked', 'text': 'التهيئة الافتتاحية مقفلة حالياً لحماية الأرصدة القديمة.'})

        customer_due_total = total_customer_dues(self.db)
        if customer_due_total > 0:
            alerts.append({'level': 'info', 'key': f'customer_dues_{int(customer_due_total)}', 'text': f'ديون الزبائن المفتوحة: {fmt_money(customer_due_total)} د.ع'})

        if self.last_autosave_status:
            alerts.append({'level': 'success', 'key': f'autosave_{self.last_autosave_status}', 'text': f'آخر حفظ تلقائي: {fmt_datetime_text(self.last_autosave_status)}'})
        elif self.last_manual_save_status:
            alerts.append({'level': 'success', 'key': f'manualsave_{self.last_manual_save_status}', 'text': f'آخر حفظ يدوي: {fmt_datetime_text(self.last_manual_save_status)}'})

        if self.last_export_status:
            alerts.append({'level': 'info', 'key': f'export_{self.last_export_status}', 'text': f'آخر تصدير Gmail: {fmt_datetime_text(self.last_export_status)}'})

        if not alerts:
            alerts.append({'level': 'success', 'key': 'all_good', 'text': 'الوضع ممتاز حالياً، لا توجد تنبيهات حرجة.'})
        return alerts

    def update_alerts_view(self, alerts):
        try:
            self.alerts_list.clear()
            danger_count = 0
            warning_count = 0
            strip_parts = []
            ordered = sorted(alerts, key=lambda row: notification_severity(row.get('level', 'info')), reverse=True)
            for row in ordered:
                level = row.get('level', 'info')
                if level == 'danger':
                    danger_count += 1
                elif level == 'warning':
                    warning_count += 1
                icon = {'danger': '⛔', 'warning': '⚠️', 'success': '✅', 'info': 'ℹ️'}.get(level, '•')
                item = QListWidgetItem(f"{icon} {row.get('text', '')}")
                bg, fg, _ = level_colors(level)
                item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
                item.setBackground(bg)
                item.setForeground(fg)
                self.alerts_list.addItem(item)
                if len(strip_parts) < 3 and row.get('text'):
                    strip_parts.append(f"{icon} {row.get('text', '')}")
            self.alert_badge.setText(str(len(alerts)))
            if danger_count:
                self.alerts_hint.setText(f'عندك {danger_count} تنبيه أحمر يحتاج تدخل سريع.')
            elif warning_count:
                self.alerts_hint.setText(f'عندك {warning_count} تنبيه أصفر يحتاج متابعة.')
            else:
                self.alerts_hint.setText('التنبيهات هنا تتحدّث تلقائياً حسب الكاش، المخزون، الديون، الحفظ، والتصدير.')
            self.alerts_strip_text.setText('   •   '.join(strip_parts) if strip_parts else 'الوضع ممتاز حالياً، لا توجد تنبيهات حرجة.')
            self.alerts_strip_text.setToolTip(self.alerts_hint.text())
            self.update_notifications_ui([{'level': row.get('level', 'info'), 'title': 'تنبيه', 'message': row.get('text', '')} for row in ordered if row.get('key') != 'all_good'])
        except Exception:
            pass

    def monitor_alert_changes(self):
        alerts = self.collect_alerts()
        self.update_alerts_view(alerts)
        current = {row.get('key'): row for row in alerts}
        critical = [row for key, row in current.items() if key not in self.alert_state and row.get('level') in ('danger', 'warning')]
        for row in critical:
            self.append_notification_log('تنبيه جديد', row.get('text', ''), row.get('level', 'warning'), source='alert_monitor')
        if critical:
            first = sorted(critical, key=lambda row: notification_severity(row.get('level', 'info')), reverse=True)[0]
            self.show_toast(first.get('text', ''), first.get('level', 'warning'), 4200, play_sound=True)
        self.alert_state = current
        self.last_alert_scan = now_str()

    def save_all(self, quiet=False):
        try:
            save_db(self.db)
            stamp = now_str()
            if quiet:
                self.last_autosave_status = stamp
                self.append_notification_log('حفظ تلقائي', 'تم حفظ البيانات تلقائياً بنجاح.', 'success', source='autosave')
            else:
                self.last_manual_save_status = stamp
                self.append_notification_log('حفظ يدوي', 'تم حفظ البيانات يدوياً بنجاح.', 'success', source='manual_save')
            self.refresh_dashboard()
            if not quiet:
                QMessageBox.information(self, 'تم', 'تم حفظ البيانات بنجاح.')
                self.show_toast('تم حفظ البيانات بنجاح.', 'success', 2500, play_sound=False)
            return True
        except Exception as e:
            if not quiet:
                QMessageBox.warning(self, 'خطأ', f'تعذر حفظ البيانات\n{e}')
                self.append_notification_log('فشل الحفظ', f'تعذر حفظ البيانات: {e}', 'danger', source='save_error')
                self.show_toast('تعذر حفظ البيانات.', 'danger', 3200, play_sound=True)
            return False

    def setup_autosave(self):
        self.autosave_timer = QTimer(self)
        interval_sec = int(self.db.get('settings', {}).get('autosave_interval_sec', 60) or 60)
        self.autosave_timer.setInterval(max(15000, interval_sec * 1000))
        self.autosave_timer.timeout.connect(lambda: self.save_all(True))
        self.autosave_timer.start()

        self.notification_timer = QTimer(self)
        self.notification_timer.setInterval(90000)
        self.notification_timer.timeout.connect(self.periodic_notification_check)
        self.notification_timer.start()

        self.update_notifications_ui()
        self.sync_notification_sound_button()
        QTimer.singleShot(1200, self.periodic_notification_check)

    def setup_alerts_monitor(self):
        self.alert_timer = QTimer(self)
        self.alert_timer.setInterval(20000)
        self.alert_timer.timeout.connect(self.monitor_alert_changes)
        self.alert_timer.start()
        self.monitor_alert_changes()

    def export_to_gmail(self):
        try:
            if self.db.get('settings', {}).get('confirm_before_export', False):
                ans = QMessageBox.question(self, 'تأكيد', 'راح يتم تجهيز تقرير وفتح Gmail. تكمل؟', QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes)
                if ans != QMessageBox.Yes:
                    return
            exports_dir = APP_DATA_DIR / 'exports'
            exports_dir.mkdir(parents=True, exist_ok=True)
            stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            backup_path = exports_dir / f'nokhba_backup_{stamp}.json'
            summary_path = exports_dir / f'nokhba_summary_{stamp}.pdf'
            save_db(self.db)
            backup_path.write_text(json.dumps(self.db, ensure_ascii=False, indent=2), encoding='utf-8')
            if self.db.get('settings', {}).get('create_backup_on_export', True):
                try:
                    create_backup_file(self.db, 'gmail_export', self.db.get('settings', {}).get('backup_keep_files', 20))
                except Exception:
                    pass

            payables_value = total_payables(self.db)
            cash_total_value = cash_balance(self.db)
            cash_value = actual_cash_on_hand(self.db)
            st = person_profit_status(self.db)
            deficit_value = st.get('total_deficit', 0)
            commitments_value = total_commitments(self.db)
            summary = (
                f"تقرير برنامج النخبة\n"
                f"تاريخ التصدير: {now_str()}\n"
                f"مبلغ القاصة: {fmt_money(cash_total_value)} د.ع\n"
                f"الكاش الفعلي: {fmt_money(cash_value)} د.ع\n"
                f"مبلغ العجز: {fmt_money(deficit_value)} د.ع\n"
                f"ذمم الموردين: {fmt_money(payables_value)} د.ع\n"
                f"الالتزامات: {fmt_money(commitments_value)} د.ع\n\n"
                f"تم إنشاء نسخة احتياطية JSON هنا:\n{backup_path}\n"
            )
            save_text_as_pdf('تقرير برنامج النخبة', summary, summary_path)

            subject = quote('تصدير بيانات برنامج النخبة')
            body = quote(
                "تم تجهيز تقرير ونسخة احتياطية من البرنامج.\n\n"
                f"ملف التقرير المحلي:\n{summary_path}\n\n"
                f"ملف النسخة الاحتياطية المحلي:\n{backup_path}\n\n"
                "ملاحظة: فتح الجيميل لا يرفق الملفات تلقائيًا، لازم ترفقها يدويًا من المسارات أعلاه."
            )
            webbrowser.open(f'https://mail.google.com/mail/?view=cm&fs=1&su={subject}&body={body}')
            self.last_export_status = now_str()
            self.append_operation_log('تصدير Gmail', f'تم تجهيز تقرير وملفات تصدير داخل {exports_dir}.', 'info')
            self.append_notification_log('تصدير Gmail', f'تم تجهيز التقرير والنسخة الاحتياطية داخل {exports_dir}.', 'info', source='gmail_export')
            self.show_toast('تم تجهيز ملفات التصدير وفتح Gmail.', 'info', 3200, play_sound=False)
            QMessageBox.information(
                self, 'تم',
                'تم إنشاء ملف التقرير والنسخة الاحتياطية وفتح صفحة إنشاء رسالة في الجيميل.\n'
                'ملاحظة: إرفاق الملفات يكون يدويًا من مجلد exports.'
            )
        except Exception as e:
            self.append_operation_log('فشل تصدير Gmail', str(e), 'danger')
            self.append_notification_log('فشل تصدير Gmail', f'تعذر التصدير للجيميل: {e}', 'danger', source='gmail_export')
            self.show_toast('تعذر التصدير للجيميل.', 'danger', 3200, play_sound=True)
            QMessageBox.warning(self, 'خطأ', f'تعذر التصدير للجيميل:\n{e}')


    def show_change_password_dialog(self):
        dlg = QDialog(self)
        dlg.setWindowTitle('تغيير كلمة السر')
        dlg.resize(420, 260)
        apply_branding(dlg)
        dlg.setStyleSheet(WINDOW_STYLE)
        lay = QVBoxLayout(dlg)
        lay.setContentsMargins(20, 20, 20, 20)
        lay.setSpacing(12)

        form = QFormLayout()
        form.setLabelAlignment(Qt.AlignRight)
        current = QLineEdit(); current.setEchoMode(QLineEdit.Password); current.setPlaceholderText('كلمة السر الحالية'); current.setMinimumHeight(42)
        new_pwd = QLineEdit(); new_pwd.setEchoMode(QLineEdit.Password); new_pwd.setPlaceholderText('كلمة السر الجديدة'); new_pwd.setMinimumHeight(42)
        confirm = QLineEdit(); confirm.setEchoMode(QLineEdit.Password); confirm.setPlaceholderText('تأكيد كلمة السر الجديدة'); confirm.setMinimumHeight(42)
        for w in [current, new_pwd, confirm]:
            w.setStyleSheet(INPUT_STYLE)
        form.addRow('الحالية', current)
        form.addRow('الجديدة', new_pwd)
        form.addRow('التأكيد', confirm)
        lay.addLayout(form)

        tip = QLabel('سيتم حفظ كلمة السر الجديدة محلياً بشكل محمي داخل app_data/pass.txt')
        tip.setWordWrap(True)
        tip.setStyleSheet(f'color:{MUTED};font-size:12px;')
        lay.addWidget(tip)

        btns = QHBoxLayout()
        save_btn = QPushButton('حفظ')
        cancel_btn = QPushButton('إلغاء')
        save_btn.setStyleSheet(BUTTON_STYLE)
        cancel_btn.setStyleSheet(SECONDARY_BUTTON)
        save_btn.setMinimumHeight(42)
        cancel_btn.setMinimumHeight(42)
        btns.addWidget(cancel_btn)
        btns.addWidget(save_btn)
        lay.addLayout(btns)

        cancel_btn.clicked.connect(dlg.reject)

        def submit():
            old_value = current.text().strip()
            new_value = new_pwd.text().strip()
            confirm_value = confirm.text().strip()
            saved = PASSWORD_FILE.read_text(encoding='utf-8').strip() if PASSWORD_FILE.exists() else ''
            if not saved:
                QMessageBox.warning(dlg, 'تنبيه', 'لا توجد كلمة سر محفوظة حالياً.')
                return
            if not verify_password_input(old_value, saved):
                QMessageBox.warning(dlg, 'خطأ', 'كلمة السر الحالية غير صحيحة.')
                return
            if len(new_value) < 4:
                QMessageBox.warning(dlg, 'تنبيه', 'كلمة السر الجديدة يجب أن تكون 4 أحرف على الأقل.')
                return
            if new_value != confirm_value:
                QMessageBox.warning(dlg, 'تنبيه', 'تأكيد كلمة السر غير مطابق.')
                return
            try:
                PASSWORD_FILE.write_text(password_record_for_storage(new_value), encoding='utf-8')
                self.append_operation_log('تغيير كلمة السر', 'تم تغيير كلمة السر بنجاح.', 'success')
                self.append_notification_log('تغيير كلمة السر', 'تم تحديث كلمة السر بنجاح.', 'success', source='security')
                self.show_toast('تم تغيير كلمة السر بنجاح.', 'success', 2200, play_sound=False)
                QMessageBox.information(dlg, 'تم', 'تم تغيير كلمة السر بنجاح.')
                dlg.accept()
            except Exception as e:
                QMessageBox.warning(dlg, 'خطأ', f'تعذر حفظ كلمة السر الجديدة\n{e}')

        save_btn.clicked.connect(submit)
        dlg.exec()

    def logout(self):
        if self.db.get('settings', {}).get('confirm_before_logout', True):
            ans = QMessageBox.question(self, 'تأكيد تسجيل الخروج', 'هل تريد حفظ البيانات ثم تسجيل الخروج؟', QMessageBox.Yes | QMessageBox.No, QMessageBox.Yes)
            if ans != QMessageBox.Yes:
                return
        self.save_all(True)
        self.append_operation_log('تسجيل الخروج', 'تم حفظ البيانات ثم تسجيل الخروج من النظام.', 'info')
        self.append_notification_log('تسجيل الخروج', 'تم حفظ البيانات قبل تسجيل الخروج.', 'info', source='logout')
        self.show_toast('جارٍ حفظ البيانات وتسجيل الخروج...', 'info', 1600, play_sound=False)
        try:
            for win in list(self.windows.values()):
                try:
                    win.close()
                except Exception:
                    pass
            self.windows = {}
        except Exception:
            pass
        self.login_window = LoginWindow()
        self.login_window.show()
        self.close()

    def closeEvent(self, event):
        if self.db.get('settings', {}).get('confirm_before_app_exit', True):
            ans = QMessageBox.question(
                self,
                'تأكيد الإغلاق',
                'هل تريد حفظ البيانات قبل الإغلاق؟',
                QMessageBox.Yes | QMessageBox.No | QMessageBox.Cancel,
                QMessageBox.Yes
            )

            if ans == QMessageBox.Yes:
                self.save_all(True)
                self.append_operation_log('إغلاق البرنامج', 'تم إغلاق البرنامج بعد الحفظ.', 'info')
            elif ans == QMessageBox.Cancel:
                event.ignore()
                return
            else:
                self.append_operation_log('إغلاق البرنامج', 'تم إغلاق البرنامج بدون حفظ.', 'warning')
        else:
            self.append_operation_log('إغلاق البرنامج', 'تم إغلاق البرنامج.', 'info')
        event.accept()

    
    def refresh_dashboard(self):
        payables_value = total_payables(self.db)
        customer_dues_value = total_customer_dues(self.db)
        cash_total_value = cash_balance(self.db)
        agents_custody_value = total_agents_custody(self.db)
        st = person_profit_status(self.db)
        deficit_value = st['total_deficit']
        commitments_value = total_commitments(self.db)
        net_after_commitments = net_cash_after_commitments(self.db)

        self.stat_labels['cash_total'].setText(f"{fmt_money(cash_total_value)} د.ع")
        self.stat_labels['total_deficit'].setText(f"{fmt_money(deficit_value)} د.ع")
        if 'agents_custody' in self.stat_labels:
            self.stat_labels['agents_custody'].setText(f"{fmt_money(agents_custody_value)} د.ع")
        if 'customer_dues' in self.stat_labels:
            self.stat_labels['customer_dues'].setText(f"{fmt_money(customer_dues_value)} د.ع")
        if 'payables' in self.stat_labels:
            self.stat_labels['payables'].setText(f"{fmt_money(payables_value)} د.ع")

        autosave_text = f" | آخر حفظ: {self.last_autosave_status}" if self.last_autosave_status else ''
        self.cash_chip.setText(f"مبلغ القاصة: {fmt_money(cash_total_value)} د.ع{autosave_text}")
        self.update_notifications_ui()

        self.hero_cash_card.set_value(f"{fmt_money(cash_total_value)} د.ع")
        if commitments_value > 0:
            self.hero_cash_card.set_note(f"بعد الالتزامات: {fmt_money(net_after_commitments)} د.ع | بعد تنزيل العهدة: {fmt_money(cash_balance_after_custody(self.db))} د.ع")
        else:
            self.hero_cash_card.set_note(f'بعد تنزيل عهدة المندوبين: {fmt_money(cash_balance_after_custody(self.db))} د.ع')

        self.hero_dues_card.set_value(f"{fmt_money(customer_dues_value)} د.ع")
        self.hero_dues_card.set_note('')
        self.hero_payables_card.set_value(f"{fmt_money(payables_value)} د.ع")
        self.hero_payables_card.set_note('')

        if deficit_value > 0:
            self.hero_alert_card.set_value('تنبيه أحمر')
            self.hero_alert_card.set_note(f"عجز {fmt_money(deficit_value)} د.ع يحتاج معالجة")
            self.focus_right.setText('يوجد عجز يحتاج متابعة')
            self.focus_right.setStyleSheet('background:#7f1d1d;color:#ffffff;border-radius:14px;padding:10px 12px;font-size:12px;font-weight:900;')
        elif payables_value > 0 or customer_dues_value > 0:
            self.hero_alert_card.set_value('متابعة مطلوبة')
            self.hero_alert_card.set_note(f"ديون زبائن {fmt_money(customer_dues_value)} | موردين {fmt_money(payables_value)}")
            self.focus_right.setText('تشغيل مستقر مع ذمم مفتوحة')
            self.focus_right.setStyleSheet('background:#8a6b16;color:#ffffff;border-radius:14px;padding:10px 12px;font-size:12px;font-weight:900;')
        else:
            self.hero_alert_card.set_value('جاهز')
            self.hero_alert_card.set_note('لا توجد تنبيهات حرجة حالياً')
            self.focus_right.setText('تشغيل مستقر')
            self.focus_right.setStyleSheet(f'background:{ACCENT};color:{TEXT_ON_ACCENT};border-radius:14px;padding:10px 12px;font-size:12px;font-weight:900;')

        self.card_sales.set_subtitle(f"{fmt_money(customer_dues_value)} د.ع ديون مفتوحة | {len(self.db.get('sales', []))} حركة بيع")
        self.card_sales.set_meta('فتح البيع والتحصيل ←')
        low_stock_count = sum(1 for item in self.db.get('items', []) if _safe_int(item.get('qty', 0)) <= _safe_int(self.db.get('settings', {}).get('low_stock_threshold', 5)))
        self.card_items.set_subtitle(f"{len(self.db.get('items', []))} صنف | منخفض المخزون: {low_stock_count}")
        self.card_items.set_meta('فتح الأصناف والمخزن ←')
        self.card_profit.set_subtitle(f"عجز هضاب {fmt_money(st.get('hidab_balance', 0))} | عجز مصطفى {fmt_money(st.get('mostafa_balance', 0))}")
        self.card_profit.set_meta('فتح الأرباح والتوزيع ←')
        self.card_cash.set_subtitle(f"القاصة {fmt_money(cash_total_value)} | بعد تنزيل العهدة {fmt_money(cash_balance_after_custody(self.db))}")
        self.card_cash.set_meta('فتح الصندوق والمطابقة ←')
        self.focus_left.setText(f'الالتزامات الحالية {fmt_money(commitments_value)} د.ع | عهدة المندوبين {fmt_money(agents_custody_value)} د.ع | القاصة بعد العهدة {fmt_money(cash_balance_after_custody(self.db))} د.ع')

        if deficit_value > 0:
            self.quick_status.setText(f"تنبيه: العجز {fmt_money(deficit_value)} د.ع | ذمم الموردين {fmt_money(payables_value)} د.ع | افتح مركز التنبيهات")
            self.quick_status.setStyleSheet('background:#7f1d1d;color:#ffffff;border-radius:14px;padding:10px;font-weight:bold;')
        else:
            save_label = self.last_autosave_status or self.last_manual_save_status
            save_text = fmt_datetime_text(save_label) if save_label else 'لم يتم بعد'
            self.quick_status.setText(f"القاصة {fmt_money(cash_total_value)} د.ع | بعد العهدة {fmt_money(cash_balance_after_custody(self.db))} د.ع | عهدة المندوبين {fmt_money(agents_custody_value)} د.ع | آخر حفظ {save_text}")
            self.quick_status.setStyleSheet(f'background:{ACCENT};color:{TEXT_ON_ACCENT};border-radius:14px;padding:10px;font-weight:bold;')

        self.update_alerts_view(self.collect_alerts())

    def show_win(self, key, cls, *args):
        if key not in self.windows:
            self.windows[key] = cls(self, *args)
        else:
            try: self.windows[key].refresh_table()
            except Exception: pass
            try: self.windows[key].refresh_view()
            except Exception: pass
            try: self.windows[key].refresh_combos()
            except Exception: pass
            try: self.windows[key].refresh_tables()
            except Exception: pass
        try:
            self.windows[key].apply_theme()
        except Exception:
            apply_theme_to_widget(self.windows[key])
        for other_key, other_win in self.windows.items():
            if other_key != key:
                try:
                    other_win.hide()
                except Exception:
                    pass
        self.hide()
        self.windows[key].show(); self.windows[key].raise_(); self.windows[key].activateWindow()
    def open_items(self): self.show_win('items', ItemsWindow)
    def open_warehouse(self): self.show_win('warehouse', WarehouseWindow)
    def open_customers(self): self.show_win('customers', PeopleWindow, 'customers', 'الزبائن', '👥')
    def open_suppliers(self): self.show_win('suppliers', PeopleWindow, 'suppliers', 'الموردين', '🚚')
    def open_funders(self): self.show_win('funders', FundersWindow)
    def open_inbound(self): self.show_win('inbound', InboundWindow)
    def open_sales(self): self.show_win('sales', SalesWindow)
    def open_returns(self): self.show_win('returns', ReturnWindow)
    def open_damaged(self): self.show_win('damaged', DamagedWindow)
    def open_cash(self): self.show_win('cash', CashWindow)
    def open_withdrawals(self): self.show_win('withdrawals', WithdrawalsWindow)
    def open_expenses(self): self.show_win('expenses', ExpensesWindow)
    def open_opening(self): self.show_win('opening', OpeningBalancesWindow)
    def open_customer_dues(self): self.show_win('customer_dues', DuesWindow, 'customers')
    def open_supplier_dues(self): self.show_win('supplier_dues', DuesWindow, 'suppliers')
    def open_profit(self): self.show_win('profit', ProfitWindow)
    def open_agents_custody(self): self.show_win('agents_custody', AgentsCustodyWindow)
    def open_reconciliation(self): self.show_win('reconciliation', ReconciliationWindow)


class SplashScreen(QWidget):
    def __init__(self):
        super().__init__()
        self.setWindowFlags(Qt.FramelessWindowHint | Qt.WindowStaysOnTopHint)
        self.setAttribute(Qt.WA_TranslucentBackground, True)
        self.resize(520, 520)
        apply_branding(self)

        outer = QVBoxLayout(self)
        outer.setContentsMargins(0, 0, 0, 0)
        outer.setAlignment(Qt.AlignCenter)

        self.card = QFrame()
        self.card.setObjectName('splashCard')
        self.card.setFixedSize(420, 420)
        self.card.setStyleSheet(f"""
            QFrame#splashCard {{
                background:qlineargradient(x1:0,y1:0,x2:1,y2:1,stop:0 #0b1730, stop:1 #13294b);
                border:1px solid rgba(255,255,255,0.08);
                border-radius:28px;
            }}
        """)
        outer.addWidget(self.card, alignment=Qt.AlignCenter)

        layout = QVBoxLayout(self.card)
        layout.setContentsMargins(28, 28, 28, 28)
        layout.setSpacing(8)
        layout.addStretch()

        self.logo = QLabel()
        self.logo.setAlignment(Qt.AlignCenter)
        self.logo.setFixedSize(220, 220)
        layout.addWidget(self.logo, alignment=Qt.AlignCenter)

        self.title = QLabel('النخبة')
        self.title.setAlignment(Qt.AlignCenter)
        self.title.setStyleSheet('color:#7dc5ff;font-size:24px;font-weight:800;letter-spacing:3px;background:transparent;')
        self.title.setVisible(False)
        layout.addSpacing(10)
        layout.addWidget(self.title, alignment=Qt.AlignCenter)

        self.message = QLabel('مرحباً بك في برنامج النخبة')
        self.message.setAlignment(Qt.AlignCenter)
        self.message.setStyleSheet('color:rgba(255,255,255,0.82);font-size:14px;background:transparent;')
        self.message.setVisible(False)
        layout.addWidget(self.message, alignment=Qt.AlignCenter)
        layout.addStretch()

        self.title_opacity = QGraphicsOpacityEffect(self.title)
        self.title.setGraphicsEffect(self.title_opacity)
        self.title_opacity.setOpacity(0.0)

        self.message_opacity = QGraphicsOpacityEffect(self.message)
        self.message.setGraphicsEffect(self.message_opacity)
        self.message_opacity.setOpacity(0.0)

        self.original_logo_geom = self.logo.geometry()
        self._prepare_logo()
        QTimer.singleShot(80, self.start_animation)

    def _prepare_logo(self):
        size = 170
        if LOGO_PNG.exists():
            pm = QPixmap(str(LOGO_PNG))
            if not pm.isNull():
                self.logo.setPixmap(pm.scaled(size, size, Qt.KeepAspectRatio, Qt.SmoothTransformation))
                return
        self.logo.setText('ن')
        self.logo.setStyleSheet('color:#7dc5ff;font-size:120px;font-weight:900;background:transparent;')

    def start_animation(self):
        center = self.logo.geometry().center()
        start_rect = self.logo.geometry()
        start_rect.setWidth(90)
        start_rect.setHeight(90)
        start_rect.moveCenter(center)
        end_rect = self.logo.geometry()
        self.logo.setGeometry(start_rect)

        zoom = QPropertyAnimation(self.logo, b'geometry', self)
        zoom.setDuration(850)
        zoom.setStartValue(start_rect)
        zoom.setEndValue(end_rect)
        zoom.setEasingCurve(QEasingCurve.OutBack)

        self.title.setVisible(True)
        self.message.setVisible(True)
        title_fade = QPropertyAnimation(self.title_opacity, b'opacity', self)
        title_fade.setDuration(650)
        title_fade.setStartValue(0.0)
        title_fade.setEndValue(1.0)

        message_fade = QPropertyAnimation(self.message_opacity, b'opacity', self)
        message_fade.setDuration(650)
        message_fade.setStartValue(0.0)
        message_fade.setEndValue(1.0)

        self.anim_group = QParallelAnimationGroup(self)
        self.anim_group.addAnimation(zoom)
        self.anim_group.addAnimation(title_fade)
        self.anim_group.addAnimation(message_fade)
        QTimer.singleShot(450, lambda: title_fade.start())
        QTimer.singleShot(700, lambda: message_fade.start())
        zoom.start()
        QTimer.singleShot(2600, self.close)

    def show_centered(self):
        screen = QApplication.primaryScreen().availableGeometry()
        self.move(screen.center() - self.rect().center())
        self.show()



class LoginWindow(QWidget):
    def __init__(self):
        super().__init__(); self.main_window = None
        db = load_db()
        set_theme(db.get('settings', {}).get('current_theme', 'dark_lux'))
        self.setWindowTitle('تسجيل الدخول | النخبة'); self.resize(760, 600); self.setMinimumSize(700, 560)
        apply_branding(self)
        layout = QVBoxLayout(self); layout.setAlignment(Qt.AlignCenter); layout.setContentsMargins(90, 44, 90, 44); layout.setSpacing(18); layout.addStretch(1)
        card = QFrame(); card.setObjectName('loginCard'); card.setMaximumWidth(520); card.setMinimumHeight(400)
        form = QVBoxLayout(card); form.setContentsMargins(38, 34, 38, 34); form.setSpacing(14)
        badge = QLabel('')
        badge.hide()
        logo = make_logo_label(116)
        title = QLabel('تسجيل الدخول'); title.setAlignment(Qt.AlignCenter); title.setStyleSheet('font-size:30px;font-weight:900;')
        sub = QLabel('')
        sub.hide()
        self.input = QLineEdit(); self.input.setEchoMode(QLineEdit.Password); self.input.setPlaceholderText('كلمة المرور'); self.input.returnPressed.connect(self.login); self.input.setMinimumHeight(58)
        self.login_btn = QPushButton('دخول'); self.login_btn.setMinimumHeight(58); self.login_btn.clicked.connect(self.login)
        self.error_label = QLabel(''); self.error_label.setAlignment(Qt.AlignCenter); self.error_label.setStyleSheet('font-size:12px;color:#ff90b6;font-weight:800;')
        note = QLabel('')
        note.hide()
        for w in [badge, logo, title, sub, self.input, self.login_btn, self.error_label, note]: form.addWidget(w)
        layout.addWidget(card, alignment=Qt.AlignCenter); layout.addStretch(1)
        self.login_card = card
        self.apply_theme()
    def paintEvent(self, event):
        paint_app_background(self, event)
    def _sync_scroll_slider(self, minimum, maximum):
        try:
            self.scroll_slider.blockSignals(True)
            self.scroll_slider.setRange(int(minimum), int(maximum if maximum > minimum else minimum + 1))
            self.scroll_slider.setPageStep(max(40, int(self.scroll.verticalScrollBar().pageStep())))
            self.scroll_slider.setValue(int(self.scroll.verticalScrollBar().value()))
        finally:
            self.scroll_slider.blockSignals(False)

    def apply_theme(self):
        apply_theme_to_widget(self)
        self.login_card.setStyleSheet(f"QFrame#loginCard{{background:qlineargradient(x1:0,y1:0,x2:1,y2:1, stop:0 {rgba_from_hex(CARD,0.98)}, stop:1 {rgba_from_hex(CARD,0.76)});border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:34px;}}")
        self.input.setStyleSheet(f"QLineEdit{{background-color:{rgba_from_hex('#ffffff',0.08)};border:1px solid {rgba_from_hex(TEXT,0.08)};border-radius:20px;color:{TEXT};padding:0 18px;font-size:15px;font-weight:700;}} QLineEdit:focus{{border:1px solid {rgba_from_hex(ACCENT,0.50)};}}")
        self.login_btn.setStyleSheet(f"QPushButton{{background-color:{ACCENT};color:#06101b;border:none;border-radius:20px;font-size:16px;font-weight:900;}} QPushButton:hover{{background-color:#8ce3ff;}} QPushButton:pressed{{background-color:#58c9f2;}}")
    def login(self):
        pwd = self.input.text().strip()
        if not pwd: return QMessageBox.warning(self,'تنبيه','دخل كلمة المرور')
        try:
            INVOICES_DIR.mkdir(exist_ok=True)
            ATTACHMENTS_DIR.mkdir(exist_ok=True)
        except Exception:
            pass
        if not PASSWORD_FILE.exists():
            PASSWORD_FILE.write_text(password_record_for_storage(pwd), encoding='utf-8')
        saved = PASSWORD_FILE.read_text(encoding='utf-8').strip()
        if not verify_password_input(pwd, saved):
            self.error_label.setText('كلمة المرور غير صحيحة')
            return QMessageBox.warning(self,'خطأ','كلمة المرور غير صحيحة')
        if saved and not saved.startswith('sha256$'):
            try: PASSWORD_FILE.write_text(password_record_for_storage(pwd), encoding='utf-8')
            except Exception: pass
        self.main_window = MainWindow(); self.main_window.showMaximized(); self.close()


def main():
    for folder in [INVOICES_DIR, ATTACHMENTS_DIR, BACKUPS_DIR]:
        try:
            folder.mkdir(parents=True, exist_ok=True)
        except Exception:
            pass

    app = QApplication(sys.argv)
    app.setQuitOnLastWindowClosed(False)
    try:
        if APP_ICON.exists():
            app.setWindowIcon(QIcon(str(APP_ICON)))
    except Exception:
        pass

    splash = SplashScreen()
    login = LoginWindow()

    # Keep references alive for the full app lifetime.
    app._splash = splash
    app._login = login

    def show_login():
        login.show()
        login.activateWindow()
        login.raise_()
        app.setQuitOnLastWindowClosed(True)

    QTimer.singleShot(2700, show_login)
    splash.show_centered()
    sys.exit(app.exec())




# ===== أدوات الفحص والتدقيق والاستقرار =====
def _sum_float_rows(rows, key):
    total = 0.0
    for row in rows:
        try:
            total += float(row.get(key, 0) or 0)
        except Exception:
            pass
    return total


def build_health_check_report(db):
    issues = []
    notes = []
    critical = 0
    warning = 0

    # basic structure
    required_keys = ['items', 'customers', 'suppliers', 'funders', 'inbound', 'sales', 'cash', 'returns', 'damaged', 'expenses', 'operations_log', 'opening']
    for key in required_keys:
        if key not in db:
            critical += 1
            issues.append(f'حرج: المفتاح الأساسي مفقود من قاعدة البيانات: {key}')

    # duplicate ids across major ledgers
    for list_key in ['sales', 'inbound', 'cash', 'returns', 'damaged', 'expenses', 'profit_distributions']:
        seen = set()
        dup = 0
        missing = 0
        for row in db.get(list_key, []):
            rid = str(row.get('id', '') or '').strip()
            if not rid:
                missing += 1
                continue
            if rid in seen:
                dup += 1
            seen.add(rid)
        if missing:
            warning += 1
            issues.append(f'مراجعة: يوجد {missing} سجل داخل {list_key} بدون id')
        if dup:
            critical += 1
            issues.append(f'حرج: يوجد {dup} id مكرر داخل {list_key}')

    # orphaned references
    item_names = {str(x.get('name', '')).strip() for x in db.get('items', [])}
    missing_item_refs = 0
    for list_key in ['inbound', 'sales', 'returns', 'damaged', 'inventory_ledger']:
        for row in db.get(list_key, []):
            nm = str(row.get('item', '') or row.get('item_name', '') or '').strip()
            if nm and nm not in item_names:
                missing_item_refs += 1
    if missing_item_refs:
        warning += 1
        issues.append(f'مراجعة: يوجد {missing_item_refs} حركة مرتبطة بصنف غير موجود بجدول الأصناف')

    # cash receipt numbers
    receipt_missing = sum(1 for x in db.get('cash', []) if (x.get('source') in ('customer_payment', 'supplier_payment', 'return_cash_payout')) and not str(x.get('receipt_no', '')).strip())
    if receipt_missing:
        warning += 1
        issues.append(f'مراجعة: يوجد {receipt_missing} حركة نقدية مهمة بدون رقم وصل')

    # negative stock
    try:
        neg = 0
        for row in item_stock_rows(db):
            qty = float(row.get('qty', 0) or 0)
            if qty < 0:
                neg += 1
        if neg:
            critical += 1
            issues.append(f'حرج: يوجد {neg} صنف بكميات سالبة بالمخزن')
    except Exception as e:
        warning += 1
        issues.append(f'مراجعة: تعذر فحص كميات المخزن: {e}')

    if not issues:
        status = 'سليم'
        notes.append('✔️ لا توجد مؤشرات حرجة ظاهرة في فحص السلامة.')
    elif critical:
        status = 'حرج'
    else:
        status = 'يحتاج مراجعة'

    lines = [
        '🩺 فحص سلامة النظام',
        '====================',
        f'الحالة: {status}',
        f'المشاكل الحرجة: {critical}',
        f'ملاحظات المراجعة: {warning}',
        '',
    ]
    if issues:
        lines.extend(f'- {x}' for x in issues)
    if notes:
        lines.append('')
        lines.extend(notes)
    return {'status': status, 'critical': critical, 'warning': warning, 'lines': lines}


def build_accounting_audit_report(db):
    issues = []
    critical = 0
    warning = 0

    # cash consistency
    try:
        cb = cash_breakdown(db)
        final_cash = float(cb.get('final_cash', 0) or 0)
        cash_rows_total = _sum_float_rows([x for x in db.get('cash', []) if str(x.get('type', '')) == 'إيراد'], 'amount') - _sum_float_rows([x for x in db.get('cash', []) if str(x.get('type', '')) == 'مصروف'], 'amount')
        # This is not the same formula as the official cash formula, so we use it as a signal only
        if abs(cash_rows_total) < 0.01 and abs(final_cash) > 0.01:
            warning += 1
            issues.append('مراجعة: سجل الصندوق يوحي بحركة صفر تقريبًا بينما القاصة النهائية ليست صفرًا. راجع التهيئة الافتتاحية.')
    except Exception as e:
        warning += 1
        issues.append(f'مراجعة: تعذر تدقيق القاصة: {e}')

    # customer dues sanity
    try:
        for row in customer_due_summary(db):
            if float(row.get('due', 0) or 0) < -0.01:
                critical += 1
                issues.append(f"حرج: دين زبون سالب بشكل غير منطقي: {row.get('name', '')}")
    except Exception as e:
        warning += 1
        issues.append(f'مراجعة: تعذر تدقيق ديون الزبائن: {e}')

    # supplier dues sanity
    try:
        for row in supplier_due_summary(db):
            if float(row.get('due', 0) or 0) < -0.01:
                critical += 1
                issues.append(f"حرج: ذمة مورد سالبة بشكل غير منطقي: {row.get('name', '')}")
    except Exception as e:
        warning += 1
        issues.append(f'مراجعة: تعذر تدقيق الموردين: {e}')

    # return credit sanity
    bad_returns = 0
    for row in db.get('returns', []):
        credit_amount = float(row.get('credit_amount', 0) or 0)
        credit_used = float(row.get('credit_used', 0) or 0)
        cash_paid_out = float(row.get('cash_paid_out', 0) or 0)
        if credit_used + cash_paid_out - credit_amount > 0.01:
            bad_returns += 1
    if bad_returns:
        critical += 1
        issues.append(f'حرج: يوجد {bad_returns} مرتجع مجموع استخدامه وصرفه أكبر من أصل الرصيد')

    # فحص منطق العجز لهضاب/مصطفى
    try:
        snap = profit_ui_snapshot(db)
        pb = profit_breakdown(db)
        partnership_period_profit = float(pb.get('distributable_profit', 0) or 0) + float(pb.get('owner_capital_profit', 0) or 0)
        period_share_each = round(partnership_period_profit / 2.0, 2)
        hidab_expected = round(max(0.0, float(snap.get('hidab_opening_balance', 0) or 0) + float(snap.get('hidab_withdrawals', 0) or 0) + owner_profit_payment_sum(db, 'هضاب') - period_share_each), 2)
        mostafa_expected = round(max(0.0, float(snap.get('mostafa_opening_balance', 0) or 0) + float(snap.get('mostafa_withdrawals', 0) or 0) + owner_profit_payment_sum(db, 'مصطفى') - period_share_each), 2)
        person = person_profit_status(db)
        if round(float(person.get('hidab_deficit', 0) or 0), 2) != hidab_expected:
            critical += 1
            issues.append(f'حرج: معادلة عجز هضاب غير مطابقة للتوقع الحسابي ({fmt_money(hidab_expected)} د.ع)')
        if round(float(person.get('mostafa_deficit', 0) or 0), 2) != mostafa_expected:
            critical += 1
            issues.append(f'حرج: معادلة عجز مصطفى غير مطابقة للتوقع الحسابي ({fmt_money(mostafa_expected)} د.ع)')
    except Exception as e:
        warning += 1
        issues.append(f'مراجعة: تعذر تدقيق منطق الأرباح والعجز: {e}')

    # orphan cash parties
    known_customers = {str(x.get('name', '')).strip() for x in db.get('customers', [])}
    known_suppliers = {str(x.get('name', '')).strip() for x in db.get('suppliers', [])}
    orphan = 0
    for row in db.get('cash', []):
        src = row.get('source')
        party = str(row.get('party', '')).strip()
        if src == 'customer_payment' and party and party not in known_customers:
            orphan += 1
        if src == 'supplier_payment' and party and party not in known_suppliers:
            orphan += 1
    if orphan:
        warning += 1
        issues.append(f'مراجعة: يوجد {orphan} حركة صندوق مرتبطة بطرف غير موجود')

    # invoice/report consistency
    try:
        seen_invoices = set()
        invoice_mismatches = []
        for sale in db.get('sales', []):
            gid = sale.get('invoice_group_id') or f"single-{ensure_invoice_no(db, sale)}"
            if gid in seen_invoices:
                continue
            seen_invoices.add(gid)
            lines = get_invoice_sales(db, sale)
            invoice_total = round(sum(_safe_float(x.get('total', 0)) for x in lines), 2)
            current_due = round(max(0.0, current_sale_due(db, gid)), 2)
            current_paid = round(max(0.0, invoice_total - current_due), 2)
            if current_due - invoice_total > 0.01:
                invoice_mismatches.append(f'مراجعة: الفاتورة #{ensure_invoice_no(db, sale)} متبقيها الحالي أكبر من إجماليها.')
            if current_paid < -0.01:
                invoice_mismatches.append(f'مراجعة: الفاتورة #{ensure_invoice_no(db, sale)} تعطي مقبوضاً سالباً وهو غير منطقي.')
        if invoice_mismatches:
            warning += len(invoice_mismatches)
            issues.extend(invoice_mismatches)
    except Exception as e:
        warning += 1
        issues.append(f'مراجعة: تعذر تدقيق الفواتير المعروضة: {e}')

    # profit sanity
    try:
        op_profit = float(operating_profit(db) or 0)
        paid_profit = _sum_float_rows(db.get('profit_distributions', []), 'amount')
        if paid_profit - op_profit > 0.01:
            warning += 1
            issues.append('مراجعة: إجمالي الأرباح المدفوعة أكبر من الربح التشغيلي الحالي')
    except Exception as e:
        warning += 1
        issues.append(f'مراجعة: تعذر تدقيق الأرباح: {e}')

    if not issues:
        status = 'سليم'
    elif critical:
        status = 'حرج'
    else:
        status = 'يحتاج مراجعة'

    lines = [
        '🧪 التدقيق المحاسبي النهائي',
        '===========================',
        f'الحالة: {status}',
        f'المشاكل الحرجة: {critical}',
        f'ملاحظات المراجعة: {warning}',
        '',
    ]
    if issues:
        lines.extend(f'- {x}' for x in issues)
    else:
        lines.append('✔️ لم يظهر خلل محاسبي واضح بالفحص الحالي.')
    return {'status': status, 'critical': critical, 'warning': warning, 'lines': lines}


def build_stability_report(db):
    issues = []
    critical = 0
    warning = 0

    cash_now = float(cash_balance(db) or 0)
    sales_count = len(db.get('sales', []))
    inbound_count = len(db.get('inbound', []))
    returns_count = len(db.get('returns', []))
    expenses_count = len(db.get('expenses', []))

    try:
        customer_due_total = sum(float(x.get('due', 0) or 0) for x in customer_due_summary(db))
    except Exception:
        customer_due_total = 0.0
        warning += 1
        issues.append('مراجعة: تعذر حساب إجمالي ديون الزبائن')

    try:
        supplier_due_total = sum(float(x.get('due', 0) or 0) for x in supplier_due_summary(db))
    except Exception:
        supplier_due_total = 0.0
        warning += 1
        issues.append('مراجعة: تعذر حساب إجمالي ديون الموردين')

    try:
        stock_rows = item_stock_rows(db)
        stock_count = len(stock_rows)
        stock_value = sum(float(x.get('qty', 0) or 0) * float(x.get('buy_price', 0) or 0) for x in stock_rows)
    except Exception:
        stock_count = 0
        stock_value = 0.0
        warning += 1
        issues.append('مراجعة: تعذر حساب قيمة المخزن')

    try:
        op_profit = float(operating_profit(db) or 0)
    except Exception:
        op_profit = 0.0
        warning += 1
        issues.append('مراجعة: تعذر حساب الربح التشغيلي')

    # very light stability checks
    if sales_count and inbound_count == 0:
        warning += 1
        issues.append('مراجعة: توجد مبيعات بدون واردات مسجلة بالنظام')
    if stock_count == 0 and (sales_count or inbound_count):
        warning += 1
        issues.append('مراجعة: لا توجد أصناف نهائية مع وجود حركات بيع/وارد')
    if cash_now < -0.01:
        critical += 1
        issues.append('حرج: القاصة سالبة حاليًا')

    if not issues:
        status = 'سليم'
    elif critical:
        status = 'حرج'
    else:
        status = 'يحتاج مراجعة'

    lines = [
        '🧱 تقرير الاستقرار النهائي',
        '==========================',
        f'الحالة: {status}',
        '',
        'الملخص الرقمي:',
        f'- القاصة الحالية: {cash_now:,.0f} د.ع',
        f'- قيمة المخزن التقديرية: {stock_value:,.0f} د.ع',
        f'- عدد الأصناف الفعالة: {stock_count}',
        f'- ديون الزبائن: {customer_due_total:,.0f} د.ع',
        f'- ذمم الموردين: {supplier_due_total:,.0f} د.ع',
        f'- الربح التشغيلي: {op_profit:,.0f} د.ع',
        f'- عدد فواتير البيع: {sales_count}',
        f'- عدد الوارد: {inbound_count}',
        f'- عدد المرتجعات: {returns_count}',
        f'- عدد المصاريف: {expenses_count}',
        '',
    ]
    if issues:
        lines.append('الملاحظات:')
        lines.extend(f'- {x}' for x in issues)
    else:
        lines.append('✔️ المؤشرات العامة مستقرة حسب الفحص الحالي.')
    return {'status': status, 'critical': critical, 'warning': warning, 'lines': lines}


def export_text_report(title, lines):
    stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
    reports_dir = APP_DATA_DIR / 'reports'
    reports_dir.mkdir(parents=True, exist_ok=True)
    safe_title = ''.join(ch if ch.isalnum() or ch in ('_', '-') else '_' for ch in title)
    path = reports_dir / f'{safe_title}_{stamp}.pdf'
    save_lines_pdf(title, lines, path)
    return path


def _show_text_report_dialog(parent, title, report, export_name):
    dlg = QDialog(parent)
    dlg.setWindowTitle(title)
    dlg.resize(820, 620)
    apply_branding(dlg)
    layout = QVBoxLayout(dlg)

    status = QLabel(f"الحالة: {report.get('status', '')}")
    status.setAlignment(Qt.AlignRight)
    status.setStyleSheet('font-size:18px;font-weight:800;')
    layout.addWidget(status)

    box = QTextEdit()
    box.setReadOnly(True)
    box.setPlainText('\n'.join(report.get('lines', [])))
    layout.addWidget(box, 1)

    btns = QHBoxLayout()
    export_btn = QPushButton('📄 تصدير PDF')
    close_btn = QPushButton('إغلاق')

    def do_export():
        try:
            path = export_text_report(export_name, report.get('lines', []))
            QMessageBox.information(dlg, 'تم', f'تم حفظ التقرير هنا:\n{path}')
        except Exception as e:
            QMessageBox.warning(dlg, 'خطأ', f'تعذر تصدير التقرير:\n{e}')

    export_btn.clicked.connect(do_export)
    close_btn.clicked.connect(dlg.accept)
    btns.addWidget(export_btn)
    btns.addWidget(close_btn)
    layout.addLayout(btns)
    dlg.exec()


def _mw_show_health_check_dialog(self):
    try:
        report = build_health_check_report(self.db)
        self.db.setdefault('operations_log', []).append({
            'id': safe_new_id('health'),
            'date': today_str(),
            'created_at': now_str(),
            'action': 'system_health_check',
            'details': f"الحالة: {report.get('status', '')}",
        })
        save_db(self.db)
        _show_text_report_dialog(self, '🩺 فحص سلامة النظام', report, 'system_health_check')
    except Exception as e:
        QMessageBox.warning(self, 'خطأ', f'تعذر تشغيل فحص سلامة النظام\n{e}')


def _mw_show_accounting_audit_dialog(self):
    try:
        report = build_accounting_audit_report(self.db)
        self.db.setdefault('operations_log', []).append({
            'id': safe_new_id('audit'),
            'date': today_str(),
            'created_at': now_str(),
            'action': 'accounting_audit',
            'details': f"الحالة: {report.get('status', '')}",
        })
        save_db(self.db)
        _show_text_report_dialog(self, '🧪 التدقيق المحاسبي النهائي', report, 'accounting_audit')
    except Exception as e:
        QMessageBox.warning(self, 'خطأ', f'تعذر تشغيل التدقيق المحاسبي\n{e}')


def _mw_show_stability_report_dialog(self):
    try:
        report = build_stability_report(self.db)
        self.db.setdefault('operations_log', []).append({
            'id': safe_new_id('stability'),
            'date': today_str(),
            'created_at': now_str(),
            'action': 'stability_report',
            'details': f"الحالة: {report.get('status', '')}",
        })
        save_db(self.db)
        _show_text_report_dialog(self, '🧱 تقرير الاستقرار النهائي', report, 'stability_report')
    except Exception as e:
        QMessageBox.warning(self, 'خطأ', f'تعذر تشغيل تقرير الاستقرار\n{e}')


def _mw_open_app_data_folder(self):
    try:
        APP_DATA_DIR.mkdir(parents=True, exist_ok=True)
        if sys.platform.startswith('win'):
            os.startfile(str(APP_DATA_DIR))
        else:
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(APP_DATA_DIR)))
    except Exception as e:
        QMessageBox.warning(self, 'خطأ', f'تعذر فتح مجلد البيانات:\n{e}')


# ربطها داخل MainWindow بشكل فعلي
MainWindow.show_health_check_dialog = _mw_show_health_check_dialog
MainWindow.show_accounting_audit_dialog = _mw_show_accounting_audit_dialog
MainWindow.show_stability_report_dialog = _mw_show_stability_report_dialog
MainWindow.open_app_data_folder = _mw_open_app_data_folder
# ===== نهاية الأدوات =====

try:
    from tod_exact_patch import apply_patch as _apply_tod_exact_patch
    _apply_tod_exact_patch(globals())
except Exception:
    pass

if __name__ == '__main__':
    main()